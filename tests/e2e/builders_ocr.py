"""Input builders and output readers for the OCR-lane E2E tests (tests/e2e/test_ocr.py).

Nothing from ``doc2mark`` is imported here: inputs are made with PyMuPDF/Pillow and outputs are
read the way a consumer would, by parsing the Markdown the CLI wrote (Python-Markdown for
rendering, lxml for the HTML tables).
"""

import glob
import io
import re
import shutil
import subprocess
from pathlib import Path
from typing import List, Optional, Sequence

import markdown as python_markdown
import pymupdf
from bs4 import BeautifulSoup, Comment, ProcessingInstruction
from lxml import html as lxml_html
from PIL import Image, ImageDraw, ImageFont

from tests.e2e import pdfgen

# --------------------------------------------------------------------------- #
# Inputs                                                                      #
# --------------------------------------------------------------------------- #


def scan_pdf(path: Path, pages: pdfgen.Pages = "SCANNED PAGE") -> Path:
    """Image-only PDF (one full-page picture per page, no text layer): the loader routes it to
    whole-page OCR, so every page produces exactly one OCR request."""
    return pdfgen.image_pdf(path, pages)


def docx_picture_in_cell(path: Path) -> Path:
    """DOCX with a 2x2 table (``Item`` | ``Picture``, ``Logo`` | <picture>): the picture's OCR
    text goes into its cell as ``[Image: <text>]``."""
    from docx import Document
    from docx.shared import Inches

    picture = Image.new("RGB", (240, 120), "white")
    ImageDraw.Draw(picture).rectangle((20, 30, 220, 90), outline="black", width=6)
    buffer = io.BytesIO()
    picture.save(buffer, format="PNG")
    document = Document()
    document.add_paragraph("Parts list")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Picture"
    table.cell(1, 0).text = "Logo"
    table.cell(1, 1).paragraphs[0].add_run().add_picture(io.BytesIO(buffer.getvalue()), width=Inches(0.8))
    document.save(str(path))
    return Path(path)


def cjk_font(script: str = "TC", size: int = 110) -> ImageFont.FreeTypeFont:
    """The Noto Sans CJK face for ``script`` ("TC" Traditional, "SC" Simplified, "JP", "KR"),
    from the system fonts (``fonts-noto-cjk`` in the E2E image). Raises if it is not installed."""
    candidates = sorted(glob.glob("/usr/share/fonts/**/NotoSansCJK*-Regular.ttc", recursive=True))
    if not candidates and shutil.which("fc-match"):
        found = subprocess.run(["fc-match", "-f", "%{file}", "Noto Sans CJK TC"], capture_output=True, text=True)
        if found.stdout.endswith(".ttc"):
            candidates = [found.stdout]
    for path in candidates:
        for index in range(12):
            try:
                font = ImageFont.truetype(path, size, index=index)
            except OSError:
                break
            if font.getname()[0].endswith(f"CJK {script}"):
                return font
    raise FileNotFoundError(f"Noto Sans CJK {script} is not installed (fonts-noto-cjk)")


def text_scan_pdf(path: Path, lines: Sequence[str], font: ImageFont.FreeTypeFont) -> Path:
    """One-page image-only PDF showing ``lines`` in ``font`` (black on white, A4 at 200 dpi)."""
    width, height = pdfgen.A4_PIXELS_200DPI
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    y = pdfgen.MARGIN_PIXELS
    for line in lines:
        right, bottom = draw.textbbox((pdfgen.MARGIN_PIXELS, y), line, font=font)[2:]
        if right > width - pdfgen.MARGIN_PIXELS or bottom > height - pdfgen.MARGIN_PIXELS:
            raise ValueError(f"line does not fit the page: {line!r}")
        draw.text((pdfgen.MARGIN_PIXELS, y), line, fill="black", font=font)
        y = bottom + font.size // 2
    buffer = io.BytesIO()
    image.save(buffer, format="PNG", dpi=(200, 200))
    doc = pymupdf.open()
    page = doc.new_page(width=pdfgen.A4_POINTS[0], height=pdfgen.A4_POINTS[1])
    page.insert_image(page.rect, stream=buffer.getvalue())
    doc.save(str(path))
    doc.close()
    return Path(path)


# --------------------------------------------------------------------------- #
# Output readers                                                              #
# --------------------------------------------------------------------------- #

_WS = re.compile(r"\s+")
_ACTIVE_TAGS = {
    "script", "img", "svg", "iframe", "object", "embed", "style", "link", "meta", "base", "form",
    "input", "button", "math", "textarea", "xmp", "noscript", "video", "audio", "source", "frame",
}


def squash(text: str) -> str:
    """``text`` with every whitespace run removed (CJK OCR output may space characters apart)."""
    return _WS.sub("", text)


def normalize(text: str) -> str:
    """``text`` with whitespace runs collapsed to single spaces."""
    return _WS.sub(" ", text).strip()


def render(markdown: str) -> BeautifulSoup:
    """The CLI's Markdown rendered to HTML (Python-Markdown with tables), as a consumer would."""
    return BeautifulSoup(python_markdown.markdown(markdown, extensions=["tables"]), "html.parser")


def visible_text(markdown: str) -> str:
    """What a reader sees: the rendered Markdown's text, whitespace-normalized."""
    return normalize(render(markdown).get_text(" "))


def active_html(markdown: str) -> List[str]:
    """Live markup in the rendered Markdown: script-capable elements, event-handler or
    ``javascript:`` attributes, and comments/PIs that hide markup. Empty means inert."""
    found = []
    soup = render(markdown)
    for element in soup.find_all(True):
        if element.name in _ACTIVE_TAGS:
            found.append(f"<{element.name}>")
        for name, value in element.attrs.items():
            text = " ".join(value) if isinstance(value, list) else str(value)
            if name.lower().startswith("on") or "javascript:" in text.lower():
                found.append(f"<{element.name} {name}={text!r}>")
    for node in soup.find_all(string=lambda s: isinstance(s, (Comment, ProcessingInstruction))):
        if "<" in node or "script" in node.lower() or "[if" in node.lower():
            found.append(f"comment {str(node)[:60]!r}")
    return found


def html_tables(markdown: str, *, nested: bool = False) -> List[List[List[Optional[str]]]]:
    """Every ``<table>`` in the Markdown as a grid laid out the way a browser does it (rowspan and
    colspan honoured, ``rowspan="0"`` spanning to the end of its row group, per the HTML spec).

    A grid is a list of rows; a row lists the text of the cell that *starts* at each column
    (``<br>`` read as a space, whitespace normalized), or ``None`` where a cell from an earlier
    row or column covers the position. Top-level tables only, unless ``nested=True`` (then every
    table, outer ones first). A row's length is the number of columns it occupies.
    """
    fragment = lxml_html.fragment_fromstring(f"<div>{markdown}</div>")
    tables = [t for t in fragment.iter("table") if nested or not any(a.tag == "table" for a in t.iterancestors())]
    return [_grid(table) for table in tables]


def _rows(children) -> list:
    """The rows among these children. Cells placed directly in a table or row group form a
    row of their own, as a browser's parser gives them one (``<table><td>`` gets a ``<tr>``)."""
    rows, cells = [], []
    for child in children:
        if not isinstance(child.tag, str):  # comments
            continue
        if child.tag in ("td", "th"):
            cells.append(child)
            continue
        if cells:
            rows.append(cells)
            cells = []
        if child.tag == "tr":
            rows.append(child)
    if cells:
        rows.append(cells)
    return rows


def _row_groups(table) -> List[list]:
    groups, loose = [], []
    for child in table:
        if child.tag in ("thead", "tbody", "tfoot"):
            if loose:
                groups.append(_rows(loose))
                loose = []
            groups.append(_rows(child))
        else:
            loose.append(child)
    if loose:
        groups.append(_rows(loose))
    return groups


def _cell_text(cell) -> str:
    copy = lxml_html.fromstring(lxml_html.tostring(cell, with_tail=False))  # the document stays intact
    for element in copy.iter("br", "td", "th", "tr"):
        if element is not copy:
            element.tail = " " + (element.tail or "")
    return normalize(copy.text_content())


def _span(cell, name: str) -> int:
    value = (cell.get(name) or "1").strip()
    return int(value) if value.isdigit() else 1


def _grid(table) -> List[List[Optional[str]]]:
    grid: List[List[Optional[str]]] = []
    for group in _row_groups(table):
        covered = {}  # (row index in group, column) -> True
        for r, row in enumerate(group):
            cells, col = {}, 0
            for cell in (c for c in row if c.tag in ("td", "th")):
                while (r, col) in covered:
                    col += 1
                # HTML table model: rowspan 0 and rowspans past the group end at the group's
                # last row; colspan 0 is 1 and colspan is capped at 1000.
                rowspan = min(_span(cell, "rowspan") or len(group), len(group) - r)
                colspan = min(max(1, _span(cell, "colspan")), 1000)
                cells[col] = _cell_text(cell)
                for dr in range(rowspan):
                    for dc in range(colspan):
                        covered[(r + dr, col + dc)] = True
                col += colspan
            width = max([c + 1 for (rr, c) in covered if rr == r], default=0)
            grid.append([cells.get(c) for c in range(width)])
    return grid


_MD_ESCAPED_PUNCT = re.compile(r"\\([!\"#$%&'()*+,\-./:;<=>?@\[\\\]^_`{|}~])")


def pipe_table_rows(markdown: str) -> List[List[str]]:
    """Header and body rows of the GFM pipe tables in the Markdown (delimiter rows left out),
    split the way GFM/markdown-it do: a ``|`` right after a backslash belongs to the cell, then
    backslash escapes of ASCII punctuation are resolved."""
    rows = []
    for line in markdown.splitlines():
        stripped = line.strip()
        if not stripped.startswith("|"):
            continue
        body = stripped[1:]
        if body.endswith("|") and not body.endswith("\\|"):
            body = body[:-1]
        cells, current, last = [], "", ""
        for ch in body:
            if ch == "|" and last != "\\":
                cells.append(current)
                current = ""
            elif ch == "|":
                current = current[:-1] + "|"
            else:
                current += ch
            last = ch
        cells.append(current)
        cells = [normalize(_MD_ESCAPED_PUNCT.sub(r"\1", c)) for c in cells]
        if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            continue
        rows.append(cells)
    return rows
