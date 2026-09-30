"""Input builders for the pdftext E2E tests (``tests/e2e/test_pdftext.py``).

Each builder ports a reviewer probe (``~/code/.executors/d2m-rv-headings-probes`` and
``~/code/.executors/d2m-rv-tables-probes``) into a document made at test time with
PyMuPDF, python-docx and, for Word-exported PDFs, LibreOffice. Nothing from
``doc2mark`` is imported. Every builder saves to ``path`` and returns it as a ``Path``.

Body text carries unique ``BODY<page>x<line>`` tokens so a test can check that every
body line survives exactly once.
"""

import random
import shutil
import subprocess
from pathlib import Path
from typing import Callable, Dict, List, Optional, Sequence, Tuple

import pymupdf

A4 = (595, 842)
SLIDE = (960, 540)
WIDE_SLIDE = (1440, 810)

WORDS = "alpha beta gamma delta sigma omega river stone cloud metal paper glass forest ocean".split()

# Page-number and running-footer forms from the review (p03_header_footer.py, p03b, p03e).
PAGE_NUMBER_FORMS: Dict[str, Callable[[int, int], str]] = {
    "bare": lambda p, n: f"{p}",
    "Page N": lambda p, n: f"Page {p}",
    "Page N of M": lambda p, n: f"Page {p} of {n}",
    "- N -": lambda p, n: f"- {p} -",
    "N/M": lambda p, n: f"{p}/{n}",
    "N | ACME": lambda p, n: f"{p} | ACME Corp",
    "ACME | Page N": lambda p, n: f"ACME Corp | Page {p}",
    "p. N": lambda p, n: f"p. {p}",
    "roman": lambda p, n: ["i", "ii", "iii", "iv", "v", "vi", "vii", "viii"][p - 1],
    "PAGE N": lambda p, n: f"PAGE {p}",
}

EUROPEAN_PAGE_NUMBER_FORMS: Dict[str, Callable[[int, int], str]] = {
    "Seite N von M": lambda p, n: f"Seite {p} von {n}",
    "Página N de M": lambda p, n: f"Página {p} de {n}",
    "Page N sur M": lambda p, n: f"Page {p} sur {n}",
    "Pagina N di M": lambda p, n: f"Pagina {p} di {n}",
}

CJK_PAGE_NUMBER_FORMS: Dict[str, Callable[[int, int], str]] = {
    "第N頁": lambda p, n: f"第 {p} 頁",
    "第N頁共M頁": lambda p, n: f"第{p}頁，共{n}頁",
    "共M頁第N頁": lambda p, n: f"共 {n} 頁 第 {p} 頁",
    "fullwidth N": lambda p, n: f"－{chr(0xFF10 + p)}－",
}


def body_lines(page: int, count: int = 10, words: int = 11) -> List[str]:
    """``count`` lines of seeded filler words, each starting with a unique ``BODY<page>x<i>`` token."""
    rnd = random.Random(page * 7919 + count)
    return [f"BODY{page}x{i} " + " ".join(rnd.choice(WORDS) for _ in range(words)) for i in range(count)]


def _text(page, x: float, y: float, text: str, size: float = 10.5, font: str = "helv", **kwargs) -> None:
    """``insert_text`` with (x, y) as the baseline origin of the first line."""
    page.insert_text((x, y), text, fontsize=size, fontname=font, **kwargs)


def _save(doc, path: Path) -> Path:
    doc.save(str(path))
    doc.close()
    return Path(path)


# --------------------------------------------------------------------------------------
# Word-exported documents (python-docx -> LibreOffice -> PDF)
# --------------------------------------------------------------------------------------

def docx_to_pdf(docx_path: Path) -> Path:
    """Convert ``docx_path`` to PDF next to it with LibreOffice headless (private profile per call)."""
    docx_path = Path(docx_path)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        raise RuntimeError("soffice is not on PATH")
    profile = docx_path.parent / f".lo-profile-{docx_path.stem}"
    subprocess.run(
        [soffice, f"-env:UserInstallation={profile.resolve().as_uri()}", "--headless",
         "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
        check=True, capture_output=True, timeout=240,
    )
    pdf = docx_path.with_suffix(".pdf")
    if not pdf.exists():
        raise RuntimeError(f"LibreOffice did not write {pdf}")
    return pdf


def word_tables_with_captions(path: Path) -> Path:
    """Port of ``mk_docx_tables_adjacent.py``: ruled 'Table Grid' tables with a caption paragraph
    directly above and a note paragraph directly below (default spacing, and zero spacing for C).

    Cell tokens are unique per table (``AR1C0`` ... ``CR3C2``); header rows read Region/Revenue/Growth.
    Returns the .docx path; convert it with :func:`docx_to_pdf`.
    """
    from docx import Document
    from docx.shared import Pt

    def table(doc, label):
        t = doc.add_table(rows=4, cols=3)
        t.style = "Table Grid"
        for r in range(4):
            for c in range(3):
                t.cell(r, c).text = ["Region", "Revenue", "Growth"][c] if r == 0 else f"{label}R{r}C{c}"
        return t

    d = Document()
    d.add_heading("Quarterly Results", level=1)
    d.add_paragraph("CAPTION-A Table 1: Revenue by region (USD m)")
    table(d, "A")
    d.add_paragraph("NOTE-A Figures are unaudited and subject to change.")
    d.add_paragraph("BODY-A The next section discusses costs in more detail for each segment of the business.")
    d.add_heading("HEADING-B Cost Analysis", level=2)
    table(d, "B")
    d.add_paragraph("NOTE-B Costs exclude one-off restructuring charges.")
    p = d.add_paragraph("CAPTION-C Revenue split by channel")
    p.paragraph_format.space_after = Pt(0)
    table(d, "C")
    p = d.add_paragraph("NOTE-C Source: internal management accounts.")
    p.paragraph_format.space_before = Pt(0)
    d.add_paragraph("BODY-C Closing remarks for the quarter follow in the outlook section of this report.")
    d.save(str(path))
    return Path(path)


def _word_field(paragraph, instr: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    fld.append(run)
    paragraph._p.append(fld)


def word_report_with_chapter_headers(path: Path) -> Path:
    """Port of ``mk_docx_report2b.py``: four chapters with per-chapter running headers
    ("Chapter N - Title"), a "Page X of Y" footer, unique ``Paragraph N:`` body paragraphs, and a
    borderless table (Date / Description / Amount / Balance) that starts at the top of a page and
    repeats its header row on every page it spans. Compared with the probe, chapters without the table
    get 12 paragraphs (not 9) so each spans at least two pages, and the table has 240 rows (not 110) so
    its header row sits at the top of more than half of the pages whatever LibreOffice's pagination.

    Returns the .docx path; convert it with :func:`docx_to_pdf`.
    """
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    words = ("growth margin customer region product platform investment strategy capital supply "
             "demand pricing inventory logistics partner contract service quality risk audit "
             "compliance digital cloud security revenue cost forecast budget market share").split()

    def para(i):
        rnd = random.Random(i)
        return f"Paragraph {i}: " + " ".join(rnd.choice(words) for _ in range(70)) + "."

    d = Document()
    n = 0
    for ci, chapter in enumerate(["Introduction", "Market Review", "Financial Statements", "Risk Management"], 1):
        sec = d.sections[0] if ci == 1 else d.add_section(WD_SECTION.NEW_PAGE)
        sec.header.is_linked_to_previous = False
        sec.header.paragraphs[0].text = f"Chapter {ci} - {chapter}"
        sec.footer.is_linked_to_previous = False
        footer = sec.footer.paragraphs[0]
        footer.add_run("Page ")
        _word_field(footer, "PAGE")
        footer.add_run(" of ")
        _word_field(footer, "NUMPAGES")
        d.add_heading(f"{ci} {chapter}", level=1)
        if chapter == "Financial Statements":
            d.add_paragraph(para(n))
            n += 1
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)  # the table starts at the top of a page
            t = d.add_table(rows=1, cols=4)  # default style: no borders
            header = t.rows[0]
            tr_pr = header._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
            for cell, text in zip(header.cells, ["Date", "Description", "Amount", "Balance"]):
                cell.text = text
            balance = 1000
            for r in range(240):
                amount = (r * 37) % 500 + 10
                balance += amount
                row = t.add_row().cells
                row[0].text = f"2024-{(r % 12) + 1:02d}-{(r % 28) + 1:02d}"
                row[1].text = f"Invoice INV-{7000 + r} payment"
                row[2].text = f"{amount}"
                row[3].text = f"{balance}"
            d.add_paragraph(para(n))
            n += 1
        else:
            for _ in range(12):  # 2+ pages per chapter, so every running header repeats
                d.add_paragraph(para(n))
                n += 1
    d.save(str(path))
    return Path(path)


# --------------------------------------------------------------------------------------
# Table suppression (H-F1 / T5)
# --------------------------------------------------------------------------------------

def ruled_table(page, x0: float, y0: float, col_widths: Sequence[float], row_heights: Sequence[float],
                rows: Sequence[Sequence[str]], fontsize: float = 9, fontname: str = "helv") -> pymupdf.Rect:
    """Draw a ruled grid (one outlined rectangle per cell) with ``rows[r][c]`` written in each cell."""
    xs = [x0]
    for width in col_widths:
        xs.append(xs[-1] + width)
    ys = [y0]
    for height in row_heights:
        ys.append(ys[-1] + height)
    for r in range(len(row_heights)):
        for c in range(len(col_widths)):
            rect = pymupdf.Rect(xs[c], ys[r], xs[c + 1], ys[r + 1])
            page.draw_rect(rect, color=(0, 0, 0), width=0.6)
            text = rows[r][c] if r < len(rows) and c < len(rows[r]) else ""
            if text:
                page.insert_text((rect.x0 + 3, rect.y0 + fontsize + 2), text, fontname=fontname, fontsize=fontsize)
    return pymupdf.Rect(xs[0], ys[0], xs[-1], ys[-1])


def _logo(page, x0: float = 40, y0: float = 30, brand: str = "ACME") -> None:
    """Logo built like the reviewed deck's: two adjacent filled rectangles with white brand text."""
    page.draw_rect(pymupdf.Rect(x0, y0, x0 + 30, y0 + 30), color=None, fill=(0.1, 0.3, 0.7))
    page.draw_rect(pymupdf.Rect(x0 + 30, y0, x0 + 130, y0 + 30), color=None, fill=(0.1, 0.3, 0.7))
    page.insert_text((x0 + 38, y0 + 20), brand, fontsize=12, color=(1, 1, 1))


def logo_and_title(path: Path, title: str, title_origin: Tuple[float, float]) -> Path:
    """Port of ``probe_suppress_logo2.py``: a logo made of filled rectangles at the top left (which
    PyMuPDF's ``find_tables`` reports as a 1x2 table) and a page title next to it that does not touch
    the rectangles, then a text-heavy body so the page takes the text route."""
    doc = pymupdf.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    _logo(page)
    page.insert_text(title_origin, title, fontsize=16)
    page.insert_textbox(pymupdf.Rect(40, 110, 550, 700),
                        " ".join(body_lines(1, count=12, words=12)), fontsize=10)
    return _save(doc, path)


def slide_with_logo_block(path: Path, lines: Sequence[Tuple[Tuple[float, float], str, float]]) -> Path:
    """Port of the reviewed deck's slide 21 (``samples/deck_suppression.py``): a wide slide whose logo
    (filled rectangles plus white "AI") sits in the same PyMuPDF text block as the slide label, the
    title and the subtitle. ``lines`` is ``[((x, y), text, fontsize), ...]`` in Traditional Chinese."""
    doc = pymupdf.open()
    page = doc.new_page(width=WIDE_SLIDE[0], height=WIDE_SLIDE[1])
    page.draw_rect(pymupdf.Rect(52, 33, 92, 74), color=None, fill=(0.15, 0.2, 0.55))
    page.draw_rect(pymupdf.Rect(92, 33, 194, 74), color=None, fill=(0.15, 0.2, 0.55))
    page.insert_text((60, 62), "AI", fontsize=24, color=(1, 1, 1))
    for origin, text, size in lines:
        page.insert_text(origin, text, fontsize=size, fontname="china-t")
    return _save(doc, path)


def caption_just_above_table(path: Path, caption: str, gap: float) -> Path:
    """Port of ``probe_suppress.py`` S1: a 10pt caption whose baseline is ``gap`` points above the
    top rule of a ruled 3x2 table."""
    doc = pymupdf.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    page.insert_text((50, 100 - gap), caption, fontsize=10)
    ruled_table(page, 50, 100, [80, 80], [18] * 3, [["H1", "H2"], ["1", "2"], ["3", "4"]])
    page.insert_textbox(pymupdf.Rect(50, 200, 550, 400), " ".join(body_lines(1, count=4)), fontsize=10)
    return _save(doc, path)


# --------------------------------------------------------------------------------------
# Rotated pages (T6)
# --------------------------------------------------------------------------------------

# Where the rotated-space table bbox lands in unrotated space, per /Rotate value, for the
# table drawn by rotated_table_with_note (probe_rot_loss.py, loss mode).
ROTATION_NOTE_RECTS = {
    0: (440, 150, 580, 360),
    90: (440, 150, 580, 360),
    180: (150, 450, 460, 530),
    270: (360, 236, 580, 322),
}


def rotated_table_with_note(path: Path, rotation: int, cells: Sequence[Sequence[str]], note: str) -> Path:
    """Port of ``probe_rot_loss.py``: a ruled 3x2 table at (150, 330) plus a note paragraph placed
    where the table's rotated (display-space) bbox falls in unrotated coordinates, then
    ``/Rotate rotation``."""
    doc = pymupdf.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    ruled_table(page, 150, 330, [100, 100], [20] * 3, cells, fontsize=10)
    page.insert_textbox(pymupdf.Rect(*ROTATION_NOTE_RECTS[rotation]), note, fontsize=9)
    page.set_rotation(rotation)
    return _save(doc, path)


def _shift_content(doc, page, dx: float, dy: float) -> None:
    """Move everything drawn on ``page`` by (dx, dy) in PDF units (a ``cm`` around the content)."""
    page.wrap_contents()
    content = b"".join(doc.xref_stream(xref) for xref in page.get_contents())
    for xref in page.get_contents():
        doc.update_stream(xref, b"")
    doc.update_stream(page.get_contents()[0], f"q 1 0 0 1 {dx:g} {dy:g} cm\n".encode() + content + b"\nQ")


def rotated_cropped_table(path: Path, rotation: int, cells: Sequence[Sequence[str]], note: str, cropbox: str,
                          mediabox: str = "[0 0 595 842]", indirect_crop_x0: bool = False,
                          parent_cropbox: Optional[str] = None, hidden: Optional[str] = None,
                          invisible: Optional[str] = None) -> Path:
    """A ruled table with a note right beside it on a ``/Rotate rotation`` page whose raw /MediaBox
    and /CropBox entries are ``mediabox`` and ``cropbox``, written unvalidated: a CropBox may reach
    outside the MediaBox or list its corners in any order, as in real files. The content is drawn
    relative to the MediaBox's lower-left corner. With ``indirect_crop_x0`` the CropBox's first
    number is an indirect object (``[12 0 R 50 570 800]``), which PDF allows in any array.
    ``parent_cropbox`` is written as the /Pages node's raw /CropBox. ``hidden`` is a line printed
    near the bottom edge of the unrotated page, where a CropBox such as ``[30 42 570 792]`` hides
    it; ``body_lines(1, 8)`` sit well inside it. ``invisible`` is a line of invisible text (render
    mode 3) over nothing the page shows, away from the table: hidden text, left out of the text
    (PR #18), which makes the text path read the page's tables from a copy without it.

    PyMuPDF's find_tables() reports the table box of such a page shifted by the crop, over the note.
    """
    doc = pymupdf.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    ruled_table(page, 150, 330, [100, 100], [20] * 3, cells, fontsize=10)
    page.insert_textbox(pymupdf.Rect(360, 332, 440, 400), note, fontsize=8)
    if hidden is not None:
        page.insert_text((200, 830), hidden, fontsize=8)
        page.insert_text((72, 450), "\n".join(line[:40] for line in body_lines(1, 8)), fontsize=10)
    if invisible is not None:
        page.insert_text((72, 640), invisible, fontsize=10, render_mode=3)
    page.set_rotation(rotation)
    x0, y0 = (float(value) for value in mediabox.strip("[]").split()[:2])
    if x0 or y0:
        _shift_content(doc, page, x0, y0)
    doc.xref_set_key(page.xref, "MediaBox", mediabox)
    if indirect_crop_x0:
        numbers = cropbox.strip("[]").split()
        number = doc.get_new_xref()
        doc.update_object(number, numbers[0])
        cropbox = f"[{number} 0 R {' '.join(numbers[1:])}]"
    doc.xref_set_key(page.xref, "CropBox", cropbox)
    if parent_cropbox is not None:
        doc.xref_set_key(int(doc.xref_get_key(page.xref, "Parent")[1].split()[0]), "CropBox", parent_cropbox)
    return _save(doc, path)


def sideways_landscape_table(path: Path, rotation: int, cells: Sequence[Sequence[str]]) -> Path:
    """Port of ``probe_rot2.py`` R3: a landscape table drawn sideways on a portrait page so that
    ``/Rotate rotation`` (90 or 270) displays it upright, as pdflscape or rotated scans do."""
    doc = pymupdf.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    x0, y0, row_h, col_w = 100, 400, 18, 80
    for r in range(3):
        for c in range(3):
            if rotation == 90:
                rect = pymupdf.Rect(x0 + r * row_h, y0 - (c + 1) * col_w, x0 + (r + 1) * row_h, y0 - c * col_w)
                page.insert_text((rect.x0 + 13, rect.y1 - 3), cells[r][c], fontsize=9, rotate=90)
            else:
                rect = pymupdf.Rect(x0 + (2 - r) * row_h, y0 + c * col_w, x0 + (3 - r) * row_h, y0 + (c + 1) * col_w)
                page.insert_text((rect.x0 + 5, rect.y0 + 3), cells[r][c], fontsize=9, rotate=270)
            page.draw_rect(rect, width=0.6)
    page.set_rotation(rotation)
    return _save(doc, path)


# --------------------------------------------------------------------------------------
# Running headers / footers (H-F2, H-F10)
# --------------------------------------------------------------------------------------

Decorate = Callable[[pymupdf.Page, int, int], None]


def paged_document(path: Path, pages: int, decorate: Optional[Decorate] = None, *, body_y: float = 150,
                   body_count: int = 10, cjk_body: bool = False, size: Tuple[int, int] = A4) -> Path:
    """``pages`` pages of unique body text (``BODY<page>x<line>`` lines from ``body_y``), each then
    passed to ``decorate(page, page_number, pages)`` to add headers, footers or other page furniture."""
    doc = pymupdf.open()
    for p in range(1, pages + 1):
        page = doc.new_page(width=size[0], height=size[1])
        if cjk_body:
            rnd = random.Random(p)
            pool = "本公司於年度持續投資核心平台並維持營運成本穩定管理階層認為現行策略能使企業具備競爭優勢"
            text = "\n".join(f"BODY{p}x{i} " + "".join(rnd.choice(pool) for _ in range(24)) + "。"
                             for i in range(body_count))
            page.insert_text((72, body_y), text, fontsize=10.5, fontname="china-t")
        else:
            page.insert_text((72, body_y), "\n".join(body_lines(p, body_count)), fontsize=10.5, fontname="helv")
        if decorate is not None:
            decorate(page, p, pages)
    return _save(doc, path)


def footer_decorator(form: Callable[[int, int], str], *, font: str = "helv", y: float = 800,
                     x: float = 280, size: float = 9) -> Decorate:
    """A running footer (or header, with a small ``y``) written with ``form(page, pages)``."""
    def decorate(page, p, n):
        _text(page, x, y, form(p, n), size, font)
    return decorate


ROTATED_STATEMENT_ROWS = [("Revenue", "4,812", "4,377"), ("Cost of sales", "(2,905)", "(2,648)"),
                          ("Gross profit", "1,907", "1,729"), ("Selling expenses", "(611)", "(583)"),
                          ("Administrative expenses", "(402)", "(388)"), ("Operating profit", "894", "758")]


def rotated_cropped_statement(path: Path, rotation: int, cropbox: str, note: str) -> Path:
    """A borderless statement (a bold title, a label column and two right-aligned number columns set
    with tab stops, no rules: a table only PyMuPDF's text strategy finds) with ``note`` under it, on
    a landscape page stored sideways (``/Rotate rotation``, as word processors store one) whose own
    raw /CropBox is ``cropbox``."""
    upright = pymupdf.open()
    page = upright.new_page(width=A4[1], height=A4[0])
    page.insert_text((60, 70), "Statement of income (USD thousands)", fontsize=12, fontname="hebo")
    y = 100
    for label, *values in ROTATED_STATEMENT_ROWS:
        page.insert_text((60, y), label, fontsize=10)
        for right, value in zip((340, 440), values):
            page.insert_text((right - pymupdf.get_text_length(value, fontsize=10), y), value, fontsize=10)
        y += 16
    page.insert_text((60, y + 20), note, fontsize=9)
    doc = pymupdf.open()
    stored = doc.new_page(width=A4[0], height=A4[1])
    stored.show_pdf_page(stored.rect, upright, 0, rotate=rotation)
    stored.set_rotation(rotation)
    doc.xref_set_key(stored.xref, "CropBox", cropbox)
    upright.close()
    return _save(doc, path)


def label_and_number_footer(page, p: int, n: int) -> None:
    """A footer row holding a per-page label ("Lesson · N", left) and the bare page number (right)."""
    _text(page, 72, 812, f"Lesson · {p}", 9)
    _text(page, 500, 812, f"{p}", 9)


def equation_number_page(path: Path, label: str, uncovered: int, pages: int = 8) -> Path:
    """``pages`` pages of body text with a bare page number at the bottom, except page ``uncovered``,
    which has no page number: its last text line is the number ``label`` of an equation drawn as
    vector graphics (as MathType exports do) at the right margin, near the bottom of the page, and
    a sentence citing it just above."""
    doc = pymupdf.open()
    for p in range(1, pages + 1):
        page = doc.new_page(width=A4[0], height=A4[1])
        page.insert_text((72, 120), "\n".join(body_lines(p, 12)), fontsize=10.5, lineheight=14 / 10.5)
        if p == uncovered:
            page.draw_rect(pymupdf.Rect(200, 766, 380, 784), color=(0, 0, 0), width=0.8)
            page.draw_line((210, 775), (370, 775))
            _text(page, 72, 756, f"The balance follows from equation {label} below:")
            _text(page, 520, 780, label)
        else:
            _text(page, 290, 815, f"{p}", 9)
    return _save(doc, path)


def kpi_tiles_and_chart(page, p: int, n: int) -> None:
    """Port of ``p03d_numeric_zone_regular.py``: KPI tiles (big numbers that differ per page, small
    constant labels) at the top, a bar chart with year labels near the bottom, a real page number."""
    for i, (number, label) in enumerate([(120 + p, "new customers"), (37 + p, "open tickets"), (2024, "fiscal year")]):
        _text(page, 72 + i * 170, 70, f"{number}", 22)
        _text(page, 72 + i * 170, 90, label, 8)
    for i in range(5):
        height = 40 + ((p * 13 + i * 29) % 60)
        page.draw_rect(pymupdf.Rect(100 + i * 80, 740 - height, 140 + i * 80, 740), color=(0, 0, 1), fill=(0.6, 0.7, 1))
        _text(page, 105 + i * 80, 758, f"{2019 + i}", 9)
    _text(page, 290, 815, f"{p}", 9)


def title_then_running_header(title: str, title_y: float = 80, header_from: int = 2,
                               hidden: bool = False) -> Decorate:
    """Port of ``p03f_title_as_running_header.py`` (regular font): the document title on page 1
    (20pt, baseline at ``title_y``: 80 is in the top margin band, 130 just below it) that is also
    the 9pt running header of every page from ``header_from`` on. With ``hidden`` the title is
    invisible text (render mode 3) over nothing the page shows: hidden text, which the text path
    leaves out (PR #18)."""
    def decorate(page, p, n):
        if p == 1:
            _text(page, 72, title_y, title, 20, render_mode=3 if hidden else 0)
        elif p >= header_from:
            _text(page, 72, 40, title, 9)
    return decorate


def tight_footer_document(path: Path, gaps: Sequence[Optional[float]], *, cover_raise: float = 10.0,
                          lines: int = 12) -> Path:
    """Pages whose text runs down to just above the page number, as in Word/LibreOffice exports of
    full pages (``sample_documents/fail-1.docx``). Page p has ``lines`` body lines whose last one
    ends ``gaps[p - 1]`` pt above the top of the bare page number below it (glyph extents from
    Helvetica's ascender and descender), or a normal margin when that is None; page 1 prints its
    number ``cover_raise`` pt higher than the other pages, as a first page with its own footer does."""
    doc = pymupdf.open()
    ascent, descent = 0.718, 0.207  # Helvetica cap height and descender, in em
    for p, gap in enumerate(gaps, 1):
        page = doc.new_page(width=A4[0], height=A4[1])
        number_y = 800 - (cover_raise if p == 1 else 0)
        last_y = number_y - ascent * 9 - (40 if gap is None else gap) - descent * 10.5
        page.insert_text((72, last_y - 14 * (lines - 1)), "\n".join(body_lines(p, lines)), fontsize=10.5,
                         lineheight=14 / 10.5, fontname="helv")
        _text(page, 290, number_y, f"{p}", 9)
    return _save(doc, path)


def heading_on_pages(heading: str, on_pages: Sequence[int]) -> Decorate:
    """Port of p03 H4: a 16pt bold section heading at the top of some pages only."""
    def decorate(page, p, n):
        if p in on_pages:
            _text(page, 72, 80, heading, 16, "hebo")
    return decorate


STATEMENT_HEADER = ["ACME CORPORATION", "Consolidated Balance Sheet", "(in thousands of USD)"]


def statement_pages(path: Path, pages: int = 4) -> Path:
    """A short financial statement: every page repeats the bold company and statement title and
    the unit note at the top (as statements continued over pages do), then numbered line items,
    and a page number at the bottom. The repeated header block is the document's title."""
    doc = pymupdf.open()
    for p in range(1, pages + 1):
        page = doc.new_page(width=612, height=792)
        _text(page, 230, 40, STATEMENT_HEADER[0], 11, "hebo")
        _text(page, 200, 55, STATEMENT_HEADER[1], 11, "hebo")
        _text(page, 230, 70, STATEMENT_HEADER[2], 9)
        rows = [f"BODY{p}x{i} Line item {p}-{i}            {1000 + 37 * i * p:,}" for i in range(10)]
        page.insert_text((72, 110), "\n".join(rows), fontsize=10)
        _text(page, 300, 770, f"{p}", 9)
    return _save(doc, path)


def statement_after_contents(path: Path, pages: int = 4) -> Path:
    """A contents page listing "Consolidated Balance Sheet" in its top margin band (body-sized,
    one entry per line), then the statement pages of :func:`statement_pages`, whose repeated header
    block starts with the same "Consolidated Balance Sheet" title."""
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    _text(page, 72, 40, "Contents", 14, "hebo")
    for k, entry in enumerate(["Consolidated Balance Sheet", "Income Statement", "Notes to the Accounts"]):
        _text(page, 72, 62 + 14 * k, entry, 10.5)
    page.insert_text((72, 200), "\n".join(body_lines(1, 10)), fontsize=10.5, fontname="helv")
    for p in range(2, pages + 2):
        page = doc.new_page(width=612, height=792)
        _text(page, 230, 40, STATEMENT_HEADER[0], 11, "hebo")
        _text(page, 200, 55, STATEMENT_HEADER[1], 11, "hebo")
        _text(page, 230, 70, STATEMENT_HEADER[2], 9)
        rows = [f"BODY{p}x{i} Line item {p}-{i}            {1000 + 37 * i * p:,}" for i in range(10)]
        page.insert_text((72, 110), "\n".join(rows), fontsize=10)
        _text(page, 300, 770, f"{p}", 9)
    return _save(doc, path)


def statement_with_one_string_rows(path: Path, pages: int = 6) -> Path:
    """An account statement whose rows are each drawn as one string with padded columns: a
    running header, a bold header row (Date / Description / Amount / Balance) repeated at the
    top of every page, 25 transaction rows, and "Page N of M"."""
    doc = pymupdf.open()
    n = 0
    for p in range(1, pages + 1):
        page = doc.new_page(width=A4[0], height=A4[1])
        _text(page, 72, 40, "Account Statement - ACME Bank", 9)
        _text(page, 72, 90, "Date        Description              Amount     Balance", 10, "hebo")
        y = 114
        for _ in range(25):
            n += 1
            _text(page, 72, y, f"2024-01-{(n - 1) % 28 + 1:02d}  TXN{n:04d} payment received   {n * 3:>6}   {n * 7:>7}", 10)
            y += 14
        _text(page, 280, 815, f"Page {p} of {pages}", 9)
    return _save(doc, path)


def title_on_every_page(title: str) -> Decorate:
    """A 14pt bold document title at the top of every page (a heading-sized running title)."""
    def decorate(page, p, n):
        _text(page, 72, 45, title, 14, "hebo")
    return decorate


LETTERHEAD_ROWS = [["ACME Corporation", "Doc ID: QMS-042 Rev 3"], ["Quality Manual", "Approved 2024-05-01"]]


def letterhead_table(page, p: int, n: int) -> None:
    """Port of ``mk_docx_header_table.py`` (``r8``): a letterhead drawn as a ruled 2x2 table at the
    top of every page."""
    ruled_table(page, 72, 30, [200, 200], [16, 16], LETTERHEAD_ROWS, fontsize=9)


def chaptered_document(path: Path, chapters: Sequence[Tuple[str, int]]) -> Path:
    """Synthetic book-style document: every page has the running header "Chapter k - <title>" of
    its chapter (9pt, top) and a "Page N of M" footer; the first page of a chapter also has the
    chapter heading "k <title>" (16pt bold). ``chapters`` is ``[(title, page_count), ...]``."""
    total = sum(count for _, count in chapters)
    doc = pymupdf.open()
    page_no = 0
    for k, (title, count) in enumerate(chapters, 1):
        for i in range(count):
            page_no += 1
            page = doc.new_page(width=A4[0], height=A4[1])
            _text(page, 72, 40, f"Chapter {k} - {title}", 9)
            body_y = 150
            if i == 0:
                _text(page, 72, 110, f"{k} {title}", 16, "hebo")
            page.insert_text((72, body_y), "\n".join(body_lines(page_no, 10)), fontsize=10.5, fontname="helv")
            _text(page, 260, 800, f"Page {page_no} of {total}", 9)
    return _save(doc, path)


def report_with_chrome_only_cover(path: Path, title: str, pages: int = 5) -> Path:
    """A report whose every page has the running header "ACME Corp - Annual Report" and a
    "Page N of M" footer; page 1 has nothing else in its text layer (a picture cover), and
    page 2 opens with the document title in 24pt bold."""
    doc = pymupdf.open()
    for p in range(1, pages + 1):
        page = doc.new_page(width=A4[0], height=A4[1])
        _text(page, 72, 40, "ACME Corp - Annual Report", 9)
        if p == 2:
            _text(page, 72, 130, title, 24, "hebo")
        if p >= 2:
            page.insert_text((72, 180), "\n".join(body_lines(p, 10)), fontsize=10.5, fontname="helv")
        _text(page, 260, 800, f"Page {p} of {pages}", 9)
    return _save(doc, path)


def slides_with_titles(path: Path, titles: Sequence[str]) -> Path:
    """16:9 slides with a 30pt bold title each (consecutive slides may share a title, as
    continued slides do), four bullets and a slide number."""
    doc = pymupdf.open()
    for s, title in enumerate(titles, 1):
        page = doc.new_page(width=SLIDE[0], height=SLIDE[1])
        _text(page, 48, 52, title, 30, "hebo")
        page.insert_text((48, 130), "\n".join(f"BODY{s}x{i} point about the topic number {i}" for i in range(4)),
                         fontsize=20, fontname="helv")
        _text(page, 900, 520, f"{s}", 10)
    return _save(doc, path)


def rotated_report(path: Path, rotation: int, *, cropbox: Optional[str] = None, pages: int = 5) -> Path:
    """Report pages drawn sideways so that ``/Rotate rotation`` displays them upright (a
    landscape scan or pdflscape page): running header "ACME Corp Annual Report 2024" at the
    displayed top, body lines, "Page N of M" at the displayed bottom. ``cropbox`` is written
    as each page's raw /CropBox entry (e.g. ``"[20 30 575 812]"``), unvalidated, so it may
    reach outside the MediaBox as it does in real files."""
    doc = pymupdf.open()
    for p in range(1, pages + 1):
        page = doc.new_page(width=A4[0], height=A4[1])
        page.set_rotation(rotation)
        to_page = page.derotation_matrix  # displayed -> unrotated page coordinates
        width, height = page.rect.width, page.rect.height

        def put(x, y, text, size):
            page.insert_text(pymupdf.Point(x, y) * to_page, text, fontsize=size, rotate=rotation)

        put(72, 50, "ACME Corp Annual Report 2024", 9)
        for i in range(8):
            put(72, 150 + 14 * i, f"BODY{p}x{i} alpha beta gamma delta sigma omega river stone", 10.5)
        put(width / 2 - 20, height - 30, f"Page {p} of {pages}", 9)
        if cropbox:
            doc.xref_set_key(page.xref, "CropBox", cropbox)
    return _save(doc, path)


SIDE_NOTES = [(356, 345, "SIDENOTE-R keep me"), (60, 345, "SIDENOTE-L keep me"),
              (150, 318, "ABOVE-NOTE keep me"), (150, 402, "BELOW-NOTE keep me")]


def rotated_table_with_inherited_cropbox(path: Path, rotation: int, cells: Sequence[Sequence[str]]) -> Path:
    """A ruled table with notes on all four sides on a ``/Rotate`` page whose CropBox is not on
    the page itself but inherited from the /Pages node."""
    doc = pymupdf.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    ruled_table(page, 150, 330, [100, 100], [20] * 3, cells, fontsize=10)
    for x, y, note in SIDE_NOTES:
        page.insert_text((x, y), note, fontsize=8)
    page.set_rotation(rotation)
    pages_xref = int(doc.xref_get_key(page.xref, "Parent")[1].split()[0])
    doc.xref_set_key(pages_xref, "CropBox", "[15 17 580 822]")
    return _save(doc, path)


def rotated_table_on_offset_mediabox(path: Path, cells: Sequence[Sequence[str]], note: str, mediabox: str) -> Path:
    """A ruled table and a note on a /Rotate 90 page whose MediaBox does not start at the
    origin (e.g. ``"[0 -842 595 0]"``) and which has no CropBox."""
    doc = pymupdf.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    ruled_table(page, 150, 330, [100, 100], [20] * 3, cells, fontsize=10)
    page.insert_text((150, 420), note, fontsize=9)
    page.set_rotation(90)
    x0, y0 = (float(value) for value in mediabox.strip("[]").split()[:2])
    _shift_content(doc, page, x0, y0)
    doc.xref_set_key(page.xref, "MediaBox", mediabox)
    return _save(doc, path)


def slides_with_footer_block(path: Path, slides: int) -> Path:
    """Port of the ``r6`` deck export: 16:9 slides whose footer "ACME Confidential" (left) and slide
    number (right) share one text block, with a 44pt title and 28pt bullets."""
    doc = pymupdf.open()
    for s in range(1, slides + 1):
        page = doc.new_page(width=SLIDE[0], height=SLIDE[1])
        _text(page, 200, 90, f"Slide {s} topic {WORDS[s % len(WORDS)]}", 44)
        bullets = [f"• BODY{s}x{i} {WORDS[(s + i) % len(WORDS)]} {WORDS[(s * 3 + i) % len(WORDS)]}" for i in range(3)]
        page.insert_text((43, 160), "\n".join(bullets), fontsize=28, fontname="helv")
        _text(page, 36, 517, "ACME Confidential", 10)
        _text(page, 893, 517, f"{s}", 10)
    return _save(doc, path)


# --------------------------------------------------------------------------------------
# Overprinted text (H-F6)
# --------------------------------------------------------------------------------------

CLIPPED_CELLS = [("Lot 2024-10", "10 pallets", 3.5), ("Invoice 7710", "10.00", 3.0), ("Batch no. 45", "45 kg net", 4.0)]


def clipped_spreadsheet_cells(path: Path) -> Path:
    """Spreadsheet-style rows whose first cell's text runs on under the next cell (as clipped
    cells are exported), so two different strings overlap without being overprints: each
    right cell starts a few digits before the left text ends."""
    doc = pymupdf.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    page.insert_text((72, 60), "\n".join(body_lines(1, 3)), fontsize=10.5)
    digit = pymupdf.get_text_length("0", fontsize=10)
    for k, (left, right, back) in enumerate(CLIPPED_CELLS):
        y = 140 + 30 * k
        _text(page, 72, y, left, 10)
        _text(page, 72 + pymupdf.get_text_length(left, fontsize=10) - back * digit, y, right, 10)
    return _save(doc, path)


def tall_metrics_document(path: Path, pages: int = 6, rotation: int = 0) -> Path:
    """Pages whose text font declares Word/Cambria-like metrics (FontDescriptor /Ascent 3117
    /Descent -2464, as in ``sample_documents/sample_pdf.pdf``), for which PyMuPDF 1.28 reports
    line boxes ~5.6x taller than the text: neighbouring lines' boxes overlap. Each page has a
    running header, a "Page N of M" footer and body lines; page 1 also has consecutive
    lines that repeat ("No" twice, "Total" over "Total revenue grew"), and an unrotated page 2
    a table with its caption just above the top rule. With ``rotation`` the text is drawn so
    that ``/Rotate rotation`` displays it upright."""
    doc = pymupdf.open()
    font = pymupdf.Font("helv").buffer
    for p in range(1, pages + 1):
        page = doc.new_page(width=A4[0], height=A4[1])
        page.set_rotation(rotation)
        page.insert_font(fontname="TALL", fontbuffer=font)
        to_page = page.derotation_matrix  # displayed -> unrotated page coordinates
        height = page.rect.height

        def put(x, y, text, size=10.5):
            page.insert_text(pymupdf.Point(x, y) * to_page, text, fontname="TALL", fontsize=size, rotate=rotation)

        put(72, 40, "ACME Corp - Annual Report", 9)
        y = 150
        for line in body_lines(p, 10):
            put(72, y, line)
            y += 14
        if p == 1:
            for line in ("No", "No", "Total", "Total revenue grew"):
                put(72, y, line)
                y += 14
        if p == 2 and not rotation:
            put(72, y + 20 - 2, "Table 7: Tall metrics caption")
            ruled_table(page, 72, y + 20, [100, 100], [18] * 3, [["TH1", "TH2"], ["TV11", "TV12"], ["TV21", "TV22"]])
        put(260, height - 42, f"Page {p} of {pages}", 9)
    for xref in range(1, doc.xref_length()):
        if doc.xref_get_key(xref, "Type") == ("name", "/FontDescriptor"):
            doc.xref_set_key(xref, "Ascent", "3117")
            doc.xref_set_key(xref, "Descent", "-2464")
            doc.xref_set_key(xref, "FontBBox", "[-1475 -2464 2868 3117]")
    return _save(doc, path)


def overprinted_text(path: Path) -> Path:
    """Port of ``p04_layout.py`` S5 plus the other overprint shapes seen in real decks: fake bold
    (0.3pt offset), a drop shadow, an exact CJK overprint, a highlighted phrase drawn twice over its
    own sentence, a line printed three times, and one word fake-bolded inside a sentence. It also
    holds legitimately repeated words ("No No", "10 10", and "No" on two consecutive lines)."""
    doc = pymupdf.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    page.insert_text((72, 60), "\n".join(body_lines(1, 3)), fontsize=10.5)
    for dx in (0, 0.3):
        _text(page, 72 + dx, 120, "Quarterly Revenue Summary", 14)
    _text(page, 73.5, 161.5, "Shadowed Heading Text", 18, color=(0.7, 0.7, 0.7))
    _text(page, 72, 160, "Shadowed Heading Text", 18)
    for _ in range(2):
        _text(page, 72, 200, "資料治理與流程銜接", 14, "china-t")
    sentence, phrase = "TeamSync AI focuses on data governance, and more", "data governance,"
    _text(page, 72, 240, sentence, 12)
    offset = pymupdf.get_text_length("TeamSync AI focuses on ", fontsize=12)
    for _ in range(2):
        _text(page, 72 + offset, 240, phrase, 12, color=(0.8, 0, 0))
    for _ in range(3):
        _text(page, 72, 280, "Triple printed line", 12)
    _text(page, 72, 320, "This is ", 12)
    offset = pymupdf.get_text_length("This is ", fontsize=12)
    for dx in (0, 0.3):
        _text(page, 72 + offset + dx, 320, "important", 12)
    _text(page, 72 + offset + pymupdf.get_text_length("important ", fontsize=12), 320, "text here.", 12)
    _text(page, 72, 360, "Votes: No No", 12)
    _text(page, 72, 380, "Scores: 10 10", 12)
    _text(page, 72, 410, "No", 12)
    _text(page, 72, 424, "No", 12)
    page.insert_text((72, 470), "\n".join(body_lines(2, 3)), fontsize=10.5)
    return _save(doc, path)
