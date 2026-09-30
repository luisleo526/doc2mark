"""Input builders for the reading-order E2E tests (``tests/e2e/test_order.py``).

The layouts port the review probes of finding H-F3 (``~/code/.executors/d2m-rv-headings-probes``:
``mk_docx_2col.py``, ``p04_layout.py`` S1, ``p04b_landscape_rotate.py``, ``p04c_landscape_lengths.py``,
``p11_two_column_order.py``) into documents made at test time with PyMuPDF and, for the Word
export, python-docx and LibreOffice. Nothing from ``doc2mark`` is imported.

Text is typeset the way a layout program sets it: paragraphs are broken first-fit to the column
width and flow from the bottom of one column to the top of the next, so a paragraph can start in
one column and end in the next. Every paragraph starts with ``[<tag>]`` and ends with ``[/<tag>]``,
so a test can check that paragraphs come in order and that no paragraph is read end first.
"""

import io
import random
from pathlib import Path
from typing import Dict, List, Sequence, Tuple

import pymupdf

A4 = (595, 842)
WORDS = ("alpha beta gamma delta sigma omega river stone cloud metal paper glass forest ocean "
         "harbour signal winter copper meadow lantern").split()


def tagged(tag: str, words: int, seed: int) -> str:
    """A paragraph of ``words`` seeded filler words between ``[tag]`` and ``[/tag]``."""
    rnd = random.Random(seed)
    return f"[{tag}] " + " ".join(rnd.choice(WORDS) for _ in range(words)) + f" [/{tag}]"


def wrap(text: str, width: float, size: float, font: str = "helv") -> List[str]:
    """``text`` broken first-fit into lines no wider than ``width``."""
    lines, line = [], ""
    for word in text.split():
        candidate = f"{line} {word}" if line else word
        if line and pymupdf.get_text_length(candidate, font, size) > width:
            lines.append(line)
            line = word
        else:
            line = candidate
    return lines + ([line] if line else [])


def _save(doc, path: Path) -> Path:
    doc.save(str(path))
    doc.close()
    return Path(path)


def _png(width: int, height: int, color=(70, 120, 190)) -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (width, height), color).save(buffer, "PNG")
    return buffer.getvalue()


# Blocks for the typesetter: ("h", text) is a bold heading, ("p", text) a paragraph.
Block = Tuple[str, str]


def flow(pages: Sequence, frames: Sequence[Tuple[int, float, float, float, float]], blocks: Sequence[Block],
         size: float = 9.5, leading: float = 11.5, gap: float = 7.0,
         avoid: Sequence[Tuple[int, float, float, float, float]] = ()) -> None:
    """Set ``blocks`` into ``frames`` ((page index, x0, y0, x1, y1) column boxes, in reading order),
    breaking lines first-fit to the width of the frame they land in. A paragraph continues in the
    next frame when the current one is full; a heading moves on with the line after it. Lines wrap
    around the ``avoid`` boxes ((page index, x0, y0, x1, y1), 8pt apart). Raises when the text does
    not fit."""
    frame_index, y = 0, None
    for kind, text in blocks:
        font = "hebo" if kind == "h" else "helv"
        line_size = size + 1 if kind == "h" else size
        words = text.split()
        while words:
            if frame_index >= len(frames):
                raise ValueError("text does not fit into the frames")
            page_index, x0, y0, x1, y1 = frames[frame_index]
            y = y0 + size if y is None else y
            if y > y1 or (kind == "h" and y + leading > y1):
                frame_index, y = frame_index + 1, None
                continue
            left, right = x0, x1
            for page_avoid, ax0, ay0, ax1, ay1 in avoid:
                if page_avoid == page_index and ay0 - 3 < y and y - line_size < ay1 + 3:
                    if ax0 <= left + 1 < ax1:
                        left = ax1 + 8
                    elif left < ax0 < right:
                        right = ax0 - 8
            count = 1
            while count < len(words) and pymupdf.get_text_length(
                    " ".join(words[:count + 1]), font, line_size) <= right - left:
                count += 1
            pages[page_index].insert_text((left, y), " ".join(words[:count]), fontsize=line_size, fontname=font)
            words = words[count:]
            y += leading
        y += gap


def two_column_paper_pdf(path: Path) -> Tuple[Path, List[str]]:
    """Two A4 pages of a two-column paper: a centred title and author line and a full-width
    abstract span the columns on page 1; sections flow through the left, then the right column of
    both pages (paragraphs cross from one column to the next); a footnote sits at the foot of the
    left column of page 1 and a page number at the bottom centre of each page.

    Returns the path and the paragraph tags in reading order."""
    doc = pymupdf.open()
    for _ in range(2):
        doc.new_page(width=A4[0], height=A4[1])
    pages = [doc[0], doc[1]]   # a new page invalidates the Page objects made before it
    title = "Reading Order in Multi Column Documents"
    width = pymupdf.get_text_length(title, "hebo", 18)
    pages[0].insert_text(((A4[0] - width) / 2, 70), title, fontsize=18, fontname="hebo")
    authors = "[AUTH] Ada Lovelace and Alan Turing [/AUTH]"
    width = pymupdf.get_text_length(authors, "helv", 11)
    pages[0].insert_text(((A4[0] - width) / 2, 92), authors, fontsize=11, fontname="helv")
    abstract = wrap(tagged("ABS", 70, 1), 495, 9.5)
    pages[0].insert_text((50, 122), "\n".join(abstract), fontsize=9.5, fontname="helv", lineheight=1.2)
    top = 122 + 11.4 * len(abstract) + 18
    tags, blocks = [], []
    sections = ["1 Introduction", "2 Related Work", "3 Method", "4 Results", "5 Discussion", "6 Conclusion"]
    number = 0
    for section in sections:
        blocks.append(("h", section))
        for _ in range(3):
            number += 1
            tag = f"P{number:02d}"
            tags.append(tag)
            blocks.append(("p", tagged(tag, random.Random(number).randint(45, 85), 100 + number)))
    frames = [(0, 50, top, 290, 700), (0, 305, top, 545, 790), (1, 50, 60, 290, 790), (1, 305, 60, 545, 790)]
    flow(pages, frames, blocks)
    footnote = "1 [FN1] The data set and the code are available from the authors on request. [/FN1]"
    pages[0].insert_text((50, 772), "\n".join(wrap(footnote, 240, 7.5)), fontsize=7.5, fontname="helv",
                         lineheight=1.2)
    for number, page in enumerate(pages, 1):
        page.insert_text((A4[0] / 2 - 3, 818), str(number), fontsize=9, fontname="helv")
    return _save(doc, path), tags


def three_column_pdf(path: Path) -> Tuple[Path, List[str], List[str]]:
    """One A4 page with three columns of unequal width (160, 188 and 143pt): a full-width title;
    a band of text flowing through the three columns; a full-width figure (a picture and a
    ``Figure 1`` caption) across the columns; a second band of three columns under it.

    Returns the path and the paragraph tags of the band above and of the band below the figure."""
    doc = pymupdf.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    page.insert_text((40, 62), "Three Column Newsletter", fontsize=20, fontname="hebo")
    columns = [(40, 200), (212, 400), (412, 555)]
    above = [f"A{k}" for k in range(1, 6)]
    below = [f"B{k}" for k in range(1, 7)]
    flow([page], [(0, x0, 90, x1, 330) for x0, x1 in columns],
         [("p", tagged(tag, 64, 300 + k)) for k, tag in enumerate(above)], size=9, leading=11)
    page.insert_image(pymupdf.Rect(40, 345, 555, 455), stream=_png(515, 110))
    caption = "Figure 1: [FIG] Revenue by region and quarter, all figures unaudited. [/FIG]"
    page.insert_text((40, 472), caption, fontsize=9, fontname="helv")
    flow([page], [(0, x0, 495, x1, 790) for x0, x1 in columns],
         [("p", tagged(tag, 64, 400 + k)) for k, tag in enumerate(below)], size=9, leading=11)
    return _save(doc, path), above, below


def sidebar_pdf(path: Path) -> Tuple[Path, List[str]]:
    """One A4 page: a heading and four paragraphs in a 320pt main column, and beside them a grey
    ``Key facts`` sidebar box and a large pull-quote in a 155pt margin column (review probe
    p04_layout S1). Returns the path and the main-flow tags in order."""
    doc = pymupdf.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    page.insert_text((50, 70), "Market Update", fontsize=20, fontname="hebo")
    main = [f"M{k}" for k in range(1, 5)]
    flow([page], [(0, 50, 95, 370, 780)], [("p", tagged(tag, 75, 500 + k)) for k, tag in enumerate(main)],
         size=10.5, leading=13, gap=9)
    page.draw_rect(pymupdf.Rect(390, 150, 545, 372), color=(0.6, 0.6, 0.6), fill=(0.92, 0.92, 0.92))
    page.insert_text((400, 172), "Key facts", fontsize=12, fontname="hebo")
    side = wrap(tagged("SIDE", 38, 7), 135, 9)
    page.insert_text((400, 192), "\n".join(side), fontsize=9, fontname="helv", lineheight=1.25)
    quote = wrap("[PQ] Growth came from the regions we had nearly given up on. [/PQ]", 150, 14, "heit")
    page.insert_text((392, 470), "\n".join(quote), fontsize=14, fontname="heit", lineheight=1.2)
    return _save(doc, path), main


def pull_quote_pdf(path: Path) -> Tuple[Path, List[str]]:
    """One A4 page in two columns with a pull-quote set across the gutter, the text of both columns
    wrapped around it (magazine style), under a full-width headline. Returns the paragraph tags in
    order."""
    doc = pymupdf.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    page.insert_text((50, 70), "The Harbour Towns Come Back", fontsize=22, fontname="hebo")
    quote_box = (0, 185, 330, 410, 420)
    tags = [f"Q{k}" for k in range(1, 9)]
    flow([page], [(0, 50, 100, 290, 790), (0, 305, 100, 545, 790)],
         [("p", tagged(tag, 62, 900 + k)) for k, tag in enumerate(tags)], size=10, leading=12.5, gap=8,
         avoid=[quote_box])
    _, qx0, qy0, qx1, qy1 = quote_box
    page.draw_line((qx0, qy0), (qx1, qy0), color=(0.2, 0.2, 0.2), width=1.2)
    page.draw_line((qx0, qy1), (qx1, qy1), color=(0.2, 0.2, 0.2), width=1.2)
    quote = wrap("[PQ] We rebuilt the quay with our own hands and the ships came back. [/PQ]",
                 qx1 - qx0 - 10, 14, "heit")
    page.insert_text((qx0 + 5, qy0 + 22), "\n".join(quote), fontsize=14, fontname="heit", lineheight=1.25)
    return _save(doc, path), tags


def _visual_text(page, x: float, y: float, text: str, size: float, font: str = "helv") -> None:
    """Text at displayed (rotated-page) coordinates, upright for the reader."""
    page.insert_text(pymupdf.Point(x, y) * page.derotation_matrix, text, fontsize=size, fontname=font,
                     rotate=page.rotation)


def _visual_line(page, x0: float, y0: float, x1: float, y1: float) -> None:
    matrix = page.derotation_matrix
    page.draw_line(pymupdf.Point(x0, y0) * matrix, pymupdf.Point(x1, y1) * matrix, color=(0, 0, 0), width=0.8)


ROTATIONS = (90, 180, 270, 90)


def rotated_pages_pdf(path: Path) -> Tuple[Path, Dict[int, List[str]]]:
    """Four pages built as report tools emit landscape pages: a portrait MediaBox with ``/Rotate``
    90, 180, 270 and 90, the content drawn so that it reads upright (review probes p04b, p04c).
    Each page shows a running header at the top, three one-line paragraphs of different lengths
    (5, 14 and 9 words), a table title and a ruled 3x3 table under it, a closing paragraph and a
    page number at the bottom.

    Returns the path and, per page, the tokens in the order the page shows them."""
    doc = pymupdf.open()
    order: Dict[int, List[str]] = {}
    for number, rotation in enumerate(ROTATIONS, 1):
        page = doc.new_page(width=A4[0], height=A4[1])
        page.set_rotation(rotation)
        width, height = page.rect.width, page.rect.height
        _visual_text(page, 60, 40, "ACME Corp - Confidential", 9)
        tokens = []
        for k, (y, words) in enumerate(zip((100, 130, 160), (5, 14, 9)), 1):
            tag = f"R{number}.{k}"
            _visual_text(page, 60, y, tagged(tag, words, number * 10 + k), 11)
            tokens.append(f"[{tag}]")
        title = f"[TT{number}] Table {number}: Revenue by region"
        _visual_text(page, 60, 205, title, 11, "hebo")
        tokens.append(f"[TT{number}]")
        x_edges, y_edges = (60, 200, 320, 440), (215, 235, 255, 275)
        for x in x_edges:
            _visual_line(page, x, y_edges[0], x, y_edges[-1])
        for y in y_edges:
            _visual_line(page, x_edges[0], y, x_edges[-1], y)
        for r in range(3):
            for c in range(3):
                _visual_text(page, x_edges[c] + 5, y_edges[r] + 14, f"T{number}R{r}C{c}", 9)
        tokens.append(f"T{number}R0C0")
        tag = f"R{number}.4"
        _visual_text(page, 60, 320, tagged(tag, 12, number * 10 + 4), 11)
        tokens.append(f"[{tag}]")
        _visual_text(page, width / 2 - 3, height - 25, str(number), 9)
        order[number] = tokens
    return _save(doc, path), order


def word_two_column_docx(path: Path) -> Tuple[Path, List[str]]:
    """Port of ``mk_docx_2col.py`` with Word sections: a title and an abstract in one column, a
    continuous section break into two columns (headings and tagged paragraphs, balanced by the
    next break), then one column again for a closing paragraph. Returns the .docx path (convert it
    with ``builders_pdftext.docx_to_pdf``) and the tags in reading order."""
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def columns(section, count):
        sect_pr = section._sectPr
        for old in sect_pr.findall(qn("w:cols")):
            sect_pr.remove(old)
        cols = OxmlElement("w:cols")
        cols.set(qn("w:num"), str(count))
        cols.set(qn("w:space"), "425")
        sect_pr.append(cols)

    document = Document()
    document.add_heading("Two Column Paper Title", level=0)
    tags = ["ABS"]
    document.add_paragraph(tagged("ABS", 60, 11))
    columns(document.sections[0], 1)
    body = document.add_section(WD_SECTION.CONTINUOUS)
    columns(body, 2)
    number = 0
    for heading in ("Introduction", "Related Work", "Method", "Results"):
        document.add_heading(heading, level=1)
        for _ in range(2):
            number += 1
            tag = f"P{number:02d}"
            tags.append(tag)
            document.add_paragraph(tagged(tag, random.Random(number).randint(35, 60), 200 + number))
    closing = document.add_section(WD_SECTION.CONTINUOUS)
    columns(closing, 1)
    tags.append("END")
    document.add_paragraph(tagged("END", 30, 99))
    document.save(str(path))
    return Path(path), tags


FORM_ROWS = [("Applicant name", "Jane Q. Example"), ("Date of birth", "12 March 1985"),
             ("Nationality", "Portuguese"), ("Passport number", "P1234567"),
             ("Address", "12 Harbour Road, Lisbon"), ("Phone", "+351 912 345 678"),
             ("Employer", "Acme Shipping Ltd"), ("Position", "Chief Engineer")]


def form_pdf(path: Path) -> Path:
    """A single-column form: an intro paragraph, eight label/value rows (labels at x=72, values at
    x=240, one row every 24pt, all labels drawn before the values), and a closing paragraph. Read
    row by row, as the page shows it."""
    doc = pymupdf.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    page.insert_text((72, 70), "Visa Application Summary", fontsize=16, fontname="hebo")
    intro = wrap("[INTRO] Please check the details below and tell the consulate about any error "
                 "before the interview date. [/INTRO]", 450, 10)
    page.insert_text((72, 100), "\n".join(intro), fontsize=10, fontname="helv", lineheight=1.2)
    # labels first, then values, as a form filler writes them: each is a text block of its own
    for row, (label, _) in enumerate(FORM_ROWS):
        page.insert_text((72, 150 + 24 * row), label, fontsize=10, fontname="hebo")
    for row, (_, value) in enumerate(FORM_ROWS):
        page.insert_text((240, 150 + 24 * row), value, fontsize=10, fontname="helv")
    y = 150 + 24 * len(FORM_ROWS)
    closing = wrap("[CLOSE] Bring the original documents to the interview; copies are not accepted. [/CLOSE]",
                   450, 10)
    page.insert_text((72, y + 20), "\n".join(closing), fontsize=10, fontname="helv", lineheight=1.2)
    return _save(doc, path)


def single_column_pdf(path: Path) -> Tuple[Path, List[str]]:
    """A single-column page: title, paragraphs, an indented quotation, a right-aligned date line, a
    short line and a picture beside the end of a paragraph. Returns the tokens top to bottom."""
    doc = pymupdf.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    page.insert_text((72, 70), "Annual Letter", fontsize=18, fontname="hebo")
    date = "[DATE] 30 September 2026"
    page.insert_text((523 - pymupdf.get_text_length(date, "helv", 10), 95), date, fontsize=10, fontname="helv")
    tokens = ["[DATE]"]
    y = 125
    for k in range(1, 4):
        tag = f"S{k}"
        lines = wrap(tagged(tag, 60, 700 + k), 451, 10)
        page.insert_text((72, y), "\n".join(lines), fontsize=10, fontname="helv", lineheight=1.2)
        tokens.append(f"[{tag}]")
        y += 12 * len(lines) + 14
        if k == 2:
            quote = wrap(tagged("QUOTE", 30, 9), 380, 10, "heit")
            page.insert_text((108, y), "\n".join(quote), fontsize=10, fontname="heit", lineheight=1.2)
            tokens.append("[QUOTE]")
            y += 12 * len(quote) + 14
    page.insert_text((72, y), "[SHORT] Yours sincerely,", fontsize=10, fontname="helv")
    tokens.append("[SHORT]")
    page.insert_image(pymupdf.Rect(420, y - 10, 520, y + 40), stream=_png(100, 50))
    return _save(doc, path), tokens


RUNNING_HEADER = "<Draft> ACME Corp - Internal"


def markup_running_header_pdf(path: Path) -> Path:
    """Four pages with the running header ``<Draft> ACME Corp - Internal`` at the top of each page
    and two body paragraphs per page (lane md follow-up: the kept first copy must be escaped)."""
    doc = pymupdf.open()
    for number in range(1, 5):
        page = doc.new_page(width=A4[0], height=A4[1])
        page.insert_text((72, 40), RUNNING_HEADER, fontsize=9, fontname="helv")
        for k, y in enumerate((100, 220), 1):
            lines = wrap(tagged(f"H{number}.{k}", 50, number * 3 + k), 451, 10)
            page.insert_text((72, y), "\n".join(lines), fontsize=10, fontname="helv", lineheight=1.2)
        page.insert_text((290, 815), str(number), fontsize=9, fontname="helv")
    return _save(doc, path)


BULLETS_THEN_NUMBER = ["Key points:", "• Revenue increased", "• Costs decreased",
                       "5. Outlook for the next year"]


def bullets_then_number_pdf(path: Path) -> Path:
    """A lead-in line, two bullets and a numbered line ``5. Outlook for the next year`` in one
    block, drawn with a TextWriter so the bullet glyphs survive (lane md follow-up)."""
    doc = pymupdf.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    intro = wrap("The board met on Tuesday to review the results of the third quarter and the plan for the "
                 "coming year.", 451, 10)
    page.insert_text((72, 80), "\n".join(intro), fontsize=10, fontname="helv", lineheight=1.2)
    font = pymupdf.Font("helv")
    writer = pymupdf.TextWriter(page.rect)
    for index, line in enumerate(BULLETS_THEN_NUMBER):
        writer.append((72, 140 + 12 * index), line, font=font, fontsize=10)
    writer.write_text(page)
    page.insert_text((72, 220), "The next meeting is planned for the end of January.", fontsize=10,
                     fontname="helv")
    return _save(doc, path)


NESTED_ITEMS = [(72, "4. Training and rollout"), (90, "\u2022 Plan the kick-off sessions"),
                (90, "\u2022 Run the hands-on courses"), (90, "\u2022 Track adoption after launch"),
                (72, "5. Reporting")]


def nested_bullets_pdf(path: Path) -> Path:
    """A numbered item with three nested bullets and the next numbered item (a Word proposal's
    shape): a tight list that must stay tight."""
    doc = pymupdf.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    page.insert_text((72, 80), "The work plan has five parts, the last two of which are listed here.", fontsize=10,
                     fontname="helv")
    font = pymupdf.Font("helv")
    writer = pymupdf.TextWriter(page.rect)
    for index, (x, line) in enumerate(NESTED_ITEMS):
        writer.append((x, 110 + 12 * index), line, font=font, fontsize=10)
    writer.write_text(page)
    return _save(doc, path)


# --------------------------------------------------------------------------------------
# Review round 1 (supervisor review of PR #23): layouts from ~/code/.executors/d2m-review-order
# --------------------------------------------------------------------------------------

def spanning_figure_pdf(path: Path, kind: str = "picture") -> Tuple[Path, List[str]]:
    """After the reviewer's ``L04_spanfig``: two columns above and below a figure across both
    columns (a raster picture, or with ``kind="drawing"`` a bar chart drawn with vector paths and no
    text), with a short caption under it that does not reach across the gutter. Returns the path and
    the tags in reading order (the caption opens the band under the figure)."""
    doc = pymupdf.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    rnd = random.Random(50)
    above = [f"F{k}" for k in range(1, 7)]
    below = [f"F{k}" for k in range(7, 11)]
    flow([page], [(0, 50, 60, 290, 315), (0, 305, 60, 545, 315)],
         [("p", tagged(tag, rnd.randint(30, 42), 50 + k)) for k, tag in enumerate(above, 1)])
    if kind == "picture":
        page.insert_image(pymupdf.Rect(50, 330, 545, 480), stream=_png(400, 120))
    else:
        page.draw_rect(pymupdf.Rect(50, 330, 545, 480), color=(0.3, 0.3, 0.3), width=0.8)
        for k in range(12):
            height = 20 + 9 * ((k * 7) % 12)
            page.draw_rect(pymupdf.Rect(70 + 38 * k, 470 - height, 96 + 38 * k, 470), color=None, fill=(0.2, 0.4, 0.7))
    page.insert_text((50, 495), "[CAP] Figure 1: a figure across both columns [/CAP]", fontsize=9, fontname="helv")
    flow([page], [(0, 50, 515, 290, 790), (0, 305, 515, 545, 790)],
         [("p", tagged(tag, rnd.randint(30, 42), 60 + k)) for k, tag in enumerate(below, 1)])
    return _save(doc, path), above + ["CAP"] + below


def word_picture_between_sections_docx(path: Path, caption: bool = True) -> Tuple[Path, List[str]]:
    """Port of the reviewer's ``mk_word_fig.py``: a title, a two-column section, a full-width
    picture (with or without a caption) in a one-column section, then another two-column section.
    Returns the .docx path and the tags in reading order."""
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches

    def columns(section, count):
        sect_pr = section._sectPr
        for old in sect_pr.findall(qn("w:cols")):
            sect_pr.remove(old)
        cols = OxmlElement("w:cols")
        cols.set(qn("w:num"), str(count))
        cols.set(qn("w:space"), "360")
        sect_pr.append(cols)

    rnd = random.Random(50 + (0 if caption else 1))
    document = Document()
    document.add_heading("Paper With A Wide Figure", level=0)
    columns(document.sections[0], 1)
    tags = []
    columns(document.add_section(WD_SECTION.CONTINUOUS), 2)
    for k in range(1, 5):
        tags.append(f"A{k}")
        document.add_paragraph(f"[A{k}] " + " ".join(rnd.choice(WORDS) for _ in range(rnd.randint(40, 60))) + f" [/A{k}]")
    columns(document.add_section(WD_SECTION.CONTINUOUS), 1)
    document.add_picture(io.BytesIO(_png(600, 200)), width=Inches(6))
    if caption:
        document.add_paragraph("Figure 1: Overview of the system.")
    columns(document.add_section(WD_SECTION.CONTINUOUS), 2)
    for k in range(1, 5):
        tags.append(f"B{k}")
        document.add_paragraph(f"[B{k}] " + " ".join(rnd.choice(WORDS) for _ in range(rnd.randint(40, 60))) + f" [/B{k}]")
    document.save(str(path))
    return Path(path), tags


def word_columns_docx(path: Path, seed: int) -> Tuple[Path, List[str]]:
    """Port of the reviewer's ``mk_word.py``: a title and 4-9 sections of 2-4 ragged-right paragraphs
    in a two-column Word section. Returns the .docx path and the tags in reading order."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rnd = random.Random(seed)
    document = Document()
    document.add_heading("Two Column Paper Title", level=0)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "360")
    document.sections[0]._sectPr.append(cols)
    tags, number = [], 0
    for section in range(rnd.randint(4, 9)):
        document.add_heading(f"Section {section + 1}", level=1)
        for _ in range(rnd.randint(2, 4)):
            number += 1
            tag = f"W{number:02d}"
            tags.append(tag)
            document.add_paragraph(f"[{tag}] " + " ".join(rnd.choice(WORDS) for _ in range(rnd.randint(25, 80)))
                                   + f" [/{tag}]")
    document.save(str(path))
    return Path(path), tags


def docx_to_pdfs(paths: Sequence[Path]) -> List[Path]:
    """Convert several .docx files with one LibreOffice run (private profile) and return the PDFs."""
    import shutil
    import subprocess

    paths = [Path(p) for p in paths]
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        raise RuntimeError("soffice is not on PATH")
    folder = paths[0].parent
    profile = folder / ".lo-profile-order"
    subprocess.run([soffice, f"-env:UserInstallation={profile.resolve().as_uri()}", "--headless",
                    "--convert-to", "pdf", "--outdir", str(folder), *map(str, paths)],
                   check=True, capture_output=True, timeout=300)
    pdfs = [path.with_suffix(".pdf") for path in paths]
    missing = [pdf for pdf in pdfs if not pdf.exists()]
    if missing:
        raise RuntimeError(f"LibreOffice did not write {missing}")
    return pdfs


def ragged_columns_pdf(path: Path) -> Tuple[Path, List[str]]:
    """Port of the reviewer's ``L16_chrome_2col``: three pages of two ragged-right columns under a
    running header and over a ``Page N`` footer. Lower in the left column some lines run a point or
    two further right than the lines beside the right column's text. Returns the tags in order."""
    doc = pymupdf.open()
    for _ in range(3):
        doc.new_page(width=A4[0], height=A4[1])
    pages = [doc[k] for k in range(3)]
    tags, blocks = [], []
    for k in range(1, 25):
        tags.append(f"H{k}")
        blocks.append(("p", tagged(f"H{k}", random.Random(180 + k).randint(50, 70), 180 + k)))
    for number, page in enumerate(pages, 1):
        page.insert_text((50, 35), "Journal of Ordering Vol. 3", fontsize=8, fontname="helv")
        page.insert_text((480, 820), f"Page {number}", fontsize=8, fontname="helv")
    flow(pages, [(k, x0, 60, x1, 790) for k in range(3) for x0, x1 in ((50, 290), (305, 545))], blocks)
    return _save(doc, path), tags


def rotated_two_column_pdf(path: Path, rotation: int) -> Tuple[Path, List[str]]:
    """Port of the reviewer's ``L14``/``L15``: one landscape page built as a portrait MediaBox with
    ``/Rotate`` 90 or 270, set upright in two ragged-right columns under a title. Returns the tags
    in reading order."""
    doc = pymupdf.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    page.set_rotation(rotation)
    matrix = page.derotation_matrix
    page.insert_text(pymupdf.Point(40, 45) * matrix, "[RT] Landscape report title [/RT]", fontsize=14, fontname="hebo",
                     rotate=rotation)
    tags = ["RT"]
    frames = [(40, 70, 410, 560), (430, 70, 800, 560)]
    frame, y = 0, None
    seed = 160 if rotation == 90 else 170
    for k in range(1, 11):
        tags.append(f"R{k}")
        words = tagged(f"R{k}", random.Random(seed + k).randint(40, 60), seed + k).split()
        while words:
            x0, y0, x1, y1 = frames[frame]
            y = y0 + 9.5 if y is None else y
            if y > y1:
                frame, y = frame + 1, None
                continue
            count = 1
            while count < len(words) and pymupdf.get_text_length(" ".join(words[:count + 1]), "helv", 9.5) <= x1 - x0:
                count += 1
            page.insert_text(pymupdf.Point(x0, y) * matrix, " ".join(words[:count]), fontsize=9.5, fontname="helv",
                             rotate=rotation)
            words = words[count:]
            y += 11.5
        y += 7
    return _save(doc, path), tags


def glossary_centred_terms_pdf(path: Path) -> Tuple[Path, List[Tuple[str, str]]]:
    """Port of the reviewer's ``L19_glossary_centred3``: six bold three-line terms at x=50, each
    centred vertically on its multi-line definition at x=180. Returns the (term, definition) tags."""
    doc = pymupdf.open()
    page = doc.new_page(width=A4[0], height=A4[1])
    y, pairs = 80, []
    for k in range(1, 7):
        lines = wrap(tagged(f"D{k}", 60 + k * 4, 300 + k), 360, 10)
        height = 12 * len(lines)
        page.insert_text((180, y), "\n".join(lines), fontsize=10, fontname="helv", lineheight=1.2)
        page.insert_text((50, y + height / 2 - 18), f"[G{k}] Long term\nname number\n{k} here", fontsize=10,
                         fontname="hebo", lineheight=1.2)
        y += height + 16
        pairs.append((f"G{k}", f"D{k}"))
    return _save(doc, path), pairs


STAMP = "arXiv:2609.01234v1 [cs.CL] 30 Sep 2026"
PAPER_SECTIONS = ["Abstract", "1. Introduction", "2. Related Work", "3. Method", "4. Results"]


def margin_stamp_paper_pdf(path: Path) -> Path:
    """Two pages of a two-column paper with an arXiv-style stamp set vertically (20pt, grey) in the
    left margin of page 1, larger than the 16pt title. Sections ``Abstract``, ``1. Introduction`` ..
    ``4. Results`` are bold 11pt headings; the last ones fall on page 2."""
    doc = pymupdf.open()
    for _ in range(2):
        doc.new_page(width=A4[0], height=A4[1])
    pages = [doc[0], doc[1]]
    title = "Ordering Text In Two Column Papers"
    width = pymupdf.get_text_length(title, "hebo", 16)
    pages[0].insert_text(((A4[0] - width) / 2, 70), title, fontsize=16, fontname="hebo")
    pages[0].insert_text((30, 620), STAMP, fontsize=20, fontname="tiro", color=(0.5, 0.5, 0.5), rotate=90)
    blocks = []
    for number, heading in enumerate(PAPER_SECTIONS):
        blocks.append(("h", heading))
        for k in range(3 if number else 1):
            tag = f"S{number}x{k}"
            blocks.append(("p", tagged(tag, random.Random(number * 10 + k).randint(60, 90), number * 10 + k)))
    flow(pages, [(0, 60, 100, 295, 790), (0, 310, 100, 545, 790), (1, 60, 60, 295, 790), (1, 310, 60, 545, 790)],
         blocks, size=9.5, leading=11.5)
    return _save(doc, path)
