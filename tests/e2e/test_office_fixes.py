"""E2E tests for the Word / PowerPoint / Excel fixes of the docs audit (PR #27, "Code issues found" 14, 19, 20,
21 and 23).

Every test builds its document at test time (tests/e2e/builders_office_fixes.py), runs the real ``doc2mark`` CLI
on it and asserts on the Markdown / JSON the CLI writes. OCR failures are real: ``--ocr openai`` without an API
key, the way the audit found them. ``cache_dir`` has no CLI switch, so the cache tests use the public Python API
in a subprocess, with the local stand-in for the OpenAI API (tests/e2e/fake_openai.py) answering the run that has
a key. Tesseract and LibreOffice run for real where a test needs them.
"""

import base64
import json
import os
import re
import subprocess
import sys
from typing import List
from urllib.parse import unquote

import pytest

from tests.e2e import builders_ocr
from tests.e2e import builders_office as office
from tests.e2e import builders_office_fixes as build
from tests.e2e import fake_openai as fake
from tests.e2e.fake_openai import FakeOpenAI
from tests.e2e.test_office import only_table, routed_via, sections, tables

UNAVAILABLE = "[image: OCR unavailable]"
NO_KEY = {"OPENAI_API_KEY": None, "OPENAI_BASE_URL": None}
SLIDE_NUMBER_PROMPT = chr(0x2039) + "#" + chr(0x203A)


def lines_of(markdown: str) -> List[str]:
    return [line.strip() for line in markdown.split("\n") if line.strip()]


def in_order(text: str, needles: List[str]) -> List[int]:
    """Offsets of ``needles`` in ``text``, each searched after the previous one (-1 once one is missing)."""
    offsets, start = [], 0
    for needle in needles:
        at = text.find(needle, start) if start >= 0 else -1
        offsets.append(at)
        start = at + len(needle) if at >= 0 else -1
    return offsets


def ocr_issues(result) -> dict:
    return ((result.json or {}).get("metadata") or {}).get("extra", {}).get("ocr_issues") or {}


def region(markdown: str, name: str) -> List[str]:
    """The text of every ``<!-- name -->`` ... ``<!-- /name -->`` block, in order."""
    return re.findall(rf"<!-- {name} -->(.*?)<!-- /{name} -->", markdown, re.S)


# --------------------------------------------------------------------------- 19: PowerPoint


def test_19_layout_prompts_and_field_placeholders_are_not_slide_text(run_cli, e2e_dir):
    """The layout's placeholder prompts ("Click to edit Master title style"), its date field and ``‹#›`` were
    added to every slide as captions, and so were slide-level date and slide-number FIELDS. A slide shows a
    layout placeholder only through its own placeholder, and a field's cached text is not text the slide sets;
    typed text in a footer or date placeholder is."""
    path = build.pptx_prompts_deck(e2e_dir / "prompts.pptx")

    result = run_cli(path, "--ocr", "none", fmt="both")

    assert result.exit_code == 0, result.describe()
    md = result.markdown
    assert "Click to edit" not in md, result.describe()
    assert SLIDE_NUMBER_PROMPT not in md, result.describe()
    assert "10/1/2026" not in md and "1/27/13" not in md, result.describe()  # the date fields, slide and layout
    parts = sections(md, "slide")
    assert sorted(parts) == [1, 2, 3, 4], result.describe()
    assert "2" not in [line.strip("* ") for line in lines_of(parts[2])], parts[2]  # slide 2's slide-number field
    assert "ACME CONFIDENTIAL FOOTER 7788" in parts[1] and "March 2026 edition" in parts[2], result.describe()
    for text in ("Quarterly Review 2026", "Operations team update", "Agenda", "First topic 8101", "Second topic 8102",
                 "Summary", "Closing remarks 8201"):
        assert md.count(text) == 1, f"{text!r}\n{result.describe()}"


def test_19_text_of_a_plain_shape_is_emitted_once(run_cli, e2e_dir):
    """A rectangle's text was emitted twice: by the text-frame reader and again by the fallback for "unknown"
    shape types (an AutoShape is not on its allow-list)."""
    path = build.pptx_prompts_deck(e2e_dir / "prompts.pptx")

    result = run_cli(path, "--ocr", "none")

    assert result.exit_code == 0, result.describe()
    assert result.markdown.count("UNIQUE BOX TEXT 5150") == 1, result.describe()
    assert result.markdown.count("Plain text box 5151") == 1, result.describe()


def test_19_text_the_layout_or_master_draws_on_slides_is_kept_once(run_cli, e2e_dir):
    """A static text box on a slide layout (``LAYOUT TAGLINE 6262``, on slides 2 and 4) or on the master
    (``MASTER BRAND LINE 6363``, on every slide) is shown on the slides. It is real text: kept once, on the first
    slide that shows it, like the first copy of a PDF running header (the layout copy was a caption on every
    slide; the master copy was lost)."""
    path = build.pptx_prompts_deck(e2e_dir / "prompts.pptx")

    result = run_cli(path, "--ocr", "none")

    assert result.exit_code == 0, result.describe()
    parts = sections(result.markdown, "slide")
    assert result.markdown.count("LAYOUT TAGLINE 6262") == 1 and "LAYOUT TAGLINE 6262" in parts[2], result.describe()
    assert result.markdown.count("MASTER BRAND LINE 6363") == 1 and "MASTER BRAND LINE 6363" in parts[1], (
        result.describe())


def test_19_slide_count_is_the_number_of_slides(run_cli, e2e_dir, sample_documents_dir):
    """``slide_count`` counted the text "Slide " in the Markdown (the 5-slide sample gave 1)."""
    built = run_cli(build.pptx_prompts_deck(e2e_dir / "prompts.pptx"), "--ocr", "none", fmt="json")
    sample = run_cli(sample_documents_dir / "sample_presentation.pptx", "--ocr", "none", fmt="json")

    for result, slides in ((built, 4), (sample, 5)):
        assert result.exit_code == 0, result.describe()
        assert result.json["metadata"]["slide_count"] == slides, result.json["metadata"]
        assert result.json["metadata"]["page_count"] == slides, result.json["metadata"]


def test_19_tracked_sample_deck_has_no_prompts_and_keeps_its_text(run_cli, sample_documents_dir):
    result = run_cli(sample_documents_dir / "sample_presentation.pptx", "--ocr", "none")

    assert result.exit_code == 0, result.describe()
    md = result.markdown
    assert "Click to edit" not in md and SLIDE_NUMBER_PROMPT not in md and "6/13/25" not in md, result.describe()
    for text in ("Sample PPTX Presentation", "Generated on 2025-06-13", "Sample Content with Image",
                 "This slide demonstrates image embedding", "Structured Table Example",
                 "Unstructured Table (Merged Cells)", "Vertical Merging Example"):
        assert md.count(text) == 1, f"{text!r}\n{result.describe()}"


def test_19_converted_legacy_presentation_has_no_layout_text_or_duplicates(run_cli, require_tool, sample_documents_dir):
    """A .ppt goes through LibreOffice, whose layouts carry ``Footer`` and ``<#>`` placeholders and whose text
    boxes become AutoShapes: both leaked (the placeholders on every slide, the AutoShape text twice)."""
    require_tool("soffice")

    result = run_cli(sample_documents_dir / "sample_legacy_presentation.ppt", "--ocr", "none", fmt="both")

    assert result.exit_code == 0, result.describe()
    md = result.markdown
    assert "Footer" not in lines_of(md) and "*Footer*" not in md, result.describe()
    assert "&lt;#>" not in md and "<#>" not in md, result.describe()
    for text in ("Sample PPTX Presentation", "Structured Table Example", "Unstructured Table (Merged Cells)",
                 "Vertical Merging Example"):
        assert md.count(text) == 1, f"{text!r}\n{result.describe()}"
    assert result.json["metadata"]["slide_count"] == 5, result.json["metadata"]


def test_19_deck_taken_through_the_image_route_reports_its_slides(run_cli, require_tool, e2e_dir):
    require_tool("tesseract")
    require_tool("soffice")
    path = office.pptx_background_deck(e2e_dir / "backdrops.pptx", ["BACKDROP 7140", "BACKDROP 7241", "BACKDROP 7342"])

    result = run_cli(path, "--ocr", "tesseract", "--ocr-images", fmt="json")

    assert result.exit_code == 0, result.describe()
    assert routed_via(result) == "pdf", result.json["metadata"]
    assert result.json["metadata"]["slide_count"] == 3, result.json["metadata"]


# --------------------------------------------------------------------------- 20: Word


BOX_ORDER = ["Intro paragraph before the boxes 1000.", "Anchor paragraph 1100.", "TEXTBOX ALPHA 1001",
             "TEXTBOX ALPHA LINE TWO 1002", "Middle paragraph 1200.", "Shape anchor 1250.", "RECTANGLE TEXT 1251",
             "GROUP BOX ONE 2001", "GROUP BOX TWO 2002", "Legacy anchor 1300.", "LEGACY VML BOX 3001",
             "Closing paragraph 1400."]


def test_20_text_boxes_and_shapes_are_read_once_in_document_order(run_cli, e2e_dir):
    """Text in text boxes and shapes (DrawingML with its VML fallback, groups, VML-only boxes, a box in a table
    cell, a box in the header) was not extracted at all. Each box is read once (never from both the DrawingML
    and the VML copy), right after the text of the paragraph it is anchored in."""
    path = build.docx_text_boxes(e2e_dir / "boxes.docx")

    result = run_cli(path, "--ocr", "none", fmt="both")

    assert result.exit_code == 0, result.describe()
    md = result.markdown
    for text in BOX_ORDER + ["CELL BOX 4001", "HEADER BOX 5001"]:
        assert md.count(text) == 1, f"{text!r} x{md.count(text)}\n{result.describe()}"
    offsets = in_order(md, BOX_ORDER)
    assert -1 not in offsets, list(zip(BOX_ORDER, offsets))
    table = only_table(md, result.describe)
    assert table.grid == [["Key", "Value"], ["boxed", "Cell text 4000\nCELL BOX 4001"]], table
    assert -1 not in in_order(md, ["Header text 5000", "HEADER BOX 5001"]), result.describe()


HEADERS_AND_FOOTERS = ["FIRST PAGE LETTERHEAD 3310", "HEADER LINE 4471", "FOOTER LINE 5582",
                       "SECOND SECTION HEADER 6093"]


def test_20_headers_and_footers_are_emitted_once_per_section_with_their_tables(run_cli, e2e_dir):
    """Header and footer paragraphs were only ``text:header`` / ``text:footer`` JSON items (never in the Markdown)
    and header tables were lost. Each header and footer a section shows is written once, in a marked block:
    headers where the section starts, footers where it ends; a footer the next section inherits is not
    repeated."""
    path = build.docx_headers_and_footers(e2e_dir / "headers.docx")

    result = run_cli(path, "--ocr", "none", fmt="both")

    assert result.exit_code == 0, result.describe()
    md = result.markdown
    for text in HEADERS_AND_FOOTERS:
        assert md.count(text) == 1, f"{text!r}\n{result.describe()}"
    order = ["<!-- page 1 -->", "<!-- header -->", "FIRST PAGE LETTERHEAD 3310", "HEADER LINE 4471", "QR-7731",
             "<!-- /header -->", "Section one body 1001.", "Section one closing 1002.", "<!-- footer -->",
             "FOOTER LINE 5582", "<!-- /footer -->", "<!-- page 2 -->", "<!-- header -->",
             "SECOND SECTION HEADER 6093", "<!-- /header -->", "Section two body 2001."]
    offsets = in_order(md, order)
    assert -1 not in offsets, f"{list(zip(order, offsets))}\n{result.describe()}"
    assert [table.grid for table in tables(md)] == [[["Ref", "QR-7731"]]], result.describe()
    assert len(region(md, "header")) == 2 and len(region(md, "footer")) == 1, result.describe()
    items = result.json["json_content"]
    assert not [item for item in items if item["type"] in ("text:header", "text:footer")], items
    where = {item["content"]: item.get("region") for item in items if item["type"] != "table"}
    assert [where.get(text) for text in HEADERS_AND_FOOTERS] == ["header", "header", "footer", "header"], items
    assert where.get("Section one closing 1002.") is None, items


def test_20_header_picture_is_read_from_the_header(run_cli, e2e_dir):
    """A header picture's relationship id was looked up in the body's relationships: the Markdown got the bytes
    of another part (here the document's custom XML) as a PNG. It is the header's own picture, in the header."""
    logo = build.label_picture(build.HEADER_LOGO)
    path = build.docx_headers_and_footers(e2e_dir / "headers.docx", logo=logo)

    result = run_cli(path, "--ocr", "none", "--extract-images")

    assert result.exit_code == 0, result.describe()
    pictures = re.findall(r"data:image/png;base64,([A-Za-z0-9+/=]+)", result.markdown)
    assert [base64.b64decode(data) for data in pictures] == [logo], result.describe()[:2000]
    assert "data:image/png;base64," in region(result.markdown, "header")[0], result.describe()[:2000]


def test_20_header_picture_is_ocrd_in_the_header(run_cli, require_tool, e2e_dir):
    require_tool("tesseract")
    path = build.docx_headers_and_footers(e2e_dir / "headers.docx")

    result = run_cli(path, "--ocr", "tesseract", "--ocr-images", fmt="both")

    assert result.exit_code == 0, result.describe()
    assert routed_via(result) != "pdf", result.json["metadata"]
    assert result.markdown.count("9921") == 1, result.describe()
    assert "9921" in region(result.markdown, "header")[0], result.describe()


def test_20_only_caption_shaped_paragraphs_become_captions(run_cli, e2e_dir):
    """Any paragraph starting with Table/Figure/Chart/Image... was an italic caption ("Tablets are popular ...").
    A caption is a paragraph in a caption style, or a caption word with a number and a separator
    ("Figure 2: ...")."""
    path = build.docx_captions(e2e_dir / "captions.docx")

    result = run_cli(path, "--ocr", "none", fmt="both")

    assert result.exit_code == 0, result.describe()
    lines = lines_of(result.markdown)
    for text in build.NOT_CAPTIONS:
        assert text in lines, f"{text!r} is not a plain paragraph\n{result.describe()}"
    for text in build.CAPTIONS + [build.STYLED_CAPTION]:
        assert f"*{text}*" in lines, f"{text!r} is not a caption\n{result.describe()}"
    types = {item["content"]: item["type"] for item in result.json["json_content"]}
    assert {types[text] for text in build.NOT_CAPTIONS} == {"text:normal"}, types
    assert {types[text] for text in build.CAPTIONS + [build.STYLED_CAPTION]} == {"text:caption"}, types


def test_20_bold_italic_and_links_are_markdown_and_the_text_stays_verbatim(run_cli, e2e_dir):
    """Bold, italics and link targets were dropped. Runs are written **bold**, *italic* and [text](url) where the
    markers keep words whole; link targets other than http(s) and mailto, and links to bookmarks, keep only their
    text; headings and table cells stay plain; the JSON items keep the verbatim text."""
    path = build.docx_formatting(e2e_dir / "formatting.docx")

    result = run_cli(path, "--ocr", "none", fmt="both")

    assert result.exit_code == 0, result.describe()
    lines = lines_of(result.markdown)
    for expected in [
        "Plain **bold** and *italic* then [link text](https://example.com/docs?id=7).",
        "***Both styles*** here.",
        "superscript stays whole.",
        "unsafe link and see section and [email us](mailto:sales@example.com).",
        "A [field link](https://example.org/field) and a [complex link](https://example.org/complex).",
        "Markup **&lt;b>not html&lt;/b>** stays text.",
        "- **Important** item text",
        "## Bold heading",
        "的資料治理、流程",
    ]:
        assert expected in lines, f"{expected!r}\n{result.describe()}"
    page = builders_ocr.render(result.markdown)
    links = [(a.get_text(), unquote(a["href"])) for a in page.find_all("a")]
    assert links == [("link text", "https://example.com/docs?id=7"), ("email us", "mailto:sales@example.com"),
                     ("Python", build.WIKI_URL), ("field link", "https://example.org/field"),
                     ("complex link", "https://example.org/complex")], links
    assert [b.get_text() for b in page.find_all("strong")] == ["bold", "Both styles", "<b>not html</b>", "Important"]
    assert builders_ocr.active_html(result.markdown) == [], result.describe()
    assert only_table(result.markdown, result.describe).grid == [["Label", "Bold cell"]]
    assert "**Bold cell**" not in result.markdown, result.describe()
    contents = [item["content"] for item in result.json["json_content"]]
    assert "Plain bold and italic then link text." in contents, contents
    assert "Read about Python online." in contents, contents


# --------------------------------------------------------------------------- 21: Excel


def test_21_workbook_reports_its_sheet_names_and_filled_cells(run_cli, e2e_dir, sample_documents_dir):
    """``sheet_names`` and ``total_cells`` were only set by the basic fallback converter."""
    import openpyxl

    built = run_cli(build.xlsx_for_metadata(e2e_dir / "meta.xlsx"), "--ocr", "none", fmt="json")
    sample_path = sample_documents_dir / "sample_spreadsheet.xlsx"
    sample = run_cli(sample_path, "--ocr", "none", fmt="json")

    assert built.exit_code == 0 and sample.exit_code == 0, built.describe() + sample.describe()
    metadata = built.json["metadata"]
    assert metadata["sheet_names"] == ["Sales", "Notes", "Empty"], metadata
    assert metadata["total_cells"] == 10, metadata
    assert sample.json["metadata"]["sheet_names"] == openpyxl.load_workbook(sample_path).sheetnames
    assert sample.json["metadata"]["total_cells"] > 0, sample.json["metadata"]


# --------------------------------------------------------------------------- 23 and 14: pictures whose OCR failed


OFFICE_PICTURE_CASES = {
    "docx": (build.docx_with_pictures, [{"issue": "failed", "image": 1, "page": 1},
                                        {"issue": "failed", "image": 2, "page": 2}]),
    "pptx": (build.pptx_with_pictures, [{"issue": "failed", "image": 1, "slide": 1},
                                        {"issue": "failed", "image": 2, "slide": 2}]),
    "xlsx": (build.xlsx_with_pictures, [{"issue": "failed", "image": 1, "sheet": "Parts"},
                                        {"issue": "failed", "image": 2, "sheet": "Parts"}]),
}


@pytest.mark.parametrize("kind", sorted(OFFICE_PICTURE_CASES))
def test_23_office_pictures_whose_ocr_failed_are_marked_and_recorded(run_cli, e2e_dir, kind):
    """A picture whose OCR request raised (no API key) left the literal text ``OCR failed`` in the Markdown
    (``[Image: OCR failed]`` in a cell), as if the picture said so, and ``ocr_issues`` counted no failure. It gets
    the PDF path's marker, and the failure is counted with its page, slide or sheet."""
    builder, locations = OFFICE_PICTURE_CASES[kind]
    path = builder(e2e_dir / f"pictures.{kind}")

    result = run_cli(path, "--ocr", "openai", "--ocr-images", fmt="both", env=NO_KEY)

    assert result.exit_code == 0, result.describe()
    assert routed_via(result) != "pdf", result.json["metadata"]
    md = result.markdown
    assert "OCR failed" not in md, result.describe()
    assert md.count(UNAVAILABLE) == 2, result.describe()
    if kind in ("docx", "xlsx"):
        grid = only_table(sections(md, "sheet")[2] if kind == "xlsx" else md, result.describe).grid
        assert grid[1] == ["Logo", UNAVAILABLE], grid
    issues = ocr_issues(result)
    assert issues.get("failed") == 2, issues
    assert issues.get("locations") == locations, issues
    assert any("API key" in error for error in issues.get("errors", [])), issues


def test_23_image_file_whose_ocr_failed_is_marked_and_recorded(run_cli, e2e_dir):
    path = build.picture_file(e2e_dir / "lot.png")

    result = run_cli(path, "--ocr", "openai", "--ocr-images", fmt="both", env=NO_KEY)

    assert result.exit_code == 0, result.describe()
    assert "OCR extraction failed" not in result.markdown and "API key" not in result.markdown, result.describe()
    assert result.markdown.count(UNAVAILABLE) == 1, result.describe()
    assert ocr_issues(result).get("failed") == 1, result.json["metadata"]


CACHE_SCRIPT = (
    "import json, sys\n"
    "from doc2mark import UnifiedDocumentLoader\n"
    "path, cache, key = sys.argv[1:4]\n"
    "first = UnifiedDocumentLoader(ocr_provider='openai', cache_dir=cache).load(path, ocr_images=True)\n"
    "second = UnifiedDocumentLoader(ocr_provider='openai', api_key=key, cache_dir=cache).load(path, ocr_images=True)\n"
    "print(json.dumps({'first': first.content, 'second': second.content,\n"
    "                  'issues': (first.metadata.extra or {}).get('ocr_issues')}))\n"
)

CACHE_CASES = {
    "docx": build.docx_with_pictures,
    "pptx": build.pptx_with_pictures,
    "xlsx": build.xlsx_with_pictures,
    "png": build.picture_file,
}


@pytest.mark.parametrize("kind", sorted(CACHE_CASES))
def test_14_failed_ocr_is_not_cached_and_the_run_with_a_key_reads_the_pictures(e2e_dir, kind):
    """A Word / PowerPoint / Excel or image file whose OCR request raised (no API key) was stored in ``cache_dir``
    with the error text and replayed after the key was fixed: the error only reached ``ocr_issues["errors"]``, and
    ``failed`` stayed 0. The failure counts, the document is not stored, and the next run reads the pictures."""
    path = CACHE_CASES[kind](e2e_dir / f"pictures.{kind}")
    with FakeOpenAI() as server:
        server.script(structured=[fake.page("LOT 4471 PART")], free_form=[fake.text("LOT 4471 PART")])
        env = {key: value for key, value in os.environ.items() if key != "OPENAI_API_KEY"}
        env["OPENAI_BASE_URL"] = server.env["OPENAI_BASE_URL"]
        proc = subprocess.run([sys.executable, "-c", CACHE_SCRIPT, str(path), str(e2e_dir / "document-cache"),
                               fake.API_KEY], cwd=e2e_dir, capture_output=True, text=True, encoding="utf-8",
                              timeout=600, env=env)
        requests = len(server.requests_of("structured"))

    assert proc.returncode == 0, f"exit {proc.returncode}\n{proc.stdout}\n{proc.stderr}"
    output = json.loads(proc.stdout.strip().splitlines()[-1])
    assert "OCR failed" not in output["first"] and "OCR extraction failed" not in output["first"], output
    assert UNAVAILABLE in output["first"], output
    assert (output["issues"] or {}).get("failed", 0) >= 1, output
    assert requests >= 1, output
    assert "LOT 4471 PART" in output["second"] and UNAVAILABLE not in output["second"], output
