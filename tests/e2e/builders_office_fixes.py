"""Fixture builders for tests/e2e/test_office_fixes.py (docs-audit issues 14, 19, 20, 21 and 23), made at test time.

Documents are written with python-docx, python-pptx and openpyxl, plus raw OOXML where those libraries cannot
express a construct: Word text boxes and shapes (DrawingML ``wps:txbx`` with its VML fallback, groups, VML-only
text boxes), hyperlink fields, and slide-level date / footer / slide-number placeholders. The markup follows what
Word and PowerPoint write. Nothing from ``doc2mark`` is imported. Every builder saves to ``path`` and returns it as
a ``Path``.
"""

import io
from pathlib import Path
from typing import Optional, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.shared import Inches
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu

from tests.e2e import builders_office as office

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NAMESPACES = (
    f'xmlns:w="{W_NS}" xmlns:r="{R_NS}" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office"'
)
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

# Filler that keeps a document with small pictures on the native (text) path of the Office OCR route.
FILLER = ("The maintenance team replaced the seals on every pump, logged the pressure tests and "
          "closed the open audit items before the end of the quarter.")


def label_picture(text: str, size=(900, 260)) -> bytes:
    """A small PNG showing ``text`` (distinct text gives distinct picture bytes)."""
    return office.text_picture(text, size=size, font_size=90)


def _w_paragraphs(lines: Sequence[str]) -> str:
    return "".join(f'<w:p><w:r><w:t xml:space="preserve">{line}</w:t></w:r></w:p>' for line in lines)


# --------------------------------------------------------------------------- DOCX text boxes


class _Ids:
    """Unique drawing ids (``wp:docPr/@id``) and VML shape ids within one document."""

    def __init__(self):
        self.next = 100

    def take(self) -> int:
        self.next += 1
        return self.next


def _wsp(shape_id: int, lines: Sequence[str], text_box: bool, in_group: bool = False) -> str:
    """A DrawingML shape with text (``wps:wsp`` + ``wps:txbx``): a text box, or a rectangle with text."""
    non_visual = f'<wps:cNvPr id="{shape_id}" name="Shape {shape_id}"/>' if in_group else ""
    kind = ' txBox="1"' if text_box else ""
    return (
        f'<wps:wsp>{non_visual}<wps:cNvSpPr{kind}/>'
        '<wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="2286000" cy="457200"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></wps:spPr>'
        f'<wps:txbx><w:txbxContent>{_w_paragraphs(lines)}</w:txbxContent></wps:txbx>'
        '<wps:bodyPr rot="0" vert="horz" wrap="square" lIns="91440" tIns="45720" rIns="91440" bIns="45720" '
        'anchor="t" anchorCtr="0"><a:noAutofit/></wps:bodyPr></wps:wsp>'
    )


def _anchor(ids: _Ids, graphic_data: str, uri: str, cy: int = 457200) -> str:
    doc_id = ids.take()
    return (
        '<w:drawing><wp:anchor distT="0" distB="0" distL="114300" distR="114300" simplePos="0" '
        f'relativeHeight="{251658240 + doc_id}" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="column"><wp:posOffset>3200400</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="2286000" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:wrapSquare wrapText="bothSides"/><wp:docPr id="{doc_id}" name="Drawing {doc_id}"/>'
        '<wp:cNvGraphicFramePr/>'
        f'<a:graphic><a:graphicData uri="{uri}">{graphic_data}</a:graphicData></a:graphic>'
        '</wp:anchor></w:drawing>'
    )


def _vml_shape(ids: _Ids, lines: Sequence[str], top_pt: int = 0) -> str:
    return (
        f'<v:shape id="_x0000_s{1024 + ids.take()}" type="#_x0000_t202" '
        f'style="position:absolute;margin-left:252pt;margin-top:{top_pt}pt;width:180pt;height:36pt;z-index:1">'
        f'<v:textbox><w:txbxContent>{_w_paragraphs(lines)}</w:txbxContent></v:textbox></v:shape>'
    )


def text_box_run(ids: _Ids, lines: Sequence[str], text_box: bool = True) -> str:
    """A run holding a floating text box (or a rectangle with text: ``text_box=False``) the way Word 2010+ saves
    it: ``mc:AlternateContent`` with the DrawingML shape in ``mc:Choice`` and the same text again in a VML
    ``mc:Fallback`` for older readers."""
    shape = _wsp(0, lines, text_box)
    return (
        f'<w:r {NAMESPACES}><mc:AlternateContent><mc:Choice Requires="wps">'
        f'{_anchor(ids, shape, "http://schemas.microsoft.com/office/word/2010/wordprocessingShape")}'
        f'</mc:Choice><mc:Fallback><w:pict>{_vml_shape(ids, lines)}</w:pict></mc:Fallback>'
        '</mc:AlternateContent></w:r>'
    )


def group_run(ids: _Ids, boxes: Sequence[Sequence[str]]) -> str:
    """A run holding a group (``wpg:wgp``) of text boxes, with its VML ``v:group`` fallback."""
    shapes = "".join(_wsp(ids.take(), lines, True, in_group=True) for lines in boxes)
    group = (
        '<wpg:wgp><wpg:cNvGrpSpPr/><wpg:grpSpPr><a:xfrm><a:off x="0" y="0"/>'
        f'<a:ext cx="2286000" cy="{457200 * len(boxes)}"/><a:chOff x="0" y="0"/>'
        f'<a:chExt cx="2286000" cy="{457200 * len(boxes)}"/></a:xfrm></wpg:grpSpPr>{shapes}</wpg:wgp>'
    )
    fallback = "".join(_vml_shape(ids, lines, top_pt=36 * index) for index, lines in enumerate(boxes))
    return (
        f'<w:r {NAMESPACES}><mc:AlternateContent><mc:Choice Requires="wpg">'
        f'{_anchor(ids, group, "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup", 457200 * len(boxes))}'
        f'</mc:Choice><mc:Fallback><w:pict><v:group id="_x0000_g{ids.take()}" '
        'style="position:absolute;margin-left:252pt;width:180pt;height:72pt" coordsize="3600,1440">'
        f'{fallback}</v:group></w:pict></mc:Fallback></mc:AlternateContent></w:r>'
    )


def vml_text_box_run(ids: _Ids, lines: Sequence[str]) -> str:
    """A run holding a VML-only text box (``w:pict``), as Word 2003-era files have them."""
    return f'<w:r {NAMESPACES}><w:pict>{_vml_shape(ids, lines)}</w:pict></w:r>'


def _append_run(paragraph, run_xml: str) -> None:
    paragraph._p.append(parse_xml(run_xml))


def docx_text_boxes(path: Path) -> Path:
    """Body paragraphs with a text box, a rectangle with text, a group of two text boxes, a VML-only text box, a
    table cell holding a text box, and a header holding a text box. Every box text is unique; the paragraphs
    around them say where the boxes are anchored."""
    ids = _Ids()
    document = Document()
    document.add_paragraph("Intro paragraph before the boxes 1000.")
    anchor = document.add_paragraph("Anchor paragraph 1100.")
    _append_run(anchor, text_box_run(ids, ["TEXTBOX ALPHA 1001", "TEXTBOX ALPHA LINE TWO 1002"]))
    document.add_paragraph("Middle paragraph 1200.")
    shape = document.add_paragraph("Shape anchor 1250.")
    _append_run(shape, text_box_run(ids, ["RECTANGLE TEXT 1251"], text_box=False))
    grouped = document.add_paragraph()
    _append_run(grouped, group_run(ids, [["GROUP BOX ONE 2001"], ["GROUP BOX TWO 2002"]]))
    legacy = document.add_paragraph("Legacy anchor 1300.")
    _append_run(legacy, vml_text_box_run(ids, ["LEGACY VML BOX 3001"]))
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Key"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "boxed"
    cell = table.cell(1, 1).paragraphs[0]
    cell.add_run("Cell text 4000")
    _append_run(cell, text_box_run(ids, ["CELL BOX 4001"]))
    document.add_paragraph("Closing paragraph 1400.")
    header = document.sections[0].header.paragraphs[0]
    header.add_run("Header text 5000")
    _append_run(header, text_box_run(ids, ["HEADER BOX 5001"]))
    document.save(str(path))
    return Path(path)


# --------------------------------------------------------------------------- DOCX headers and footers


HEADER_LOGO = "ACME 9921"


def docx_headers_and_footers(path: Path, logo: Optional[bytes] = None) -> Path:
    """Two sections (the second starts on a new page).

    Section 1 has a different first page: its first-page header ``FIRST PAGE LETTERHEAD 3310``; its default
    header ``HEADER LINE 4471``, a picture (``logo``, by default one showing ``ACME 9921``) and a table
    (``Ref`` | ``QR-7731``); its footer ``FOOTER LINE 5582``. Section 2 has its own header
    ``SECOND SECTION HEADER 6093`` and inherits (links to) the footer of section 1."""
    document = Document()
    first = document.sections[0]
    first.different_first_page_header_footer = True
    first.first_page_header.paragraphs[0].text = "FIRST PAGE LETTERHEAD 3310"
    header = first.header
    header.paragraphs[0].text = "HEADER LINE 4471"
    header.paragraphs[0].add_run().add_picture(io.BytesIO(logo or label_picture(HEADER_LOGO)), width=Inches(1.2))
    reference = header.add_table(rows=1, cols=2, width=Inches(4))
    reference.cell(0, 0).text = "Ref"
    reference.cell(0, 1).text = "QR-7731"
    first.footer.paragraphs[0].text = "FOOTER LINE 5582"
    document.add_paragraph("Section one body 1001. " + FILLER)
    document.add_paragraph("Section one closing 1002.")
    second = document.add_section(WD_SECTION.NEW_PAGE)
    second.different_first_page_header_footer = False
    second.header.is_linked_to_previous = False
    second.header.paragraphs[0].text = "SECOND SECTION HEADER 6093"
    document.add_paragraph("Section two body 2001. " + FILLER)
    document.save(str(path))
    return Path(path)


# --------------------------------------------------------------------------- DOCX captions


NOT_CAPTIONS = [
    "Tablets are popular devices this year.",
    "Table of contents follows below.",
    "Figure 1 shows the revenue growing.",
    "Charter members met on Monday.",
    "Images of the site are attached.",
]
CAPTIONS = [
    "Figure 2: Revenue by quarter",
    "Table 3. Regional totals",
]
STYLED_CAPTION = "Photo of the new plant"


def docx_captions(path: Path) -> Path:
    """Body paragraphs that start with a caption word but are not captions, two caption-shaped paragraphs in
    the Normal style, and one paragraph in Word's Caption style."""
    document = Document()
    for text in NOT_CAPTIONS + CAPTIONS:
        document.add_paragraph(text)
    document.add_paragraph(STYLED_CAPTION, style="Caption")
    document.save(str(path))
    return Path(path)


# --------------------------------------------------------------------------- DOCX bold, italics and links


def _hyperlink(paragraph, text: str, url: Optional[str] = None, anchor: Optional[str] = None,
               bold: bool = False) -> None:
    attributes = ""
    if url is not None:
        rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        attributes += f' r:id="{rid}"'
    if anchor is not None:
        attributes += f' w:anchor="{anchor}"'
    run_props = "<w:rPr><w:rStyle w:val=\"Hyperlink\"/>" + ("<w:b/>" if bold else "") + "</w:rPr>"
    paragraph._p.append(parse_xml(
        f'<w:hyperlink {NAMESPACES}{attributes}><w:r>{run_props}<w:t xml:space="preserve">{text}</w:t></w:r>'
        '</w:hyperlink>'))


def _field_hyperlink(paragraph, text: str, url: str, simple: bool) -> None:
    if simple:
        paragraph._p.append(parse_xml(
            f'<w:fldSimple {NAMESPACES} w:instr=" HYPERLINK &quot;{url}&quot; "><w:r><w:t>{text}</w:t></w:r>'
            '</w:fldSimple>'))
        return
    for xml in ('<w:r {ns}><w:fldChar w:fldCharType="begin"/></w:r>',
                '<w:r {ns}><w:instrText xml:space="preserve"> HYPERLINK "{url}" \\o "tooltip" </w:instrText></w:r>',
                '<w:r {ns}><w:fldChar w:fldCharType="separate"/></w:r>',
                '<w:r {ns}><w:t>{text}</w:t></w:r>',
                '<w:r {ns}><w:fldChar w:fldCharType="end"/></w:r>'):
        paragraph._p.append(parse_xml(xml.format(ns=NAMESPACES, url=url, text=text)))


WIKI_URL = "https://en.wikipedia.org/wiki/Python_(programming_language)"


def docx_formatting(path: Path) -> Path:
    """Paragraphs with bold, italic and bold-italic runs, hyperlinks (external, internal bookmark, mailto,
    ``javascript:``, a URL with parentheses, simple and complex HYPERLINK fields), a bold run inside a word,
    markup-like text in bold, a bulleted item and a heading with bold runs, a table cell with a bold run, and a
    CJK sentence whose bold part ends with punctuation."""
    document = Document()
    p = document.add_paragraph("Plain ")
    p.add_run("bold").bold = True
    p.add_run(" and ")
    p.add_run("italic").italic = True
    p.add_run(" then ")
    _hyperlink(p, "link text", url="https://example.com/docs?id=7")
    p.add_run(".")

    p = document.add_paragraph()
    run = p.add_run("Both styles")
    run.bold = run.italic = True
    p.add_run(" here.")

    p = document.add_paragraph("super")
    p.add_run("script").bold = True
    p.add_run(" stays whole.")

    p = document.add_paragraph()
    _hyperlink(p, "unsafe link", url="javascript:alert(1)")
    p.add_run(" and ")
    _hyperlink(p, "see section", anchor="_Toc0001")
    p.add_run(" and ")
    _hyperlink(p, "email us", url="mailto:sales@example.com")
    p.add_run(".")

    p = document.add_paragraph("Read about ")
    _hyperlink(p, "Python", url=WIKI_URL)
    p.add_run(" online.")

    p = document.add_paragraph("A ")
    _field_hyperlink(p, "field link", "https://example.org/field", simple=True)
    p.add_run(" and a ")
    _field_hyperlink(p, "complex link", "https://example.org/complex", simple=False)
    p.add_run(".")

    p = document.add_paragraph("Markup ")
    p.add_run("<b>not html</b>").bold = True
    p.add_run(" stays text.")

    p = document.add_paragraph(style="List Bullet")
    p.add_run("Important").bold = True
    p.add_run(" item text")

    heading = document.add_heading("Bold ", level=2)
    heading.add_run("heading").bold = True

    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Label"
    table.cell(0, 1).paragraphs[0].add_run("Bold cell").bold = True

    p = document.add_paragraph("的")
    p.add_run("資料治理、").bold = True
    p.add_run("流程")
    document.save(str(path))
    return Path(path)


# --------------------------------------------------------------------------- PPTX placeholders and shapes


def _layout(prs, name: str):
    return next(layout for layout in prs.slide_layouts if layout.name == name)


def _placeholder_idx(layout, ph_type: str) -> str:
    for shape in layout.placeholders:
        ph = shape._element.find(f".//{{{P_NS}}}ph")
        if ph is not None and ph.get("type") == ph_type:
            return ph.get("idx")
    raise ValueError(f"layout {layout.name!r} has no {ph_type} placeholder")


def _add_footer_placeholder(slide, ph_type: str, paragraph_xml: str) -> None:
    """Add a slide-level date (``dt``), footer (``ftr``) or slide-number (``sldNum``) placeholder, as
    PowerPoint does when "Header & Footer" is applied to a slide (python-pptx does not copy these)."""
    tree = slide.shapes._spTree
    shape_id = max(int(e.get("id")) for e in tree.iter(f"{{{P_NS}}}cNvPr")) + 1
    idx = _placeholder_idx(slide.slide_layout, ph_type)
    tree.append(etree.fromstring(
        f'<p:sp xmlns:p="{P_NS}" xmlns:a="{A_NS}"><p:nvSpPr>'
        f'<p:cNvPr id="{shape_id}" name="{ph_type} Placeholder {shape_id}"/>'
        '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
        f'<p:nvPr><p:ph type="{ph_type}" sz="quarter" idx="{idx}"/></p:nvPr></p:nvSpPr><p:spPr/>'
        f'<p:txBody><a:bodyPr/><a:lstStyle/>{paragraph_xml}</p:txBody></p:sp>'))


def _field(field_type: str, text: str) -> str:
    return (f'<a:p><a:fld id="{{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}}" type="{field_type}">'
            f'<a:rPr lang="en-US"/><a:t>{text}</a:t></a:fld><a:endParaRPr lang="en-US"/></a:p>')


def _typed(text: str) -> str:
    return f'<a:p><a:r><a:rPr lang="en-US"/><a:t>{text}</a:t></a:r></a:p>'


def _add_static_text(shapes_owner, text: str, top: int) -> None:
    """A non-placeholder text box on a slide layout or master: drawn on every slide that uses it."""
    tree = shapes_owner.shapes._spTree
    shape_id = max(int(e.get("id")) for e in tree.iter(f"{{{P_NS}}}cNvPr")) + 1
    tree.append(etree.fromstring(
        f'<p:sp xmlns:p="{P_NS}" xmlns:a="{A_NS}"><p:nvSpPr>'
        f'<p:cNvPr id="{shape_id}" name="Tagline {shape_id}"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="457200" y="{top}"/><a:ext cx="4572000" cy="369332"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
        f'<p:txBody><a:bodyPr wrap="none"/><a:lstStyle/>{_typed(text)}</p:txBody></p:sp>'))


PPTX_REAL_TEXT = ["Quarterly Review 2026", "Operations team update", "Agenda", "First topic 8101",
                  "Second topic 8102", "UNIQUE BOX TEXT 5150", "Plain text box 5151", "Summary",
                  "Closing remarks 8201", "ACME CONFIDENTIAL FOOTER 7788", "March 2026 edition",
                  "LAYOUT TAGLINE 6262", "MASTER BRAND LINE 6363"]


def pptx_prompts_deck(path: Path) -> Path:
    """Four slides on python-pptx's default template (whose layouts carry PowerPoint's prompt texts:
    "Click to edit Master title style", a date field, ``‹#›``).

    1. Title Slide: title, subtitle, and slide-level placeholders: a date FIELD, a footer with typed text
       ``ACME CONFIDENTIAL FOOTER 7788`` and a slide-number FIELD.
    2. Title and Content: title ``Agenda``, two body lines, a slide-number field showing ``2`` and a date
       placeholder with typed text ``March 2026 edition``.
    3. Title Only with its title left empty, a rectangle with ``UNIQUE BOX TEXT 5150`` and a text box.
    4. Title and Content again: ``Summary`` / ``Closing remarks 8201``.

    The Title and Content layout also draws a static text box ``LAYOUT TAGLINE 6262`` (slides 2 and 4) and the
    slide master one with ``MASTER BRAND LINE 6363`` (every slide)."""
    prs = Presentation()
    content_layout = _layout(prs, "Title and Content")
    _add_static_text(content_layout, "LAYOUT TAGLINE 6262", 6172200)
    _add_static_text(prs.slide_master, "MASTER BRAND LINE 6363", 6400800)

    slide = prs.slides.add_slide(_layout(prs, "Title Slide"))
    slide.shapes.title.text = "Quarterly Review 2026"
    slide.placeholders[1].text = "Operations team update"
    _add_footer_placeholder(slide, "dt", _field("datetime1", "10/1/2026"))
    _add_footer_placeholder(slide, "ftr", _typed("ACME CONFIDENTIAL FOOTER 7788"))
    _add_footer_placeholder(slide, "sldNum", _field("slidenum", "‹#›"))

    slide = prs.slides.add_slide(content_layout)
    slide.shapes.title.text = "Agenda"
    slide.placeholders[1].text_frame.text = "First topic 8101"
    slide.placeholders[1].text_frame.add_paragraph().text = "Second topic 8102"
    _add_footer_placeholder(slide, "sldNum", _field("slidenum", "2"))
    _add_footer_placeholder(slide, "dt", _typed("March 2026 edition"))

    slide = prs.slides.add_slide(_layout(prs, "Title Only"))
    rectangle = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(914400), Emu(1828800), Emu(3657600), Emu(914400))
    rectangle.text_frame.text = "UNIQUE BOX TEXT 5150"
    box = slide.shapes.add_textbox(Emu(914400), Emu(3200400), Emu(3657600), Emu(914400))
    box.text_frame.text = "Plain text box 5151"

    slide = prs.slides.add_slide(content_layout)
    slide.shapes.title.text = "Summary"
    slide.placeholders[1].text = "Closing remarks 8201"
    prs.save(str(path))
    return Path(path)


# --------------------------------------------------------------------------- OCR failures (no API key)


def docx_with_pictures(path: Path, cell_label: str = "CELL LOGO 2601", body_label: str = "BODY PHOTO 2602") -> Path:
    """Page 1: a paragraph and a table whose cell holds a picture; a page break; page 2: a paragraph holding a
    second picture. Small pictures and enough text, so the document stays on the native path."""
    document = Document()
    document.add_paragraph("Parts list for the pump stations. " + FILLER)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Picture"
    table.cell(1, 0).text = "Logo"
    table.cell(1, 1).paragraphs[0].add_run().add_picture(io.BytesIO(label_picture(cell_label)), width=Inches(1.0))
    document.add_paragraph("Last line of page one.").add_run().add_break(WD_BREAK.PAGE)
    photo = document.add_paragraph("Site photo: ")
    photo.add_run().add_picture(io.BytesIO(label_picture(body_label)), width=Inches(1.5))
    document.add_paragraph("Page two text. " + FILLER)
    document.save(str(path))
    return Path(path)


def pptx_with_pictures(path: Path, labels: Sequence[str] = ("SLIDE PHOTO 2701", "SLIDE PHOTO 2702")) -> Path:
    """One text slide per label, each with a small picture showing that label (text-dominant: the Office
    pipeline itself OCRs the pictures)."""
    prs = Presentation()
    for number, label in enumerate(labels, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Emu(457200), Emu(457200), Emu(7315200), Emu(1828800))
        box.text_frame.word_wrap = True
        box.text_frame.text = f"Slide {number}: {FILLER}"
        slide.shapes.add_picture(io.BytesIO(label_picture(label)), Emu(5486400), Emu(4114800), Emu(2743200),
                                 Emu(792480))
    prs.save(str(path))
    return Path(path)


def xlsx_with_pictures(path: Path, cell_label: str = "CELL PART 2801", outside_label: str = "LOOSE PART 2802") -> Path:
    """Sheet ``Summary`` (text only) and sheet ``Parts`` with a picture anchored in table cell B2 and another
    anchored at E9, outside the table."""
    return office.workbook(path, [
        {"title": "Summary", "rows": [["Total", 1]]},
        {"title": "Parts", "rows": [["Item", "Picture"], ["Logo", None], ["Valve", 3]],
         "images": {"B2": label_picture(cell_label), "E9": label_picture(outside_label)}},
    ])


def picture_file(path: Path, label: str = "LOT 4471") -> Path:
    """A PNG image file showing ``label``."""
    Path(path).write_bytes(label_picture(label, size=(1200, 300)))
    return Path(path)


# --------------------------------------------------------------------------- XLSX metadata


def xlsx_for_metadata(path: Path) -> Path:
    """Sheets ``Sales`` (8 filled cells in a 3 x 3 grid plus a merged note row: 9), ``Notes`` (1) and
    ``Empty`` (none): 10 filled cells."""
    return office.workbook(path, [
        {"title": "Sales", "rows": [["Region", "Q1", "Q2"], ["North", 10, None], ["South", 30, 40], [],
                                    ["All figures in EUR", None, None]],
         "merges": ["A5:C5"]},
        {"title": "Notes", "rows": [["Reviewed by finance"]]},
        {"title": "Empty", "rows": []},
    ])
