"""Office format processors (DOCX, XLSX, PPTX)."""

import logging
import zipfile
import base64
from pathlib import Path
from typing import Any, Dict, List, Optional, Union, Tuple

from doc2mark.core.base import (
    BaseProcessor,
    DocumentFormat,
    DocumentMetadata,
    ProcessedDocument,
    ProcessingError
)
from doc2mark.ocr.base import BaseOCR

# Import the advanced pipeline loader
try:
    from doc2mark.pipelines.office_advanced_pipeline import (
        DocxLoader, PptxLoader, XlsxLoader, UniversalOfficeLoader
    )
    ADVANCED_PIPELINE_AVAILABLE = True
except ImportError:
    ADVANCED_PIPELINE_AVAILABLE = False
    logging.warning("Advanced Office pipeline not available. Using basic processing.")

logger = logging.getLogger(__name__)

_P_NS = 'http://schemas.openxmlformats.org/presentationml/2006/main'
_A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
_W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
_WP_NS = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
_MC_NS = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
_W_T = f'{{{_W_NS}}}t'
_WP_INLINE = f'{{{_WP_NS}}}inline'
_WP_ANCHOR = f'{{{_WP_NS}}}anchor'
_WP_EXTENT = f'{{{_WP_NS}}}extent'
_PIC_PIC = '{http://schemas.openxmlformats.org/drawingml/2006/picture}pic'
# Subtrees whose content the rendered document does not show (deleted revisions)
# or shows a second copy of (mc:Fallback duplicates the mc:Choice drawing).
_DOCX_UNRENDERED = {f'{{{_W_NS}}}del', f'{{{_W_NS}}}moveFrom', f'{{{_MC_NS}}}Fallback'}


def _int_attr(element, name: str, default: int = 0) -> int:
    try:
        return int(element.get(name, default))
    except (TypeError, ValueError):
        return default


def _union_area(rects, width: float, height: float) -> float:
    """Area covered by the union of ``(x0, y0, x1, y1)`` rectangles clipped to the page."""
    clipped = []
    for x0, y0, x1, y1 in rects:
        x0, x1 = sorted((x0, x1))
        y0, y1 = sorted((y0, y1))
        x0, y0, x1, y1 = max(x0, 0.0), max(y0, 0.0), min(x1, width), min(y1, height)
        if x1 > x0 and y1 > y0:
            clipped.append((x0, y0, x1, y1))
    xs = sorted({x for rect in clipped for x in (rect[0], rect[2])})
    area = 0.0
    for left, right in zip(xs, xs[1:]):
        spans = sorted((r[1], r[3]) for r in clipped if r[0] <= left and r[2] >= right)
        covered, start, end = 0.0, None, None
        for top, bottom in spans:
            if end is None or top > end:
                covered += 0.0 if end is None else end - start
                start, end = top, bottom
            else:
                end = max(end, bottom)
        covered += 0.0 if end is None else end - start
        area += covered * (right - left)
    return area


def _pptx_background_picture(slide):
    """The ``a:blipFill`` of the background shown on ``slide`` (its own, else its
    layout's, else its master's), or None when that background is not a picture."""
    layout = slide.slide_layout
    for owner in (slide, layout, layout.slide_master):
        background = owner._element.find(f'{{{_P_NS}}}cSld/{{{_P_NS}}}bg')
        if background is not None:
            return background.find(f'{{{_P_NS}}}bgPr/{{{_A_NS}}}blipFill')
    return None


def _pptx_inherited_shapes(slide) -> list:
    """Shape trees of the layout and master that are drawn on ``slide``."""
    if slide._element.get('showMasterSp') in ('0', 'false'):
        return []
    layout = slide.slide_layout
    trees = [layout.shapes]
    if layout._element.get('showMasterSp') not in ('0', 'false'):
        trees.append(layout.slide_master.shapes)
    return trees


def _pptx_hidden(shape) -> bool:
    c_nv_pr = shape._element.find(f'./*/{{{_P_NS}}}cNvPr')
    return c_nv_pr is not None and c_nv_pr.get('hidden') in ('1', 'true')


def _pptx_walk_shapes(shapes, transform=(1.0, 0.0, 1.0, 0.0)):
    """Yield ``(shape, (x0, y0, x1, y1))`` for every visible shape, recursing into
    groups; rectangles are in slide coordinates (group child offsets/extents applied)."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    sx, dx, sy, dy = transform
    for shape in shapes:
        if _pptx_hidden(shape):
            continue
        try:
            left, top = float(shape.left or 0), float(shape.top or 0)
            width, height = float(shape.width or 0), float(shape.height or 0)
        except (AttributeError, TypeError, ValueError):
            left = top = width = height = 0.0
        yield shape, (left * sx + dx, top * sy + dy, (left + width) * sx + dx, (top + height) * sy + dy)
        try:
            is_group = shape.shape_type == MSO_SHAPE_TYPE.GROUP
        except (AttributeError, NotImplementedError):
            is_group = False
        if not is_group:
            continue
        xfrm = shape._element.find(f'{{{_P_NS}}}grpSpPr/{{{_A_NS}}}xfrm')
        child = (1.0, 0.0, 1.0, 0.0)
        if xfrm is not None:
            off, ext = xfrm.find(f'{{{_A_NS}}}off'), xfrm.find(f'{{{_A_NS}}}ext')
            ch_off, ch_ext = xfrm.find(f'{{{_A_NS}}}chOff'), xfrm.find(f'{{{_A_NS}}}chExt')
            if None not in (off, ext, ch_off, ch_ext):
                csx = _int_attr(ext, 'cx') / _int_attr(ch_ext, 'cx') if _int_attr(ch_ext, 'cx') else 1.0
                csy = _int_attr(ext, 'cy') / _int_attr(ch_ext, 'cy') if _int_attr(ch_ext, 'cy') else 1.0
                child = (csx, _int_attr(off, 'x') - _int_attr(ch_off, 'x') * csx,
                         csy, _int_attr(off, 'y') - _int_attr(ch_off, 'y') * csy)
        yield from _pptx_walk_shapes(shape.shapes, (child[0] * sx, child[1] * sx + dx, child[2] * sy, child[3] * sy + dy))


def _pptx_is_picture(shape) -> bool:
    """A picture shape (placeholder pictures included) or a shape filled with a picture."""
    element = shape._element
    if element.tag == f'{{{_P_NS}}}pic':
        return True
    sp_pr = element.find(f'{{{_P_NS}}}spPr')
    return sp_pr is not None and sp_pr.find(f'{{{_A_NS}}}blipFill') is not None


def _pptx_text_length(shape) -> int:
    length = 0
    if getattr(shape, 'has_text_frame', False):
        length += len(shape.text_frame.text)
    if getattr(shape, 'has_table', False):
        length += sum(len(cell.text) for row in shape.table.rows for cell in row.cells)
    return length


def _docx_rendered_elements(root):
    """Every element under ``root`` in document order, minus deleted revisions and
    ``mc:Fallback`` duplicates."""
    stack = [iter(root)]
    while stack:
        for element in stack[-1]:
            if not isinstance(element.tag, str) or element.tag in _DOCX_UNRENDERED:
                continue
            yield element
            stack.append(iter(element))
            break
        else:
            stack.pop()


class OfficeProcessor(BaseProcessor):
    """Processor for modern Office formats (DOCX, XLSX, PPTX) with advanced features."""

    def __init__(self, ocr: Optional[BaseOCR] = None, table_style: Optional[str] = None):
        """Initialize Office processor.
        
        Args:
            ocr: OCR provider for image extraction
            table_style: Output style for complex tables:
                - 'minimal_html': Clean HTML with only rowspan/colspan (default)
                - 'markdown_grid': Markdown with merge annotations  
                - 'styled_html': Full HTML with inline styles (legacy)
        """
        self.ocr = ocr
        self.table_style = table_style
        self._docx = None
        self._openpyxl = None
        self._pptx = None

    @property
    def python_docx(self):
        """Lazy load python-docx."""
        if self._docx is None:
            try:
                import docx
                self._docx = docx
            except ImportError:
                raise ImportError(
                    "python-docx is not installed. "
                    "Install it with: pip install python-docx"
                )
        return self._docx

    @property
    def openpyxl(self):
        """Lazy load openpyxl."""
        if self._openpyxl is None:
            try:
                import openpyxl
                self._openpyxl = openpyxl
            except ImportError:
                raise ImportError(
                    "openpyxl is not installed. "
                    "Install it with: pip install openpyxl"
                )
        return self._openpyxl

    @property
    def python_pptx(self):
        """Lazy load python-pptx."""
        if self._pptx is None:
            try:
                import pptx
                self._pptx = pptx
            except ImportError:
                raise ImportError(
                    "python-pptx is not installed. "
                    "Install it with: pip install python-pptx"
                )
        return self._pptx

    def can_process(self, file_path: Union[str, Path]) -> bool:
        """Check if this processor can handle the file."""
        file_path = Path(file_path)
        extension = file_path.suffix.lower().lstrip('.')
        return extension in ['docx', 'xlsx', 'pptx']

    def process(
            self,
            file_path: Union[str, Path],
            **kwargs
    ) -> ProcessedDocument:
        """Process Office document using advanced pipeline if available."""
        file_path = Path(file_path)
        extension = file_path.suffix.lower().lstrip('.')

        # Get file size
        file_size = file_path.stat().st_size

        # Image-dominant office docs (docx/pptx that are mostly pictures with no real
        # text layer) have no faithful OOXML text to extract — route them through the
        # PDF image strategy (whole-page render OCR + page_markdown synthesis), the
        # same content-based decision the PDF pipeline already makes. Text/table office
        # docs stay on the native path; any failure falls back to native extraction
        # and is recorded in ``metadata.extra`` (``routed_via``/``route_error``).
        route_info: Dict[str, Any] = {}
        routed = self._maybe_route_image_dominant(file_path, file_size, route_info=route_info, **kwargs)
        if routed is not None:
            return routed

        # Use advanced pipeline if available
        json_content = None
        if ADVANCED_PIPELINE_AVAILABLE:
            content, metadata, images, json_content = self._process_with_advanced_pipeline(file_path, **kwargs)
        else:
            # Fallback to basic processing
            if extension == 'docx':
                content, metadata, images = self._process_docx_basic(file_path, **kwargs)
                doc_format = DocumentFormat.DOCX
            elif extension == 'xlsx':
                content, metadata, images = self._process_xlsx_basic(file_path, **kwargs)
                doc_format = DocumentFormat.XLSX
            elif extension == 'pptx':
                content, metadata, images = self._process_pptx_basic(file_path, **kwargs)
                doc_format = DocumentFormat.PPTX
            else:
                raise ProcessingError(f"Unsupported Office format: {extension}")

        # Determine format
        if extension == 'docx':
            doc_format = DocumentFormat.DOCX
        elif extension == 'xlsx':
            doc_format = DocumentFormat.XLSX
        elif extension == 'pptx':
            doc_format = DocumentFormat.PPTX

        # Build metadata
        doc_metadata = DocumentMetadata(
            filename=file_path.name,
            format=doc_format,
            size_bytes=file_size,
            **metadata
        )
        if route_info:
            doc_metadata.extra.update(route_info)

        return ProcessedDocument(
            content=content,
            metadata=doc_metadata,
            images=images,
            json_content=json_content
        )

    # ------------------------------------------------------------------ #
    # Office image-dominance route — delegates to the PDF image strategy   #
    # ------------------------------------------------------------------ #
    def _maybe_route_image_dominant(
        self, file_path: Path, file_size: int, route_info: Optional[Dict[str, Any]] = None, **kwargs
    ) -> Optional[ProcessedDocument]:
        """Return a ProcessedDocument via the PDF image strategy when this office
        doc is image-dominant; else None so the caller continues with native
        extraction. Gated to docx/pptx with OCR enabled; never raises (any failure
        — including no LibreOffice — falls back to native).

        The OOXML signals are only a cheap pre-filter. A document they flag is
        converted to PDF, and the converted PDF's own route decision is final: when
        it routes ``text`` the document goes back to native extraction. Whenever
        routing was attempted but the document stays native, ``route_info`` gets
        ``routed_via='native'`` plus ``route_reason`` (the PDF routes text) or
        ``route_error`` (conversion or PDF processing failed), so the caller can
        record it in ``metadata.extra`` instead of falling back silently.
        """
        ext = file_path.suffix.lower().lstrip('.')
        if ext not in ('docx', 'pptx'):                 # xlsx (data grids) never routes
            return None
        ocr_requested = bool(kwargs.get('ocr_images') and kwargs.get('extract_images', False))
        if self.ocr is None or not ocr_requested:
            return None
        try:
            if not self._is_image_dominant(file_path):
                return None
            return self._process_as_image_dominant(file_path, file_size, route_info=route_info, **kwargs)
        except Exception as e:
            logger.warning(
                f"Office image-dominance route unavailable ({e}); using native extraction"
            )
            if route_info is not None:
                route_info['routed_via'] = 'native'
                route_info['route_error'] = str(e).strip()[:500] or type(e).__name__
            return None

    def _is_image_dominant(self, file_path: Path) -> bool:
        """Two-signal image/text decision from the OOXML structure (no rendering)."""
        from doc2mark.core.strategy import decide_doc_strategy
        ext = file_path.suffix.lower().lstrip('.')
        if ext == 'pptx':
            mean_cov, mean_text = self._pptx_image_signals(file_path)
        elif ext == 'docx':
            mean_cov, mean_text = self._docx_image_signals(file_path)
        else:
            return False
        strategy = decide_doc_strategy(mean_cov, mean_text)
        logger.info(
            f"📑 Office strategy: {strategy} (mean coverage {mean_cov:.2f}, "
            f"mean text {mean_text:.0f} chars)"
        )
        return strategy == 'image'

    def _pptx_image_signals(self, file_path: Path) -> Tuple[float, float]:
        """Mean picture coverage + mean text chars per slide, counted the way the
        converted PDF shows them.

        Coverage is the union area, clipped to the slide, of every visible picture:
        picture shapes at any group depth, pictures in placeholders, shapes filled
        with a picture, the background picture fill that shows on the slide (its
        own, else its layout's, else its master's) and non-placeholder pictures the
        layout and master draw on it. Text counts every visible text frame and
        table cell, grouped ones included. Hidden shapes count for neither.
        """
        from pptx import Presentation
        prs = Presentation(str(file_path))
        width, height = float(prs.slide_width or 0), float(prs.slide_height or 0)
        slide_area = (width * height) or 1.0
        covs: List[float] = []
        texts: List[int] = []
        for slide in prs.slides:
            rects: List[Tuple[float, float, float, float]] = []
            if _pptx_background_picture(slide) is not None:
                rects.append((0.0, 0.0, width, height))
            for shapes in _pptx_inherited_shapes(slide):
                rects.extend(rect for shape, rect in _pptx_walk_shapes(shapes)
                             if not shape.is_placeholder and _pptx_is_picture(shape))
            txt = 0
            for shape, rect in _pptx_walk_shapes(slide.shapes):
                if _pptx_is_picture(shape):
                    rects.append(rect)
                txt += _pptx_text_length(shape)
            covs.append(min(_union_area(rects, width, height) / slide_area, 1.0))
            texts.append(txt)
        n = len(covs) or 1
        return sum(covs) / n, sum(texts) / n

    def _docx_image_signals(self, file_path: Path) -> Tuple[float, float]:
        """Total picture coverage (vs one page) + total text chars, counted the way
        the converted PDF shows them: text and pictures in the body (paragraphs,
        tables, content controls, text boxes) and in every header and footer.
        Deleted revisions and ``mc:Fallback`` copies of drawings are skipped.

        Totals (not per-page) suffice for this pre-filter: a real multi-page text doc
        easily exceeds the 200-char text limit, and a document it flags is converted
        so the converted PDF decides.
        """
        import docx
        d = docx.Document(str(file_path))
        roots = [d.element.body]
        for rel in d.part.rels.values():
            if rel.reltype.endswith(('/header', '/footer')) and not rel.is_external:
                roots.append(rel.target_part.element)
        text_len = 0
        img_area = 0.0
        for root in roots:
            for el in _docx_rendered_elements(root):
                if el.tag == _W_T:
                    text_len += len(el.text or '')
                elif el.tag in (_WP_INLINE, _WP_ANCHOR) and el.find(f'.//{_PIC_PIC}') is not None:
                    extent = el.find(_WP_EXTENT)
                    if extent is not None:
                        img_area += _int_attr(extent, 'cx') * float(_int_attr(extent, 'cy'))
        sect = d.sections[0]
        page_area = float((sect.page_width or 0) * (sect.page_height or 0)) or 1.0
        return min(img_area / page_area, 1.0), float(text_len)

    def _converted_pdf_route(self, pdf_path: Path) -> Optional[str]:
        """The PDF pipeline's document route (``"image"``/``"text"``) for the
        converted file, or None when it cannot be evaluated."""
        from doc2mark.pipelines.pymupdf_advanced_pipeline import PDFLoader
        loader = PDFLoader(pdf_path, ocr=self.ocr, table_style=self.table_style)
        try:
            decide = getattr(loader, '_document_image_strategy', None)
            return decide() if decide is not None else None
        finally:
            loader.close()

    def _process_as_image_dominant(
        self, file_path: Path, file_size: int, route_info: Optional[Dict[str, Any]] = None, **kwargs
    ) -> Optional[ProcessedDocument]:
        """Convert the office doc to PDF and, when the converted PDF routes
        ``image``, process it with the PDF image strategy (whole-page render OCR +
        page_markdown synthesis), then restore the office identity in the metadata.

        Only the converted PDF's decision counts. When it routes ``text`` (the OOXML
        pre-filter over-estimated the pictures) this returns None and the caller
        uses native extraction, which keeps the OOXML tables and structure; the PDF
        pipeline evaluates the same rule on the same file, so an ``image`` decision
        here cannot turn into a PDF text-route extraction.
        """
        import tempfile
        from doc2mark.utils.libreoffice import convert_office_to
        from doc2mark.formats.pdf import PDFProcessor
        ext = file_path.suffix.lower().lstrip('.')
        doc_format = DocumentFormat.DOCX if ext == 'docx' else DocumentFormat.PPTX
        with tempfile.TemporaryDirectory() as tmp:
            pdf_path = convert_office_to(file_path, 'pdf', tmp, timeout=300)
            if self._converted_pdf_route(pdf_path) == 'text':
                logger.info(f"📑 {file_path.name}: the converted PDF routes text; using native extraction")
                if route_info is not None:
                    route_info['routed_via'] = 'native'
                    route_info['route_reason'] = 'converted PDF routes text'
                return None
            pdf_proc = PDFProcessor(ocr=self.ocr, table_style=self.table_style)
            result = pdf_proc.process(
                pdf_path,
                extract_images=kwargs.get('extract_images', True),
                use_ocr=bool(kwargs.get('ocr_images', True)),
                show_progress=kwargs.get('show_progress', False),
            )
        # Restore the original office identity; record the routing for downstream.
        result.metadata.format = doc_format
        result.metadata.filename = file_path.name
        result.metadata.size_bytes = file_size
        if result.metadata.extra is None:
            result.metadata.extra = {}
        result.metadata.extra['routed_via'] = 'pdf'
        return result

    def _process_with_advanced_pipeline(
        self,
        file_path: Path,
        **kwargs
    ) -> Tuple[str, dict, List[Dict[str, Any]], Optional[List[Dict[str, Any]]]]:
        """Process document using advanced pipeline with all features."""
        try:
            # Configure options
            extract_images = kwargs.get('extract_images', False)
            ocr_images = kwargs.get('ocr_images', False) and extract_images and self.ocr is not None

            # Use the advanced pipeline
            json_data = UniversalOfficeLoader.load(
                file_path,
                extract_images=extract_images,
                ocr_images=ocr_images,
                show_progress=kwargs.get('show_progress', False),
                ocr=self.ocr,
                table_style=kwargs.get('table_style', self.table_style)
            )

            # Use office_to_markdown for content (includes page/slide/sheet markers)
            from doc2mark.pipelines.office_advanced_pipeline import office_to_markdown
            content = office_to_markdown(json_data)

            # Extract images from json_data
            images = []
            for item in json_data["content"]:
                if item["type"] == "image":
                    images.append({
                        'data': base64.b64decode(item["content"]),
                        'page': item.get('page', 1)
                    })

            json_content = json_data.get('content', [])
            
            # Build metadata
            metadata = {
                'page_count': json_data.get('pages', 1),
                'image_count': len(images),
            }
            
            # Add format-specific metadata
            if file_path.suffix.lower() == '.docx':
                metadata['word_count'] = len(content.split())
            elif file_path.suffix.lower() == '.xlsx':
                # XLSX specific metadata is already included in json_data
                pass
            elif file_path.suffix.lower() == '.pptx':
                # Count slides from content
                slide_count = content.count('Slide ')
                metadata['slide_count'] = max(slide_count, 1)
            
            return content, metadata, images, json_content

        except Exception as e:
            logger.warning(f"Advanced pipeline processing failed: {e}, falling back to basic processing")
            # Fallback to basic processing
            extension = file_path.suffix.lower().lstrip('.')
            if extension == 'docx':
                c, m, i = self._process_docx_basic(file_path, **kwargs)
                return c, m, i, None
            elif extension == 'xlsx':
                c, m, i = self._process_xlsx_basic(file_path, **kwargs)
                return c, m, i, None
            elif extension == 'pptx':
                c, m, i = self._process_pptx_basic(file_path, **kwargs)
                return c, m, i, None
            else:
                raise ProcessingError(f"Unsupported Office format: {extension}")

    def _process_docx_basic(self, file_path: Path, **kwargs) -> Tuple[str, dict, List[Dict[str, Any]]]:
        """Basic DOCX processing (fallback when advanced pipeline not available)."""
        try:
            doc = self.python_docx.Document(str(file_path))

            # Extract text content
            markdown_parts = []

            # Process paragraphs
            for para in doc.paragraphs:
                if para.text.strip():
                    # Check style for headings
                    if para.style.name.startswith('Heading'):
                        level = int(para.style.name[-1]) if para.style.name[-1].isdigit() else 1
                        markdown_parts.append(f"{'#' * level} {para.text}")
                    else:
                        # Handle text formatting
                        text = self._format_docx_paragraph(para)
                        markdown_parts.append(text)
                    markdown_parts.append("")

            # Process tables
            for table in doc.tables:
                table_md = self._convert_docx_table_to_markdown_basic(table)
                markdown_parts.append(table_md)
                markdown_parts.append("")

            # Extract images if requested
            images = []
            if kwargs.get('extract_images', False) and self.ocr:
                images = self._extract_docx_images_basic(file_path)

            # Metadata
            metadata = {
                'page_count': len(doc.element.xpath('//w:sectPr')),
                'word_count': sum(len(para.text.split()) for para in doc.paragraphs),
                'author': doc.core_properties.author,
                'title': doc.core_properties.title,
                'creation_date': str(doc.core_properties.created) if doc.core_properties.created else None,
                'modification_date': str(doc.core_properties.modified) if doc.core_properties.modified else None,
                'image_count': len(images),
            }

            return '\n'.join(markdown_parts), metadata, images

        except Exception as e:
            logger.error(f"Failed to process DOCX: {e}")
            raise ProcessingError(f"DOCX processing failed: {str(e)}")

    def _process_xlsx_basic(self, file_path: Path, **kwargs) -> Tuple[str, dict, List[Dict[str, Any]]]:
        """Basic XLSX processing (fallback when advanced pipeline not available)."""
        try:
            wb = self.openpyxl.load_workbook(str(file_path), data_only=True)

            markdown_parts = []
            total_cells = 0

            # Process each sheet
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                
                # Find actual max column with data (not Excel's 16384 limit)
                actual_max_col = 0
                actual_max_row = 0
                
                # Check first 100 rows to find actual data extent
                for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 100)), 1):
                    for col_idx, cell in enumerate(row, 1):
                        if cell.value is not None:
                            actual_max_col = max(actual_max_col, col_idx)
                            actual_max_row = max(actual_max_row, row_idx)
                
                # Check rest of rows if needed
                if sheet.max_row > 100:
                    for row in sheet.iter_rows(min_row=101, max_row=sheet.max_row, max_col=actual_max_col):
                        if any(cell.value is not None for cell in row):
                            actual_max_row = sheet.max_row
                            break

                if actual_max_row > 0 and actual_max_col > 0:
                    markdown_parts.append(f"## {sheet_name}")
                    markdown_parts.append("")

                    # Convert sheet to markdown table
                    table_data = []
                    for row in sheet.iter_rows(min_row=1, max_row=actual_max_row, 
                                               min_col=1, max_col=actual_max_col, 
                                               values_only=True):
                        # Filter out completely empty rows
                        if any(cell is not None for cell in row):
                            table_data.append([str(cell) if cell is not None else "" for cell in row])
                            total_cells += actual_max_col

                    if table_data:
                        table_md = self._convert_list_to_markdown_table(table_data)
                        markdown_parts.append(table_md)
                        markdown_parts.append("")

            # Extract images if requested
            images = []
            if kwargs.get('extract_images', False) and self.ocr:
                images = self._extract_xlsx_images_basic(file_path)
                
                # Add extracted image text to content if available
                if images:
                    image_texts = [img['text'] for img in images if img.get('text', '').strip()]
                    if image_texts:
                        markdown_parts.append("## Extracted Images")
                        markdown_parts.append("")
                        for i, text in enumerate(image_texts, 1):
                            markdown_parts.append(f"### Image {i}")
                            markdown_parts.append("```xml")
                            markdown_parts.append("<ocr_result>")
                            markdown_parts.append(text)
                            markdown_parts.append("</ocr_result>")
                            markdown_parts.append("```")
                            markdown_parts.append("")

            # Metadata
            metadata = {
                'page_count': len(wb.sheetnames),  # Using page_count to represent number of sheets
                'sheet_names': wb.sheetnames,
                'total_cells': total_cells,
                'image_count': len(images),
            }

            return '\n'.join(markdown_parts), metadata, images

        except Exception as e:
            logger.error(f"Failed to process XLSX: {e}")
            raise ProcessingError(f"XLSX processing failed: {str(e)}")

    def _process_pptx_basic(self, file_path: Path, **kwargs) -> Tuple[str, dict, List[Dict[str, Any]]]:
        """Basic PPTX processing (fallback when advanced pipeline not available)."""
        try:
            prs = self.python_pptx.Presentation(str(file_path))

            markdown_parts = []
            image_count = 0

            # Process each slide
            for i, slide in enumerate(prs.slides, 1):
                markdown_parts.append(f"## Slide {i}")
                markdown_parts.append("")

                # Extract text from shapes
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        # Handle title shapes
                        if shape == slide.shapes.title:
                            markdown_parts.append(f"### {shape.text}")
                        else:
                            markdown_parts.append(shape.text)
                        markdown_parts.append("")

                    # Count images
                    if shape.shape_type == 13:  # Picture
                        image_count += 1

                # Process tables
                for shape in slide.shapes:
                    if shape.has_table:
                        table_md = self._convert_pptx_table_to_markdown_basic(shape.table)
                        markdown_parts.append(table_md)
                        markdown_parts.append("")

            # Extract images if requested
            images = []
            if kwargs.get('extract_images', False) and self.ocr:
                images = self._extract_pptx_images_basic(file_path)
                
                # Add extracted image text to content if available
                if images:
                    image_texts = [img['text'] for img in images if img.get('text', '').strip()]
                    if image_texts:
                        markdown_parts.append("## Extracted Images")
                        markdown_parts.append("")
                        for i, text in enumerate(image_texts, 1):
                            markdown_parts.append(f"### Image {i}")
                            markdown_parts.append("```xml")
                            markdown_parts.append("<ocr_result>")
                            markdown_parts.append(text)
                            markdown_parts.append("</ocr_result>")
                            markdown_parts.append("```")
                            markdown_parts.append("")

            # Metadata
            metadata = {
                'page_count': len(prs.slides),
                'slide_count': len(prs.slides),
                'image_count': image_count,
                'title': prs.core_properties.title,
                'author': prs.core_properties.author,
            }

            return '\n'.join(markdown_parts), metadata, images

        except Exception as e:
            logger.error(f"Failed to process PPTX: {e}")
            raise ProcessingError(f"PPTX processing failed: {str(e)}")

    def _format_docx_paragraph(self, paragraph) -> str:
        """Format DOCX paragraph with inline formatting."""
        if not paragraph.runs:
            return paragraph.text

        formatted_text = []
        for run in paragraph.runs:
            text = run.text
            if run.bold:
                text = f"**{text}**"
            if run.italic:
                text = f"*{text}*"
            formatted_text.append(text)

        return ''.join(formatted_text)

    def _convert_docx_table_to_markdown_basic(self, table) -> str:
        """Convert DOCX table to markdown (basic version)."""
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(cells)

        return self._convert_list_to_markdown_table(rows)

    def _convert_pptx_table_to_markdown_basic(self, table) -> str:
        """Convert PPTX table to markdown (basic version)."""
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(cells)

        return self._convert_list_to_markdown_table(rows)

    def _convert_list_to_markdown_table(self, data: List[List[str]]) -> str:
        """Convert list of lists to markdown table."""
        if not data:
            return ""

        # Determine column widths
        col_widths = [0] * len(data[0])
        for row in data:
            for i, cell in enumerate(row):
                col_widths[i] = max(col_widths[i], len(cell))

        # Build table
        lines = []

        # Header
        header = "| " + " | ".join(cell.ljust(col_widths[i]) for i, cell in enumerate(data[0])) + " |"
        lines.append(header)

        # Separator
        separator = "|" + "|".join("-" * (w + 2) for w in col_widths) + "|"
        lines.append(separator)

        # Data rows
        for row in data[1:]:
            row_str = "| " + " | ".join(cell.ljust(col_widths[i]) for i, cell in enumerate(row)) + " |"
            lines.append(row_str)

        return '\n'.join(lines)

    def _extract_docx_images_basic(self, file_path: Path) -> List[Dict[str, Any]]:
        """Extract images from DOCX file using basic method."""
        images = []

        try:
            with zipfile.ZipFile(file_path, 'r') as zip_file:
                # Collect all images for batch processing
                image_batch = []
                image_info = []

                # Find image files in the media folder
                for file_info in zip_file.filelist:
                    if file_info.filename.startswith('word/media/'):
                        image_data = zip_file.read(file_info.filename)
                        image_batch.append(image_data)
                        image_info.append({
                            'filename': file_info.filename,
                            'size': len(image_data)
                        })

                # Batch process OCR if we have images and OCR is available
                if image_batch and self.ocr:
                    try:
                        # Prepare OCR kwargs with language configuration if available
                        ocr_kwargs = {}
                        if hasattr(self.ocr, 'config') and self.ocr.config and self.ocr.config.language:
                            ocr_kwargs['language'] = self.ocr.config.language
                            logger.info(
                                f"🌍 Passing language configuration to Office batch OCR: {self.ocr.config.language}")

                        logger.info(f"🚀 Processing {len(image_batch)} DOCX images with batch OCR")
                        ocr_results = self.ocr.batch_process_images(image_batch, **ocr_kwargs)

                        # Combine results
                        for info, ocr_result in zip(image_info, ocr_results):
                            text = ocr_result.text if hasattr(ocr_result, 'text') else str(ocr_result)
                            info['text'] = text
                            images.append(info)

                        logger.info(f"✅ DOCX batch OCR completed successfully")

                    except Exception as e:
                        logger.warning(f"Failed to batch OCR DOCX images: {e}")
                        # Add images without OCR text
                        for info in image_info:
                            info['text'] = ''
                            images.append(info)
                else:
                    # No OCR available, add images without text
                    for info in image_info:
                        info['text'] = ''
                        images.append(info)

        except Exception as e:
            logger.warning(f"Failed to extract images from DOCX: {e}")

        return images

    def _extract_xlsx_images_basic(self, file_path: Path) -> List[Dict[str, Any]]:
        """Extract images from XLSX file using basic method."""
        images = []

        try:
            with zipfile.ZipFile(file_path, 'r') as zip_file:
                # Collect all images for batch processing
                image_batch = []
                image_info = []

                # Find image files in the media folder
                for file_info in zip_file.filelist:
                    if file_info.filename.startswith('xl/media/'):
                        image_data = zip_file.read(file_info.filename)
                        image_batch.append(image_data)
                        image_info.append({
                            'filename': file_info.filename,
                            'size': len(image_data)
                        })

                # Batch process OCR if we have images and OCR is available
                if image_batch and self.ocr:
                    try:
                        # Prepare OCR kwargs with language configuration if available
                        ocr_kwargs = {}
                        if hasattr(self.ocr, 'config') and self.ocr.config and self.ocr.config.language:
                            ocr_kwargs['language'] = self.ocr.config.language
                            logger.info(
                                f"🌍 Passing language configuration to Office batch OCR: {self.ocr.config.language}")

                        logger.info(f"🚀 Processing {len(image_batch)} XLSX images with batch OCR")
                        ocr_results = self.ocr.batch_process_images(image_batch, **ocr_kwargs)

                        # Combine results
                        for info, ocr_result in zip(image_info, ocr_results):
                            text = ocr_result.text if hasattr(ocr_result, 'text') else str(ocr_result)
                            info['text'] = text
                            images.append(info)

                        logger.info(f"✅ XLSX batch OCR completed successfully")

                    except Exception as e:
                        logger.warning(f"Failed to batch OCR XLSX images: {e}")
                        # Add images without OCR text
                        for info in image_info:
                            info['text'] = ''
                            images.append(info)
                else:
                    # No OCR available, add images without text
                    for info in image_info:
                        info['text'] = ''
                        images.append(info)

        except Exception as e:
            logger.warning(f"Failed to extract images from XLSX: {e}")

        return images

    def _extract_pptx_images_basic(self, file_path: Path) -> List[Dict[str, Any]]:
        """Extract images from PPTX file using basic method."""
        images = []

        try:
            with zipfile.ZipFile(file_path, 'r') as zip_file:
                # Collect all images for batch processing
                image_batch = []
                image_info = []

                # Find image files in the media folder
                for file_info in zip_file.filelist:
                    if file_info.filename.startswith('ppt/media/'):
                        image_data = zip_file.read(file_info.filename)
                        image_batch.append(image_data)
                        image_info.append({
                            'filename': file_info.filename,
                            'size': len(image_data)
                        })

                # Batch process OCR if we have images and OCR is available
                if image_batch and self.ocr:
                    try:
                        # Prepare OCR kwargs with language configuration if available
                        ocr_kwargs = {}
                        if hasattr(self.ocr, 'config') and self.ocr.config and self.ocr.config.language:
                            ocr_kwargs['language'] = self.ocr.config.language
                            logger.info(
                                f"🌍 Passing language configuration to Office batch OCR: {self.ocr.config.language}")

                        logger.info(f"🚀 Processing {len(image_batch)} PPTX images with batch OCR")
                        ocr_results = self.ocr.batch_process_images(image_batch, **ocr_kwargs)

                        # Combine results
                        for info, ocr_result in zip(image_info, ocr_results):
                            text = ocr_result.text if hasattr(ocr_result, 'text') else str(ocr_result)
                            info['text'] = text
                            images.append(info)

                        logger.info(f"✅ PPTX batch OCR completed successfully")

                    except Exception as e:
                        logger.warning(f"Failed to batch OCR PPTX images: {e}")
                        # Add images without OCR text
                        for info in image_info:
                            info['text'] = ''
                            images.append(info)
                else:
                    # No OCR available, add images without text
                    for info in image_info:
                        info['text'] = ''
                        images.append(info)

        except Exception as e:
            logger.warning(f"Failed to extract images from PPTX: {e}")

        return images
