import base64
import bisect
import hashlib
import json
import logging
import re
import unicodedata
import zipfile
from contextlib import contextmanager
from pathlib import Path
from typing import Dict, List, Any, Union, Optional, Tuple


def _image_hash(data: bytes) -> str:
    """Compute a stable, collision-resistant hash for image deduplication."""
    return hashlib.sha256(data).hexdigest()


def _safe_lxml_parser():
    """Return an lxml parser hardened against XXE / entity-expansion attacks.

    ``resolve_entities=False`` is the load-bearing flag: it blocks file:// XXE
    and entity-expansion attacks when parsing XML extracted from untrusted
    .docx/.pptx archives.
    """
    from lxml import etree
    return etree.XMLParser(
        resolve_entities=False,
        no_network=True,
        load_dtd=False,
        dtd_validation=False,
        huge_tree=False,
    )


from doc2mark.core.table import TableStyle, TableRenderer, TableData
from doc2mark.ocr.schema import plain_ocr_text  # noqa: E402
from doc2mark.utils.markdown import (  # noqa: E402
    emphasis_fits, escape_inline_pieces, escape_line_start, escape_markdown_text, normalize_control_chars,
    wrap_inline,
)
from doc2mark.utils.number_format import format_cell_value  # noqa: E402

# What a picture whose OCR failed shows, as on the PDF path: it was not read (no API key, an outage, an
# answer the provider flags failed). The failure itself is in metadata.extra["ocr_issues"].
OCR_UNAVAILABLE = "[image: OCR unavailable]"

# Office document libraries
try:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.shape import InlineShape
    from docx.enum.shape import WD_INLINE_SHAPE
    from docx.oxml.ns import qn
except ImportError:
    raise ImportError("python-docx is required. Install with: pip install python-docx")

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    raise ImportError("python-pptx is required. Install with: pip install python-pptx")

try:
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
except ImportError:
    raise ImportError("openpyxl is required. Install with: pip install openpyxl")

# Import VisionAgent for OCR functionality (optional)
try:
    from doc2mark.ocr.openai import VisionAgent

    OCR_AVAILABLE = True
except ImportError:
    OCR_AVAILABLE = False
    logging.warning("OCR functionality not available. Install VisionAgent to enable OCR.")

logger = logging.getLogger(__name__)


# --------------------------------------------------------------------------- #
# WordprocessingML text walking                                                #
# --------------------------------------------------------------------------- #
# python-docx only reads runs that are direct children of a paragraph (or of a
# hyperlink) and paragraphs/tables that are direct children of the body or a cell.
# Text inside content controls, tracked insertions, simple fields, smart tags,
# custom XML and nested tables was silently dropped, so DOCX text is read from
# the XML tree instead.

_W_P, _W_R, _W_T = qn('w:p'), qn('w:r'), qn('w:t')
_W_TAB, _W_PTAB, _W_BR, _W_CR = qn('w:tab'), qn('w:ptab'), qn('w:br'), qn('w:cr')
_W_NO_BREAK_HYPHEN = qn('w:noBreakHyphen')
_W_TYPE, _W_VAL = qn('w:type'), qn('w:val')
_W_TBL, _W_TR, _W_TC = qn('w:tbl'), qn('w:tr'), qn('w:tc')
_W_TR_PR, _W_TC_PR = qn('w:trPr'), qn('w:tcPr')
_W_GRID_BEFORE, _W_GRID_AFTER, _W_GRID_SPAN = qn('w:gridBefore'), qn('w:gridAfter'), qn('w:gridSpan')
_W_V_MERGE, _W_H_MERGE = qn('w:vMerge'), qn('w:hMerge')
_W_TBL_GRID, _W_GRID_COL = qn('w:tblGrid'), qn('w:gridCol')
_W_SDT, _W_SDT_CONTENT, _W_CUSTOM_XML = qn('w:sdt'), qn('w:sdtContent'), qn('w:customXml')
_W_DEL, _W_MOVE_FROM, _W_DRAWING = qn('w:del'), qn('w:moveFrom'), qn('w:drawing')
_W_PPR, _W_P_STYLE, _W_NUM_PR = qn('w:pPr'), qn('w:pStyle'), qn('w:numPr')
_W_NUM_ID, _W_ILVL, _W_OUTLINE_LVL = qn('w:numId'), qn('w:ilvl'), qn('w:outlineLvl')
_M_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
_M_OMATH, _M_OMATH_PARA, _M_T = f'{{{_M_NS}}}oMath', f'{{{_M_NS}}}oMathPara', f'{{{_M_NS}}}t'
_MC_NS = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
_MC_ALTERNATE_CONTENT = f'{{{_MC_NS}}}AlternateContent'
_MC_CHOICE, _MC_FALLBACK = f'{{{_MC_NS}}}Choice', f'{{{_MC_NS}}}Fallback'
_W_RPR, _W_R_STYLE, _W_B, _W_I = qn('w:rPr'), qn('w:rStyle'), qn('w:b'), qn('w:i')
_W_HYPERLINK, _W_ANCHOR, _R_ID = qn('w:hyperlink'), qn('w:anchor'), qn('r:id')
_W_FLD_SIMPLE, _W_INSTR = qn('w:fldSimple'), qn('w:instr')
_W_FLD_CHAR, _W_FLD_CHAR_TYPE, _W_INSTR_TEXT = qn('w:fldChar'), qn('w:fldCharType'), qn('w:instrText')
_W_TXBX_CONTENT, _W_LAST_RENDERED_BREAK = qn('w:txbxContent'), qn('w:lastRenderedPageBreak')
_W_SECT_PR, _W_TITLE_PG, _W_EVEN_AND_ODD = qn('w:sectPr'), qn('w:titlePg'), qn('w:evenAndOddHeaders')
_W_HEADER_REF, _W_FOOTER_REF = qn('w:headerReference'), qn('w:footerReference')

# Paragraph-level elements whose content the reader does not see: properties,
# deleted revisions and the moved-from copy of moved text.
_DOCX_UNSEEN = frozenset({
    _W_PPR, _W_RPR, _W_DEL, _W_MOVE_FROM, qn('w:sdtPr'), qn('w:sdtEndPr'),
})


def _docx_paragraph_items(element, link=None):
    """Yield ``(run, link)`` for the runs (``w:r``) and math zones of a paragraph in reading
    order; ``link`` is the innermost ``w:hyperlink`` or ``w:fldSimple`` holding the run, or None.

    Descends into hyperlinks, content controls (``w:sdt``), tracked insertions
    (``w:ins``) and moves (``w:moveTo``), simple fields (``w:fldSimple``), smart
    tags and custom XML; skips deleted text (``w:del``) and ``w:moveFrom``.
    """
    for child in element:
        tag = child.tag
        if not isinstance(tag, str) or tag in _DOCX_UNSEEN:
            continue
        if tag == _W_R or tag in (_M_OMATH, _M_OMATH_PARA):
            yield child, link
        elif tag == _MC_ALTERNATE_CONTENT:
            branch = child.find(_MC_CHOICE)
            if branch is None:
                branch = child.find(_MC_FALLBACK)
            if branch is not None:
                yield from _docx_paragraph_items(branch, link)
        else:
            yield from _docx_paragraph_items(child, child if tag in (_W_HYPERLINK, _W_FLD_SIMPLE) else link)


def _docx_paragraph_content(element):
    """Yield the runs (``w:r``) and math zones of a paragraph in reading order (see
    ``_docx_paragraph_items``)."""
    for run, _ in _docx_paragraph_items(element):
        yield run


def _docx_text_boxes(element):
    """Yield the contents (``w:txbxContent``) of the text boxes and shapes with text drawn in
    ``element`` (a run), outermost first, in document order: DrawingML text boxes and shapes
    (``wps:txbx``, also inside groups) and the VML text boxes of older files (``v:textbox``). Of an
    ``mc:AlternateContent`` only the branch a reader shows is read (``mc:Choice``, else
    ``mc:Fallback``), so a box saved both ways, as Word saves it, counts once. Boxes nested in a
    box are read with that box's paragraphs."""
    for child in element:
        tag = child.tag
        if not isinstance(tag, str) or tag in (_W_DEL, _W_MOVE_FROM):
            continue
        if tag == _W_TXBX_CONTENT:
            yield child
        elif tag == _MC_ALTERNATE_CONTENT:
            branch = child.find(_MC_CHOICE)
            if branch is None:
                branch = child.find(_MC_FALLBACK)
            if branch is not None:
                yield from _docx_text_boxes(branch)
        else:
            yield from _docx_text_boxes(child)


def _docx_paragraph_text_boxes(p_el):
    """The text boxes anchored in a paragraph's runs (see ``_docx_text_boxes``)."""
    for run in _docx_paragraph_content(p_el):
        if run.tag == _W_R:
            yield from _docx_text_boxes(run)


def _docx_flag(element) -> bool:
    """An on/off property element (``w:titlePg``, ``w:evenAndOddHeaders``): on when present,
    unless its ``w:val`` says off."""
    return element is not None and element.get(_W_VAL, 'true').lower() not in ('0', 'false', 'off')


def _docx_toggle(props, tag) -> Optional[bool]:
    """A toggle property (``w:b``, ``w:i``) of run properties: True / False when set, None when not."""
    element = props.find(tag) if props is not None else None
    return None if element is None else _docx_flag(element)


_HYPERLINK_FIELD = re.compile(r'\s*HYPERLINK\b(.*)', re.IGNORECASE | re.DOTALL)
_FIELD_TOKEN = re.compile(r'"([^"]*)"|(\\\S+)|(\S+)')


def _hyperlink_field_target(instruction: str) -> Optional[str]:
    """The address a ``HYPERLINK`` field instruction points to (``HYPERLINK "https://…" \\o "tip"``),
    with its ``\\l`` bookmark as a fragment; None for a bookmark in the document alone
    (``HYPERLINK \\l "_Toc1"``) or another field."""
    match = _HYPERLINK_FIELD.match(instruction or '')
    if not match:
        return None
    address = anchor = None
    switch = None
    for token in _FIELD_TOKEN.finditer(match.group(1)):
        quoted, flag, bare = token.groups()
        if flag is not None:
            switch = flag.lower() if flag.lower() in ('\\l', '\\o', '\\t') else None
            continue
        value = quoted if quoted is not None else bare
        if switch == '\\l':
            anchor = value
        elif switch is None and address is None:
            address = value
        switch = None
    if not address:
        return None
    return f"{address}#{anchor}" if anchor else address


_LINK_SCHEMES = ('http://', 'https://', 'mailto:')
_LINK_UNSAFE = frozenset(' "<>()\\`')


def _link_destination(target: Optional[str]) -> Optional[str]:
    """``target`` written as a Markdown link destination, or None when it is not a web or mail
    address (``javascript:``, ``file:``, a relative path): such a link keeps only its text.
    Spaces, parentheses, angle brackets, quotes and backslashes are percent-encoded so the
    destination cannot end early or change the link."""
    target = (target or '').strip()
    if not target.lower().startswith(_LINK_SCHEMES):
        return None
    return ''.join(f'%{ord(char):02X}' if char in _LINK_UNSAFE or ord(char) < 0x20 or ord(char) == 0x7f
                   else char for char in target)


def _strip_pieces(pieces):
    """``pieces`` (``(text, ...)`` tuples) without the whitespace at the start and end of their
    joined text, as ``str.strip`` removes it."""
    text = ''.join(piece[0] for piece in pieces)
    start, end = len(text) - len(text.lstrip()), len(text.rstrip())
    stripped, offset = [], 0
    for piece in pieces:
        low, high = max(offset, start), min(offset + len(piece[0]), end)
        if low < high:
            stripped.append((text[low:high],) + tuple(piece[1:]))
        offset += len(piece[0])
    return stripped


# Characters that may stand right next to a link's brackets: brackets, quotes and sentence
# punctuation. Anything else glued to a link's text (``https://`` before a link on ``www.``...) would
# split the text around it for lexical retrieval.
_LINK_NEIGHBOURS = frozenset('([{<"\'\u00ab\u201c\u2018)]}>\u00bb\u201d\u2019.,;:!?')
_PAGE_NUMBER_FIELDS = frozenset({'PAGE', 'NUMPAGES', 'SECTIONPAGES'})


def _sets_off_link(char: str) -> bool:
    """Whether ``char`` (next to a link's text; "" at a line edge) leaves the link's brackets
    between words: the line edge, a space, a bracket, quote or sentence punctuation, or a CJK
    character or CJK punctuation (CJK text has no spaces between words)."""
    if not char or char.isspace() or char in _LINK_NEIGHBOURS:
        return True
    if ord(char) < 0x2E80 or (0xFF10 <= ord(char) <= 0xFF5A and char.isalnum()):
        return False  # a letter or digit of a word (fullwidth Latin letters and digits too)
    return char.isalnum() or unicodedata.category(char).startswith('P')


def _markup_edges(markdown: str) -> Tuple[str, str, str]:
    """``(leading spaces, core, trailing spaces)`` of a piece of inline Markdown."""
    core = markdown.strip()
    if not core:
        return markdown, '', ''
    start = len(markdown) - len(markdown.lstrip())
    return markdown[:start], core, markdown[start + len(core):]


def _markdown_line(pieces) -> str:
    """Inline Markdown for one line of styled runs (``(text, bold, italic, link target)``, no line
    breaks). The line is escaped at once (``escape_inline_pieces``, so a ``<`` in one run sees the
    letter in the next); then runs are written ``**bold**`` / ``*italic*`` and the runs of a web or
    mail link ``[text](address)``, only where that changes nothing a reader sees:

    - emphasis only where its markers keep words whole and CommonMark can close them
      (``emphasis_fits``), never on a line holding a literal ``*`` (it would pair with the markers),
      and never touching the markers of the run before it;
    - a link only when its text has no bracket, it is set off from the text around it and not
      preceded by ``!`` (which would make it an image);
    - nothing on a line holding a backtick (a code span could swallow the markup);
    - a backslash right before inserted markup is doubled, so it does not escape the markup.

    Otherwise the text is kept as it is; runs of one style and one written link are merged."""
    text = ''.join(piece[0] for piece in pieces)
    escaped = escape_inline_pieces(text)
    if '`' in text:
        return ''.join(escaped)

    def grouped(segments: List[List[Any]]) -> List[List[Any]]:
        merged: List[List[Any]] = []
        for segment in segments:
            if merged and merged[-1][3] == segment[3] and (merged[-1][2] == segment[2] or not segment[1].strip()):
                merged[-1][0] += segment[0]
                merged[-1][1] += segment[1]
            else:
                merged.append(list(segment))
        return merged

    segments: List[List[Any]] = []   # [markdown, plain text, (bold, italic), link target]
    offset = 0
    for piece_text, bold, italic, target, *_ in pieces:
        segments.append([''.join(escaped[offset:offset + len(piece_text)]), piece_text, (bold, italic), target])
        offset += len(piece_text)
    segments = grouped(segments)

    def neighbours(first: int, last: int) -> Tuple[str, str]:
        """The characters just outside segments ``first``..``last`` ("" at the line edge)."""
        head, tail = segments[first][1][:1], segments[last][1][-1:]
        before = head if head.isspace() else (segments[first - 1][1][-1:] if first else '')
        after = tail if tail.isspace() else (segments[last + 1][1][:1] if last + 1 < len(segments) else '')
        return before, after

    def spans() -> List[Tuple[int, int]]:
        found, first = [], 0
        while first < len(segments):
            last = first
            while last + 1 < len(segments) and segments[last + 1][3] == segments[first][3]:
                last += 1
            found.append((first, last))
            first = last + 1
        return found

    # A link that is not written keeps its text: its runs join the runs around them.
    for first, last in spans():
        if segments[first][3] is None:
            continue
        plain = ''.join(segment[1] for segment in segments[first:last + 1])
        before, after = neighbours(first, last)
        if not (_link_destination(segments[first][3]) and plain.strip() and '[' not in plain and ']' not in plain
                and before != '!' and _sets_off_link(before) and _sets_off_link(after)):
            for segment in segments[first:last + 1]:
                segment[3] = None
    segments = grouped(segments)

    tokens: List[Tuple[bool, str]] = []   # (inserted markup, text)
    emphasis = '*' not in text
    marked_before = False   # the previous run ends with an emphasis marker
    for first, last in spans():
        destination = _link_destination(segments[first][3]) if segments[first][3] is not None else None
        inner: List[Tuple[bool, str]] = []
        for position in range(first, last + 1):
            markdown, plain, (bold, italic), _ = segments[position]
            marker = '***' if bold and italic else '**' if bold else '*' if italic else ''
            lead, core, trail = _markup_edges(markdown)
            left, right = neighbours(position, position)
            if (emphasis and marker and any(char.isalnum() for char in plain)
                    and emphasis_fits(left, plain.strip(), right) and not (marked_before and not lead)):
                inner += [(False, lead), (True, marker), (False, core), (True, marker), (False, trail)]
                marked_before = not trail
            else:
                inner.append((False, markdown))
                marked_before = False
        if destination:
            # The spaces around the link's text stay outside its brackets.
            head, tail = inner[0][1], inner[-1][1]
            lead = head[:len(head) - len(head.lstrip())]
            trail = tail[len(tail.rstrip()):]
            if lead:
                inner[0] = (False, head[len(lead):])
            if trail:
                inner[-1] = (False, inner[-1][1][:len(inner[-1][1]) - len(trail)])
            inner = [(False, lead), (True, '[')] + inner + [(True, f']({destination})'), (False, trail)]
            marked_before = False
        tokens += inner

    out: List[str] = []
    backslashes = 0   # backslashes at the end of the output so far
    for markup, chunk in tokens:
        if not chunk:
            continue
        if markup:
            if backslashes % 2:
                out.append('\\')   # a backslash right before inserted markup would escape it
            out.append(chunk)
            backslashes = 0
        else:
            out.append(chunk)
            kept = chunk.rstrip('\\')
            backslashes = backslashes + len(chunk) if not kept else len(chunk) - len(kept)
    return ''.join(out)


def _docx_inline_markdown(pieces) -> str:
    """Markdown for a paragraph's text from its styled runs (see ``_markdown_line``): one line per
    line break, block syntax at a line start escaped as for body text."""
    lines: List[List[Tuple[str, bool, bool, Optional[str]]]] = [[]]
    for text, bold, italic, target, *_ in pieces:
        for number, part in enumerate(normalize_control_chars(text).split('\n')):
            if number:
                lines.append([])
            if part:
                lines[-1].append((part, bold, italic, target))
    return '\n'.join(escape_line_start(_markdown_line(line)) for line in lines)


def _docx_run_text(run) -> str:
    """Text of a run as python-docx reads it (``w:t``, tabs, line breaks, non-breaking
    hyphens; field instructions and deleted text excluded), or of a math zone."""
    if run.tag != _W_R:
        return ''.join(t.text or '' for t in run.iter(_M_T))
    parts = []
    for child in run:
        tag = child.tag
        if tag == _W_T:
            parts.append(child.text or '')
        elif tag in (_W_TAB, _W_PTAB):
            parts.append('\t')
        elif tag == _W_BR:
            parts.append('\n' if child.get(_W_TYPE, 'textWrapping') == 'textWrapping' else '')
        elif tag == _W_CR:
            parts.append('\n')
        elif tag == _W_NO_BREAK_HYPHEN:
            parts.append('-')
    return ''.join(parts)


def _docx_blocks(container):
    """Yield the paragraphs and tables of a body, cell or content control in document
    order, looking inside block-level content controls and custom XML."""
    for child in container:
        tag = child.tag
        if tag == _W_P or tag == _W_TBL:
            yield child
        elif tag == _W_SDT:
            content = child.find(_W_SDT_CONTENT)
            if content is not None:
                yield from _docx_blocks(content)
        elif tag == _W_CUSTOM_XML:
            yield from _docx_blocks(child)


def _docx_children(element, tag):
    """``tag`` children (rows of a table, cells of a row), including those wrapped in
    content controls or custom XML."""
    for child in element:
        if child.tag == tag:
            yield child
        elif child.tag == _W_SDT:
            content = child.find(_W_SDT_CONTENT)
            if content is not None:
                yield from _docx_children(content, tag)
        elif child.tag == _W_CUSTOM_XML:
            yield from _docx_children(child, tag)


def _docx_int(element, default: int = 0) -> int:
    """``w:val`` of ``element`` as an int; ``default`` when absent or malformed."""
    if element is None:
        return default
    try:
        return int(element.get(_W_VAL, default))
    except (TypeError, ValueError):
        return default


def _docx_rendered(root):
    """Every element under ``root`` in document order, minus deleted revisions,
    moved-from copies and ``mc:Fallback`` duplicates."""
    unseen = (_W_DEL, _W_MOVE_FROM, _MC_FALLBACK)
    stack = [iter(root)]
    while stack:
        for element in stack[-1]:
            if not isinstance(element.tag, str) or element.tag in unseen:
                continue
            yield element
            stack.append(iter(element))
            break
        else:
            stack.pop()


def _attr_int(element, name: str, default: int = 0) -> int:
    try:
        return int(element.get(name, default))
    except (TypeError, ValueError):
        return default


def _roman(number: int) -> str:
    numerals = ((1000, 'M'), (900, 'CM'), (500, 'D'), (400, 'CD'), (100, 'C'), (90, 'XC'),
                (50, 'L'), (40, 'XL'), (10, 'X'), (9, 'IX'), (5, 'V'), (4, 'IV'), (1, 'I'))
    out = []
    for value, symbol in numerals:
        count, number = divmod(number, value)
        out.append(symbol * count)
    return ''.join(out)


_CJK_DIGITS = '〇一二三四五六七八九'
_CJK_COUNTING = ('零一二三四五六七八九', '千百十', True)
_CJK_LEGAL_TRADITIONAL = ('零壹貳參肆伍陸柒捌玖', '仟佰拾', False)
_CJK_LEGAL_SIMPLIFIED = ('零壹贰叁肆伍陆柒捌玖', '仟佰拾', False)


def _cjk_counting(number: int, numerals=_CJK_COUNTING) -> str:
    """一, 十, 十一, 二十, 一百零一 (or 壹, 壹拾, 壹拾壹 ... in the legal forms) for
    numbers below 10,000; larger ones stay Arabic."""
    if not 0 < number < 10000:
        return str(number)
    digits, units, bare_ten = numerals
    out, zero_pending = [], False
    for unit_value, unit in ((1000, units[0]), (100, units[1]), (10, units[2]), (1, '')):
        digit, number = divmod(number, unit_value)
        if digit:
            if zero_pending:
                out.append(digits[0])
                zero_pending = False
            leading_ten = bare_ten and unit_value == 10 and digit == 1 and not out
            out.append(('' if leading_ten else digits[digit]) + unit)
        elif out:
            zero_pending = True
    return ''.join(out)


def _format_list_number(number: int, fmt: str) -> str:
    """One list counter in a Word ``w:numFmt``."""
    if fmt == 'decimalZero':
        return f'{number:02d}'
    if fmt in ('upperRoman', 'lowerRoman'):
        text = _roman(number) if number > 0 else str(number)
        return text if fmt == 'upperRoman' else text.lower()
    if fmt in ('upperLetter', 'lowerLetter'):
        if number <= 0:
            return str(number)
        letter = chr(ord('A') + (number - 1) % 26) * ((number - 1) // 26 + 1)
        return letter if fmt == 'upperLetter' else letter.lower()
    if fmt in ('decimalFullWidth', 'decimalFullWidth2'):
        return ''.join(chr(ord(ch) + 0xFEE0) for ch in str(number))
    if fmt.startswith('decimalEnclosedCircle') and 0 < number <= 20:
        return chr(0x2460 + number - 1)
    if fmt in ('chineseCounting', 'chineseCountingThousand', 'taiwaneseCounting',
               'taiwaneseCountingThousand', 'japaneseCounting'):
        return _cjk_counting(number)
    if fmt == 'ideographLegalTraditional':
        return _cjk_counting(number, _CJK_LEGAL_TRADITIONAL)
    if fmt == 'chineseLegalSimplified':
        return _cjk_counting(number, _CJK_LEGAL_SIMPLIFIED)
    if fmt in ('ideographDigital', 'taiwaneseDigital', 'japaneseDigitalTenThousand'):
        return ''.join(_CJK_DIGITS[int(ch)] for ch in str(number))
    if fmt == 'ideographTraditional' and 0 < number <= 10:
        return '甲乙丙丁戊己庚辛壬癸'[number - 1]
    if fmt == 'ideographZodiac' and 0 < number <= 12:
        return '子丑寅卯辰巳午未申酉戌亥'[number - 1]
    return str(number)


class _DocxStructure:
    """Paragraph structure Word shows but python-docx does not expose: list numbers
    and bullets (``word/numbering.xml``) and outline (heading) levels.

    ``list_marker`` must be called once per paragraph in document order: counters are
    kept per abstract numbering definition, deeper levels restart when a shallower
    level advances, and a ``w:startOverride`` restarts a level when its list instance
    is first used.
    """

    _MAX_STYLE_DEPTH = 20

    def __init__(self, document):
        self._styles: Dict[str, Any] = {}
        self._default_style: Optional[str] = None
        self._nums: Dict[int, Tuple[Optional[int], Dict[int, int]]] = {}
        self._levels: Dict[int, Dict[int, Tuple[str, Optional[str], int, bool]]] = {}
        self._level_styles: Dict[int, Dict[str, int]] = {}
        self._style_links: Dict[int, str] = {}
        self._counters: Dict[int, Dict[int, int]] = {}
        self._started: set = set()
        try:
            for style in document.styles.element.findall(qn('w:style')):
                style_id = style.get(qn('w:styleId'))
                self._styles[style_id] = style
                if style.get(qn('w:type')) == 'paragraph' and style.get(qn('w:default')) in ('1', 'true', 'on'):
                    self._default_style = style_id
        except Exception as e:
            logger.debug(f"DOCX styles unavailable: {e}")
        try:
            numbering = document.part.numbering_part.element
        except Exception:
            numbering = None  # the package has no numbering part
        if numbering is None:
            return
        for abstract in numbering.findall(qn('w:abstractNum')):
            abstract_id = _attr_int(abstract, qn('w:abstractNumId'))
            levels, level_styles = {}, {}
            for lvl in abstract.findall(qn('w:lvl')):
                ilvl = _attr_int(lvl, _W_ILVL)
                fmt = lvl.find(qn('w:numFmt'))
                text = lvl.find(qn('w:lvlText'))
                picture_bullet = lvl.find(qn('w:lvlPicBulletId')) is not None
                levels[ilvl] = (
                    'bullet' if picture_bullet else (fmt.get(_W_VAL, 'decimal') if fmt is not None else 'decimal'),
                    text.get(_W_VAL) if text is not None else None,
                    _docx_int(lvl.find(qn('w:start')), 1) if lvl.find(qn('w:start')) is not None else 1,
                    lvl.find(qn('w:isLgl')) is not None,
                )
                p_style = lvl.find(_W_P_STYLE)
                if p_style is not None:
                    level_styles[p_style.get(_W_VAL)] = ilvl
            self._levels[abstract_id] = levels
            self._level_styles[abstract_id] = level_styles
            link = abstract.find(qn('w:numStyleLink'))
            if link is not None:
                self._style_links[abstract_id] = link.get(_W_VAL)
        for num in numbering.findall(qn('w:num')):
            abstract_ref = num.find(qn('w:abstractNumId'))
            overrides = {}
            for override in num.findall(qn('w:lvlOverride')):
                start = override.find(qn('w:startOverride'))
                if start is not None:
                    overrides[_attr_int(override, _W_ILVL)] = _docx_int(start, 1)
            self._nums[_attr_int(num, qn('w:numId'))] = (
                _docx_int(abstract_ref, -1) if abstract_ref is not None else None, overrides)

    def _style_chain(self, p_el):
        ppr = p_el.find(_W_PPR)
        style_ref = ppr.find(_W_P_STYLE) if ppr is not None else None
        style_id = style_ref.get(_W_VAL) if style_ref is not None else self._default_style
        seen = set()
        while style_id and style_id not in seen and len(seen) < self._MAX_STYLE_DEPTH:
            seen.add(style_id)
            style = self._styles.get(style_id)
            if style is None:
                return
            yield style_id, style
            based_on = style.find(qn('w:basedOn'))
            style_id = based_on.get(_W_VAL) if based_on is not None else None

    @contextmanager
    def separate_numbering(self):
        """Count list numbers apart while a header, footer or text box (a story of its own) is read, so
        its lists neither take nor move the body's numbers; the body's counters are restored after."""
        saved = (self._counters, self._started)
        self._counters, self._started = {}, set()
        try:
            yield
        finally:
            self._counters, self._started = saved

    def run_emphasis(self, run) -> Tuple[bool, bool]:
        """(bold, italic) of a run: its own ``w:b`` / ``w:i``, else those of its character style
        (``w:rStyle`` and the styles it is based on, e.g. Strong, Emphasis). Paragraph styles are not
        consulted: a heading or a quote styled bold as a whole is its style, not emphasis in its text."""
        props = run.find(_W_RPR)
        bold, italic = _docx_toggle(props, _W_B), _docx_toggle(props, _W_I)
        style_ref = props.find(_W_R_STYLE) if props is not None else None
        style_id = style_ref.get(_W_VAL) if style_ref is not None else None
        seen = set()
        while (bold is None or italic is None) and style_id and style_id not in seen \
                and len(seen) < self._MAX_STYLE_DEPTH:
            seen.add(style_id)
            style = self._styles.get(style_id)
            if style is None:
                break
            style_props = style.find(_W_RPR)
            bold = _docx_toggle(style_props, _W_B) if bold is None else bold
            italic = _docx_toggle(style_props, _W_I) if italic is None else italic
            based_on = style.find(qn('w:basedOn'))
            style_id = based_on.get(_W_VAL) if based_on is not None else None
        return bool(bold), bool(italic)

    def outline_level(self, p_el) -> Optional[int]:
        """The outline level set on the paragraph or, failing that, the nearest style of its
        chain: 0-8 for heading levels 1-9, 9 for body text; None when none is set."""
        ppr = p_el.find(_W_PPR)
        if ppr is not None and ppr.find(_W_OUTLINE_LVL) is not None:
            return min(max(_docx_int(ppr.find(_W_OUTLINE_LVL), 9), 0), 9)
        for _, style in self._style_chain(p_el):
            style_ppr = style.find(_W_PPR)
            if style_ppr is not None and style_ppr.find(_W_OUTLINE_LVL) is not None:
                return min(max(_docx_int(style_ppr.find(_W_OUTLINE_LVL), 9), 0), 9)
        return None

    def _num_pr(self, p_el) -> Tuple[Optional[int], Optional[int], Optional[str]]:
        """(numId, ilvl, id of the style that supplied the numbering)."""
        num_id = ilvl = None
        source_style = None
        ppr = p_el.find(_W_PPR)
        num_pr = ppr.find(_W_NUM_PR) if ppr is not None else None
        if num_pr is not None:
            if num_pr.find(_W_NUM_ID) is not None:
                num_id = _docx_int(num_pr.find(_W_NUM_ID), 0)
            if num_pr.find(_W_ILVL) is not None:
                ilvl = _docx_int(num_pr.find(_W_ILVL), 0)
        for style_id, style in self._style_chain(p_el):
            if num_id is not None and ilvl is not None:
                break
            style_ppr = style.find(_W_PPR)
            style_num = style_ppr.find(_W_NUM_PR) if style_ppr is not None else None
            if style_num is None:
                continue
            if num_id is None and style_num.find(_W_NUM_ID) is not None:
                num_id = _docx_int(style_num.find(_W_NUM_ID), 0)
                source_style = style_id
            if ilvl is None and style_num.find(_W_ILVL) is not None:
                ilvl = _docx_int(style_num.find(_W_ILVL), 0)
        return num_id, ilvl, source_style

    def _abstract_for(self, num_id: int) -> Optional[int]:
        abstract_id = self._nums.get(num_id, (None, {}))[0]
        # A list that only links to a numbering style takes that style's definition.
        link = self._style_links.get(abstract_id)
        if link and not self._levels.get(abstract_id):
            style = self._styles.get(link)
            style_ppr = style.find(_W_PPR) if style is not None else None
            style_num = style_ppr.find(_W_NUM_PR) if style_ppr is not None else None
            linked = _docx_int(style_num.find(_W_NUM_ID), 0) if style_num is not None else 0
            if linked and linked != num_id:
                abstract_id = self._nums.get(linked, (None, {}))[0]
        return abstract_id

    def list_marker(self, p_el) -> Optional[Tuple[str, int]]:
        """(marker, level) for a list paragraph: ``"1."``, ``"a)"``, ``"1.2."`` from
        the level's text and number format, ``"-"`` for bullets; None otherwise."""
        num_id, ilvl, source_style = self._num_pr(p_el)
        if not num_id or num_id not in self._nums:
            return None
        abstract_id = self._abstract_for(num_id)
        levels = self._levels.get(abstract_id)
        if not levels:
            return None
        if ilvl is None:
            ilvl = self._level_styles.get(abstract_id, {}).get(source_style, 0)
        ilvl = max(0, min(ilvl, 8))
        if ilvl not in levels:
            return None
        counters = self._counters.setdefault(abstract_id, {})
        if num_id not in self._started:
            self._started.add(num_id)
            for level, start in self._nums[num_id][1].items():
                counters[level] = start - 1
        fmt, text, start, _ = levels[ilvl]
        counters[ilvl] = counters[ilvl] + 1 if ilvl in counters else start
        for deeper in [level for level in counters if level > ilvl]:
            del counters[deeper]
        if fmt == 'bullet':
            return '-', ilvl
        if fmt == 'none' or not text:  # no number shown; any literal level text still is
            literal = re.sub(r'%[1-9]', '', text or '').strip()
            return (literal, ilvl) if literal else None

        def counter(match):
            level = int(match.group(1)) - 1
            level_fmt, _, level_start, _ = levels.get(level, ('decimal', None, 1, False))
            value = counters.get(level, level_start)
            return _format_list_number(value, 'decimal' if levels[ilvl][3] else level_fmt)

        marker = re.sub(r'%([1-9])', counter, text).strip()
        return (marker, ilvl) if marker else None


# A caption by its shape: a caption word, then a number (3, 2.1, 4-2, A1, 3a, IV, A) and a separator
# (":", ".", ")", a dash) or the end of the text.
_CAPTION_TEXT = re.compile(
    r'(?i:figure|fig\.?|table|tbl\.?|chart|graph|image|plate|scheme|exhibit)\s+'
    r'(?:[A-Z]?\d+(?:[.\-–]\d+)*[a-z]?|[IVXLC]{1,7}|[A-Z])'
    r'(?:\s*[:.)\-–—]|\s*$)')


def _style_list_depth(style_name: str) -> int:
    """Nesting depth implied by Word's numbered list styles ("List Bullet 2" -> 1)."""
    match = re.fullmatch(r'list(?: (?:bullet|number|continue))? ([2-9])', (style_name or '').strip().lower())
    return int(match.group(1)) - 1 if match else 0


def _soft_breaks(text: str) -> str:
    """python-pptx returns a soft line break (``a:br``) as a vertical tab; make it a newline."""
    return text.replace('\x0b', '\n') if text else text


# --------------------------------------------------------------------------- #
# SpreadsheetML package reading (what openpyxl does not expose)                #
# --------------------------------------------------------------------------- #
_CONTROL_CHARS = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')
_SML_NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
_SML_C, _SML_F, _SML_V = f'{{{_SML_NS}}}c', f'{{{_SML_NS}}}f', f'{{{_SML_NS}}}v'
_PKG_REL_NS = 'http://schemas.openxmlformats.org/package/2006/relationships'
_OFFICE_REL_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
_RD_NS = 'http://schemas.microsoft.com/office/spreadsheetml/2017/richdata'
_RVREL_NS = 'http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel'


def _formula_uncached(cell) -> bool:
    """A formula cell saved without a cached result: no ``<v>``, or an empty one on a
    cell that is not a string result (an empty string is a real cached result)."""
    value = cell.find(_SML_V)
    if value is None:
        return True
    return not (value.text or '').strip() and cell.get('t') != 'str'


# A formula element (with or without a namespace prefix: openpyxl writes <f>, the Open XML
# SDK <x:f>) or a value-metadata attribute (vm="1" or vm='1'; XML allows either quote).
_FORMULA_OR_RICH_VALUE = re.compile(rb'<(?:[A-Za-z_][\w.-]*:)?f[\s/>]|\svm\s*=\s*["\']')
_ZIP_SCAN_OVERLAP = 256  # longer than any tag prefix the pattern has to see whole


def _zip_member_matches(archive, name: str, pattern) -> bool:
    """Whether a zip member matches ``pattern`` anywhere (streamed in 1 MiB chunks)."""
    tail = b''
    with archive.open(name) as stream:
        while True:
            chunk = stream.read(1 << 20)
            if not chunk:
                return False
            window = tail + chunk
            if pattern.search(window):
                return True
            tail = window[-_ZIP_SCAN_OVERLAP:]


def _part_path(base_dir: str, target: str) -> str:
    """Resolve a relationship target against the directory of its source part."""
    if target.startswith('/'):
        return target.lstrip('/')
    parts = [piece for piece in base_dir.split('/') if piece]
    for piece in target.split('/'):
        if piece == '..':
            if parts:
                parts.pop()
        elif piece and piece != '.':
            parts.append(piece)
    return '/'.join(parts)


def _read_xml(archive, name: str):
    from lxml import etree
    return etree.fromstring(archive.read(name), parser=_safe_lxml_parser())


def _relationships(archive, source_part: str) -> Dict[str, Tuple[str, str]]:
    """{Id: (type, part name)} for the internal relationships of ``source_part``
    (``""`` for the package itself)."""
    base, _, filename = source_part.rpartition('/')
    rels_name = f"{base + '/' if base else ''}_rels/{filename}.rels"
    try:
        root = _read_xml(archive, rels_name)
    except KeyError:
        return {}
    return {rel.get('Id'): (rel.get('Type', ''), _part_path(base, rel.get('Target', '')))
            for rel in root.findall(f'{{{_PKG_REL_NS}}}Relationship')
            if rel.get('TargetMode') != 'External'}


def _xlsx_sheet_parts(archive) -> Dict[str, str]:
    """{sheet title: worksheet part name}."""
    workbook = next((target for kind, target in _relationships(archive, '').values()
                     if kind.endswith('/officeDocument')), 'xl/workbook.xml')
    rels = _relationships(archive, workbook)
    parts = {}
    for sheet in _read_xml(archive, workbook).iter(f'{{{_SML_NS}}}sheet'):
        rel = rels.get(sheet.get(f'{{{_OFFICE_REL_NS}}}id'))
        if rel:
            parts[sheet.get('name')] = rel[1]
    return parts


def _xlsx_rich_value_images(archive, names) -> Dict[int, str]:
    """{cell ``vm`` index (1-based): media part} for pictures placed in cells.

    Excel stores such a cell as the error ``#VALUE!`` with a value-metadata index; the
    picture is found through ``metadata.xml`` (valueMetadata -> XLRICHVALUE future
    metadata) -> ``rdrichvalue.xml`` (a rich value whose structure has
    ``_rvRel:LocalImageIdentifier``) -> ``richValueRel.xml`` -> its relationship target.
    """
    def member(suffix: str) -> Optional[str]:
        return next((name for name in sorted(names) if name.lower().endswith(suffix)), None)

    metadata_name = 'xl/metadata.xml' if 'xl/metadata.xml' in names else member('/metadata.xml')
    values_name, structures_name = member('/rdrichvalue.xml'), member('/rdrichvaluestructure.xml')
    rel_name = member('/richvaluerel.xml')
    if not (metadata_name and values_name and structures_name and rel_name):
        return {}
    metadata = _read_xml(archive, metadata_name)
    types = [kind.get('name') for kind in metadata.iter(f'{{{_SML_NS}}}metadataType')]
    future = next((f for f in metadata.iter(f'{{{_SML_NS}}}futureMetadata') if f.get('name') == 'XLRICHVALUE'), None)
    value_metadata = next(metadata.iter(f'{{{_SML_NS}}}valueMetadata'), None)
    if future is None or value_metadata is None:
        return {}
    rich_of_future = []
    for block in future.findall(f'{{{_SML_NS}}}bk'):
        rvb = next(block.iter(f'{{{_RD_NS}}}rvb'), None)
        rich_of_future.append(_attr_int(rvb, 'i', -1) if rvb is not None else -1)
    structures = [[key.get('n') for key in structure.findall(f'{{{_RD_NS}}}k')]
                  for structure in _read_xml(archive, structures_name).findall(f'{{{_RD_NS}}}s')]
    rel_of_rich = {}
    for index, value in enumerate(_read_xml(archive, values_name).findall(f'{{{_RD_NS}}}rv')):
        structure = _attr_int(value, 's', -1)
        keys = structures[structure] if 0 <= structure < len(structures) else []
        entries = value.findall(f'{{{_RD_NS}}}v')
        if '_rvRel:LocalImageIdentifier' in keys:
            entry = keys.index('_rvRel:LocalImageIdentifier')
            if entry < len(entries) and (entries[entry].text or '').strip().isdigit():
                rel_of_rich[index] = int(entries[entry].text)
    rel_ids = [rel.get(f'{{{_OFFICE_REL_NS}}}id') for rel in _read_xml(archive, rel_name).iter(f'{{{_RVREL_NS}}}rel')]
    targets = _relationships(archive, rel_name)
    images = {}
    for vm, block in enumerate(value_metadata.findall(f'{{{_SML_NS}}}bk'), start=1):
        record = block.find(f'{{{_SML_NS}}}rc')
        if record is None:
            continue
        kind, future_index = _attr_int(record, 't') - 1, _attr_int(record, 'v', -1)
        if not (0 <= kind < len(types) and types[kind] == 'XLRICHVALUE' and 0 <= future_index < len(rich_of_future)):
            continue
        rel_index = rel_of_rich.get(rich_of_future[future_index])
        if rel_index is not None and rel_index < len(rel_ids) and rel_ids[rel_index] in targets:
            images[vm] = targets[rel_ids[rel_index]][1]
    return images


def _drawing_blip_rid(drawing) -> Optional[str]:
    """``r:embed`` of the picture a ``wp:inline`` / ``wp:anchor`` shows (a ``pic:pic`` directly in
    its graphic data), or None."""
    for graphic in drawing:
        if not graphic.tag.endswith('}graphic'):
            continue
        for data in graphic:
            if not data.tag.endswith('}graphicData'):
                continue
            for pic in data:
                if not pic.tag.endswith('}pic'):
                    continue
                for fill in pic:
                    if not fill.tag.endswith('}blipFill'):
                        continue
                    for blip in fill:
                        if blip.tag.endswith('}blip'):
                            for name, value in blip.attrib.items():
                                if name.endswith('}embed'):
                                    return value
    return None


def _issue_location(location: Optional[Dict[str, Any]]) -> Dict[str, Any]:
    """Where a picture sits, for the loader's OCR issue record: its slide, sheet or (Word) page."""
    location = location or {}
    if 'slide' in location:
        return {"slide": location['slide']}
    if 'sheet' in location:
        return {"sheet": location.get('sheet_name') or location['sheet']}
    if 'page' in location:
        return {"page": location['page']}
    return {}


_P_NS = 'http://schemas.openxmlformats.org/presentationml/2006/main'
_A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'


def _pptx_hidden(shape) -> bool:
    """A shape hidden on its slide (``p:cNvPr hidden="1"``)."""
    c_nv_pr = shape._element.find(f'./*/{{{_P_NS}}}cNvPr')
    return c_nv_pr is not None and c_nv_pr.get('hidden') in ('1', 'true')


def _pptx_drawn_masters(slide) -> list:
    """The slide layout and slide master whose shapes are drawn on ``slide`` (``showMasterSp``)."""
    if slide._element.get('showMasterSp') in ('0', 'false'):
        return []
    layout = slide.slide_layout
    owners = [layout]
    if layout._element.get('showMasterSp') not in ('0', 'false'):
        owners.append(layout.slide_master)
    return owners


def _pptx_visible_shapes(shapes):
    """Every visible shape of a shape tree, groups opened (a hidden group hides its shapes)."""
    for shape in shapes:
        if _pptx_hidden(shape):
            continue
        yield shape
        try:
            is_group = shape.shape_type == MSO_SHAPE_TYPE.GROUP
        except (AttributeError, NotImplementedError):
            is_group = False
        if is_group:
            yield from _pptx_visible_shapes(shape.shapes)


def _pptx_typed_text(paragraph) -> str:
    """A paragraph's text without its fields (``a:fld``: the slide number, the date), whose cached
    text is not text the slide sets."""
    parts = []
    for child in paragraph._p:
        if child.tag == f'{{{_A_NS}}}r':
            parts.append(''.join(t.text or '' for t in child.iter(f'{{{_A_NS}}}t')))
        elif child.tag == f'{{{_A_NS}}}br':
            parts.append('\n')
    return ''.join(parts)


class BaseOfficeLoader:
    """Base class for Office document loaders"""

    def __init__(self, file_path: Union[str, Path], ocr=None, table_style: Union[str, TableStyle] = None):
        self.file_path = Path(file_path)
        if not self.file_path.exists():
            raise FileNotFoundError(f"File not found: {self.file_path}")
        self.doc = None
        self.ocr = ocr  # Store the OCR instance
        # Each picture's OCR text for a table cell label (by image hash): plain text, which the
        # table renderer escapes once (the OCR result's own text is escaped Markdown).
        self._ocr_cell_texts: Dict[str, str] = {}
        
        # Set table output style
        if table_style is None:
            self.table_style = TableStyle.default()
        elif isinstance(table_style, str):
            self.table_style = TableStyle(table_style)
        else:
            self.table_style = table_style

        # Log OCR configuration if available
        if self.ocr:
            logger.info(f"📷 OCR configured for {self.__class__.__name__}: {type(self.ocr).__name__}")
            if hasattr(self.ocr, 'config') and self.ocr.config:
                if hasattr(self.ocr.config, 'language') and self.ocr.config.language:
                    logger.info(f"🌍 OCR Language setting: {self.ocr.config.language}")
        else:
            logger.warning(f"⚠️  No OCR instance provided to {self.__class__.__name__}")

    def convert_to_json(self,
                        extract_images: bool = True,
                        ocr_images: bool = False,
                        show_progress: bool = True) -> Dict[str, Any]:
        """Convert document to JSON format"""
        raise NotImplementedError("Subclasses must implement this method")

    def _convert_table_to_markdown(self, table_data: Union[List[List[str]], Any],
                                   extract_images: bool = True, ocr_images: bool = False,
                                   ocr_results_map: Optional[Dict[str, str]] = None) -> str:
        """Render a table given as rows of cell values.

        Merged cells are never guessed from blank cells: a blank cell is a blank cell.
        Loaders that know real merges (``w:gridSpan``/``w:vMerge`` in DOCX,
        ``gridSpan``/``vMerge`` in PPTX, ``merged_cells.ranges`` in XLSX) build their
        spans from the file and render them themselves.
        """
        if not table_data or not any(table_data):
            return ""
        col_count = max(len(row) for row in table_data)
        rows = [["" if cell is None else str(cell) for cell in row] + [""] * (col_count - len(row))
                for row in table_data]
        return TableRenderer(self.table_style).render(TableData.from_raw(rows, {'is_complex': False}))

    def _extract_image_as_base64(self, image_data: bytes, image_format: str = 'png') -> str:
        """Convert image bytes to base64 string"""
        return base64.b64encode(image_data).decode('utf-8')

    @staticmethod
    def _ocr_description(ocr_text: Optional[str], **fields) -> Optional[Dict[str, Any]]:
        """The ``text:image_description`` item for a picture's OCR text (inside the internal
        ``<image_ocr_result>`` provenance wrapper that the Markdown render strips), or None
        when OCR found no text in the picture: an empty description is never emitted, as on
        the PDF path, which skips images that OCR to nothing. ``ocr_text`` None means the OCR
        failed: the picture gets the PDF path's marker, ``[image: OCR unavailable]``."""
        text = OCR_UNAVAILABLE if ocr_text is None else ocr_text.strip()
        if not text:
            return None
        return {"type": "text:image_description", "content": f"<image_ocr_result>{text}</image_ocr_result>", **fields}

    @staticmethod
    def _ocr_failed(result) -> bool:
        """No answer, or an answer the provider flags ``failed`` (a timeout, an error): the picture
        was not read, which is not the same as a picture without text."""
        return result is None or bool((getattr(result, 'metadata', None) or {}).get('failed'))

    def _label_ocr_issues(self, locations: List[Optional[Dict[str, Any]]]) -> None:
        """Tell the loader's OCR issue record where the pictures of the last OCR call sit."""
        label_issues = getattr(self.ocr, "label_last_batch", None)
        if callable(label_issues) and any(locations):
            label_issues([_issue_location(location) for location in locations])

    def _ocr_image(self, image_bytes: bytes, location: Optional[Dict[str, Any]] = None) -> Optional[str]:
        """OCR one picture on its own (one the document's batch did not cover): its text ("" when
        OCR found none), or None when it could not be read (no image data or OCR, a request that
        raised, an answer the provider flags ``failed``). ``location`` is where the picture sits
        (``{'slide': n}``, ``{'sheet': n, 'sheet_name': name}``, ``{'page': n}``), for the loader's
        OCR issue record (see ``_issue_location``)."""
        if not image_bytes or not self.ocr:
            return None
        # Use the configured OCR instance with language configuration if available
        kwargs = {}
        if hasattr(self.ocr, 'config') and self.ocr.config and self.ocr.config.language:
            kwargs['language'] = self.ocr.config.language
        try:
            result = self.ocr.process_image(image_bytes, **kwargs)
        except Exception as e:
            logger.error(f"OCR failed: {e}; the picture is marked {OCR_UNAVAILABLE}")
            self._label_ocr_issues([location])
            return None
        self._label_ocr_issues([location])
        if self._ocr_failed(result):
            return None
        if hasattr(result, 'text'):
            self._ocr_cell_texts[_image_hash(image_bytes)] = plain_ocr_text(
                result.text, getattr(result, 'document', None))
            return result.text
        return str(result)

    def _classify_text_type(self, text: str, style_name: str) -> str:
        """Classify text type based on style and content"""
        if not text:
            return "text:normal"

        # Check style-based classification first
        if style_name:
            style_lower = style_name.lower()

            # Check for title/heading styles ('subtitle' first: it contains 'title')
            if 'subtitle' in style_lower:
                return "text:section"
            elif 'title' in style_lower:
                return "text:title"
            elif 'heading' in style_lower:
                level = re.search(r'\d+', style_lower)
                return "text:title" if level and level.group(0) == '1' else "text:section"
            elif 'caption' in style_lower:
                return "text:caption"
            elif any(x in style_lower for x in ['list', 'bullet']):
                return "text:list"

        # Content-based classification as fallback
        # Check for list patterns
        if re.match(r'^[\u2022•\-\*\d]+[\.\)]\s+', text):
            return "text:list"

        # Caption-shaped text: a caption word, a number and a separator ("Figure 2: ...",
        # "Table 3. ..."), or a source / note line. "Tablets are ...", "Table of contents" and
        # "Figure 1 shows ..." are sentences.
        if _CAPTION_TEXT.match(text) or re.match(r'(Source|Note|Notes)[\.:)]', text, re.IGNORECASE):
            return "text:caption"

        # Default to normal text
        return "text:normal"

    def _batch_ocr_images(self, images_info: List[Dict[str, Any]]) -> Dict[Any, Optional[str]]:
        """OCR the document's pictures in one batch call.

        Args:
            images_info: List of dictionaries containing:
                - 'id': Unique identifier for the image
                - 'data': Image bytes
                - 'location': where the picture sits (slide, sheet or page), for the loader's
                  OCR issue record

        Returns:
            ``{image hash: OCR text}`` (and ``{('_fallback_ocr', index): text}`` for XLSX pictures
            read from the package); the text is None for a picture that was not read: an answer
            the provider flags ``failed``, or every picture of a call that raised (no API key, an
            outage). Those are not sent again one by one: a provider that raised fails them again,
            and the next run asks again (such a document is not cached). A picture missing from
            the map is OCR'd on its own where it is used.
        """
        if not images_info or not self.ocr:
            return {}

        image_data_list = [info['data'] for info in images_info]
        language_info = getattr(self.ocr.config, 'language', 'auto') if hasattr(self.ocr,
                                                                                'config') and self.ocr.config else 'auto'
        logger.info(f"Processing {len(image_data_list)} images with configured OCR (language: {language_info})...")

        # Pass language configuration if available
        kwargs = {}
        if hasattr(self.ocr, 'config') and self.ocr.config and self.ocr.config.language:
            kwargs['language'] = self.ocr.config.language
            logger.info(f"🌍 Passing language configuration to OCR: {self.ocr.config.language}")

        try:
            ocr_results = self.ocr.batch_process_images(image_data_list, **kwargs)
        except Exception as e:
            logger.error(f"Batch OCR processing failed: {e}; {len(images_info)} picture(s) marked {OCR_UNAVAILABLE}")
            ocr_results = [None] * len(images_info)
        # Tell the loader's OCR issue record where each picture is (slide, sheet or page).
        self._label_ocr_issues([info.get('location') for info in images_info])

        results_map: Dict[Any, Optional[str]] = {}
        failed = 0
        for image_info, ocr_result in zip(images_info, ocr_results):
            img_hash = _image_hash(image_info['data'])
            if self._ocr_failed(ocr_result):
                ocr_text = None
                failed += 1
            else:
                ocr_text = ocr_result.text if hasattr(ocr_result, 'text') else str(ocr_result)
                if hasattr(ocr_result, 'text'):
                    self._ocr_cell_texts[img_hash] = plain_ocr_text(ocr_text, getattr(ocr_result, 'document', None))
            # The same picture sent twice keeps an answer over a failure.
            if ocr_text is not None or img_hash not in results_map:
                results_map[img_hash] = ocr_text
            # XLSX pictures read from the package are also found by their index
            if image_info.get('location', {}).get('source') == 'zip_fallback':
                results_map[('_fallback_ocr', image_info['location']['img_idx'])] = ocr_text

        logger.info(f"OCR answered {len(ocr_results) - failed} of {len(images_info)} picture(s)")
        return results_map

    def _collect_all_images(self) -> List[Dict[str, Any]]:
        """Collect all images from the document for batch processing
        
        This method should be overridden by subclasses
        
        Returns:
            List of dictionaries containing:
                - 'id': Unique identifier for the image
                - 'data': Image bytes
                - 'location': Location info (page/slide/sheet number, etc.)
        """
        raise NotImplementedError("Subclasses must implement _collect_all_images")


class DocxLoader(BaseOfficeLoader):
    """Loader for DOCX (Word) documents"""

    def __init__(self, file_path: Union[str, Path], ocr=None, table_style: Union[str, TableStyle] = None):
        super().__init__(file_path, ocr, table_style)
        self._open_document()
        self._structure = _DocxStructure(self.doc)
        # The part whose relationships resolve the r:id of pictures and links: the document, or the
        # header or footer being read (a header's r:id means nothing in the document's relationships).
        self._part = self.doc.part
        # Where the picture being read sits ({'page': n}), for the loader's OCR issue record.
        self._location: Optional[Dict[str, Any]] = None
        self._layout_cache = None

    def _open_document(self):
        """Open DOCX document with error handling and configuration logging"""
        try:
            self.doc = Document(self.file_path)

            # Log DOCX configuration
            logger.info("=" * 60)
            logger.info(f"DOCX Configuration for: {self.file_path.name}")
            logger.info("=" * 60)
            logger.info(f"File path: {self.file_path}")
            logger.info(f"File size: {self.file_path.stat().st_size / (1024 * 1024):.2f} MB")

            # Count paragraphs and tables
            para_count = len(list(self.doc.paragraphs))
            table_count = len(self.doc.tables)
            logger.info(f"Total paragraphs: {para_count}")
            logger.info(f"Total tables: {table_count}")

            # Count images (approximate - includes inline shapes)
            inline_shape_count = len(self.doc.inline_shapes) if hasattr(self.doc, 'inline_shapes') else 0
            logger.info(f"Inline shapes (includes images): {inline_shape_count}")

            # Count sections
            section_count = len(self.doc.sections)
            logger.info(f"Total sections: {section_count}")

            # Document properties
            core_props = self.doc.core_properties
            if core_props:
                logger.info("Document Properties:")
                if core_props.title:
                    logger.info(f"  Title: {core_props.title}")
                if core_props.author:
                    logger.info(f"  Author: {core_props.author}")
                if core_props.subject:
                    logger.info(f"  Subject: {core_props.subject}")
                if core_props.created:
                    logger.info(f"  Created: {core_props.created}")
                if core_props.modified:
                    logger.info(f"  Modified: {core_props.modified}")
                if core_props.last_modified_by:
                    logger.info(f"  Last modified by: {core_props.last_modified_by}")

            logger.info("=" * 60)

        except Exception as e:
            logger.error(f"Failed to open DOCX: {e}")
            raise

    def convert_to_json(self,
                        extract_images: bool = True,
                        ocr_images: bool = False,
                        show_progress: bool = True) -> Dict[str, Any]:
        """Convert DOCX to JSON format.

        Body paragraphs and tables come in document order, each tagged with the page it starts on
        (counted from page and section breaks, see ``_layout``), the text boxes anchored in a
        paragraph right after its text. The headers and footers a section shows are added once
        each (a part shared with an earlier section is not repeated): headers before the section's
        first block, footers after its last, their items marked ``"region": "header"`` /
        ``"footer"``. Footnotes and endnotes come last."""
        result = {
            "filename": self.file_path.name,
            "pages": 1,  # DOCX doesn't have fixed pages
            "content": []
        }

        if show_progress:
            logging.info(f"Processing DOCX: {self.file_path.name}")

        # One OCR batch for every picture the document shows, before the text is laid out
        ocr_images = bool(extract_images and ocr_images and self.ocr)
        ocr_results_map = {}
        if ocr_images:
            if show_progress:
                logger.info("Collecting all images for batch OCR processing...")
            all_images_info = self._collect_all_images()
            if all_images_info:
                if show_progress:
                    logger.info(f"Processing {len(all_images_info)} images with batch OCR...")
                ocr_results_map = self._batch_ocr_images(all_images_info)

        blocks, pages, stories, _ = self._layout()
        content = result["content"]
        options = (extract_images, ocr_images, ocr_results_map, set())  # the set: pictures already emitted
        section, last_page = -1, 1
        for block, page, block_section in blocks:
            while section < block_section:
                if section >= 0:
                    self._add_story_items(stories[section], "footer", last_page, content, *options)
                section += 1
                self._add_story_items(stories[section], "header", page, content, *options)
            start = len(content)
            self._location = {'page': page}
            self._add_block(block, content, *options)
            for item in content[start:]:
                item["page"] = page
            last_page = page
        # The last section's footers; a section without blocks (after a trailing section break)
        # still shows its headers and footers.
        while section < len(stories) - 1:
            if section >= 0:
                self._add_story_items(stories[section], "footer", last_page, content, *options)
            section += 1
            self._add_story_items(stories[section], "header", last_page, content, *options)
        if section >= 0:
            self._add_story_items(stories[section], "footer", last_page, content, *options)
        result["pages"] = pages

        # Extract footnotes and endnotes
        try:
            footnotes = self._load_footnotes()
            for note_id, note_text in sorted(footnotes.items(), key=lambda x: int(x[0]) if x[0].isdigit() else 0):
                content.append({
                    "type": "text:footnote",
                    "content": f"[^{note_id}]: {note_text}",
                    "page": pages,  # footnotes at end of doc
                })
        except Exception as e:
            logger.debug(f"Failed to extract footnotes: {e}")

        return result

    def _layout(self):
        """``(blocks, pages, stories, section_pages)``: the body's top-level paragraphs and tables in
        document order as ``(element, page, section)``, the page count, for each section the header
        and footer parts it shows (``_section_stories``) and ``{section: (first page, last page)}``.

        Pages follow Word's pagination only approximately: an explicit page break or Word's last
        rendered page break in a paragraph puts it on the next page, and a paragraph ending a
        next/odd/even-page section puts the following ones on the next page (the first rendered
        break after it is that same break). A paragraph carrying ``w:sectPr`` ends its section."""
        if self._layout_cache is not None:
            return self._layout_cache
        body = self.doc.element.body
        blocks: List[Tuple[Any, int, int]] = []
        section_breaks = []
        page, section, skip_rendered = 1, 0, False
        for block in _docx_blocks(body):
            if block.tag != _W_P:
                blocks.append((block, page, section))
                continue
            # Detect page breaks before this paragraph (at most one increment)
            has_break = False
            for run_el in block.findall(_W_R):
                br = run_el.find(_W_BR)
                if br is not None and br.get(_W_TYPE) == 'page':
                    has_break = True
                    break
                if run_el.find(_W_LAST_RENDERED_BREAK) is not None:
                    if skip_rendered:
                        skip_rendered = False
                    else:
                        has_break = True
                    break
            ppr = block.find(_W_PPR)
            sect_pr = ppr.find(_W_SECT_PR) if ppr is not None else None
            if sect_pr is not None:
                section_breaks.append(sect_pr)
            if not has_break and sect_pr is not None:
                sect_type = sect_pr.find(_W_TYPE)
                if (sect_type.get(_W_VAL, 'nextPage') if sect_type is not None else 'nextPage') in (
                        'nextPage', 'oddPage', 'evenPage'):
                    # This paragraph ENDS the section: it stays on this page, the next one starts a new page.
                    blocks.append((block, page, section))
                    page += 1
                    skip_rendered = True
                    section += 1
                    continue
            if has_break:
                page += 1
            blocks.append((block, page, section))
            if sect_pr is not None:
                section += 1
        stories = self._section_stories(section_breaks + [body.find(_W_SECT_PR)])
        section_pages: Dict[int, Tuple[int, int]] = {}
        for _, block_page, block_section in blocks:
            section_pages[block_section] = (section_pages.get(block_section, (block_page,))[0], block_page)
        self._layout_cache = (blocks, page, stories, section_pages)
        return self._layout_cache

    def _section_stories(self, sect_prs) -> List[Dict[str, list]]:
        """For each section (its ``w:sectPr``; None when missing), the header and footer parts Word
        shows in it: ``{"header": [parts], "footer": [parts]}``, in the order first page (with
        ``w:titlePg``), default, even pages (with ``w:evenAndOddHeaders`` in the settings). A section
        without a reference of a kind inherits the previous section's. A part is listed only for the
        first section that shows it, so a header or footer shared by several sections is written
        once."""
        try:
            even_and_odd = _docx_flag(self.doc.settings.element.find(_W_EVEN_AND_ODD))
        except Exception:
            even_and_odd = False
        related = self.doc.part.related_parts
        references: Dict[Tuple[str, str], Optional[str]] = {}
        seen = set()   # part names and contents already listed
        stories = []
        for sect_pr in sect_prs:
            shown: Dict[str, list] = {"header": [], "footer": []}
            if sect_pr is not None:
                for region, tag in (("header", _W_HEADER_REF), ("footer", _W_FOOTER_REF)):
                    for reference in sect_pr.findall(tag):
                        references[(region, reference.get(_W_TYPE, 'default'))] = reference.get(_R_ID)
                kinds = (['first'] if _docx_flag(sect_pr.find(_W_TITLE_PG)) else []) + ['default'] + (
                    ['even'] if even_and_odd else [])
                for region in ("header", "footer"):
                    for kind in kinds:
                        rid = references.get((region, kind))
                        part = related.get(rid) if rid else None
                        if part is None or part.partname in seen:
                            continue
                        seen.add(part.partname)
                        signature = self._story_signature(part)
                        if signature in seen:
                            continue  # the same text and pictures as a header or footer already listed
                        seen.add(signature)
                        shown[region].append(part)
            stories.append(shown)
        return stories

    def _story_signature(self, part) -> Tuple[str, ...]:
        """What a header or footer part shows: its text (whitespace collapsed) and its pictures. Sections
        that each have their own part with the same content (unlinked headers, merged documents) show
        one header, written once, like the first copy of a PDF running header."""
        texts, pictures = [], []
        for element in _docx_rendered(part.element):
            if element.tag == _W_T:
                texts.append(element.text or '')
            elif element.tag.endswith(('}inline', '}anchor')):
                rid = _drawing_blip_rid(element)
                target = part.related_parts.get(rid) if rid else None
                if target is not None and getattr(target, 'blob', None):
                    pictures.append(_image_hash(target.blob))
        return ('story', ' '.join(''.join(texts).split())) + tuple(pictures)

    def _add_story_items(self, stories: Dict[str, list], region: str, page: int, content: List[Dict],
                         extract_images: bool, ocr_images: bool, ocr_results_map: Dict[Any, Optional[str]],
                         processed_image_hashes: set) -> None:
        """Add the paragraphs and tables of one section's ``region`` ("header" or "footer") parts,
        tagged with ``page`` and ``region``."""
        for part in stories[region]:
            start = len(content)
            previous, self._part = self._part, part
            self._location = {'page': page}
            try:
                with self._structure.separate_numbering():
                    for block in _docx_blocks(part.element):
                        self._add_block(block, content, extract_images, ocr_images, ocr_results_map,
                                        processed_image_hashes)
            except Exception as e:
                logger.warning(f"Failed to read a {region} ({part.partname}): {e}")
            finally:
                self._part = previous
            for item in content[start:]:
                item["page"] = page
                item["region"] = region
                if item["type"] in ("text:title", "text:section"):
                    # A heading-styled header line is not a heading of the document's outline.
                    item["type"] = "text:normal"
                    item.pop("level", None)
                    marker = item.pop("marker", None)
                    if marker:
                        item["markdown"] = escape_markdown_text(f"{marker} {item['content']}")

    def _add_block(self, block, content: List[Dict], extract_images: bool, ocr_images: bool,
                   ocr_results_map: Dict[Any, Optional[str]], processed_image_hashes: set) -> None:
        """Add the items of a paragraph (with the text boxes anchored in it) or of a table."""
        if block.tag == _W_P:
            self._process_paragraph(Paragraph(block, self.doc), content, extract_images, ocr_images,
                                    ocr_results_map, processed_image_hashes)
            return
        try:
            table_md = self._convert_table_to_markdown(Table(block, self.doc), extract_images, ocr_images,
                                                       ocr_results_map)
        except Exception as e:
            # One malformed table must not send the whole document to the basic
            # converter (which reorders content): keep its text, lose its structure.
            logger.warning(f"DOCX table conversion failed ({e}); keeping the table text without merges")
            table_md = self._table_text_fallback(block)
        if table_md:
            content.append({"type": "table", "content": table_md})

    def _load_footnotes(self) -> Dict[str, str]:
        """Extract footnotes and endnotes from the DOCX ZIP.

        python-docx has no native API for footnotes, so we parse the raw XML
        from ``word/footnotes.xml`` and ``word/endnotes.xml`` directly.

        Returns:
            Mapping of note id -> text content.
        """
        import zipfile
        from lxml import etree

        notes: Dict[str, str] = {}
        nsmap = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

        try:
            with zipfile.ZipFile(str(self.file_path), "r") as zf:
                for xml_path in ("word/footnotes.xml", "word/endnotes.xml"):
                    if xml_path not in zf.namelist():
                        continue
                    tree = etree.parse(zf.open(xml_path), parser=_safe_lxml_parser())
                    root = tree.getroot()
                    is_endnote = "endnote" in xml_path
                    tag_local = "endnote" if is_endnote else "footnote"
                    for note_el in root.findall(f"w:{tag_local}", nsmap):
                        note_id = note_el.get(f"{{{nsmap['w']}}}id", "")
                        # Skip separator/continuation notes (id 0 and -1)
                        if note_id in ("0", "-1"):
                            continue
                        # Prefix endnote ids to avoid collision with footnotes
                        key = f"en{note_id}" if is_endnote else note_id
                        # Collect all paragraph text
                        paragraphs = note_el.findall(".//w:p", nsmap)
                        texts = []
                        for p in paragraphs:
                            runs = p.findall(".//w:r/w:t", nsmap)
                            texts.append("".join(r.text or "" for r in runs))
                        note_text = " ".join(t for t in texts if t.strip())
                        if note_text.strip():
                            notes[key] = note_text.strip()
        except (zipfile.BadZipFile, etree.XMLSyntaxError, KeyError) as e:
            logger.debug(f"Could not parse footnotes XML: {e}")

        return notes

    def _process_paragraph(self, paragraph: Paragraph, content: List[Dict], extract_images: bool,
                           ocr_images: bool = False, ocr_results_map: Optional[Dict[Any, Optional[str]]] = None,
                           processed_image_hashes: set = None):
        """Add a paragraph's pictures, its text item and the items of the text boxes anchored in it.

        The text item's ``content`` is the paragraph's text as written; when its runs carry bold,
        italics or web / mail links, a normal paragraph or a list item also gets ``markdown``: the
        same text as inline Markdown (``**bold**``, ``*italic*``, ``[text](address)``, escaped like
        body text), which the Markdown output uses."""
        if ocr_results_map is None:
            ocr_results_map = {}
        if processed_image_hashes is None:
            processed_image_hashes = set()
        p_el = paragraph._p

        # Extract images from runs first
        if extract_images:
            for run in _docx_paragraph_content(p_el):
                if run.tag == _W_R:
                    content.extend(self._extract_run_images(run, ocr_images, ocr_results_map, processed_image_hashes))

        # List numbering advances for every list paragraph, empty ones included, as in Word.
        numbering = self._structure.list_marker(p_el)

        # Then extract text. In a header or footer, a paragraph that only shows a page number
        # (PAGE, NUMPAGES, SECTIONPAGES fields) is not text of the document: its number is whatever
        # page Word last showed it on, as the slide number of a PowerPoint footer.
        pieces = _strip_pieces(self._paragraph_pieces(p_el))
        text = ''.join(piece[0] for piece in pieces)
        numbered_by = [field for text_part, *_, field in pieces if any(char.isalnum() for char in text_part)]
        if self._part is not self.doc.part and numbered_by and all(field in _PAGE_NUMBER_FIELDS for field in numbered_by):
            text = ''
        if text:
            # Footnote/endnote references in the paragraph
            refs = self._note_references(p_el)
            suffix = " " + " ".join(f"[^{ref}]" for ref in refs) if refs else ""
            item = self._paragraph_item(paragraph, text + suffix, numbering)
            if item["type"] == "text:normal" or (item["type"] == "text:list" and item.get("marker")):
                markdown = _docx_inline_markdown(pieces) + suffix
                if markdown != escape_markdown_text(item["content"]):
                    item["markdown"] = markdown
            content.append(item)

        # Text boxes and shapes anchored in the paragraph, after its own text, in document order.
        # A text box is a story of its own: its lists do not move the numbers of the body's.
        for box in _docx_paragraph_text_boxes(p_el):
            with self._structure.separate_numbering():
                for block in _docx_blocks(box):
                    self._add_block(block, content, extract_images, ocr_images, ocr_results_map,
                                    processed_image_hashes)

    @staticmethod
    def _note_references(p_el) -> List[str]:
        """Ids of the footnotes (``3``) and endnotes (``en3``) referenced in a paragraph, in order."""
        refs = []
        for run in _docx_paragraph_content(p_el):
            for tag, prefix in ((qn('w:footnoteReference'), ''), (qn('w:endnoteReference'), 'en')):
                reference = run.find(tag)
                if reference is not None:
                    ref_id = reference.get(qn('w:id'), '')
                    if ref_id and ref_id not in ('0', '-1'):
                        refs.append(f"{prefix}{ref_id}")
        return refs

    def _paragraph_pieces(self, p_el) -> List[Tuple[str, bool, bool, Optional[str], Optional[str]]]:
        """``(text, bold, italic, link target, field)`` for the runs of a paragraph in reading order;
        their texts joined are the paragraph's text. A run's link target comes from its
        ``w:hyperlink``, its ``w:fldSimple`` HYPERLINK field, or the HYPERLINK field whose result it
        shows (``w:fldChar`` begin, ``w:instrText``, separate, result runs, end); ``field`` names the
        field whose result the run shows (``PAGE``, ``HYPERLINK``...), or is None."""
        pieces = []
        fields: List[List[Any]] = []   # open fields: [instruction parts, showing its result]
        for run, link in _docx_paragraph_items(p_el):
            if run.tag == _W_R:
                for child in run:
                    if child.tag == _W_FLD_CHAR:
                        kind = child.get(_W_FLD_CHAR_TYPE)
                        if kind == 'begin':
                            fields.append([[], False])
                        elif kind == 'separate' and fields:
                            fields[-1][1] = True
                        elif kind == 'end' and fields:
                            fields.pop()
                    elif child.tag == _W_INSTR_TEXT and fields and not fields[-1][1]:
                        fields[-1][0].append(child.text or '')
            text = _docx_run_text(run)
            if not text:
                continue
            target = self._link_target(link)
            instruction = ''.join(fields[-1][0]) if fields and fields[-1][1] else None
            if instruction is None and link is not None and link.tag == _W_FLD_SIMPLE:
                instruction = link.get(_W_INSTR) or ''
            if target is None and instruction:
                target = _hyperlink_field_target(instruction)
            field = (instruction.split() or [None])[0].upper() if instruction else None
            bold, italic = self._structure.run_emphasis(run) if run.tag == _W_R else (False, False)
            pieces.append((text, bold, italic, target, field))
        return pieces

    def _link_target(self, link) -> Optional[str]:
        """The address a ``w:hyperlink`` (its external relationship, plus its ``w:anchor`` bookmark as
        a fragment) or a ``w:fldSimple`` HYPERLINK field points to; None for a bookmark in the
        document alone, another field, or no link."""
        if link is None:
            return None
        if link.tag == _W_FLD_SIMPLE:
            return _hyperlink_field_target(link.get(_W_INSTR))
        rid = link.get(_R_ID)
        relationship = self._part.rels.get(rid) if rid else None
        if relationship is None or not relationship.is_external:
            return None
        anchor = link.get(_W_ANCHOR)
        return f"{relationship.target_ref}#{anchor}" if anchor else relationship.target_ref

    def _paragraph_item(self, paragraph: Paragraph, text: str,
                        numbering: Optional[Tuple[str, int]]) -> Dict[str, Any]:
        """Typed content item for a paragraph's text.

        Headings (outline level, or the Title/Subtitle/Heading N styles) carry
        ``level`` (1-6); list paragraphs carry ``marker`` (``"1."``, ``"-"``...) and
        ``list_level``. The structure lives beside ``content`` so the text stays verbatim.
        """
        try:
            style_name = paragraph.style.name if paragraph.style is not None else ""
        except (AttributeError, KeyError, ValueError):
            style_name = ""
        outline = self._structure.outline_level(paragraph._p)
        level = self._heading_level(outline, style_name)
        if level:
            item = {"type": "text:title" if level == 1 else "text:section", "content": text, "level": min(level, 6)}
            if numbering and numbering[0] != '-':
                item["marker"] = numbering[0]
            return item
        if numbering:
            marker, depth = numbering
            return {"type": "text:list", "content": text, "marker": marker,
                    "list_level": max(depth, _style_list_depth(style_name))}
        text_type = self._classify_text_type(text, style_name)
        if outline == 9 and text_type in ("text:title", "text:section"):
            text_type = "text:normal"  # explicitly body text, whatever the style is called
        return {"type": text_type, "content": text}

    @staticmethod
    def _heading_level(outline: Optional[int], style_name: str) -> Optional[int]:
        """1-9 for headings, None for body text. The outline level Word uses for its
        navigation pane decides (9 is body text, even in a style named "Heading 2"); the
        built-in Title (1) and Subtitle (2) styles, which have none, go by name, and so do
        Heading N styles when no outline level is set anywhere."""
        if outline is not None and outline <= 8:
            return outline + 1
        name = (style_name or "").strip().lower()
        if name == "title":
            return 1
        if name == "subtitle":
            return 2
        match = re.fullmatch(r'heading\s*([1-9])', name)
        return int(match.group(1)) if match and outline is None else None

    def _extract_run_images(self, run, ocr_images: bool = False,
                            ocr_results_map: Optional[Dict[Any, Optional[str]]] = None,
                            processed_image_hashes: set = None) -> List[Dict[str, Any]]:
        """Items for the pictures drawn in a run (``w:drawing`` inline or anchored), see
        ``_drawing_item``."""
        if ocr_results_map is None:
            ocr_results_map = {}
        if processed_image_hashes is None:
            processed_image_hashes = set()
        items = []
        try:
            for child in getattr(run, '_element', run):
                if child.tag != _W_DRAWING:
                    continue
                for drawing in child:
                    if drawing.tag.endswith(('}inline', '}anchor')):
                        item = self._drawing_item(drawing, ocr_images, ocr_results_map, processed_image_hashes)
                        if item:
                            items.append(item)
        except Exception as e:
            logger.warning(f"Failed to extract images from run: {e}")
        return items

    def _drawing_item(self, drawing, ocr_images: bool, ocr_results_map: Dict[Any, Optional[str]],
                      processed_image_hashes: set) -> Optional[Dict[str, Any]]:
        """The item for the picture a drawing shows: its OCR text (the ``[image: OCR unavailable]``
        marker when its OCR failed; nothing when OCR found no text in it) or, without OCR, the image
        as base64. A picture already emitted in this document (the same bytes) gives nothing."""
        image_bytes = self._drawing_image_bytes(drawing)
        if not image_bytes:
            return None
        img_hash = _image_hash(image_bytes)
        if img_hash in processed_image_hashes:
            return None
        processed_image_hashes.add(img_hash)
        if not ocr_images:
            return {"type": "image", "content": base64.b64encode(image_bytes).decode('utf-8')}
        if img_hash in ocr_results_map:
            return self._ocr_description(ocr_results_map[img_hash])
        logger.warning("A picture was not in the OCR batch; OCR'ing it on its own")
        return self._ocr_description(self._ocr_image(image_bytes, self._location))

    def _drawing_image_bytes(self, drawing) -> Optional[bytes]:
        """Bytes of the picture a ``wp:inline`` / ``wp:anchor`` shows, resolved through the
        relationships of the part being read (the document, or a header or footer)."""
        rid = _drawing_blip_rid(drawing)
        try:
            part = self._part.related_parts.get(rid) if rid else None
        except Exception as e:
            logger.warning(f"Failed to get image by rId {rid}: {e}")
            return None
        return getattr(part, 'blob', None) if part is not None else None

    def _collect_all_images(self) -> List[Dict[str, Any]]:
        """Every picture the document shows, once per distinct image, in document order: in body
        paragraphs and tables at any depth, content controls and text boxes (the ``mc:Choice``
        branch), then in the headers and footers; deleted revisions and ``mc:Fallback`` copies are
        skipped. Each carries the page it is on (a header or footer picture: the page its block is
        written at)."""
        blocks, _, stories, section_pages = self._layout()
        last_page = blocks[-1][1] if blocks else 1
        images_info: List[Dict[str, Any]] = []
        seen = set()

        def add_drawings(root, location: Dict[str, Any]) -> None:
            for element in _docx_rendered(root):
                if element.tag != _W_DRAWING:
                    continue
                for drawing in element:
                    if not drawing.tag.endswith(('}inline', '}anchor')):
                        continue
                    data = self._drawing_image_bytes(drawing)
                    if not data or _image_hash(data) in seen:
                        continue
                    seen.add(_image_hash(data))
                    images_info.append({'id': f'docx_image_{len(images_info) + 1}', 'data': data,
                                        'location': dict(location)})

        for block, page, _ in blocks:
            add_drawings(block, {'type': 'body', 'page': page})
        for section, shown in enumerate(stories):
            first, last = section_pages.get(section, (last_page, last_page))
            for region, page in (("header", first), ("footer", last)):
                for part in shown[region]:
                    previous, self._part = self._part, part
                    try:
                        add_drawings(part.element, {'type': region, 'page': page})
                    except Exception as e:
                        logger.warning(f"Failed to collect the pictures of a {region}: {e}")
                    finally:
                        self._part = previous

        logger.info(f"Collected {len(images_info)} unique images from DOCX")
        return images_info

    # ------------------------------------------------------------------ #
    # Tables: read from the w:tbl XML                                     #
    # ------------------------------------------------------------------ #
    def _convert_table_to_markdown(self, table_data, extract_images: bool = True, ocr_images: bool = False,
                                   ocr_results_map: Optional[Dict[str, str]] = None) -> str:
        """Render a python-docx ``Table`` (merges from ``w:gridSpan``/``w:vMerge``, rows
        offset by ``w:gridBefore``/``w:gridAfter``); other inputs go to the base class."""
        tbl = getattr(table_data, '_tbl', None)
        if tbl is None:
            return super()._convert_table_to_markdown(table_data, extract_images, ocr_images, ocr_results_map)
        ocr_results_map = ocr_results_map or {}
        texts, spans = self._docx_table_grid(
            tbl, lambda tc: "\n".join(self._tc_lines(tc, extract_images, ocr_images, ocr_results_map)))
        if not texts:
            return ""
        logger.debug(f"DOCX Table: {len(texts)}x{len(texts[0])}, {len(spans)} merged cells")
        table_obj = TableData.from_raw(texts, {'is_complex': bool(spans), 'cell_spans': spans})
        return TableRenderer(self.table_style).render(table_obj)

    @staticmethod
    def _docx_table_grid(tbl, cell_text) -> Tuple[List[List[str]], Dict[Tuple[int, int], Tuple[int, int]]]:
        """Lay a ``w:tbl`` out on its column grid.

        Each row starts at its ``w:gridBefore`` column; a cell covers ``w:gridSpan``
        columns; ``w:vMerge`` continuation cells extend the cell above that starts in
        the same column with the same width (legacy ``w:hMerge`` continuations extend the
        cell to their left). Text found in a continuation cell is appended to the merged
        cell rather than dropped. Deleted (tracked) rows are skipped.

        Returns ``(texts, spans)``: a rectangular grid of cell texts (``""`` for empty and
        covered positions) and ``{(row, col): (rowspan, colspan)}`` for merged cells.
        """
        grid_cols = tbl.find(_W_TBL_GRID)
        width = len(grid_cols.findall(_W_GRID_COL)) if grid_cols is not None else 0
        rows = []
        for tr in _docx_children(tbl, _W_TR):
            tr_pr = tr.find(_W_TR_PR)
            if tr_pr is not None and tr_pr.find(_W_DEL) is not None:
                continue
            col = max(0, _docx_int(tr_pr.find(_W_GRID_BEFORE) if tr_pr is not None else None, 0))
            cells = []
            for tc in _docx_children(tr, _W_TC):
                tc_pr = tc.find(_W_TC_PR)
                span = max(1, _docx_int(tc_pr.find(_W_GRID_SPAN) if tc_pr is not None else None, 1))
                v_merge = tc_pr.find(_W_V_MERGE) if tc_pr is not None else None
                h_merge = tc_pr.find(_W_H_MERGE) if tc_pr is not None else None
                continues_above = v_merge is not None and v_merge.get(_W_VAL, 'continue') != 'restart'
                continues_left = h_merge is not None and h_merge.get(_W_VAL, 'continue') != 'restart'
                if continues_left and cells:
                    start, previous_span, above, _, texts_in = cells[-1]
                    cells[-1] = (start, previous_span + span, above, tc, texts_in + [tc])
                else:
                    cells.append((col, span, continues_above, tc, [tc]))
                col += span
            col += max(0, _docx_int(tr_pr.find(_W_GRID_AFTER) if tr_pr is not None else None, 0))
            width = max(width, col)
            rows.append(cells)

        texts = [[""] * width for _ in rows]
        spans: Dict[Tuple[int, int], Tuple[int, int]] = {}
        origin_of: Dict[Tuple[int, int], Tuple[int, int]] = {}
        for r, cells in enumerate(rows):
            for col, span, continues_above, _, tcs in cells:
                span = min(span, width - col)
                if span <= 0:
                    continue
                text = "\n".join(t for t in (cell_text(tc) for tc in tcs) if t)
                origin = origin_of.get((r - 1, col)) if continues_above else None
                if origin is not None and spans.get(origin, (1, 1))[1] == span:
                    rowspan, colspan = spans.get(origin, (1, span))
                    spans[origin] = (rowspan + 1, colspan)
                    origin_of[(r, col)] = origin
                    if text:
                        texts[origin[0]][origin[1]] = "\n".join(t for t in (texts[origin[0]][origin[1]], text) if t)
                    continue
                texts[r][col] = text
                origin_of[(r, col)] = (r, col)
                if span > 1:
                    spans[(r, col)] = (1, span)
        return texts, spans

    def _tc_lines(self, container, extract_images: bool, ocr_images: bool,
                  ocr_results_map: Dict[Any, Optional[str]]) -> List[str]:
        """Lines of a cell in document order: one per non-empty paragraph, then the lines of
        the text boxes anchored in it, and one per row of a nested table (its cells joined
        with ``" | "``)."""
        lines = []
        for block in _docx_blocks(container):
            if block.tag == _W_P:
                text = self._cell_paragraph_text(block, extract_images, ocr_images, ocr_results_map)
                if text:
                    lines.append(text)
                for box in _docx_paragraph_text_boxes(block):
                    with self._structure.separate_numbering():
                        lines.extend(self._tc_lines(box, extract_images, ocr_images, ocr_results_map))
            else:
                texts, _ = self._docx_table_grid(
                    block, lambda tc: " ".join(self._tc_lines(tc, extract_images, ocr_images, ocr_results_map)))
                lines.extend(" | ".join(t for t in row if t) for row in texts if any(row))
        return lines

    def _cell_paragraph_text(self, p_el, extract_images: bool, ocr_images: bool,
                             ocr_results_map: Dict[Any, Optional[str]]) -> str:
        """Text of a paragraph in a table cell, image markers in place and its list
        number or bullet in front."""
        parts = []
        for run in _docx_paragraph_content(p_el):
            if extract_images and run.tag == _W_R:
                parts.extend(self._cell_image_markers(run, ocr_images, ocr_results_map))
            parts.append(_docx_run_text(run))
        numbering = self._structure.list_marker(p_el)
        text = "".join(parts).strip()
        return f"{numbering[0]} {text}" if text and numbering else text

    def _cell_image_markers(self, run, ocr_images: bool, ocr_results_map: Dict[Any, Optional[str]]) -> List[str]:
        """``[Image]`` (or ``[Image: <OCR text>]``, ``[image: OCR unavailable]`` when its OCR failed)
        for each picture drawn in ``run``."""
        markers = []
        for child in run:
            if child.tag != _W_DRAWING:
                continue
            for drawing in child:
                if not drawing.tag.endswith(('}inline', '}anchor')):
                    continue
                image_bytes = self._drawing_image_bytes(drawing)
                if not image_bytes:
                    continue
                if not ocr_images:
                    markers.append("[Image]")
                    continue
                img_hash = _image_hash(image_bytes)
                ocr_text = ocr_results_map[img_hash] if img_hash in ocr_results_map else self._ocr_image(
                    image_bytes, self._location)
                if ocr_text is None:
                    markers.append(OCR_UNAVAILABLE)
                    continue
                # The cell's text is escaped by the table renderer: give it the plain OCR text.
                ocr_text = self._ocr_cell_texts.get(img_hash, ocr_text).strip()
                markers.append(f"[Image: {ocr_text}]" if ocr_text else "[Image]")
        return markers

    def _table_text_fallback(self, tbl) -> str:
        """Every cell's text on a plain grid (no merges), for a table whose structure
        could not be read; the text itself is never dropped (and text a reader does not
        see, such as the ``mc:Fallback`` copy of a text box, is not added)."""
        def shown_text(element) -> str:
            return "".join(t.text or "" for t in _docx_rendered(element) if t.tag == _W_T)

        try:
            rows = [[shown_text(tc) for tc in _docx_children(tr, _W_TC)] for tr in _docx_children(tbl, _W_TR)]
            width = max((len(row) for row in rows), default=0)
            rows = [row + [""] * (width - len(row)) for row in rows if any(cell.strip() for cell in row)]
            if rows:
                return TableRenderer(self.table_style).render(TableData.from_raw(rows, {'is_complex': False}))
        except Exception as e:
            logger.warning(f"DOCX table text fallback failed ({e}); emitting the raw text")
        return "\n".join(t.text for t in _docx_rendered(tbl) if t.tag == _W_T and t.text)


class PptxLoader(BaseOfficeLoader):
    """Loader for PPTX (PowerPoint) documents"""

    def __init__(self, file_path: Union[str, Path], ocr=None, table_style: Union[str, TableStyle] = None):
        super().__init__(file_path, ocr, table_style)
        self._open_document()
        self._inherited_seen: set = set()

    def _open_document(self):
        """Open PPTX document"""
        try:
            self.doc = Presentation(self.file_path)
            logger.info(f"Opened PPTX: {self.file_path.name}")
            logger.info(f"Total slides: {len(self.doc.slides)}")
        except Exception as e:
            logger.error(f"Failed to open PPTX: {e}")
            raise

    def convert_to_json(self,
                        extract_images: bool = True,
                        ocr_images: bool = False,
                        show_progress: bool = True) -> Dict[str, Any]:
        """Convert PPTX to JSON format"""
        document = {
            "filename": self.file_path.name,
            "pages": len(self.doc.slides),
            "content": []
        }
        # Shapes of layouts and masters whose text a slide already showed (see _inherited_text).
        self._inherited_seen = set()

        # Batch OCR processing if requested
        ocr_images = bool(extract_images and ocr_images and self.ocr)
        ocr_results_map = {}
        if ocr_images:
            if show_progress:
                logger.info("Collecting all images for batch OCR processing...")

            all_images_info = self._collect_all_images()

            if all_images_info:
                if show_progress:
                    logger.info(f"Processing {len(all_images_info)} images with batch OCR...")

                try:
                    # Use the configured OCR instance from BaseOfficeLoader
                    ocr_results_map = self._batch_ocr_images(all_images_info)

                    if show_progress:
                        logger.info(f"Successfully processed {len(ocr_results_map)} images with OCR")

                except Exception as e:
                    logger.error(f"Batch OCR processing failed: {e}")
                    ocr_images = False  # Fall back to base64 extraction

        # Process each slide
        for slide_idx, slide in enumerate(self.doc.slides):
            if show_progress:
                logger.info(f"Processing slide {slide_idx + 1}/{len(self.doc.slides)}")

            slide_content = self._process_slide(slide, slide_idx + 1, extract_images, ocr_images, ocr_results_map)
            document["content"].extend(slide_content)

            # Extract notes from the slide if present
            notes_content = self._extract_slide_notes(slide, slide_idx + 1)
            if notes_content:
                document["content"].append(notes_content)

        return document

    def _collect_all_images(self) -> List[Dict[str, Any]]:
        """Collect all images from PPTX presentation for batch processing"""
        images_info = []
        image_counter = 0

        for slide_idx, slide in enumerate(self.doc.slides):
            background = self._slide_background_image(slide)
            if background is not None:
                image_counter += 1
                images_info.append({
                    'id': f'pptx_slide{slide_idx + 1}_background_{image_counter}',
                    'data': background,
                    'location': {'slide': slide_idx + 1, 'type': 'background'}
                })

            # Check placeholders
            for placeholder in slide.placeholders:
                if hasattr(placeholder, 'image'):
                    try:
                        image_data = placeholder.image.blob
                        image_counter += 1
                        images_info.append({
                            'id': f'pptx_slide{slide_idx + 1}_placeholder_{image_counter}',
                            'data': image_data,
                            'location': {'slide': slide_idx + 1, 'type': 'placeholder'}
                        })
                    except (AttributeError, ValueError) as e:
                        logger.debug(f"Failed to extract placeholder image: {e}")

            # Check all shapes
            for shape_idx, shape in enumerate(slide.shapes):
                if hasattr(shape, 'shape_type') and shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        image_data = shape.image.blob
                        image_counter += 1
                        images_info.append({
                            'id': f'pptx_slide{slide_idx + 1}_shape{shape_idx}_{image_counter}',
                            'data': image_data,
                            'location': {'slide': slide_idx + 1, 'type': 'shape', 'shape_idx': shape_idx}
                        })
                    except (AttributeError, ValueError) as e:
                        logger.debug(f"Failed to extract shape image: {e}")

                # Check grouped shapes
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    group_images = self._collect_images_from_group(shape, slide_idx + 1, image_counter)
                    images_info.extend(group_images)
                    image_counter += len(group_images)

        logger.info(f"Collected {len(images_info)} images from PPTX")
        return images_info

    def _collect_images_from_group(self, group_shape, slide_num: int, counter_start: int) -> List[Dict[str, Any]]:
        """Recursively collect images from grouped shapes"""
        images = []
        counter = counter_start

        try:
            for shape in group_shape.shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    # Recursive call for nested groups
                    nested_images = self._collect_images_from_group(shape, slide_num, counter)
                    images.extend(nested_images)
                    counter += len(nested_images)
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        image_data = shape.image.blob
                        counter += 1
                        images.append({
                            'id': f'pptx_slide{slide_num}_group_{counter}',
                            'data': image_data,
                            'location': {'slide': slide_num, 'type': 'grouped'}
                        })
                    except (AttributeError, ValueError) as e:
                        logger.debug(f"Failed to extract grouped image: {e}")
        except (AttributeError, TypeError) as e:
            logger.debug(f"Failed to iterate group shapes: {e}")

        return images

    def _process_slide(self, slide, slide_num: int, extract_images: bool, ocr_images: bool,
                       ocr_results_map: Optional[Dict[str, str]] = None) -> List[Dict[str, Any]]:
        """Process a single slide"""
        if ocr_results_map is None:
            ocr_results_map = {}
        content_items = []

        logger.info(f"Processing slide {slide_num}...")

        # A picture set as this slide's own background (Format Background > Picture fill)
        # is slide content; layout/master backgrounds are shared decoration and skipped.
        if extract_images:
            background = self._slide_background_image(slide)
            if background is not None:
                item = self._image_item(background, slide_num, ocr_images, ocr_results_map)
                if item:
                    item["_top"] = -1
                    item["_left"] = -1
                    content_items.append(item)

        # First extract from placeholders (most structured content)
        placeholder_content = self._extract_from_placeholders(slide, slide_num, extract_images, ocr_images,
                                                              ocr_results_map)
        content_items.extend(placeholder_content)
        logger.info(f"  Extracted {len(placeholder_content)} items from placeholders")

        # Then extract from regular shapes (including grouped shapes)
        shape_content = self._extract_from_shapes(slide, slide_num, extract_images, ocr_images, ocr_results_map)
        content_items.extend(shape_content)
        logger.info(f"  Extracted {len(shape_content)} items from shapes")

        # Text the slide's layout and master draw on it (their placeholders are prompts, shown only
        # through the slide's own placeholders)
        try:
            content_items.extend(self._inherited_text(slide, slide_num))
        except Exception as e:
            logger.warning(f"Error extracting layout text: {e}")

        # Sort by position to maintain reading order
        content_items.sort(key=lambda x: (x.get("_top") or 0, x.get("_left") or 0))

        # Remove internal position markers
        for item in content_items:
            item.pop("_top", None)
            item.pop("_left", None)

        logger.info(f"  Total items extracted from slide {slide_num}: {len(content_items)}")
        return content_items

    def _inherited_text(self, slide, slide_num: int) -> List[Dict[str, Any]]:
        """Text the slide's layout and master draw on it: their visible non-placeholder shapes (a
        tagline, a company line), each kept once per presentation, on the first slide that shows
        it, as a PDF keeps the first copy of a running header. Their placeholders are prompts
        ("Click to edit Master title style") and date / footer / slide-number fields, which a slide
        shows only through its own placeholders."""
        items = []
        for owner in _pptx_drawn_masters(slide):
            for shape in _pptx_visible_shapes(owner.shapes):
                if shape.is_placeholder or not getattr(shape, 'has_text_frame', False):
                    continue
                key = (str(owner.part.partname), shape.shape_id)
                if key in self._inherited_seen:
                    continue
                self._inherited_seen.add(key)
                # Fields (a slide number, a date) show the layout's cached text, not the slide's.
                lines = [_soft_breaks(_pptx_typed_text(paragraph)).strip() for paragraph in shape.text_frame.paragraphs]
                text = "\n".join(line for line in lines if line)
                if text:
                    items.append({"type": "text:normal", "content": text, "page": slide_num,
                                  "_top": shape.top or 0, "_left": shape.left or 0})
        return items

    def _extract_from_placeholders(self, slide, slide_num: int, extract_images: bool, ocr_images: bool,
                                   ocr_results_map: Optional[Dict[str, str]] = None) -> List[
        Dict[str, Any]]:
        """Extract content from slide placeholders"""
        if ocr_results_map is None:
            ocr_results_map = {}
        content_items = []

        try:
            # Import placeholder type enum
            from pptx.enum.shapes import PP_PLACEHOLDER

            placeholder_count = len(slide.placeholders) if hasattr(slide, 'placeholders') else 0
            logger.debug(f"Slide {slide_num}: Found {placeholder_count} placeholders")

            for idx, placeholder in enumerate(slide.placeholders):
                try:
                    # Get placeholder type for better classification
                    ph_type = placeholder.placeholder_format.type
                    logger.debug(f"  Placeholder {idx}: type={ph_type}")

                    # Handle text placeholders
                    if placeholder.has_text_frame:
                        text_content = self._extract_text_from_placeholder(placeholder, ph_type)
                        if text_content:
                            text_content["page"] = slide_num
                            text_content["_top"] = placeholder.top if hasattr(placeholder, 'top') else 0
                            text_content["_left"] = placeholder.left if hasattr(placeholder, 'left') else 0
                            content_items.append(text_content)
                            logger.debug(
                                f"    Extracted text: {text_content['type']} - {len(text_content['content'])} chars")

                    # Handle table placeholders
                    if hasattr(placeholder, 'has_table') and placeholder.has_table:
                        table_content = self._extract_table_from_shape(placeholder, slide_num)
                        if table_content:
                            table_content["_top"] = placeholder.top if hasattr(placeholder, 'top') else 0
                            table_content["_left"] = placeholder.left if hasattr(placeholder, 'left') else 0
                            content_items.append(table_content)
                            logger.debug(f"    Extracted table")

                    # Pictures in placeholders: a picture placeholder, or a picture inserted
                    # into a content (object/body) placeholder, which becomes a p:pic too
                    if extract_images and (ph_type == PP_PLACEHOLDER.PICTURE or hasattr(placeholder, 'image')):
                        if hasattr(placeholder, 'image'):
                            image_content = self._extract_image_from_placeholder(placeholder, slide_num, ocr_images,
                                                                                 ocr_results_map)
                            if image_content:
                                image_content["_top"] = placeholder.top if hasattr(placeholder, 'top') else 0
                                image_content["_left"] = placeholder.left if hasattr(placeholder, 'left') else 0
                                content_items.append(image_content)
                                logger.debug(f"    Extracted image")

                except Exception as e:
                    logger.warning(f"Failed to process placeholder {idx}: {e}")

        except ImportError as e:
            logger.warning(f"PP_PLACEHOLDER enum not available: {e}, using basic extraction")
            # Fallback: extract from all placeholders without type information
            try:
                for idx, placeholder in enumerate(slide.placeholders):
                    if placeholder.has_text_frame:
                        text = _soft_breaks(placeholder.text_frame.text).strip()
                        if text:
                            content_items.append({
                                "type": self._classify_text_type(text, ""),
                                "content": text,
                                "page": slide_num,
                                "_top": placeholder.top if hasattr(placeholder, 'top') else 0,
                                "_left": placeholder.left if hasattr(placeholder, 'left') else 0
                            })
            except Exception as e:
                logger.error(f"Fallback placeholder extraction failed: {e}")

        logger.debug(f"Slide {slide_num}: Extracted {len(content_items)} items from placeholders")
        return content_items

    def _extract_text_from_placeholder(self, placeholder, ph_type) -> Optional[Dict[str, Any]]:
        """Extract text from a placeholder with proper type classification. A date, footer, header
        or slide-number placeholder gives only the text the slide sets in it: its fields (the slide
        number, an automatic date) are not text of the slide."""
        text_parts = []
        try:
            from pptx.enum.shapes import PP_PLACEHOLDER
            marginal = ph_type in (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.HEADER,
                                   PP_PLACEHOLDER.SLIDE_NUMBER)
        except ImportError:
            marginal = False

        for paragraph in placeholder.text_frame.paragraphs:
            # Don't strip yet - check raw text first
            para_text = _soft_breaks(_pptx_typed_text(paragraph) if marginal else paragraph.text)
            # Only skip if truly empty or just whitespace
            if para_text and not para_text.isspace():
                # Now we can strip for storage
                text_parts.append(para_text.strip())

        if text_parts:
            full_text = "\n".join(text_parts)

            # Don't skip short text
            if not full_text:
                return None

            # Import placeholder types
            try:
                from pptx.enum.shapes import PP_PLACEHOLDER

                # Map placeholder types to our text types
                if ph_type in [PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE]:
                    text_type = "text:title"
                elif ph_type == PP_PLACEHOLDER.SUBTITLE:
                    text_type = "text:section"
                elif ph_type in [PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT,
                                 PP_PLACEHOLDER.VERTICAL_BODY, PP_PLACEHOLDER.VERTICAL_OBJECT]:
                    # Further classify body text
                    text_type = self._classify_text_type(full_text, "")
                elif ph_type in [PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER,
                                 PP_PLACEHOLDER.HEADER, PP_PLACEHOLDER.SLIDE_NUMBER]:
                    text_type = "text:caption"
                else:
                    text_type = "text:normal"

                # Log very short text from placeholders
                if len(full_text) <= 5:
                    logger.debug(f"    Found short text in placeholder: '{full_text}' (type: {ph_type})")

            except (AttributeError, IndexError) as e:
                logger.debug(f"Failed to classify placeholder: {e}")
                # Fallback classification
                text_type = self._classify_text_type(full_text, "")

            return {
                "type": text_type,
                "content": full_text
            }

        return None

    def _extract_from_shapes(self, slide, slide_num: int, extract_images: bool, ocr_images: bool,
                             ocr_results_map: Optional[Dict[str, str]] = None) -> List[
        Dict[str, Any]]:
        """Extract content from regular shapes (non-placeholders)"""
        if ocr_results_map is None:
            ocr_results_map = {}
        content_items = []

        shape_count = len(slide.shapes) if hasattr(slide, 'shapes') else 0
        logger.debug(f"Slide {slide_num}: Processing {shape_count} shapes")

        for idx, shape in enumerate(slide.shapes):
            try:
                # Skip placeholders as they're already processed
                if hasattr(shape, 'is_placeholder') and shape.is_placeholder:
                    logger.debug(f"  Shape {idx}: Skipping (is placeholder)")
                    continue

                # Process grouped shapes recursively
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    logger.debug(f"  Shape {idx}: Processing group shape")
                    group_content = self._extract_from_group_shape(shape, slide_num, extract_images, ocr_images,
                                                                   ocr_results_map)
                    content_items.extend(group_content)

                # Tables - extract full table structure (skip text extraction to avoid duplicates)
                if shape.has_table:
                    table_content = self._extract_table_from_shape(shape, slide_num)
                    if table_content:
                        table_content["_top"] = shape.top if hasattr(shape, 'top') else 0
                        table_content["_left"] = shape.left if hasattr(shape, 'left') else 0
                        content_items.append(table_content)
                        logger.debug(f"  Shape {idx}: Extracted table")
                else:
                    # Extract ALL text from non-table shapes using comprehensive method
                    all_text_items = self._extract_all_text_from_shape(shape, slide_num)
                    for text_item in all_text_items:
                        text_item["page"] = slide_num
                        text_item["_top"] = shape.top if hasattr(shape, 'top') else 0
                        text_item["_left"] = shape.left if hasattr(shape, 'left') else 0
                        content_items.append(text_item)

                # Pictures
                if extract_images and hasattr(shape, 'shape_type') and shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    image_content = self._extract_image_from_shape(shape, slide_num, ocr_images, ocr_results_map)
                    if image_content:
                        image_content["_top"] = shape.top if hasattr(shape, 'top') else 0
                        image_content["_left"] = shape.left if hasattr(shape, 'left') else 0
                        content_items.append(image_content)
                        logger.debug(f"  Shape {idx}: Extracted image")

                # Log unknown shape types for debugging
                if hasattr(shape, 'shape_type'):
                    shape_type_name = shape.shape_type
                    if shape_type_name not in [MSO_SHAPE_TYPE.GROUP, MSO_SHAPE_TYPE.PICTURE,
                                               MSO_SHAPE_TYPE.PLACEHOLDER, MSO_SHAPE_TYPE.TEXT_BOX]:
                        logger.debug(
                            f"  Shape {idx}: Type={shape_type_name}, Name={shape.name if hasattr(shape, 'name') else 'unnamed'}")

                        # Try to extract text from any shape as a last resort (a shape with a text
                        # frame, such as an AutoShape, was read above)
                        if not getattr(shape, 'has_text_frame', False) and getattr(shape, 'text', '').strip():
                            content_items.append({
                                "type": "text:normal",
                                "content": _soft_breaks(shape.text).strip(),
                                "page": slide_num,
                                "_top": shape.top if hasattr(shape, 'top') else 0,
                                "_left": shape.left if hasattr(shape, 'left') else 0
                            })
                            logger.debug(f"    Extracted text from unknown shape type: {shape.text[:50]}...")

            except Exception as e:
                logger.warning(f"Failed to process shape {idx}: {e}")

        logger.debug(f"Slide {slide_num}: Extracted {len(content_items)} items from shapes")
        return content_items

    def _extract_from_group_shape(self, group_shape, slide_num: int, extract_images: bool, ocr_images: bool,
                                  ocr_results_map: Optional[Dict[str, str]] = None) -> List[
        Dict[str, Any]]:
        """Recursively extract content from grouped shapes"""
        if ocr_results_map is None:
            ocr_results_map = {}
        content_items = []

        try:
            for shape in group_shape.shapes:
                # Recursively handle nested groups
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    nested_content = self._extract_from_group_shape(shape, slide_num, extract_images, ocr_images,
                                                                    ocr_results_map)
                    content_items.extend(nested_content)

                # Extract text
                elif shape.has_text_frame:
                    text_content = self._extract_text_from_shape(shape)
                    if text_content:
                        text_content["page"] = slide_num
                        text_content["_top"] = shape.top if hasattr(shape, 'top') else 0
                        text_content["_left"] = shape.left if hasattr(shape, 'left') else 0
                        content_items.append(text_content)

                # Extract tables
                elif shape.has_table:
                    table_content = self._extract_table_from_shape(shape, slide_num)
                    if table_content:
                        table_content["_top"] = shape.top if hasattr(shape, 'top') else 0
                        table_content["_left"] = shape.left if hasattr(shape, 'left') else 0
                        content_items.append(table_content)

                # Extract images
                elif extract_images and shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    image_content = self._extract_image_from_shape(shape, slide_num, ocr_images, ocr_results_map)
                    if image_content:
                        image_content["_top"] = shape.top if hasattr(shape, 'top') else 0
                        image_content["_left"] = shape.left if hasattr(shape, 'left') else 0
                        content_items.append(image_content)

        except Exception as e:
            logger.warning(f"Failed to process group shape: {e}")

        return content_items

    def _extract_text_from_shape(self, shape) -> Optional[Dict[str, Any]]:
        """Extract text from a shape"""
        text_parts = []

        for paragraph in shape.text_frame.paragraphs:
            # Don't strip yet - check raw text first
            para_text = _soft_breaks(paragraph.text)
            # Only skip if truly empty or just whitespace
            if para_text and not para_text.isspace():
                # Now we can strip for storage
                text_parts.append(para_text.strip())

        if text_parts:
            full_text = "\n".join(text_parts)

            # Don't skip short text - it might be important (like page numbers, labels, etc.)
            if not full_text:
                return None

            # Try to get more context for classification
            shape_name = shape.name.lower() if hasattr(shape, 'name') else ""

            # Check if it's likely a title based on shape name
            if 'title' in shape_name:
                text_type = "text:title"
            elif 'subtitle' in shape_name:
                text_type = "text:section"
            else:
                text_type = self._classify_text_type(full_text, "")

            # Log very short text for debugging
            if len(full_text) <= 5:
                logger.debug(f"    Found short text: '{full_text}' in shape: {shape_name}")

            return {
                "type": text_type,
                "content": full_text
            }

        return None

    def _extract_slide_notes(self, slide, slide_num: int) -> Optional[Dict[str, Any]]:
        """Extract notes from a slide"""
        try:
            if slide.has_notes_slide:
                notes_slide = slide.notes_slide
                notes_text = _soft_breaks(notes_slide.notes_text_frame.text).strip()

                if notes_text:
                    return {
                        "type": "text:normal",
                        "content": f"[Slide {slide_num} Notes]\n{notes_text}",
                        "page": slide_num
                    }
        except Exception as e:
            logger.warning(f"Failed to extract slide notes: {e}")

        return None

    def _extract_table_from_shape(self, shape, slide_num: int) -> Optional[Dict[str, Any]]:
        """Extract table from a shape with enhanced merged cell detection for PPTX"""
        table = shape.table

        # Get table dimensions
        num_rows = len(table.rows)
        num_cols = len(table.columns)

        logger.debug(f"Extracting table from slide {slide_num}: {num_rows}x{num_cols}")

        # Initialize table data with proper dimensions
        table_data = [['' for _ in range(num_cols)] for _ in range(num_rows)]
        merged_cells_info = []

        # Track which cells we've already processed as part of merges
        processed_cells = set()

        # Process each cell to detect merges using gridSpan and vMerge properties
        for row_idx in range(num_rows):
            for col_idx in range(num_cols):
                if (row_idx, col_idx) in processed_cells:
                    continue

                try:
                    # Get the cell at this position
                    cell = table.cell(row_idx, col_idx)
                    tc = cell._tc  # Get the table cell element

                    # Get span properties
                    gridSpan = tc.gridSpan  # Horizontal span (>1 means merged horizontally)
                    vMerge = tc.vMerge  # Vertical merge (True for continuation cells)

                    # Get cell text
                    cell_text = _soft_breaks(cell.text).strip() if hasattr(cell, 'text') else ''

                    # Check if this is a vertically merged cell continuation
                    if vMerge:
                        # This cell is part of a vertical merge from above
                        # Find the origin cell by looking upward
                        origin_row = row_idx - 1
                        while origin_row >= 0:
                            origin_cell = table.cell(origin_row, col_idx)
                            if not origin_cell._tc.vMerge:
                                # Found the origin of the vertical merge
                                break
                            origin_row -= 1

                        # Mark this cell as processed (it's empty in a vertical merge)
                        processed_cells.add((row_idx, col_idx))
                        logger.debug(f"Cell({row_idx},{col_idx}) is part of vertical merge from row {origin_row}")
                        continue

                    # Determine the span of this cell
                    colspan = gridSpan if gridSpan > 1 else 1
                    rowspan = 1

                    # If this cell has text and cells below are empty with vMerge=True, it's a vertical merge origin
                    if row_idx < num_rows - 1:
                        # Check cells below for vertical merge
                        check_row = row_idx + 1
                        while check_row < num_rows:
                            check_cell = table.cell(check_row, col_idx)
                            if check_cell._tc.vMerge:
                                rowspan += 1
                                check_row += 1
                            else:
                                break

                    # Place text in the current cell
                    table_data[row_idx][col_idx] = cell_text

                    # Mark all cells covered by this merge as processed
                    for r in range(row_idx, row_idx + rowspan):
                        for c in range(col_idx, col_idx + colspan):
                            processed_cells.add((r, c))
                            # Clear cells that are part of the merge (except origin)
                            if (r, c) != (row_idx, col_idx):
                                if r < num_rows and c < num_cols:
                                    table_data[r][c] = ''

                    # Record merge info if this is a merged cell
                    if rowspan > 1 or colspan > 1:
                        merged_cells_info.append({
                            'row': row_idx,
                            'col': col_idx,
                            'rowspan': rowspan,
                            'colspan': colspan
                        })
                        logger.debug(
                            f"Found merged cell at ({row_idx}, {col_idx}) with span {rowspan}x{colspan}, text: '{cell_text[:30]}...'")

                except Exception as e:
                    logger.warning(f"Error processing cell at ({row_idx}, {col_idx}): {e}")
                    table_data[row_idx][col_idx] = ''
                    processed_cells.add((row_idx, col_idx))

        # Log the extraction result
        logger.info(f"PPTX Table: Extracted {num_rows}x{num_cols} table with {len(merged_cells_info)} merged cells")

        # Debug: Log first few rows of extracted data
        if logger.isEnabledFor(logging.DEBUG):
            for i, row in enumerate(table_data[:3]):
                logger.debug(f"  Row {i}: {[cell[:20] + '...' if cell and len(cell) > 20 else cell for cell in row]}")

        # For PowerPoint tables, we'll always use HTML format for better structure preservation
        if table_data:
            # Create table info for HTML conversion
            # Build cell_spans and render via TableData
            cell_spans = {}
            for merge_info in merged_cells_info:
                key = (merge_info['row'], merge_info['col'])
                cell_spans[key] = (merge_info['rowspan'], merge_info['colspan'])

            table_obj = TableData.from_raw(table_data, {
                'is_complex': True,  # Always treat PowerPoint tables as complex
                'cell_spans': cell_spans,
            })
            renderer = TableRenderer(self.table_style)
            html_table = renderer._render_html(table_obj)

            return {
                "type": "table",
                "content": html_table,
                "page": slide_num
            }

        return None

    @staticmethod
    def _slide_background_image(slide) -> Optional[bytes]:
        """Bytes of the picture fill of the slide's own background (``p:bg/p:bgPr/a:blipFill``), or None."""
        p_ns = 'http://schemas.openxmlformats.org/presentationml/2006/main'
        blip = slide._element.find(f'{{{p_ns}}}cSld/{{{p_ns}}}bg/{{{p_ns}}}bgPr/'
                                   '{http://schemas.openxmlformats.org/drawingml/2006/main}blipFill/'
                                   '{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
        rid = blip.get(qn('r:embed')) if blip is not None else None
        if not rid:
            return None
        try:
            return slide.part.related_part(rid).blob
        except (KeyError, AttributeError) as e:
            logger.debug(f"Slide background picture unavailable: {e}")
            return None

    def _image_item(self, image_data: bytes, slide_num: int, ocr_images: bool,
                    ocr_results_map: Optional[Dict[Any, Optional[str]]] = None) -> Optional[Dict[str, Any]]:
        """Content item for one picture: its OCR text when OCR runs (None when OCR finds no
        text; the ``[image: OCR unavailable]`` marker when its OCR failed), else the image as
        base64."""
        if ocr_images and self.ocr:
            img_hash = _image_hash(image_data)
            if img_hash in (ocr_results_map or {}):
                ocr_text = ocr_results_map[img_hash]
            else:
                logger.warning(f"OCR result not found for a picture on slide {slide_num}, using fallback OCR")
                ocr_text = self._ocr_image(image_data, {'slide': slide_num})
            return self._ocr_description(ocr_text, page=slide_num)
        return {"type": "image", "content": self._extract_image_as_base64(image_data), "page": slide_num}

    def _extract_image_from_placeholder(self, placeholder, slide_num: int, ocr_images: bool,
                                        ocr_results_map: Optional[Dict[Any, Optional[str]]] = None) -> Optional[
        Dict[str, Any]]:
        """Extract image from a picture placeholder"""
        try:
            return self._image_item(placeholder.image.blob, slide_num, ocr_images, ocr_results_map)
        except Exception as e:
            logger.warning(f"Failed to extract image from placeholder: {e}")
            return None

    def _extract_image_from_shape(self, shape, slide_num: int, ocr_images: bool,
                                  ocr_results_map: Optional[Dict[Any, Optional[str]]] = None) -> Optional[Dict[str, Any]]:
        """Extract image from a shape"""
        try:
            return self._image_item(shape.image.blob, slide_num, ocr_images, ocr_results_map)
        except Exception as e:
            logger.warning(f"Failed to extract image: {e}")
            return None

    def _extract_all_text_from_shape(self, shape, slide_num: int) -> List[Dict[str, Any]]:
        """Extract all possible text from any shape type"""
        text_items = []

        try:
            # 1. Regular text frame
            if hasattr(shape, 'has_text_frame') and shape.has_text_frame:
                text_content = self._extract_text_from_shape(shape)
                if text_content:
                    text_items.append(text_content)

            # 2. Chart text elements (removed table cell extraction - tables handled separately)
            if hasattr(shape, 'has_chart') and shape.has_chart:
                chart = shape.chart

                # Chart title
                if hasattr(chart, 'has_title') and chart.has_title:
                    try:
                        title_text = _soft_breaks(chart.chart_title.text_frame.text).strip()
                        if title_text:
                            text_items.append({
                                "type": "text:caption",
                                "content": f"Chart: {title_text}"
                            })
                    except (AttributeError, ValueError) as e:
                        logger.debug(f"Failed to extract chart title: {e}")

                # Axis titles
                try:
                    if hasattr(chart, 'category_axis') and chart.category_axis.has_title:
                        axis_text = _soft_breaks(chart.category_axis.axis_title.text_frame.text).strip()
                        if axis_text:
                            text_items.append({
                                "type": "text:caption",
                                "content": f"X-axis: {axis_text}"
                            })
                except (AttributeError, ValueError) as e:
                    logger.debug(f"Failed to extract x-axis title: {e}")

                try:
                    if hasattr(chart, 'value_axis') and chart.value_axis.has_title:
                        axis_text = _soft_breaks(chart.value_axis.axis_title.text_frame.text).strip()
                        if axis_text:
                            text_items.append({
                                "type": "text:caption",
                                "content": f"Y-axis: {axis_text}"
                            })
                except (AttributeError, ValueError) as e:
                    logger.debug(f"Failed to extract y-axis title: {e}")

            # 3. SmartArt text (often in grouped shapes)
            # SmartArt is typically a group shape with specific properties
            if hasattr(shape, 'shape_type') and shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                # Check if it might be SmartArt by looking at shape name
                shape_name = shape.name.lower() if hasattr(shape, 'name') else ""
                if 'diagram' in shape_name or 'smart' in shape_name:
                    logger.debug(f"  Possible SmartArt detected: {shape.name}")

            # 4. Shape alt text (often contains descriptions)
            if hasattr(shape, 'alt_text') and shape.alt_text:
                text_items.append({
                    "type": "text:caption",
                    "content": f"[Alt text]: {shape.alt_text}"
                })

            # 5. Connector text (lines with labels)
            if hasattr(shape, 'connector_type'):
                if hasattr(shape, 'text_frame') and shape.text_frame:
                    connector_text = _soft_breaks(shape.text_frame.text).strip()
                    if connector_text:
                        text_items.append({
                            "type": "text:caption",
                            "content": connector_text
                        })

            # 6. OLE objects might have accessible text
            if hasattr(shape, 'ole_format'):
                logger.debug(f"  Found OLE object: {shape.name if hasattr(shape, 'name') else 'unnamed'}")
                # OLE objects are embedded files, limited text extraction

        except Exception as e:
            logger.warning(f"Error extracting text from shape: {e}")

        return text_items

class XlsxLoader(BaseOfficeLoader):
    """Loader for XLSX (Excel) documents

    Every sheet is read cell by cell with openpyxl (cached values). Values are
    rendered the way the spreadsheet displays them (number formats), merged cells
    come only from ``merged_cells.ranges`` (never guessed from blank cells), fully
    empty rows and columns are dropped, leading single-cell title rows become text
    above the table, a formula saved without a cached value shows its formula text,
    and pictures anchored inside the table (or placed in a cell) are marked in their
    cell (``[Image]``, or ``[Image: <OCR text>]`` when OCR runs).
    """

    def __init__(self, file_path: Union[str, Path], ocr=None, table_style: Union[str, TableStyle] = None):
        super().__init__(file_path, ocr, table_style)
        self._open_document()

    def _open_document(self):
        """Open XLSX document"""
        try:
            self.doc = openpyxl.load_workbook(self.file_path, data_only=True)
            logger.info(f"Opened XLSX: {self.file_path.name}")
            logger.info(f"Total sheets: {len(self.doc.worksheets)}")
        except Exception as e:
            logger.error(f"Failed to open XLSX: {e}")
            raise
        self._scans: Optional[Dict[str, Dict[str, Dict[Tuple[int, int], str]]]] = None
        self._media_cache: Dict[str, bytes] = {}
        self._fallback_media: Dict[int, str] = {}

    def convert_to_json(self,
                        extract_images: bool = True,
                        ocr_images: bool = False,
                        show_progress: bool = True) -> Dict[str, Any]:
        """Convert XLSX to JSON format. Besides ``content``, the result carries ``sheet_names``
        (every sheet, in workbook order) and ``total_cells`` (the cells that show a value, over
        all sheets)."""
        document = {
            "filename": self.file_path.name,
            "pages": len(self.doc.worksheets),
            "content": [],
            "sheet_names": list(self.doc.sheetnames),
            "total_cells": 0,
        }
        ocr_images = bool(ocr_images and self.ocr)

        # Image data is read once (openpyxl's Image._data() can only be read once) and
        # OCR'd in one batch before the sheets are laid out, so a picture's OCR text can
        # go into the cell it is anchored to.
        ocr_results_map = {}
        image_data_cache = {}
        if extract_images:
            if show_progress:
                logger.info("Collecting all images...")
            all_images_info = self._collect_all_images()
            for info in all_images_info:
                location = info['location']
                image_data_cache[(location['sheet_name'], location['img_idx'])] = info['data']
                if location.get('source') == 'zip_fallback':
                    image_data_cache[('_fallback', location['img_idx'])] = info['data']
                    self._fallback_media[location['img_idx']] = info.get('media_file', '')
            if ocr_images and all_images_info:
                if show_progress:
                    logger.info(f"Processing {len(all_images_info)} images with batch OCR...")
                try:
                    ocr_results_map = self._batch_ocr_images(all_images_info)
                except Exception as e:
                    logger.error(f"Batch OCR processing failed: {e}")
                    ocr_images = False  # Fall back to base64 extraction

        for sheet_idx, sheet_name in enumerate(self.doc.sheetnames):
            if show_progress:
                logger.info(f"Processing sheet {sheet_idx + 1}/{len(self.doc.sheetnames)}: {sheet_name}")
            sheet = self.doc[sheet_name]
            document["content"].append({
                "type": "text:title",
                "content": f"Sheet: {sheet_name}",
                "page": sheet_idx + 1
            })
            items, embedded = self._extract_sheet_content(
                sheet, sheet_idx + 1, extract_images, ocr_images, ocr_results_map, image_data_cache)
            document["content"].extend(items)
            document["total_cells"] += self._filled_cells
            if extract_images:
                document["content"].extend(self._extract_images_from_sheet(
                    sheet, sheet_idx + 1, ocr_images, ocr_results_map, image_data_cache, embedded))

        return document

    # ------------------------------------------------------------------ #
    # Sheet layout                                                         #
    # ------------------------------------------------------------------ #
    def _extract_sheet_content(self, sheet, sheet_num: int, extract_images: bool, ocr_images: bool,
                               ocr_results_map: Dict[Any, str], image_data_cache: Dict[tuple, bytes]
                               ) -> Tuple[List[Dict[str, Any]], set]:
        """Content items of one sheet (title rows, then the table) and the pictures
        embedded in its cells (``("anchor", index)`` / ``("media", part name)``). Sets
        ``self._filled_cells``: how many of its cells show a value."""
        embedded: set = set()
        self._filled_cells = 0
        if not hasattr(sheet, 'merged_cells'):  # a chartsheet has no cells
            return [], embedded
        scan = self._sheet_scans().get(sheet.title, {})
        formulas, pictures = scan.get('formulas', {}), scan.get('pictures', {})

        texts: Dict[Tuple[int, int], str] = {}
        for (row, col), cell in self._sheet_cells(sheet):
            if (row, col) in pictures and cell.data_type == 'e':
                continue  # a picture placed in the cell; Excel stores #VALUE! as its fallback
            text = self._cell_text(cell) if cell.value is not None else formulas.get((row, col), "")
            if text:
                texts[(row, col)] = text
        self._filled_cells = len(texts)

        merges = [(r.min_row, r.min_col, r.max_row, r.max_col) for r in sheet.merged_cells.ranges
                  if texts.get((r.min_row, r.min_col))]
        # Decided on the text alone: a logo beside a title must not make the title a table row.
        table_start = self._table_start(texts, merges)

        cells = dict(texts)
        if extract_images and (texts or pictures):
            marks = self._picture_marks(sheet, texts, merges, pictures, ocr_images, ocr_results_map,
                                        image_data_cache, embedded)
            for position, labels in marks.items():
                cells[position] = " ".join(([cells[position]] if cells.get(position) else []) + labels)
        if not cells:
            return [], embedded
        if table_start is None:
            table_start = min(row for row, _ in cells)

        by_row: Dict[int, Dict[int, str]] = {}
        for (row, col), text in cells.items():
            by_row.setdefault(row, {})[col] = text
        # Rows above the table (titles, and pictures placed between them and the table) are
        # text, one paragraph per row with its cells in column order.
        items = [{"type": "text:normal", "content": " ".join(text for _, text in sorted(by_row[row].items())),
                  "page": sheet_num}
                 for row in sorted(row for row in by_row if row < table_start)]
        table_rows = sorted(row for row in by_row if row >= table_start)
        cols = sorted({col for row in table_rows for col in by_row[row]})
        row_pos = {row: i for i, row in enumerate(table_rows)}
        col_pos = {col: j for j, col in enumerate(cols)}
        grid = [[by_row[row].get(col, "") for col in cols] for row in table_rows]
        # A merge spans the rows and columns it covers that are still in the table.
        spans = {}
        for r1, c1, r2, c2 in merges:
            if r1 not in row_pos or c1 not in col_pos:
                continue
            rowspan = bisect.bisect_right(table_rows, r2) - row_pos[r1]
            colspan = bisect.bisect_right(cols, c2) - col_pos[c1]
            if rowspan > 1 or colspan > 1:
                spans[(row_pos[r1], col_pos[c1])] = (rowspan, colspan)
        if grid:
            table_obj = TableData.from_raw(grid, {'is_complex': bool(spans), 'cell_spans': spans})
            items.append({
                "type": "table",
                "content": TableRenderer(self.table_style).render(table_obj),
                "page": sheet_num
            })
        return items, embedded

    @staticmethod
    def _table_start(texts: Dict[Tuple[int, int], str], merges) -> Optional[int]:
        """First row of the sheet's table; the text rows above it are titles.

        Only rows above the first row with two or more cells can be titles, and only on
        strong evidence: the row's single cell is merged across the table's width, or a
        blank row separates it from the table ("ACME Corp - Sales Report FY2025", blank row,
        header row). A lone cell right above the table ("Customer" over two-column data) may
        be a header and stays in the table. None when the sheet has no text.
        """
        rows = sorted({row for row, _ in texts})
        if not rows:
            return None
        cols_by_row: Dict[int, List[int]] = {}
        for row, col in texts:
            cols_by_row.setdefault(row, []).append(col)
        first_multi = next((row for row in rows if len(cols_by_row[row]) >= 2), None)
        if first_multi is None:
            return rows[0]
        table_cols = [col for row, col in texts if row >= first_multi]
        left, right = min(table_cols), max(table_cols)
        start = first_multi
        for row in reversed([row for row in rows if row < first_multi]):
            col = cols_by_row[row][0]
            across = any((r1, c1) == (row, col) and c1 <= left and c2 >= right and r2 < start
                         for r1, c1, r2, c2 in merges)
            if across or start - row > 1:
                return start
            start = row
        return start

    @staticmethod
    def _sheet_cells(sheet):
        """((row, col), cell) for every stored cell, in row-major order."""
        cells = getattr(sheet, '_cells', None)
        if isinstance(cells, dict):
            return sorted(cells.items())
        return (((cell.row, cell.column), cell) for row in sheet.iter_rows() for cell in row)

    @staticmethod
    def _cell_text(cell) -> str:
        """The cell as displayed: number formats applied, line breaks normalised,
        control characters removed."""
        try:
            text = format_cell_value(cell.value, cell.number_format)
        except Exception:
            text = str(cell.value)
        text = text.replace('\r\n', '\n').replace('\r', '\n').replace('\x0b', '\n').replace('\x0c', '\n')
        return _CONTROL_CHARS.sub('', text).strip()

    def _picture_marks(self, sheet, texts, merges, pictures, ocr_images: bool, ocr_results_map,
                       image_data_cache, embedded: set) -> Dict[Tuple[int, int], List[str]]:
        """``[Image]`` / ``[Image: <OCR text>]`` labels for pictures inside the sheet's
        table: anchored pictures whose top-left cell lies within the used range (a cell
        inside a merge maps to the merge's top-left cell) and pictures placed in cells."""
        marks: Dict[Tuple[int, int], List[str]] = {}
        if texts:
            row_lo, row_hi = min(r for r, _ in texts), max(r for r, _ in texts)
            col_lo, col_hi = min(c for _, c in texts), max(c for _, c in texts)
        else:
            row_lo = row_hi = col_lo = col_hi = None
        for index, image in enumerate(getattr(sheet, '_images', [])):
            position = self._anchor_cell(image)
            if position is None or row_lo is None:
                continue
            position = next(((r1, c1) for r1, c1, r2, c2 in merges
                             if r1 <= position[0] <= r2 and c1 <= position[1] <= c2), position)
            if not (row_lo <= position[0] <= row_hi and col_lo <= position[1] <= col_hi):
                continue
            label = self._picture_label(image_data_cache.get((sheet.title, index)), ocr_images, ocr_results_map,
                                        sheet.title)
            marks.setdefault(position, []).append(label)
            if label != "[Image]":
                embedded.add(('anchor', index))
        for position, media in pictures.items():
            label = self._picture_label(self._media_bytes(media), ocr_images, ocr_results_map, sheet.title)
            marks.setdefault(position, []).append(label)
            if label != "[Image]":
                embedded.add(('media', media))
        return marks

    def _picture_label(self, data: Optional[bytes], ocr_images: bool, ocr_results_map,
                       sheet_name: Optional[str] = None) -> str:
        """``[Image: <OCR text>]`` when OCR read the picture, ``[image: OCR unavailable]`` when its
        OCR failed, else ``[Image]``. Only a picture whose OCR result went into its cell counts as
        embedded; any other picture is still emitted after the table, so extracted image data is
        never dropped. ``sheet_name``: the picture's sheet, for the loader's OCR issue record."""
        if not (ocr_images and self.ocr and data):
            return "[Image]"
        img_hash = _image_hash(data)
        if img_hash in ocr_results_map:
            text = ocr_results_map[img_hash]
        else:  # not in the batch: OCR it once, and remember it for the image item
            text = ocr_results_map[img_hash] = self._ocr_image(data, {'sheet': sheet_name} if sheet_name else None)
        if text is None:
            return OCR_UNAVAILABLE
        # The cell's text is escaped by the table renderer: give it the plain OCR text.
        text = (self._ocr_cell_texts.get(img_hash, text) or "").strip()
        return f"[Image: {text}]" if text else "[Image]"

    @staticmethod
    def _anchor_cell(image) -> Optional[Tuple[int, int]]:
        """1-based (row, col) of a picture's top-left anchor cell."""
        anchor = getattr(image, 'anchor', None)
        try:
            if isinstance(anchor, str):
                from openpyxl.utils.cell import coordinate_to_tuple
                return coordinate_to_tuple(anchor)
            start = getattr(anchor, '_from', None)
            if start is not None:
                return int(start.row) + 1, int(start.col) + 1
        except (TypeError, ValueError, AttributeError):
            pass
        return None

    def _media_bytes(self, part_name: str) -> Optional[bytes]:
        if part_name not in self._media_cache:
            try:
                with zipfile.ZipFile(self.file_path) as archive:
                    self._media_cache[part_name] = archive.read(part_name)
            except (KeyError, zipfile.BadZipFile, OSError) as e:
                logger.debug(f"XLSX picture {part_name} unavailable: {e}")
                self._media_cache[part_name] = None
        return self._media_cache[part_name]

    # ------------------------------------------------------------------ #
    # What openpyxl does not expose: uncached formulas, in-cell pictures   #
    # ------------------------------------------------------------------ #
    def _sheet_scans(self) -> Dict[str, Dict[str, Dict[Tuple[int, int], str]]]:
        """Per sheet title: ``formulas`` {(row, col): "=formula"} for formulas saved
        without a cached value (openpyxl reads them as empty), and ``pictures``
        {(row, col): media part} for pictures placed in cells (Excel's rich values)."""
        if self._scans is None:
            try:
                self._scans = self._scan_package()
            except Exception as e:
                logger.warning(f"XLSX formula/picture scan failed ({e}); uncached formulas stay empty")
                self._scans = {}
        return self._scans

    def _scan_package(self) -> Dict[str, Dict[str, Dict[Tuple[int, int], str]]]:
        from lxml import etree
        from openpyxl.utils.cell import coordinate_to_tuple
        scans: Dict[str, Dict[str, Dict[Tuple[int, int], str]]] = {}
        unresolved: Dict[str, List[Tuple[int, int]]] = {}
        value_metadata: Dict[str, Dict[Tuple[int, int], int]] = {}
        with zipfile.ZipFile(self.file_path) as archive:
            names = set(archive.namelist())
            for title, part in _xlsx_sheet_parts(archive).items():
                if part not in names or not _zip_member_matches(archive, part, _FORMULA_OR_RICH_VALUE):
                    continue
                formulas: Dict[Tuple[int, int], str] = {}
                for _, element in etree.iterparse(archive.open(part), events=('end',), tag=_SML_C,
                                                  resolve_entities=False, no_network=True, huge_tree=False):
                    reference = element.get('r')
                    if reference:
                        position = coordinate_to_tuple(reference)
                        formula = element.find(_SML_F)
                        if formula is not None and _formula_uncached(element):
                            if formula.text:
                                formulas[position] = "=" + formula.text
                            else:
                                unresolved.setdefault(title, []).append(position)
                        if element.get('vm'):
                            value_metadata.setdefault(title, {})[position] = _attr_int(element, 'vm')
                    element.clear(keep_tail=True)
                scans[title] = {'formulas': formulas, 'pictures': {}}
            if value_metadata:
                images = _xlsx_rich_value_images(archive, names)
                for title, cells in value_metadata.items():
                    scans[title]['pictures'] = {position: images[vm] for position, vm in cells.items() if vm in images}
        if unresolved:
            # Shared formulas store their text once; openpyxl translates it for each cell.
            formula_book = openpyxl.load_workbook(self.file_path, data_only=False)
            for title, positions in unresolved.items():
                for row, col in positions:
                    value = formula_book[title].cell(row, col).value
                    value = getattr(value, 'text', value)
                    if isinstance(value, str) and value:
                        scans[title]['formulas'][(row, col)] = value if value.startswith('=') else '=' + value
        return scans

    def _extract_images_from_sheet(self, sheet, sheet_num: int, ocr_images: bool,
                                   ocr_results_map: Optional[Dict[str, str]] = None,
                                   image_data_cache: Optional[Dict[tuple, bytes]] = None,
                                   embedded: Optional[set] = None) -> List[Dict[str, Any]]:
        """Extract images from an Excel sheet, except those already shown in a table
        cell (``embedded``: ``("anchor", index)`` / ``("media", part name)``)."""
        if ocr_results_map is None:
            ocr_results_map = {}
        if image_data_cache is None:
            image_data_cache = {}
        embedded = embedded or set()
        images = []
        sheet_name = sheet.title

        # First try standard image extraction
        for img_idx, image in enumerate(getattr(sheet, '_images', [])):
            if ('anchor', img_idx) in embedded:
                continue
            try:
                # Try to get cached image data first to avoid double _data() call
                cache_key = (sheet_name, img_idx)  # Use sheet_name for consistency
                if cache_key in image_data_cache:
                    image_data = image_data_cache[cache_key]
                else:
                    # Fallback to extracting if not cached (shouldn't happen)
                    try:
                        image_data = image._data()
                        # Cache it now to avoid future calls
                        image_data_cache[cache_key] = image_data
                    except Exception as e:
                        logger.warning(f"Failed to extract image data: {e}")
                        continue

                if ocr_images:
                    # Use image content hash to find OCR result
                    img_hash = _image_hash(image_data)
                    if img_hash in ocr_results_map:
                        ocr_text = ocr_results_map[img_hash]
                    else:
                        # Fallback to individual OCR
                        logger.warning(f"OCR result not found for image on sheet {sheet_name}, using fallback OCR")
                        ocr_text = self._ocr_image(image_data, {'sheet': sheet_num, 'sheet_name': sheet_name})
                    item = self._ocr_description(ocr_text, page=sheet_num)
                    if item:
                        images.append(item)
                else:
                    # Return base64 encoded image
                    base64_data = self._extract_image_as_base64(image_data)
                    images.append({
                        "type": "image",
                        "content": base64_data,
                        "page": sheet_num
                    })

            except Exception as e:
                logger.warning(f"Failed to extract image: {e}")
        
        # If no images found via standard method and this is the first sheet, check for fallback images
        if len(images) == 0 and sheet_num == 1:
            # Check if we have fallback images in cache
            fallback_idx = 0
            while True:
                fallback_key = ('_fallback', fallback_idx)
                if fallback_key in image_data_cache and ('media', self._fallback_media.get(fallback_idx)) in embedded:
                    fallback_idx += 1  # a picture placed in a cell, already shown there
                    continue
                if fallback_key in image_data_cache:
                    image_data = image_data_cache[fallback_key]
                    
                    if ocr_images:
                        # First try the special fallback OCR key for this specific image, then
                        # the hash, then an individual OCR call
                        fallback_ocr_key = ('_fallback_ocr', fallback_idx)
                        img_hash = _image_hash(image_data)
                        if fallback_ocr_key in ocr_results_map:
                            ocr_text = ocr_results_map[fallback_ocr_key]
                            logger.info(f"Using OCR result for fallback image {fallback_idx}")
                        elif img_hash in ocr_results_map:
                            ocr_text = ocr_results_map[img_hash]
                        else:
                            logger.warning(f"OCR result not found for fallback image {fallback_idx}, using fallback OCR")
                            ocr_text = self._ocr_image(image_data, {'sheet': sheet_num, 'sheet_name': sheet_name})
                        item = self._ocr_description(ocr_text, page=sheet_num)
                        if item:
                            images.append(item)
                    else:
                        # Return base64 encoded image
                        base64_data = self._extract_image_as_base64(image_data)
                        images.append({
                            "type": "image",
                            "content": base64_data,
                            "page": sheet_num
                        })
                    
                    logger.info(f"Added fallback image {fallback_idx} to sheet {sheet_name}")
                    fallback_idx += 1
                else:
                    break  # No more fallback images

        return images

    def _collect_all_images(self) -> List[Dict[str, Any]]:
        """Collect all images from XLSX workbook for batch processing"""
        images_info = []
        image_counter = 0

        # First try standard openpyxl method
        for sheet_idx, sheet_name in enumerate(self.doc.sheetnames):
            sheet = self.doc[sheet_name]

            # Extract images from sheet using standard method
            for img_idx, image in enumerate(getattr(sheet, '_images', [])):
                try:
                    image_data = image._data()
                    image_counter += 1
                    images_info.append({
                        'id': f'xlsx_sheet{sheet_idx + 1}_{sheet_name}_img{img_idx}',
                        'data': image_data,
                        'location': {'sheet': sheet_idx + 1, 'sheet_name': sheet_name, 'img_idx': img_idx}
                    })
                except Exception as e:
                    logger.warning(f"Failed to extract image from sheet {sheet_name}: {e}")

        # If no images found via standard method, try fallback ZIP extraction
        if len(images_info) == 0:
            logger.info("No images found via standard method, trying ZIP extraction fallback...")
            images_info = self._extract_images_from_zip()

        logger.info(f"Collected {len(images_info)} images from XLSX")
        return images_info
    
    def _extract_images_from_zip(self) -> List[Dict[str, Any]]:
        """Fallback method to extract images directly from XLSX ZIP structure
        
        This handles edge cases where images are embedded in non-standard ways:
        - Cell backgrounds
        - VML drawings
        - Legacy formats
        """
        images_info = []
        
        try:
            import zipfile
            import xml.etree.ElementTree as ET
            
            with zipfile.ZipFile(self.file_path, 'r') as zip_file:
                # Look for all media files
                media_files = [f for f in zip_file.namelist() if '/media/' in f and 
                             any(f.lower().endswith(ext) for ext in
                                 ['.png', '.jpg', '.jpeg', '.gif', '.bmp',
                                  '.tiff', '.tif', '.emf', '.wmf'])]
                
                # Try to map images to their cell locations via drawing files
                image_to_cell_map = {}
                drawing_files = [f for f in zip_file.namelist() if '/drawings/' in f and f.endswith('.xml')]
                
                for drawing_file in drawing_files:
                    try:
                        content = zip_file.read(drawing_file).decode('utf-8')
                        # Parse drawing XML to find image anchors
                        # This is a simplified extraction - full implementation would need proper namespace handling
                        import re
                        # Look for cell references in anchor tags
                        # Pattern to find from cell references
                        from_cells = re.findall(r'<xdr:from>.*?<xdr:col>(\d+)</xdr:col>.*?<xdr:row>(\d+)</xdr:row>', 
                                               content, re.DOTALL)
                        # Map drawing index to cell location
                        for idx, (col, row) in enumerate(from_cells):
                            if idx < len(media_files):
                                # Convert to Excel cell reference (0-based to 1-based)
                                cell_ref = f"{chr(65 + int(col))}{int(row) + 1}"  # Simple conversion for single letters
                                image_to_cell_map[media_files[idx]] = cell_ref
                    except Exception as e:
                        logger.debug(f"Could not parse drawing file {drawing_file}: {e}")
                
                for idx, media_file in enumerate(media_files):
                    try:
                        image_data = zip_file.read(media_file)
                        
                        # Try to determine which sheet this image belongs to
                        # For now, we'll associate with the first sheet as a fallback
                        sheet_idx = 0
                        sheet_name = self.doc.sheetnames[0] if self.doc.sheetnames else "Sheet1"
                        
                        # Get cell reference if available
                        cell_ref = image_to_cell_map.get(media_file, None)
                        
                        images_info.append({
                            'id': f'xlsx_fallback_{idx}_{media_file.split("/")[-1]}',
                            'data': image_data,
                            'location': {
                                'sheet': sheet_idx + 1, 
                                'sheet_name': sheet_name, 
                                'img_idx': idx, 
                                'source': 'zip_fallback',
                                'cell_ref': cell_ref  # Add cell reference if found
                            },
                            'media_file': media_file  # Keep track of which media file this is
                        })
                        
                        logger.info(f"Extracted image via ZIP fallback: {media_file} (cell: {cell_ref or 'unknown'})")
                    except Exception as e:
                        logger.warning(f"Failed to extract {media_file}: {e}")
                        
        except Exception as e:
            logger.warning(f"ZIP fallback extraction failed: {e}")
            
        return images_info


class UniversalOfficeLoader:
    """Universal loader that detects file type and uses appropriate loader"""

    @staticmethod
    def load(file_path: Union[str, Path],
             extract_images: bool = True,
             ocr_images: bool = False,
             show_progress: bool = True,
             ocr=None,
             table_style: Union[str, TableStyle] = None) -> Dict[str, Any]:
        """Load any supported Office document
        
        Args:
            file_path: Path to the Office document
            extract_images: Whether to extract images
            ocr_images: Whether to OCR images
            show_progress: Whether to show progress logs
            ocr: OCR instance for image processing
            table_style: Output style for complex tables:
                - 'minimal_html': Clean HTML with only rowspan/colspan (default)
                - 'markdown_grid': Markdown with merge annotations
                - 'styled_html': Full HTML with inline styles (legacy)
        """
        file_path = Path(file_path)
        extension = file_path.suffix.lower()

        loaders = {
            '.docx': DocxLoader,
            '.pptx': PptxLoader,
            '.xlsx': XlsxLoader
        }

        if extension not in loaders:
            raise ValueError(f"Unsupported file type: {extension}")

        loader_class = loaders[extension]
        loader = loader_class(file_path, ocr=ocr, table_style=table_style)

        try:
            return loader.convert_to_json(
                extract_images=extract_images,
                ocr_images=ocr_images,
                show_progress=show_progress
            )
        finally:
            # Cleanup if needed
            pass


# Convenience functions
def office_to_json(
        file_path: Union[str, Path],
        output_path: Optional[Union[str, Path]] = None,
        output_markdown: bool = False,
        extract_images: bool = True,
        ocr_images: bool = False,
        show_progress: bool = True,
        ocr=None
) -> Dict[str, Any]:
    """
    Convert Office document to JSON with content in reading order
    
    Args:
        file_path: Path to the Office file (DOCX, PPTX, or XLSX)
        output_path: Optional path to save JSON output
        output_markdown: Also save as markdown file
        extract_images: Extract images as base64
        ocr_images: Use OCR to convert images to text descriptions
        show_progress: Show progress messages
    
    Returns:
        JSON data with content array
    """
    json_data = UniversalOfficeLoader.load(
        file_path,
        extract_images=extract_images,
        ocr_images=ocr_images,
        show_progress=show_progress,
        ocr=ocr
    )

    if output_path:
        output_path = Path(output_path)
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(json_data, f, ensure_ascii=False, indent=2)
        logger.info(f"JSON saved to: {output_path}")

        if output_markdown:
            markdown_path = output_path.with_suffix('.md')
            markdown_content = office_to_markdown(json_data)
            with open(markdown_path, 'w', encoding='utf-8') as f:
                f.write(markdown_content)
            logger.info(f"Markdown saved to: {markdown_path}")

    return json_data


_MARKDOWN_LIST_MARKER = re.compile(r'[-+*]|\d{1,9}[.)]')


def office_to_markdown(json_data: Dict[str, Any]) -> str:
    """Convert Office JSON data to markdown string.

    Document text is escaped so it cannot turn into Markdown/HTML structure (a paragraph
    ``# of units`` stays a paragraph, ``<img …>`` stays text); see ``doc2mark.utils.markdown``.
    A paragraph or list item with a ``markdown`` key (Word bold, italics and links) is written
    from it. Items of a Word header or footer (``"region": "header"`` / ``"footer"``) are written
    between ``<!-- header -->`` and ``<!-- /header -->`` (``footer``) lines.
    """
    from doc2mark.utils.markdown import escape_heading_text, escape_list_item

    def escape_footnote(text: str) -> str:
        # "[^3]: note text" -> keep the definition label, escape the note
        match = re.match(r"(\[\^[^\]]+\]:[ \t]*)(.*)", text, re.DOTALL)
        return match.group(1) + escape_markdown_text(match.group(2)) if match else escape_markdown_text(text)

    markdown_parts = []
    current_page = None
    # Depth of the deepest list item still open in the current run of list items
    # (only markers Markdown recognises open one); -1 when none is.
    list_depth, list_page = -1, None

    # Determine marker label from filename
    filename = json_data.get("filename", "")
    if filename.lower().endswith(".pptx"):
        page_label = "slide"
    elif filename.lower().endswith(".xlsx"):
        page_label = "sheet"
    else:
        page_label = "page"

    region = None
    for item in json_data["content"]:
        # Close a header / footer block before a page marker, open one after it
        if item.get("region") != region and region:
            markdown_parts.append(f"<!-- /{region} -->\n")
            region = None
            list_depth = -1
        # Add page/slide/sheet separator if needed
        if "page" in item and item["page"] != current_page:
            if markdown_parts:
                markdown_parts.append("")
            markdown_parts.append(f"<!-- {page_label} {item['page']} -->")
            markdown_parts.append("")
            current_page = item["page"]
        if item.get("region") != region:
            region = item["region"]
            markdown_parts.append(f"<!-- {region} -->\n")
            list_depth = -1

        # Structure read from the file (list marker, heading level) sits beside the
        # verbatim ``content``; only the prefix is built here. A list item nests at
        # most one level below a preceding item of the same run whose marker is a
        # Markdown list marker ("-", "1.", "1)"), so an indented line always continues
        # a list and never starts an indented code block.
        if item["type"] == "text:list" and item.get("marker"):
            if item.get("page") != list_page:
                list_depth = -1
            depth = min(item.get("list_level", 0), list_depth + 1)
            text = item.get("markdown") or escape_markdown_text(item["content"])
            markdown_parts.append(f"{'    ' * depth}{item['marker']} {text}\n")
            if _MARKDOWN_LIST_MARKER.fullmatch(item["marker"]):
                list_depth = depth
            else:  # a paragraph ("(a) ..."): it closes list items at its own depth and deeper
                list_depth = min(list_depth, depth - 1)
            list_page = item.get("page")
            continue
        list_depth = -1
        if item["type"] in ("text:title", "text:section") and item.get("level"):
            number = f"{item['marker']} " if item.get("marker") else ""
            markdown_parts.append(f"{'#' * item['level']} {number}{escape_heading_text(item['content'])}\n")
            continue

        if item["type"] == "text:title":
            # Add titles with H1 formatting and extra spacing
            markdown_parts.append(f"# {escape_heading_text(item['content'])}\n")
        elif item["type"] == "text:section":
            # Add sections with H2 formatting
            markdown_parts.append(f"## {escape_heading_text(item['content'])}\n")
        elif item["type"] == "text:list":
            # Add list items with proper formatting
            markdown_parts.append(f"{escape_list_item(item['content'])}\n")
        elif item["type"] == "text:caption":
            # Add captions in italics, one emphasis per line
            caption_lines = escape_markdown_text(item['content'].strip()).split("\n")
            markdown_parts.append("\n".join(f"*{line.strip()}*" for line in caption_lines if line.strip()) + "\n")
        elif item["type"] == "text:normal":
            # Add normal text with paragraph spacing
            markdown_parts.append(f"{item.get('markdown') or escape_markdown_text(item['content'])}\n")
        elif item["type"] == "text:image_description":
            # OCR'd-image text — strip the internal provenance wrapper and emit
            # clean text (no code-fence / <ocr_result> noise).
            ocr_text = item['content']
            if ocr_text.startswith('<image_ocr_result>') and ocr_text.endswith('</image_ocr_result>'):
                ocr_text = ocr_text[18:-19]
            markdown_parts.append(ocr_text.strip() + "\n")
        elif item["type"] == "table":
            # Tables already have proper spacing
            markdown_parts.append(item["content"])
        elif item["type"] == "image":
            # Include image as markdown with base64 data URL
            markdown_parts.append(f'![Image](data:image/png;base64,{item["content"]})\n')
        elif item["type"] == "text:footnote":
            markdown_parts.append(f"{escape_footnote(item['content'])}\n")

    if region:
        markdown_parts.append(f"<!-- /{region} -->\n")
    return "\n".join(markdown_parts)
