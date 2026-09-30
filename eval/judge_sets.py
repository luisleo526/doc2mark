#!/usr/bin/env python
"""Build the labelled evaluation sets of doc2mark's optional judge hooks.

    python eval/judge_sets.py --out tests/data/judge            # all sets
    python eval/judge_sets.py --out tests/data/judge --only legibility
    python eval/judge_sets.py --out tests/data/judge --check    # rebuild and compare byte for byte
    python eval/judge_sets.py --deck /path/to/private_deck.pdf  # uncommitted deck slice (counts only)

It writes ``legibility.jsonl``, ``boilerplate.jsonl``, ``non_content.jsonl`` and
``non_content_ambiguous.jsonl`` (see ``tests/data/judge/README.md``) from in-repo sources
only: the committed sample documents, PDFs generated here with PyMuPDF and Pillow, and the
seed items of the two judge spikes, embedded below. Every text and context is what the
current doc2mark gives the hook: page texts come from ``pdf_routing.measure_page`` (and the
legibility judge the pipeline actually asks), repeated lines from ``pdf_to_simple_json`` with
a recording ``boilerplate_judge``. Boilerplate items are labelled with how many copies must
remain in the Markdown (``expected_copies``). Splits go by family, with near-duplicates
clustered (``assign_splits``). ``external/{legibility,boilerplate,non_content}.jsonl`` hold
the PR #22 reviewer's fresh items (embedded verbatim, split "external", never calibrated on).

``deck_items(path)`` builds the slice of a private deck at run time; it is never written
into the repository (the repository is public).
"""
import argparse
import difflib
import hashlib
import io
import json
import os
import random
import re
import sys
import tempfile
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Callable, Dict, List, Optional, Sequence, Tuple

import pymupdf
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "sample_documents"
FILES = ("legibility", "boilerplate", "non_content", "non_content_ambiguous")
BACKSLASH = chr(92)
FFFD = chr(0xFFFD)


def _doc2mark():
    """doc2mark of this checkout (an editable install elsewhere must not be measured)."""
    if str(ROOT) not in sys.path[:1]:
        sys.path.insert(0, str(ROOT))
    import doc2mark
    if ROOT not in Path(doc2mark.__file__).resolve().parents:
        raise RuntimeError(f"doc2mark imported from {doc2mark.__file__}, not from {ROOT}")
    return doc2mark


# ============================================================================ output


def _json_line(item: dict) -> str:
    """One JSON line; invisible and garbage code points (controls, format characters,
    private use, non-ASCII spaces, U+FFFD) are written as escapes, other text as is."""
    raw = json.dumps(item, ensure_ascii=False)
    out = []
    for char in raw:
        code = ord(char)
        if char == " " or (unicodedata.category(char)[0] not in "CZ" and code != 0xFFFD):
            out.append(char)
        elif code > 0xFFFF:
            code -= 0x10000
            out.append(f"{BACKSLASH}u{0xD800 + (code >> 10):04x}{BACKSLASH}u{0xDC00 + (code & 0x3FF):04x}")
        else:
            out.append(f"{BACKSLASH}u{code:04x}")
    return "".join(out)


def write_jsonl(path: Path, items: Sequence[dict]) -> None:
    """Write atomically: a temporary file in the same directory, then rename."""
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=str(path.parent), prefix=f".{path.name}.", suffix=".tmp")
    with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as handle:
        for item in items:
            handle.write(_json_line(item) + "\n")
    os.chmod(tmp, 0o644)
    os.replace(tmp, path)


def _hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


NEAR_JACCARD, NEAR_RATIO = 0.5, 0.8   # near-duplicates: char 5-gram Jaccard or difflib ratio at or above
SIMILARITY: Dict[str, dict] = {}      # per hook, filled by assign_splits: cross-split similarity report


def _sim_text(text: str) -> str:
    """What near-duplicate detection compares: NFKC, case folded, whitespace collapsed."""
    return " ".join(unicodedata.normalize("NFKC", text).casefold().split())


def _grams(text: str) -> frozenset:
    return frozenset(text[i:i + 5] for i in range(len(text) - 4)) or frozenset([text])


def _jaccard(grams_a: frozenset, grams_b: frozenset) -> float:
    return len(grams_a & grams_b) / len(grams_a | grams_b)


def _ratio(a: str, b: str, floor: float = 0.0) -> float:
    """difflib ratio (no autojunk); 0.0 when its upper bound is below ``floor`` (not computed)."""
    matcher = difflib.SequenceMatcher(None, a, b, autojunk=False)
    if matcher.real_quick_ratio() < floor or matcher.quick_ratio() < floor:
        return 0.0
    return matcher.ratio()


def _near(jaccard: float, ratio: float) -> bool:
    return jaccard >= NEAR_JACCARD or ratio >= NEAR_RATIO


def _cross_max(pairs, texts) -> Tuple[float, float, list]:
    """Exact largest Jaccard and difflib ratio over ``pairs`` (i, j, Jaccard), and the pairs above
    0.6 on either, computing a ratio only where its upper bound can still matter."""
    best_jaccard = max((jaccard for _, _, jaccard in pairs), default=0.0)
    best_ratio, above = 0.0, []
    bounds = []
    for i, j, jaccard in pairs:
        matcher = difflib.SequenceMatcher(None, texts[i], texts[j], autojunk=False)
        bounds.append((min(matcher.real_quick_ratio(), matcher.quick_ratio()), i, j, jaccard, matcher))
    for bound, i, j, jaccard, matcher in sorted(bounds, key=lambda row: (-row[0], row[1], row[2])):
        if bound <= min(best_ratio, 0.6) and jaccard <= 0.6:
            continue
        ratio = matcher.ratio()
        best_ratio = max(best_ratio, ratio)
        if ratio > 0.6 or jaccard > 0.6:
            above.append((round(jaccard, 3), round(ratio, 3), i, j))
    return best_jaccard, best_ratio, sorted(above, key=lambda row: (-max(row[:2]), row[2], row[3]))


def assign_splits(items: List[dict], hook: str, external: Sequence[dict] = ()) -> None:
    """``split`` = "train" or "test" by FAMILY: every item of one template (its translations and
    brand or name substitutions), one base text or document, or one seed item goes to one split.
    Families holding near-duplicates (compared text: the page text, answer or line text; char
    5-gram Jaccard >= 0.5 or difflib ratio >= 0.8) are merged with union-find, and every merged
    cluster goes to one split. A cluster with a near-duplicate of an ``external`` item goes to
    test (the external set stays independent of calibration); the others are placed largest first
    (ties by the SHA-256 of their name), each on the side that keeps the labels, then the (kind,
    label) strata (and for boilerplate the (language, label) and (reason, label) strata) most
    balanced, ties by the hash. Records in ``SIMILARITY[hook]`` the largest cross-split
    similarities."""
    texts = [_sim_text(item["text"]) for item in items]
    grams = [_grams(text) for text in texts]
    parent = {item["family"]: item["family"] for item in items}

    def find(family):
        while parent[family] != family:
            parent[family] = parent[parent[family]]
            family = parent[family]
        return family

    for i in range(len(items)):
        for j in range(i + 1, len(items)):
            if items[i]["family"] != items[j]["family"] and _near(
                    _jaccard(grams[i], grams[j]), _ratio(texts[i], texts[j], NEAR_RATIO)):
                a, b = sorted((find(items[i]["family"]), find(items[j]["family"])))
                parent[b] = a
    clusters: Dict[str, List[dict]] = defaultdict(list)
    for item in items:
        clusters[find(item["family"])].append(item)
    ext_texts = [_sim_text(ext["text"]) for ext in external]
    ext_grams = [_grams(text) for text in ext_texts]
    forced = {find(items[i]["family"]) for i in range(len(items)) for e in range(len(external))
              if _near(_jaccard(grams[i], ext_grams[e]), _ratio(texts[i], ext_texts[e], NEAR_RATIO))}
    counts: Dict[tuple, List[int]] = defaultdict(lambda: [0, 0])

    def strata(members):
        keys = Counter()
        for item in members:
            keys[("label", item["label"])] += 4
            keys[("kind", item["kind"], item["label"])] += 1
            if "lang" in item:
                keys[("lang", item["lang"], item["label"])] += 1
            if "context" in item:
                keys[("reason", item["context"]["reason"], item["label"])] += 1
        return keys

    order = sorted(clusters, key=lambda name: (name not in forced, -len(clusters[name]), _hash(name)))
    for name in order:
        weights = strata(clusters[name])

        def cost(side: int) -> int:
            total = 0
            for key, n in weights.items():
                sides = list(counts[key])
                sides[side] += n
                total += abs(sides[0] - sides[1])
            return total

        costs = cost(0), cost(1)
        side = 1 if name in forced else 0 if costs[0] < costs[1] else 1 if costs[1] < costs[0] \
            else int(_hash(name)[-1], 16) % 2
        for key, n in weights.items():
            counts[key][side] += n
        for item in clusters[name]:
            item["split"] = ("train", "test")[side]
    cross = [(i, j, _jaccard(grams[i], grams[j])) for i in range(len(items)) for j in range(i + 1, len(items))
             if items[i]["split"] != items[j]["split"]]
    jaccard, ratio, above = _cross_max(cross, texts)
    info = {"families": len(parent), "clusters": len(clusters),
            "largest_cluster": max(len(members) for members in clusters.values()),
            "forced_test": sum(len(clusters[name]) for name in forced),
            "max_jaccard": jaccard, "max_ratio": ratio,
            "above_0.6": [(a, b, items[i]["id"], items[j]["id"]) for a, b, i, j in above]}
    if external:
        both = texts + ext_texts
        both_grams = grams + ext_grams
        for split in ("train", "test"):
            pairs = [(i, len(items) + e, _jaccard(both_grams[i], both_grams[len(items) + e]))
                     for i in range(len(items)) if items[i]["split"] == split for e in range(len(external))]
            jaccard, ratio, above = _cross_max(pairs, both)
            ids = [item["id"] for item in items] + [ext["id"] for ext in external]
            info[f"external_{split}"] = (jaccard, ratio, [(a, b, ids[j], ids[i]) for a, b, i, j in above])
    SIMILARITY[hook] = info


def summary(name: str, items: Sequence[dict]) -> str:
    """Counts per split x label, and per kind x label x split (Markdown)."""
    lines = [f"### {name} ({len(items)} items)", "", "| split | label 0 | label 1 | total |", "|---|---|---|---|"]
    for split in [split for split in ("train", "test", "external") if any(item["split"] == split for item in items)]:
        part = [item for item in items if item["split"] == split]
        zero = sum(1 for item in part if item["label"] == 0)
        lines.append(f"| {split} | {zero} | {len(part) - zero} | {len(part)} |")
    if any(item["split"] == "external" for item in items):
        return "\n".join(lines)
    lines += ["", "| kind | label | train | test |", "|---|---|---|---|"]
    table = Counter((item["kind"], item["label"], item["split"]) for item in items)
    for kind, label in sorted({(item["kind"], item["label"]) for item in items}):
        lines.append(f"| {kind} | {label} | {table[(kind, label, 'train')]} | {table[(kind, label, 'test')]} |")
    return "\n".join(lines)


# ======================================================================= non-content
#
# (group, kind, label, text, source). label 1 = the whole answer is only a refusal, apology,
# error or "no readable text" statement; label 0 = it transcribes or describes anything.

_SB = "spikeB"                      # the Jev fit spike's OCR-answer set (items_b.py)
_RV = "review-pr19"                 # PR #19 reviewer probes (refusal_probe.py) and review rounds
_UT = "tests/test_ocr_refusal.py"
_NEW = "new"

NON_CONTENT = [
    # ---- spike B, label 1 ------------------------------------------------------------
    ("openai-classic", "refusal_en", 1, "I'm sorry, but I can't assist with that request.", _SB),
    ("identify-people", "safety_refusal", 1,
     "I'm sorry, I can't help with identifying or making assumptions about people in images.", _SB),
    ("unable-resolution", "refusal_en", 1,
     "I am unable to transcribe this image because the resolution is too low and the text is not legible.", _SB),
    ("blank-or-no-text", "no_text_en", 1, "The image appears to be blank or contains no visible text.", _SB),
    ("no-text-detected", "placeholder", 1, "No text detected.", _SB),
    ("as-an-ai", "refusal_en", 1, "As an AI language model, I cannot read the contents of this document image.", _SB),
    ("cannot-process", "error_shape", 1, "Unable to process the image. Please provide a clearer scan of the page.", _SB),
    ("polite-offer", "refusal_en", 1,
     "I can't transcribe copyrighted book pages in full, but I can summarize the general topic if you'd like.", _SB),
    ("error-json", "error_shape", 1, '{"error": "image could not be read"}', _SB),
    ("zh-tw-refusal", "refusal_ml", 1, "抱歉，我無法辨識這張圖片中的文字。", _SB),
    ("zh-tw-resolution", "refusal_ml", 1, "很抱歉，這張圖片解析度過低，無法進行文字轉錄。", _SB),
    ("zh-cn-refusal", "refusal_ml", 1, "对不起，我无法处理这张图片的内容。", _SB),
    ("zh-no-text", "no_text_ml", 1, "圖片中沒有可辨識的文字。", _SB),
    ("ja-refusal", "refusal_ml", 1, "申し訳ありませんが、この画像の文字を読み取ることができません。", _SB),
    ("ko-refusal", "refusal_ml", 1, "죄송하지만 이 이미지의 텍스트를 인식할 수 없습니다.", _SB),
    ("de-refusal", "refusal_ml", 1, "Leider kann ich den Text in diesem Bild nicht erkennen.", _SB),
    ("placeholder-no-content", "placeholder", 1, "[No content]", _SB),
    ("mostly-empty", "no_text_en", 1,
     "The page seems to be mostly empty; I could not find any content to transcribe.", _SB),
    ("sensitive-info", "safety_refusal", 1,
     "This image may contain sensitive personal information, so I won't transcribe it.", _SB),
    # ---- spike B, label 0 (deck-derived items replaced by fictional text of the same shape)
    ("slide-zh-transcription", "transcription", 0,
     "## 04 / 核心優勢\n北辰雲平台整合採購、庫存與出貨資料，讓各部門在同一介面追蹤訂單進度。",
     _SB + " (adapted: fictional text)"),
    ("slide-zh-title", "transcription", 0, "方案簡介\n北辰雲 企業資料整合平台", _SB + " (adapted: fictional text)"),
    ("prose-en", "transcription", 0,
     "Introduction\nThis is a comprehensive sample DOCX document that demonstrates various document elements "
     "including text formatting, images, and tables.", _SB),
    ("table-html", "transcription", 0,
     "<table><tr><th>Engine</th><th>MT</th><th>DSG</th></tr><tr><td>1.0 TSI/85 kW</td><td>✓</td><td>–</td></tr></table>",
     _SB),
    ("receipt", "transcription", 0, "7-ELEVEN 統一超商\n2026/09/12 14:03\n鮮奶茶 1 x 35\n御飯糰 2 x 32\n合計 NT$99\n統一編號 22555003",
     _SB),
    ("chart-description", "description", 0,
     "Bar chart titled 'Quarterly revenue (NT$M)'. Q1 120, Q2 135, Q3 160, Q4 171. Revenue grows every quarter.", _SB),
    ("photo-description", "description", 0,
     "A photograph of a factory floor with two workers inspecting a CNC machine; a sign on the wall reads "
     "'SAFETY FIRST'.", _SB),
    ("logo-caption", "description", 0, "北辰精密 logo", _SB + " (adapted: fictional text)"),
    ("sorry-we-missed-you", "notice", 0,
     "Sorry we missed you!\nWe tried to deliver your parcel today. Scan the QR code to reschedule delivery.", _SB),
    ("returns-policy", "notice", 0,
     "We cannot accept returns after 30 days. Items must be unused and in the original packaging.", _SB),
    ("grant-letter", "note_chat_letter", 0,
     "Dear Mr. Chen,\nI am sorry to inform you that your application for the 2026 grant was not successful. "
     "We received 412 proposals this year.", _SB),
    ("error-503", "ui_error", 0,
     "Error 503\nService Unavailable\nThe server is temporarily unable to service your request. "
     "Please try again later.", _SB),
    ("no-text-no-problem", "quote_or_slide", 0, "No text? No problem.\nOur OCR engine reads handwriting, stamps and faded scans.",
     _SB),
    ("zh-apology-notice", "notice", 0, "致歉聲明\n本公司因系統異常導致9月12日訂單延遲出貨，造成不便，深感抱歉。", _SB),
    ("blank-form", "name_or_form", 0,
     "Name: ________\nDate: ________\nSignature: ________\n(This page intentionally left blank for notes)", _SB),
    ("intentionally-blank", "quote_or_slide", 0, "This page intentionally left blank.", _SB),
    ("pip-install", "transcription", 0, "$ pip install doc2mark\nSuccessfully installed doc2mark-0.6.1", _SB),
    ("thanks-slide", "quote_or_slide", 0, "Thank you!\nQ&A", _SB),
    ("page-number", "quote_or_slide", 0, "- 12 -", _SB),
    ("ja-prose", "transcription", 0, "第3章 システム構成\n本システムは、データ収集層、解析層、表示層の三層で構成される。", _SB),
    ("preamble-then-content", "caveat_then_content", 0,
     "Here is the transcription of the image:\n\nMeeting agenda\n1. Budget review\n2. Hiring plan\n3. Q4 roadmap", _SB),
    ("illegible-marker", "transcription", 0,
     "Invoice No. 2026-0917\nBill to: Acme Ltd.\nAmount due: [illegible]\nDue date: 2026-10-15", _SB),
    # ---- PR #19 reviewer hard negatives: a refusal or "no text" clause, then content --
    ("caveat-handwritten", "caveat_then_content", 0,
     "I can't transcribe the handwritten text, but the printed part reads:\nInvoice No. 2026-0917\nTotal due: $450", _RV),
    ("caveat-signature", "caveat_then_content", 0,
     "Invoice No. 2026-0917\nTotal due: $450\n(The signature at the bottom is illegible, so I can't transcribe it.)", _RV),
    ("caveat-stamp", "caveat_then_content", 0,
     "Invoice 2026-0917, Total $450. The stamp is unreadable and I won't transcribe it.", _RV),
    ("caveat-caption", "caveat_then_content", 0,
     "The image is mostly blank except for a small caption: 'Figure 3 - Plant layout'.", _RV),
    ("caveat-logo", "caveat_then_content", 0, "The image is blank apart from the company logo ACME Corp.", _RV),
    ("caveat-bar-chart", "caveat_then_content", 0,
     "There is no text in this image. It shows a bar chart with revenue rising from 120 to 171 across Q1-Q4.", _RV),
    ("caveat-factory", "caveat_then_content", 0,
     "There is no readable text in this image; it is a photograph of a factory floor with two workers at a CNC machine.",
     _RV),
    ("caveat-map", "caveat_then_content", 0,
     "There are no words in the image, only a map of Taiwan with Taipei, Taichung and Kaohsiung marked.", _RV),
    ("caveat-stop-sign", "caveat_then_content", 0,
     "I don't see any text in the image, but it shows a red octagonal stop sign on a pole.", _RV),
    ("caveat-watermark", "caveat_then_content", 0,
     "I don't see any text in this image besides the watermark 'CONFIDENTIAL'.", _RV),
    ("caveat-tsmc", "caveat_then_content", 0, "No text detected in the image apart from the logo 'TSMC'.", _RV),
    ("caveat-badge", "caveat_then_content", 0,
     "As an AI, I can't identify the person in the photo, but the name badge reads 'Dr. Lee, Chief Engineer'.", _RV),
    ("caveat-board-minutes", "caveat_then_content", 0,
     "I'm sorry, I can't read the handwritten notes in this image, but the typed heading says 'Q3 Board Minutes'.", _RV),
    ("caveat-ja", "caveat_then_content", 0,
     "すみません、この画像の手書き部分は読み取れませんが、印刷部分は以下の通りです。\n請求書番号 2026-0917\n合計 45,000円", _RV),
    ("caveat-ko", "caveat_then_content", 0,
     "죄송하지만 이 이미지의 손글씨는 인식할 수 없습니다. 인쇄된 부분: 청구서 번호 2026-0917 합계 45,000원", _RV),
    ("caveat-de", "caveat_then_content", 0,
     "Leider kann ich den handschriftlichen Text im Bild nicht erkennen. Gedruckt: Rechnung Nr. 4711, Betrag 450 EUR",
     _RV),
    # ---- PR #19 reviewer: real page text that sounds like a refusal -------------------
    ("hal-9000", "quote_or_slide", 0, "I'm sorry Dave, I'm afraid I can't do that.", _RV),
    ("hal-9000", "quote_or_slide", 0, "I'm sorry, Dave. I'm afraid I can't do that.", _UT),
    ("no-refunds", "notice", 0, "Sorry, no refunds.", _RV),
    ("liability", "notice", 0, "The company cannot be held liable for any loss or damage.", _RV),
    ("meeting-friday", "note_chat_letter", 0, "I'm sorry, but I can't make it to the meeting on Friday.", _RV),
    ("refund-receipt", "note_chat_letter", 0,
     "Dear customer, I am sorry but I cannot process your refund without a receipt.", _RV),
    ("de-chat-du", "note_chat_letter", 0, "Leider kann ich das Foto nicht lesen, kannst du es nochmal schicken?", _RV),
    ("es-chat-tu", "note_chat_letter", 0, "No pude leer la foto, ¿me la mandas otra vez?", _RV),
    ("zh-product-photo-notice", "notice", 0, "本商品圖片僅供參考，實際以實物為準", _RV),
    # ---- PR #19 reviewer: refusals and "no text" statements ----------------------------
    ("error-no-text-found", "error_shape", 1, "Error: No text found.", _RV),
    ("no-text-found", "placeholder", 1, "No text found", _RV),
    ("unreadable", "placeholder", 1, "Unreadable", _RV),
    ("illegible", "placeholder", 1, "ILLEGIBLE", _RV),
    ("illegible", "placeholder", 1, "Illegible", _UT),
    ("illegible-brackets", "placeholder", 1, "[illegible]", _RV),
    ("zh-cannot-read-picture", "refusal_ml", 1, "我無法讀取這張圖片", _RV),
    ("zh-photo-no-text", "no_text_ml", 1, "這張照片沒有文字", _RV),
    ("photo-blank", "no_text_en", 1, "The photo is blank.", _RV),
    ("scan-illegible", "no_text_en", 1, "This scan is illegible.", _RV),
    ("screenshot-empty", "no_text_en", 1, "The screenshot was empty.", _RV),
    ("cannot-process-image", "refusal_en", 1, "I cannot process the image.", _RV),
    ("will-not-comply-request", "refusal_en", 1, "I will not comply with this request.", _RV),
    ("ja-photo-refusal", "refusal_ml", 1, "申し訳ありませんが、この写真の文字は読み取れません。", _RV),
    ("ko-no-text", "no_text_ml", 1, "이미지에 텍스트가 없습니다.", _RV),
    ("de-no-text", "no_text_ml", 1, "Das Bild enthält keinen Text.", _RV),
    ("cannot-provide-transcription", "refusal_en", 1, "I'm sorry, but I cannot provide a transcription of this image.", _RV),
    ("copyrighted-text", "refusal_en", 1,
     "I apologize, but I cannot provide the text from this image as it appears to contain copyrighted material.", _RV),
    ("too-blurry-because", "refusal_en", 1, "I can't read the text in this image because it's too blurry.", _RV),
    ("zh-request-refusal", "refusal_ml", 1, "抱歉，我無法協助處理這個請求。", _RV),
    ("zh-no-transcription", "refusal_ml", 1, "很抱歉，我無法提供這張圖片的轉錄。", _RV),
    ("identify-people-short", "safety_refusal", 1, "Sorry, I can't identify people in images.", _RV),
    ("no-legible-present", "no_text_en", 1, "No legible text is present in the image.", _RV),
    ("assist-courtesy", "refusal_en", 1,
     "I'm sorry, but I can't assist with that. If you have any other questions, feel free to ask!", _RV),
    # ---- tests/test_ocr_refusal.py: canonical refusals, and hard negatives the patterns keep
    ("scan-corrupted", "refusal_en", 1, "I can't read the scan. It seems to have been corrupted.", _UT),
    ("clearer-image-tail", "refusal_en", 1,
     "I'm unable to read the text in this image. It appears to be blurry or low resolution.\n"
     "If you could provide a clearer image, I'd be happy to help!", _UT),
    ("sensitive-description", "safety_refusal", 1,
     "I'm sorry, I can't provide a description of this image. It may contain sensitive content.", _UT),
    ("blank-placeholder", "placeholder", 1, "[Blank]", _UT),
    ("error-code-ui", "ui_error", 0, "Unable to read the file.\nError code 0x80070570\nContact your administrator", _UT),
    ("payment-notice", "notice", 0, "We're sorry, we are unable to process your payment at this time.", _UT),
    ("new-policy", "note_chat_letter", 0, "I'm sorry, but I can't comply with the new policy.", _UT),
    ("hi-team-attachment", "note_chat_letter", 0,
     "Hi team, I can't read the text in the attachment - can you resend it?", _UT),
    ("upload-limits", "ui_error", 0, "Unable to process the uploaded image. Please upload a JPG or PNG under 10 MB.", _UT),
    ("retake-hint", "ui_error", 0, "The photo is blurry.\nHold the camera steady and retake it.", _UT),
    ("photo-id-signature", "note_chat_letter", 0,
     "I cannot process your photo ID without a signature. Please resend the form.", _UT),
    ("resubmit-application", "notice", 0,
     "The attached photo is not readable. Please resubmit your application with a new photo.", _UT),
    ("zh-upload-outage", "notice", 0, "抱歉，系統暫時無法處理照片上傳，工程師正在修復中。", _UT),
    ("ja-upload-outage", "notice", 0, "申し訳ありませんが、現在画像のアップロードを処理できません。", _UT),
    ("search-no-text", "ui_error", 0, "No text found.\nTry a different search term.", _UT),
    ("scans-dr-lee", "note_chat_letter", 0, "I can't read the scans until Dr. Lee signs off.", _UT),
    ("claim-deadline", "note_chat_letter", 0,
     "Unfortunately, I cannot assist with processing your claim as it was filed after the deadline.", _UT),
    ("scan-corrupted-total", "caveat_then_content", 0,
     "I can't read the scan. It seems to have been corrupted. Total: $500", _UT),
    ("ask-mark", "note_chat_letter", 0, "Sorry, I can't help with that. However, I can ask Mark tomorrow.", _UT),
    ("gratitude", "quote_or_slide", 0, "There are no words to describe our gratitude.", _UT),
    ("ai-first", "quote_or_slide", 0, "As an AI-first company, we ship faster.", _UT),
    ("named-es", "named_refusal_shape", 0, "No puedo leer la imagen del recibo de Juan Pérez", _UT),
    ("named-zh", "named_refusal_shape", 0, "很抱歉，我無法辨識王小明的照片", _UT),
    ("named-ja", "named_refusal_shape", 0, "申し訳ありませんが、山田さんの写真の文字は読み取れません。", _UT),
    # ---- new: refusals and "no text" statements in other words ---------------------------
    ("new-not-transcribable", "refusal_en", 1,
     "Unfortunately, the page in this image cannot be transcribed; the text is not legible.", _NEW),
    ("new-nothing-legible", "no_text_en", 1, "Nothing legible could be extracted from this image.", _NEW),
    ("new-nothing-to-transcribe", "no_text_en", 1,
     "The provided image doesn't appear to contain any text that I can transcribe.", _NEW),
    ("new-apologies-extract", "refusal_en", 1,
     "My apologies, I wasn't able to extract any text from the image you provided.", _NEW),
    ("new-make-out", "refusal_en", 1, "I'm not able to make out any of the writing in this picture.", _NEW),
    ("new-transcription-unavailable", "error_shape", 1,
     "Transcription unavailable: the image resolution is too low to read the text.", _NEW),
    ("new-no-legible-content", "placeholder", 1, "No legible content.", _NEW),
    ("new-no-text-parens", "placeholder", 1, "(no text)", _NEW),
    ("new-blank-page-image", "no_text_en", 1,
     "The uploaded page image is entirely blank, so there is nothing to transcribe.", _NEW),
    ("new-quality-too-poor", "refusal_en", 1, "Sorry, the image quality is too poor for me to read any of the text.", _NEW),
    ("new-no-image-received", "error_shape", 1,
     "I cannot see the image you're referring to. Could you please upload it again?", _NEW),
    ("new-image-did-not-come-through", "error_shape", 1,
     "It looks like the image didn't come through. Please try attaching it again.", _NEW),
    ("new-json-no-text", "error_shape", 1, '{"text": "", "error": "no_text_detected"}', _NEW),
    ("new-json-status", "error_shape", 1, '{"status": "error", "message": "Unable to read image"}', _NEW),
    ("new-decode-error", "error_shape", 1, "Error: image could not be decoded.", _NEW),
    ("new-ocr-failed", "error_shape", 1, "OCR failed: no text regions were detected in the image.", _NEW),
    ("new-json-null", "error_shape", 1, '{"transcription": null, "reason": "image too blurry"}', _NEW),
    ("new-pii-policy", "safety_refusal", 1,
     "I can't help with extracting personal data such as ID numbers from images of documents.", _NEW),
    ("new-id-documents", "safety_refusal", 1,
     "I can't transcribe identity documents or other images that contain personal information.", _NEW),
    ("new-privacy-handwriting", "safety_refusal", 1,
     "For privacy reasons I cannot transcribe handwritten personal information.", _NEW),
    ("new-usage-policies", "safety_refusal", 1, "This content may violate our usage policies, so I can't process it.", _NEW),
    ("new-zh-tw-undetected", "no_text_ml", 1, "此圖片無法辨識，未偵測到任何文字。", _NEW),
    ("new-zh-tw-blurry", "refusal_ml", 1, "抱歉，這張圖片太模糊了，我沒辦法轉錄內容。", _NEW),
    ("new-zh-tw-extract", "refusal_ml", 1, "無法從圖片中擷取文字。", _NEW),
    ("new-zh-tw-placeholder", "placeholder", 1, "（此頁無可辨識文字）", _NEW),
    ("new-zh-cn-blurry", "refusal_ml", 1, "抱歉，图片过于模糊，无法识别其中的文字。", _NEW),
    ("new-zh-cn-undetected", "no_text_ml", 1, "图片中未检测到任何文字。", _NEW),
    ("new-ja-unclear", "refusal_ml", 1, "画像が不鮮明なため、文字を読み取ることができませんでした。", _NEW),
    ("new-ja-recognize", "refusal_ml", 1, "申し訳ございませんが、画像の内容を認識できませんでした。", _NEW),
    ("new-ja-no-text", "no_text_ml", 1, "この画像には読み取り可能なテキストが含まれていないようです。", _NEW),
    ("new-ko-blurry", "refusal_ml", 1, "이미지가 너무 흐려서 텍스트를 읽을 수 없습니다.", _NEW),
    ("new-ko-no-text", "no_text_ml", 1, "이 이미지에서는 텍스트를 찾을 수 없습니다.", _NEW),
    ("new-de-no-text", "no_text_ml", 1, "Im Bild ist kein lesbarer Text vorhanden.", _NEW),
    ("new-de-blurry", "refusal_ml", 1,
     "Es tut mir leid, aber ich kann dieses Bild nicht transkribieren, da es zu unscharf ist.", _NEW),
    ("new-de-not-found", "no_text_ml", 1, "Leider konnte ich im Bild keinen Text finden.", _NEW),
    ("new-fr-refusal", "refusal_ml", 1, "Je suis désolé, mais je ne peux pas transcrire le texte de cette image.", _NEW),
    ("new-fr-no-text", "no_text_ml", 1, "Aucun texte lisible n'a été détecté dans l'image.", _NEW),
    ("new-fr-blurry", "refusal_ml", 1, "L'image est trop floue pour être lue.", _NEW),
    ("new-es-refusal", "refusal_ml", 1, "Lo siento, no puedo transcribir el texto de esta imagen.", _NEW),
    ("new-es-no-text", "no_text_ml", 1, "No se detectó texto legible en la imagen.", _NEW),
    ("new-es-blurry", "refusal_ml", 1, "La imagen está demasiado borrosa para leer el texto.", _NEW),
    # ---- new: content that shares words with refusals ---------------------------------
    ("new-back-in-10", "note_chat_letter", 0, "Back in 10 minutes.", _NEW),
    ("new-out-of-office", "note_chat_letter", 0, "Out of office until Monday, 6 October.", _NEW),
    ("new-caveat-heading", "caveat_then_content", 0,
     "The bottom half is too blurry to read, but the heading says 'Quarterly Safety Audit - Plant 2'.", _NEW),
    ("new-caveat-cut-off", "caveat_then_content", 0,
     "Parts of the page are cut off. Visible text:\nTerms and Conditions\n1. Scope\n2. Payment", _NEW),
    ("new-caveat-zh", "caveat_then_content", 0,
     "圖片右下角的手寫字跡無法辨識，印刷部分如下：\n報價單編號 Q-2026-118\n總計 NT$48,500", _NEW),
    ("new-caveat-fr", "caveat_then_content", 0,
     "Le tampon est illisible, mais le texte imprimé indique : Facture n° 2026-044, total 1 250,00 EUR", _NEW),
    ("new-caveat-es", "caveat_then_content", 0,
     "La firma no se puede leer; el resto dice: Recibo n.º 5521, importe 320 EUR", _NEW),
    ("new-name-en", "name_or_form", 0, "Maria Gonzalez", _NEW),
    ("new-name-ja", "name_or_form", 0, "山田 太郎", _NEW),
    ("new-name-title", "name_or_form", 0, "Dr. Priya Raman, MD\nCardiology", _NEW),
    ("new-form-values", "name_or_form", 0, "Name: Chen Wei-Ting\nDate of birth: 1987-04-12\nPolicy no.: 55-201-7788", _NEW),
    ("new-form-illegible", "name_or_form", 0, "Signature: [illegible]\nDate: 2026-09-12", _NEW),
    ("new-de-sign", "multilingual_content", 0, "Bitte nicht stören", _NEW),
    ("new-fr-closed", "multilingual_content", 0, "Fermé pour travaux jusqu'au 15 octobre.", _NEW),
    ("new-zh-hours", "multilingual_content", 0, "營業時間 09:00-18:00，週日公休", _NEW),
    ("new-elevator-apology", "notice", 0,
     "We apologize for the inconvenience. The elevator is out of service until Friday.", _NEW),
    ("new-ja-full", "notice", 0, "申し訳ございません。本日は満席です。", _NEW),
    ("new-de-machine", "notice", 0, "Entschuldigung, dieser Automat nimmt keine Scheine.", _NEW),
    ("new-ko-closed", "notice", 0, "죄송합니다. 오늘은 휴무입니다.", _NEW),
    ("new-zh-sold-out", "notice", 0, "抱歉，本日商品已售完。", _NEW),
    ("new-caption-figure", "transcription", 0, "Figure 2: Throughput vs. batch size (log scale).", _NEW),
    ("new-caption-photo", "transcription", 0, "Photo: The new Hsinchu plant, June 2026.", _NEW),
    ("new-404", "ui_error", 0, "404 Not Found\nThe requested URL /reports/2026 was not found on this server.", _NEW),
    ("new-login-screenshot", "description", 0,
     "Screenshot of a login form with the fields 'Email' and 'Password' and a blue 'Sign in' button.", _NEW),
    ("new-upload-receipt", "notice", 0, "Upload a photo of your receipt to claim the refund.", _NEW),
    ("new-recognition-release", "quote_or_slide", 0, "Text recognition accuracy improved by 12% in version 2.3.", _NEW),
]

# Ambiguous by construction (not scored): nothing in the text says whether it is the model's
# non-answer or a transcribed chat line, support reply or app message.
NON_CONTENT_AMBIGUOUS = [
    ("sorry-cant-help", "Sorry, I can't help.", _RV, "generic refusal that names nothing: a chat line reads the same"),
    ("cant-help", "I can't help.", _RV, "generic refusal that names nothing"),
    ("help-you-with-that", "I'm sorry, but I can't help you with that.", _RV, "generic refusal addressed to a person"),
    ("help-with-that", "I'm sorry, I can't help with that.", _RV, "generic refusal that names nothing"),
    ("help-with-that-curly", "I’m sorry, but I can't help with that.", _RV, "generic refusal that names nothing"),
    ("sorry-help-with-that", "Sorry, I can't help with that.", _UT, "generic refusal that names nothing"),
    ("i-cant-help-with-that", "I can't help with that.", _UT, "generic refusal that names nothing"),
    ("assistance-with-that", "I can't provide assistance with that.", _RV, "generic refusal that names nothing"),
    ("comply-with-that", "I'm sorry, I can't comply with that.", _RV, "generic refusal that names nothing"),
    ("share-that", "I'm sorry, but I can't share that.", _RV, "generic refusal that names nothing"),
    ("unable-to-help", "I’m unable to help with that.", _RV, "generic refusal that names nothing"),
    ("unable-to-assist", "I'm unable to assist.", _RV, "generic refusal that names nothing"),
    ("cannot-do-that", "Sorry, I cannot do that.", _RV, "generic refusal that names nothing"),
    ("but-cant-do-that", "I'm sorry, but I can't do that.", _UT, "generic refusal that names nothing"),
    ("cant-do-that", "I can't do that.", _UT, "generic refusal that names nothing; also a slide or quote"),
    ("will-not-comply", "I will not comply.", _UT, "generic refusal that names nothing; also a quote"),
    ("fenced-assist", "```\nI'm sorry, but I can't assist with that.\n```", _RV, "generic refusal in a code fence"),
    ("wont-read-picture", "I won't read the picture.", _RV, "odd wording: model or person"),
    ("cant-read-this-scan", "I can't read this scan.", _RV, "can't-read line a person could write about a scan they got"),
    ("fr-photo-floue", "Désolé, je ne peux pas lire la photo, elle est floue.", _RV,
     "can't-read line with a quality reason; a chat reads the same"),
    ("no-text-long", "There is no readable text in this image. It seems to be a decorative background graphic.", _SB,
     "no-text statement plus a minimal description of the image (spike B labelled it 1)"),
    ("photo-blurry", "The photo is blurry.", _UT, "a quality remark: an app hint or the model"),
    ("no-content-found", "No content found", _UT, "placeholder or an app's empty state"),
    ("document-empty", "The document is empty.", _UT, "no-content statement or an app message"),
    ("no-clear-words", "There are no clear words to display.", _UT, "no-content statement or an app message"),
    ("describe-in-words", "I cannot describe this picture in words.", _UT, "refusal to describe, or a figure of speech"),
    ("read-your-handwriting", "Sorry, I can't read your handwriting.", _UT, "model answer to a handwritten page, or a note"),
    ("document-you-sent", "Sorry, I cannot read the document you sent.", _UT, "model answer or support reply"),
    ("retake-it", "I can't read the picture, it's too blurry. Can you retake it?", _UT, "model answer or support reply"),
    ("zoom-in", "Sorry, I can't read the text in this, can you zoom in?", _UT, "model answer or chat line"),
    ("text-in-your-photo", "I'm sorry, I can't help with reading the text in your photo.", _UT,
     "model answer or support reply"),
    ("clearer-photo", "Sorry, I can't read this photo. Please send a clearer photo.", _UT, "model answer or support reply"),
    ("clearer-copy", "I can't read this scan. Please upload a clearer copy.", _UT, "model answer or support reply"),
    ("fr-vous", "Je ne peux pas lire la photo que vous m'avez envoyée.", _UT, "model answer or support reply"),
    ("de-sie", "Leider kann ich das Foto nicht lesen, bitte senden Sie es erneut.", _UT, "model answer or support reply"),
    ("es-su-documento", "Lo siento, no pude leer la foto de su documento.", _UT, "model answer or support reply"),
    ("es-intentelo", "No es posible procesar la imagen. Inténtelo de nuevo.", _UT, "engine message or app message"),
    ("zh-reupload", "抱歉，無法讀取您上傳的圖片，請重新上傳。", _UT, "model answer or app message"),
    ("zh-who-in-photo", "我無法辨識照片中的人是誰，可以再傳一張嗎？", _UT, "safety refusal or chat line"),
    ("rescan-required", "Scan is illegible - rescan required", _UT, "a stamp on the page or an engine status"),
    ("driver-description", "I cannot provide a description of the driver.", _UT,
     "safety refusal about a person in the image, or a witness note"),
    ("legal-advice", "As an AI assistant, I cannot give legal advice.", _UT, "a chatbot disclaimer on the page, or the model"),
]


# Answers written as variants of one seed (translations, a derived hard negative, a shortened
# form): one family, so one split. Other groups are their own family.
_NC_FAMILY = {
    **dict.fromkeys(("zh-tw-refusal", "zh-cn-refusal", "ja-refusal", "ko-refusal", "de-refusal",
                     "zh-cannot-read-picture", "ja-photo-refusal"), "cant-recognize-text"),
    **dict.fromkeys(("zh-no-text", "ko-no-text", "de-no-text", "zh-photo-no-text", "new-ko-no-text", "new-de-no-text",
                     "new-fr-no-text", "new-es-no-text", "new-ja-no-text", "new-zh-cn-undetected",
                     "new-zh-tw-undetected", "new-de-not-found", "no-legible-present"), "no-text-in-image"),
    **dict.fromkeys(("new-zh-tw-blurry", "new-zh-cn-blurry", "new-ja-unclear", "new-ko-blurry", "new-de-blurry",
                     "new-fr-blurry", "new-es-blurry", "too-blurry-because", "new-quality-too-poor"), "too-blurry"),
    **dict.fromkeys(("new-fr-refusal", "new-es-refusal", "cannot-provide-transcription", "zh-no-transcription",
                     "new-zh-tw-extract", "new-apologies-extract", "new-ja-recognize"), "cannot-transcribe"),
    **dict.fromkeys(("caveat-handwritten", "caveat-signature", "caveat-stamp", "caveat-ja", "caveat-ko", "caveat-de",
                     "new-caveat-zh", "new-caveat-fr", "new-caveat-es", "scan-corrupted", "scan-corrupted-total",
                     "illegible-marker"), "caveat-invoice"),
    **dict.fromkeys(("chart-description", "photo-description", "caveat-bar-chart", "caveat-factory", "caveat-map",
                     "caveat-stop-sign", "caveat-watermark", "caveat-tsmc", "caveat-logo", "caveat-caption"),
                    "caveat-no-text"),
    **dict.fromkeys(("identify-people", "identify-people-short"), "identify-people"),
    **dict.fromkeys(("illegible", "illegible-brackets", "unreadable"), "illegible"),
    **dict.fromkeys(("no-text-found", "error-no-text-found", "search-no-text"), "no-text-found"),
    **dict.fromkeys(("cannot-process", "cannot-process-image"), "cannot-process"),
    **dict.fromkeys(("zh-upload-outage", "ja-upload-outage"), "upload-outage"),
    **dict.fromkeys(("new-ja-full", "new-de-machine", "new-ko-closed", "new-zh-sold-out"), "shop-apology"),
    **dict.fromkeys(("named-es", "named-zh", "named-ja"), "named-refusal"),
    **dict.fromkeys(("openai-classic", "assist-courtesy"), "cant-assist"),
}


def build_non_content() -> Tuple[List[dict], List[dict]]:
    from doc2mark.ocr.refusal import MAX_JUDGE_CHARS, _normalize, matches_non_content_pattern

    def base(prefix, group, source, kind, label, text):
        answer = _normalize(text)
        fires = matches_non_content_pattern(text)
        return {"id": None, "group": f"{prefix}-{group}", "family": f"nc:{_NC_FAMILY.get(group, group)}",
                "source": source, "kind": kind, "label": label, "split": None, "text": text, "pattern": fires,
                "judged": bool(answer) and not fires and len(answer) <= MAX_JUDGE_CHARS}

    items, seen = [], Counter()
    for group, kind, label, text, source in NON_CONTENT:
        item = base("nc", group, source, kind, label, text)
        seen[group] += 1
        item["id"] = f"nc-{group}" + (f"-{seen[group]}" if seen[group] > 1 else "")
        items.append(item)
    texts = [item["text"] for item in items]
    if len(set(texts)) != len(texts):
        raise ValueError("duplicate non-content texts")
    ambiguous = []
    for group, text, source, why in NON_CONTENT_AMBIGUOUS:
        item = base("nca", group, source, "ambiguous", None, text)
        item["id"] = f"nca-{group}"
        item["why"] = why
        ambiguous.append(item)
    if set(texts) & {item["text"] for item in ambiguous}:
        raise ValueError("a text is both scored and ambiguous")
    return items, ambiguous


# ======================================================================== legibility

# Spike A hard legible items (synthetic, real-world shapes), verbatim.
SPIKE_A_HARD = {
    'hard_financial_table': (
        'Consolidated Statement of Income (in thousands)\n'
        '2025 2024 Change\n'
        'Revenue 1,284,332 1,102,918 16.4%\n'
        'Cost of revenue (612,004) (548,221)\n'
        'Gross profit 672,328 554,697 21.2%\n'
        'R&D (201,337) (187,002)\n'
        'SG&A (144,918) (139,550)\n'
        'Operating income 326,073 228,145 42.9%\n'
        'EPS (diluted) 3.41 2.37'
    ),
    'hard_hashes': (
        'Release checksums (SHA-256)\n'
        'doc2mark-0.6.1.tar.gz  9f86d081884c7d659a2feaa0c55ad015a3bf4f1b2b0b822cd15d6c15b0f00a08\n'
        'doc2mark-0.6.1-py3-none-any.whl  60303ae22b998861bce3b28f33eec1be758a213c86c93c076dbe9f558'
        'c11c752\n'
        'GPG key: 0x4A3F9C21D0B87E55'
    ),
    'hard_part_numbers': (
        'BOM rev C\n'
        'R101 RC0402FR-0710KL 10k 1% 0402\n'
        'C12 GRM155R71C104KA88D 100nF X7R\n'
        'U3 STM32F411CEU6 UFQFPN48\n'
        'J1 USB4105-GF-A USB-C\n'
        'D2 BAT54WS-7-F Schottky'
    ),
    'hard_code': (
        'def decide(cov: float, txt: float) -> str:\n'
        '    if cov >= 0.55 and txt < 200:\n'
        "        return 'image'\n"
        "    return 'text'\n"
        '\n'
        'for k, v in cfg.items():\n'
        "    print(f'{k}={v!r}')"
    ),
    'hard_math': (
        'Theorem 2. For all x ∈ ℝⁿ, ‖Ax‖₂ ≤ σ_max(A)‖x‖₂.\n'
        'Proof. Let A = UΣVᵀ. Then ‖Ax‖₂² = xᵀVΣ²Vᵀx ≤ σ²_max ‖x‖₂². ∎\n'
        '∑_{i=1}^{n} λ_i = tr(A), ∏ λ_i = det(A)'
    ),
    'hard_bibliography': (
        '[12] K. He, X. Zhang, S. Ren, J. Sun. Deep residual learning for image recognition. CVPR, '
        'pp. 770–778, 2016.\n'
        '[13] A. Vaswani et al. Attention is all you need. NeurIPS 30, 2017.\n'
        '[14] J. Devlin, M.-W. Chang, K. Lee, K. Toutanova. BERT. NAACL-HLT, 2019.'
    ),
    'hard_japanese': (
        '第3章 システム構成\n'
        '本システムは、データ収集層、解析層、表示層の三層で構成される。収集層では各工場のセンサー値を1秒ごとに取得し、解析層で異常検知モデルを適用する。'
    ),
    'hard_korean': (
        '제2장 개인정보 처리 방침\n'
        '회사는 이용자의 개인정보를 수집 목적 범위 내에서만 처리하며, 보유 기간이 끝나면 지체 없이 파기합니다.'
    ),
    'hard_simplified_zh': (
        '第四条 甲方应于每月五日前向乙方支付上月服务费用。逾期付款的，每逾期一日按未付金额的千分之三支付违约金。'
    ),
    'hard_form_labels': (
        'Applicant name: ____________\n'
        'Date of birth (DD/MM/YYYY): __/__/____\n'
        'Passport No.: __________\n'
        '☐ Single ☐ Married ☐ Other\n'
        'Signature: ______________'
    ),
    'hard_urls_emails': (
        'Contacts\n'
        'support@typesafe-demo.io\n'
        'https://docs.example.com/v2/api/reference#auth\n'
        'sftp://files.acme-corp.net:2222/inbound/\n'
        '+886-2-2712-3456 ext. 301'
    ),
    'hard_short_labels_en': (
        'Q3 Highlights\n'
        'ARR $12.4M\n'
        'NRR 118%\n'
        'Churn 1.9%\n'
        'Headcount 84'
    ),
    'hard_german': (
        'Die Gesellschaft hat im Geschäftsjahr 2025 einen Umsatz von 48,2 Mio. € erzielt. Das Ergeb'
        'nis vor Steuern stieg gegenüber dem Vorjahr um 12 %.'
    ),
    'hard_letterspaced_title': (
        'A N N U A L  R E P O R T  2 0 2 5\n'
        'Message from the Chairman\n'
        'Dear shareholders, this year we delivered record results across all regions.'
    ),
}


# --- spike A text transforms (character-wise, so they apply span by span) -------------

def _caesar(k: int) -> Callable[[str], str]:
    def shift(s):
        out = []
        for c in s:
            if "a" <= c <= "z":
                out.append(chr((ord(c) - 97 + k) % 26 + 97))
            elif "A" <= c <= "Z":
                out.append(chr((ord(c) - 65 + k) % 26 + 65))
            else:
                out.append(c)
        return "".join(out)
    return shift


def _perm_map(seed: int) -> Dict[str, str]:
    rnd = random.Random(seed)
    a = list("abcdefghijklmnopqrstuvwxyz")
    b = a[:]
    rnd.shuffle(b)
    mapping = dict(zip(a, b))
    mapping.update({x.upper(): y.upper() for x, y in zip(a, b)})
    return mapping


def _perm(seed: int) -> Callable[[str], str]:
    mapping = _perm_map(seed)
    return lambda s: "".join(mapping.get(c, c) for c in s)


def _glyphid(s: str) -> str:
    return "".join(chr(ord(c) - 29) if c.isalpha() and ord(c) < 128 else c for c in s)


def _cjk_offset(k: int) -> Callable[[str], str]:
    return lambda s: "".join(chr(0x4E00 + (ord(c) - 0x4E00 + k) % 0x51A0) if 0x4E00 <= ord(c) <= 0x9FFF else c
                             for c in s)


def _pua(s: str) -> str:
    return "".join(chr(0xE000 + ord(c) % 0x1000) if not c.isspace() else c for c in s)


def _cid(s: str) -> str:
    return "".join(f"(cid:{ord(c) % 180 + 3})" if not c.isspace() else c for c in s)


def _fffd(rate: float, seed: int) -> Callable[[str], str]:
    rnd = random.Random(seed)
    return lambda s: "".join(FFFD if (not c.isspace() and rnd.random() < rate) else c for c in s)


def _mojibake(s: str) -> str:
    return s.encode("utf-8").decode("cp1252", errors="replace")


_NOISE_POOL = "abcdefghijklmnopqrstuvwxyz0123456789.,;:'!|/" + BACKSLASH + "~^"


def _noisy(rate: float, seed: int) -> Callable[[str], str]:
    rnd = random.Random(seed)
    return lambda s: "".join(rnd.choice(_NOISE_POOL) if (c.isalnum() and rnd.random() < rate) else c for c in s)


def _symbols(s: str) -> str:
    sym = "■□▲△▼▽◆◇○●◎★☆♠♣♥♦"
    return "".join(sym[ord(c) % len(sym)] if c.isalnum() else c for c in s)


# Spike A garbage items on non-deck bases: (spike id, kind, base key, transform factory, line-wise?)
SPIKE_A_GARBAGE = [
    ("garb_caesar_en", "garb_shifted_latin", "pdf:sample_pdf.pdf#p1", lambda: _caesar(3), False),
    ("garb_caesar_en2", "garb_shifted_latin", "pdf:test-table.pdf#p1", lambda: _caesar(7), False),
    ("garb_perm_en", "garb_substituted_latin", "pdf:sample_pdf.pdf#p2", lambda: _perm(7), False),
    ("garb_perm_en2", "garb_substituted_latin", "pdf:complex_table_test.pdf#p1", lambda: _perm(11), False),
    ("garb_perm_docx", "garb_substituted_latin", "docx:sample_document.docx", lambda: _perm(5), False),
    ("garb_glyphid_en", "garb_glyph_offset", "pdf:sample_pdf.pdf#p1", lambda: _glyphid, False),
    ("garb_cjk_offset_zh_contract", "garb_cjk_substituted", "hard:hard_simplified_zh", lambda: _cjk_offset(450), False),
    ("garb_pua_en", "garb_pua", "pdf:sample_pdf.pdf#p2", lambda: _pua, False),
    ("garb_cid_en", "garb_cid", "pdf:sample_pdf.pdf#p1", lambda: _cid, False),
    ("garb_fffd40_en", "garb_fffd", "pdf:sample_pdf.pdf#p2", lambda: _fffd(0.4, 3), False),
    ("garb_mojibake_zh2", "garb_mojibake", "hard:hard_japanese", lambda: _mojibake, False),
    ("garb_noisy_ocr35", "garb_bad_ocr_layer", "pdf:sample_pdf.pdf#p1", lambda: _noisy(0.35, 9), False),
    ("garb_noisy_ocr60", "garb_bad_ocr_layer", "pdf:test-table.pdf#p1", lambda: _noisy(0.6, 10), False),
    ("garb_symbols", "garb_symbol_font", "pdf:complex_table_test.pdf#p1", lambda: _symbols, False),
    ("garb_mixed_half", "garb_mixed_half", "pdf:sample_pdf.pdf#p1", lambda: 21, True),
]


# --- PDF writing ------------------------------------------------------------------------

A4 = (595, 842)
SLIDE = (960, 540)
_EMBEDDED = {"sans": ("FS", "helv"), "bold": ("FB", "hebo"), "serif": ("FR", "tiro"), "mono": ("FM", "cour"),
             "broken": ("F9", "tiro"), "broken-cjk": ("F8", "china-t")}
_CJK_BASE = {"zh-t": "china-t", "zh-s": "china-s", "ja": "japan", "ko": "korea"}
_FONT_CACHE: Dict[str, pymupdf.Font] = {}


def _font(name: str) -> pymupdf.Font:
    if name not in _FONT_CACHE:
        _FONT_CACHE[name] = pymupdf.Font(name)
    return _FONT_CACHE[name]


class Pdf:
    """A PDF under construction: embedded fonts (exact Unicode text layers), CJK base fonts,
    and two "broken" fonts whose ToUnicode maps :meth:`save` rewrites."""

    def __init__(self):
        self.doc = pymupdf.open()
        self.broken: Dict[str, set] = defaultdict(set)
        self._embedded: Dict[int, set] = defaultdict(set)

    def page(self, size=A4):
        return self.doc.new_page(width=size[0], height=size[1])

    def text(self, page, x, y, text, size=11.0, font="sans", render_mode=0, color=(0, 0, 0)) -> float:
        if font in _CJK_BASE:
            fontname, width = _CJK_BASE[font], len(text) * size
        else:
            fontname, builtin = _EMBEDDED[font]
            if fontname not in self._embedded[page.number]:
                page.insert_font(fontname=fontname, fontbuffer=_font(builtin).buffer)
                self._embedded[page.number].add(fontname)
            width = _font(builtin).text_length(text, fontsize=size)
            if font.startswith("broken"):
                self.broken[fontname].update(text)
        if x + width > page.rect.width - 10:
            raise ValueError(f"line wider than the page: {text!r}")
        page.insert_text((x, y), text, fontname=fontname, fontsize=size, render_mode=render_mode, color=color)
        return width

    def lines(self, page, x, y, lines, size=11.0, font="sans", leading=1.5, **kwargs) -> float:
        for line in lines:
            if line:
                self.text(page, x, y, line, size, font, **kwargs)
            y += size * leading
        return y

    def save(self, path: Path, rewrite: Optional[Dict[str, Optional[Callable[[str], str]]]] = None) -> Path:
        if any(self._embedded.values()):
            self.doc.subset_fonts()
        for resource, fn in (rewrite or {}).items():
            _garble(self.doc, resource, fn, self.broken[resource])
        self.doc.save(str(path), garbage=3, deflate=True)
        self.doc.close()
        return Path(path)


def _cmap_pairs(cmap: str):
    for block in re.findall(r"beginbfchar(.*?)endbfchar", cmap, re.S):
        for src, dst in re.findall(r"<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>", block):
            yield int(src, 16), chr(int(dst[:4], 16))
    for block in re.findall(r"beginbfrange(.*?)endbfrange", cmap, re.S):
        for lo, hi, dst in re.findall(r"<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>", block):
            for code in range(int(lo, 16), int(hi, 16) + 1):
                yield code, chr(int(dst[:4], 16) + code - int(lo, 16))


def _utf16_hex(text: str) -> str:
    return text.encode("utf-16-be").hex().upper() if text else "FFFD"


def _garble(doc, resource: str, rewrite: Optional[Callable[[str], str]], used: set) -> None:
    """Rewrite the ToUnicode map of the embedded font ``resource``: its glyphs still render,
    but extract as ``rewrite(char)`` (no ToUnicode at all when ``rewrite`` is None: U+FFFD)."""
    xrefs = sorted({font[0] for page in doc for font in page.get_fonts(full=True) if font[4] == resource})
    if not xrefs:
        raise ValueError(f"no {resource} text in the document")
    for xref in xrefs:
        if rewrite is None:
            doc.xref_set_key(xref, "ToUnicode", "null")
            continue
        kind, value = doc.xref_get_key(xref, "ToUnicode")
        if kind != "xref":
            raise ValueError(f"font {xref} has no ToUnicode stream")
        cmap_xref = int(value.split()[0])
        cmap = doc.xref_stream(cmap_xref).decode("latin-1")
        entries = sorted({(src, rewrite(char)) for src, char in _cmap_pairs(cmap) if char in used})
        blocks = []
        for start in range(0, len(entries), 100):
            chunk = entries[start:start + 100]
            blocks.append(f"{len(chunk)} beginbfchar\n"
                          + "\n".join(f"<{src:04X}> <{_utf16_hex(dst)}>" for src, dst in chunk) + "\nendbfchar")
        doc.update_stream(cmap_xref, (
            "/CIDInit /ProcSet findresource begin 12 dict begin begincmap /CMapName /X def "
            "1 begincodespacerange <0000> <FFFF> endcodespacerange\n" + "\n".join(blocks)
            + "\nendcmap CMapName currentdict /CMap defineresource pop end end").encode())


def _shift3(c: str) -> str:
    return _caesar(3)(c)


def _glyph_offset_real(c: str) -> str:
    """A subset font whose glyph IDs were taken for Unicode ("Sample" -> "6DPSOH"): printable
    ASCII 29 lower; digits and most punctuation fall below 0x20 (MuPDF then yields U+FFFD)."""
    return chr(ord(c) - 29) if c.isascii() and not c.isspace() and ord(c) >= 29 else c


def _mojibake_char(c: str) -> str:
    return c.encode("utf-8").decode("cp1252", errors="replace") if not c.isascii() else c


def _pua_char(c: str) -> str:
    return c if c.isspace() else chr(0xE000 + ord(c) % 200)


def _cjk_char(k: int) -> Callable[[str], str]:
    def sub(c):
        code = ord(c)
        if 0x4E00 <= code <= 0x9FFF:
            return chr(0x4E00 + (code - 0x4E00 + k) % 0x51A0)
        if 0xAC00 <= code <= 0xD7A3:
            return chr(0xAC00 + (code - 0xAC00 + k) % 11172)
        return c
    return sub


# --- base texts for generated pages (fictional) ----------------------------------------

BASES: Dict[str, dict] = {
    "invoice_en": {"font": "sans", "title": "INVOICE", "lines": [
        "Northwind Traders Ltd. · 18 Harbour Road · Portsmouth PO1 3AX · United Kingdom",
        "VAT Reg. No. GB 284 7710 39",
        "Invoice no.: INV-2026-0413    Invoice date: 12 September 2026",
        "Customer: Fabrikam Retail GmbH, Königstraße 21, 70173 Stuttgart",
        "Payment terms: 30 days net    Due date: 12 October 2026",
        "1  Stainless steel hinge, 80 mm       400 x £1.85     £740.00",
        "2  Cabinet handle, brushed nickel     250 x £2.40     £600.00",
        "3  Soft-close drawer runner (pair)    120 x £6.75     £810.00",
        "4  Delivery and handling                1 x £45.00     £45.00",
        "Subtotal £2,195.00    VAT 20% £439.00    Total due £2,634.00",
        "Bank: Lloyds Bank · Sort code 30-94-21 · Account 11873420",
        "IBAN GB29 LOYD 3094 2111 8734 20 · BIC LOYDGB21"]},
    "contract_en": {"font": "serif", "title": "MASTER SERVICES AGREEMENT", "lines": [
        "This Agreement is entered into on 1 March 2026 between Contoso Ltd. (the \"Client\")",
        "and Adatum Consulting LLC (the \"Provider\").",
        "7. LIMITATION OF LIABILITY",
        "7.1 Neither party shall be liable to the other for any indirect, incidental or",
        "consequential damages, including loss of profits, arising out of this Agreement.",
        "7.2 Each party's aggregate liability shall not exceed the fees paid by the Client",
        "in the twelve (12) months preceding the event giving rise to the claim.",
        "8. TERM AND TERMINATION",
        "8.1 This Agreement commences on the Effective Date and continues for an initial",
        "term of twenty-four (24) months, renewing automatically for successive one-year",
        "terms unless either party gives ninety (90) days' written notice.",
        "8.2 Either party may terminate this Agreement for material breach that remains",
        "uncured thirty (30) days after written notice describing the breach.",
        "9. GOVERNING LAW",
        "9.1 This Agreement is governed by the laws of the State of Washington."]},
    "letter_en": {"font": "sans", "title": "Adatum Consulting LLC", "lines": [
        "401 Pine Street, Suite 900, Seattle, WA 98101",
        "15 September 2026",
        "Ms. Jordan Ellis, Head of Procurement, Contoso Ltd.",
        "Dear Ms. Ellis,",
        "Thank you for meeting with our team last Thursday. As discussed, we propose to",
        "begin the data migration in two phases: the finance ledgers in November and the",
        "inventory records in January, so that the year-end close is not affected.",
        "Our estimate assumes read-only access to the legacy system and a dedicated test",
        "environment. We will send the detailed project plan and staffing by 30 September.",
        "Please let me know if you would like to review the plan together before then.",
        "Kind regards,",
        "Samuel Okafor, Engagement Manager"]},
    "cashflow_en": {"font": "sans", "title": "Consolidated Statement of Cash Flows", "lines": [
        "For the year ended 31 December 2025 (in thousands of USD)      2025        2024",
        "Net income                                                   48,210      39,877",
        "Depreciation and amortization                                12,604      11,930",
        "Share-based compensation                                      3,118       2,764",
        "Changes in working capital                                   (6,442)     (4,105)",
        "Net cash provided by operating activities                    57,490      50,466",
        "Purchases of property and equipment                         (18,773)    (15,209)",
        "Acquisitions, net of cash acquired                           (9,850)          -",
        "Net cash used in investing activities                       (28,623)    (15,209)",
        "Repayment of long-term debt                                 (10,000)     (8,000)",
        "Dividends paid                                               (7,512)     (6,988)",
        "Net increase in cash and cash equivalents                    11,355      20,269"]},
    "code_py": {"font": "mono", "size": 9.5, "title": "sensors/alerts.py", "lines": [
        "import csv",
        "from dataclasses import dataclass",
        "",
        "@dataclass",
        "class Reading:",
        "    sensor_id: str",
        "    celsius: float",
        "",
        "def load(path: str) -> list[Reading]:",
        "    with open(path, newline=\"\") as handle:",
        "        rows = csv.DictReader(handle)",
        "        return [Reading(r[\"id\"], float(r[\"temp\"])) for r in rows]",
        "",
        "def alerts(readings, limit=85.0):",
        "    for reading in readings:",
        "        if reading.celsius >= limit:",
        "            yield f\"{reading.sensor_id}: {reading.celsius:.1f} C\""]},
    "manual_en": {"font": "sans", "title": "Installation and Safety Instructions", "lines": [
        "WARNING: Disconnect the mains supply before opening the housing.",
        "1. Mount the bracket at least 30 cm above the floor using the four M6 screws.",
        "2. Feed the supply cable through the gland and tighten it to 2.5 N·m.",
        "3. Connect L to terminal 1, N to terminal 2 and the earth conductor to the",
        "   terminal marked with the earth symbol.",
        "4. Set DIP switch 3 to ON for 230 V operation or OFF for 115 V operation.",
        "5. Close the housing and restore power. The status LED flashes green twice.",
        "If the LED stays red, check the fuse (T 2 A, 5 x 20 mm) and the wiring.",
        "Dispose of the packaging in accordance with local recycling regulations."]},
    "slide_en": {"font": "sans", "size": 22, "page": SLIDE, "title": "Quarterly Business Review Q3 2026", "lines": [
        "Revenue $48.2M (+18% YoY)", "Gross margin 61.4%", "3 new regions: Nordics, Benelux, Iberia",
        "NPS 61 (target 55)", "Next: self-serve tier in Q4"]},
    "form_en": {"font": "sans", "title": "Application for Annual Leave", "lines": [
        "Employee name: ____________________",
        "Employee no.: ________    Department: ____________",
        "Leave type:  [ ] Annual   [ ] Sick   [ ] Unpaid   [ ] Other: ________",
        "First day of leave (DD/MM/YYYY): __ /__ /____",
        "Last day of leave (DD/MM/YYYY): __ /__ /____",
        "Number of working days: ____",
        "Reason (optional): ______________________________",
        "Employee signature: ______________    Date: __________",
        "Approved by (manager): ______________    Date: __________"]},
    "timetable_en": {"font": "sans", "title": "Route 42 - Weekday timetable", "lines": [
        "Central Station       06:10   06:40   07:10   07:35   08:05",
        "Market Square         06:16   06:46   07:17   07:43   08:13",
        "University Gate       06:24   06:54   07:26   07:53   08:22",
        "Riverside Park        06:31   07:01   07:34   08:02   08:30",
        "Airport Terminal 2    06:45   07:15   07:49   08:18   08:46",
        "Services every 30 minutes until 20:00, then hourly until 23:30."]},
    "invoice_de": {"font": "sans", "title": "RECHNUNG", "lines": [
        "Müller & Söhne Werkzeugbau GmbH · Industriestraße 12 · 70565 Stuttgart",
        "Rechnungsnummer: 2026-117        Rechnungsdatum: 08.09.2026",
        "Kundennummer: K-40821            Lieferdatum: 03.09.2026",
        "1  Fräswerkzeug VHM, Ø 12 mm        20 x 38,50 €      770,00 €",
        "2  Spannzange ER32, 10 mm           10 x 24,90 €      249,00 €",
        "3  Kühlschmierstoff, 20 l            2 x 110,50 €     221,00 €",
        "Zwischensumme 1.240,00 €   zzgl. 19 % MwSt. 235,60 €",
        "Gesamtbetrag 1.475,60 €",
        "Zahlbar innerhalb von 14 Tagen ohne Abzug.",
        "Bankverbindung: Volksbank Stuttgart · IBAN DE12 6009 0100 0123 4567 89"]},
    "agb_de": {"font": "serif", "title": "Allgemeine Geschäftsbedingungen", "lines": [
        "§ 3 Lieferung und Gefahrübergang",
        "(1) Die Lieferung erfolgt ab Werk Stuttgart. Teillieferungen sind zulässig,",
        "soweit sie dem Besteller zumutbar sind.",
        "(2) Die Gefahr geht mit der Übergabe der Ware an den Spediteur auf den",
        "Besteller über, auch wenn frachtfreie Lieferung vereinbart wurde.",
        "(3) Gerät der Besteller in Annahmeverzug, sind wir berechtigt, Ersatz des",
        "entstehenden Schadens einschließlich etwaiger Mehraufwendungen zu verlangen.",
        "§ 4 Gewährleistung",
        "(1) Mängel sind uns unverzüglich, spätestens jedoch innerhalb von acht Tagen",
        "nach Eingang der Ware, schriftlich anzuzeigen.",
        "(2) Die Gewährleistungsfrist beträgt zwölf Monate ab Gefahrübergang."]},
    "cgv_fr": {"font": "serif", "title": "Conditions générales de vente", "lines": [
        "Article 4 – Prix et modalités de paiement",
        "Les prix sont exprimés en euros, hors taxes. Ils s’entendent départ usine,",
        "emballage compris. Toute commande inférieure à 150 € fera l’objet de frais",
        "de traitement forfaitaires de 12 €.",
        "Le règlement s’effectue par virement à trente jours date de facture. En cas de",
        "retard, des pénalités égales à trois fois le taux d’intérêt légal seront",
        "appliquées, ainsi qu’une indemnité forfaitaire de 40 € pour frais de recouvrement.",
        "Article 5 – Réserve de propriété",
        "Le vendeur conserve la propriété des marchandises jusqu’au paiement intégral",
        "du prix. L’acheteur s’engage à les assurer contre tout risque de perte."]},
    "report_es": {"font": "serif", "title": "Informe trimestral", "lines": [
        "Durante el tercer trimestre, las ventas crecieron un 9 % respecto al mismo",
        "periodo del año anterior, impulsadas por la apertura de dos tiendas en Valencia",
        "y Málaga. El margen bruto se situó en el 38,5 %, ligeramente por encima de lo",
        "previsto, gracias a la renegociación de los contratos de transporte.",
        "¿Qué esperamos para el cuarto trimestre? La campaña de Navidad suele concentrar",
        "el 35 % de la facturación anual; por ello se ha reforzado la plantilla con",
        "cuarenta contratos temporales y se ha ampliado el almacén de Zaragoza.",
        "¡Gracias a todo el equipo por el esfuerzo de estos meses!"]},
    "quality_ja": {"font": "ja", "title": "第2章 品質管理体制", "lines": [
        "当社は、製品の設計から出荷に至るまでの全工程において、国際規格",
        "ISO 9001に基づく品質マネジメントシステムを運用している。",
        "各工場には品質保証課を設置し、受入検査、工程内検査、出荷検査の",
        "三段階で不良品の流出を防止している。",
        "2025年度の工程内不良率は0.12％であり、前年度から改善した。",
        "今後は画像検査装置の導入を進め、目視検査に依存しない体制を構築する。"]},
    "privacy_ko": {"font": "ko", "title": "제3장 개인정보의 보유 및 이용 기간", "lines": [
        "회사는 법령에 따른 개인정보 보유 및 이용 기간 또는 정보주체로부터",
        "개인정보를 수집할 때 동의받은 기간 내에서 개인정보를 처리합니다.",
        "다만, 다음의 경우에는 해당 사유가 종료될 때까지 보관합니다.",
        "1. 관계 법령 위반에 따른 수사 또는 조사가 진행 중인 경우",
        "2. 서비스 이용에 따른 채권 및 채무 관계가 남아 있는 경우",
        "계약 또는 청약철회 기록은 5년간 보관합니다."]},
    "manual_zh_cn": {"font": "zh-s", "title": "产品使用说明", "lines": [
        "一、安装前请确认电源电压为220伏，并确保插座具有可靠接地。",
        "二、首次使用前，请用清水冲洗水箱，并空烧一次以去除异味。",
        "三、加水时水位不得超过最高刻度线，以免沸腾时溢出烫伤。",
        "四、本产品具有干烧保护功能，缺水时将自动断电。",
        "五、清洁时请先拔下电源插头，切勿将主机浸入水中。",
        "售后服务热线：400-820-1234（工作日 9:00-18:00）"]},
    "notice_zh_tw": {"font": "zh-t", "title": "重大訊息公告", "lines": [
        "主旨：本公司董事會決議通過一一五年第二季合併財務報告。",
        "一、董事會決議日期：一一五年八月七日。",
        "二、合併營業收入新台幣三十二億四千萬元，較去年同期成長百分之十一。",
        "三、歸屬於母公司業主之淨利新台幣四億二千萬元，每股盈餘二點一元。",
        "四、本次財務報告業經會計師核閱，並依規定申報主管機關備查。",
        "五、其他應敘明事項：無。"]},
    "slide_zh_tw": {"font": "zh-t", "size": 24, "page": SLIDE, "title": "第三季營運重點與下一步規劃", "lines": [
        "營收年增 18%", "新增三個海外據點", "客戶滿意度 92 分", "下一步：推出訂閱方案"]},
    "form_zh_tw": {"font": "zh-t", "title": "請假申請單", "lines": [
        "申請人：＿＿＿＿＿＿ 員工編號：＿＿＿＿＿＿",
        "部門：＿＿＿＿＿＿＿＿",
        "假別：□ 事假 □ 病假 □ 特休 □ 公假",
        "請假期間：＿＿年＿＿月＿＿日 至 ＿＿年＿＿月＿＿日",
        "請假事由：＿＿＿＿＿＿＿＿＿＿＿＿",
        "代理人簽章：＿＿＿＿＿＿ 主管核准：＿＿＿＿＿＿"]},
    "invoice_ru": {"font": "serif", "title": "Счёт-фактура № 45 от 12.09.2026", "lines": [
        "Продавец: ООО «Северный ветер», ИНН 7701234567, г. Москва",
        "Покупатель: АО «Балтийские системы», ИНН 7812345678, г. Санкт-Петербург",
        "Наименование товара: серверный шкаф 42U, 2 шт. по 58 400,00 руб.",
        "Стоимость без налога: 116 800,00 руб.",
        "НДС 20 %: 23 360,00 руб.",
        "Всего к оплате: 140 160,00 руб.",
        "Руководитель организации ____________ И. П. Смирнов"]},
    "invoice_el": {"font": "serif", "title": "Τιμολόγιο Παροχής Υπηρεσιών αρ. 17", "lines": [
        "Εκδότης: Αιγαίο Λογισμικό Α.Ε., ΑΦΜ 099887766, Αθήνα",
        "Πελάτης: Κυκλάδες Ξενοδοχειακή Ε.Π.Ε., Νάξος",
        "Περιγραφή: Συντήρηση πληροφοριακού συστήματος, Σεπτέμβριος 2026",
        "Καθαρή αξία: 1.200,00 €",
        "ΦΠΑ 24%: 288,00 €",
        "Σύνολο: 1.488,00 €",
        "Τρόπος πληρωμής: Τραπεζική κατάθεση εντός 30 ημερών."]},
    "spec_mixed": {"font": "zh-t", "size": 10, "title": "Product Specification 產品規格", "lines": [
        "Model 型號: NX-2400 Industrial Gateway",
        "Input voltage 輸入電壓: 12-48 V DC",
        "Operating temperature 工作溫度: -20 to 70 C",
        "Interfaces 介面: 2 x RS-485, 1 x CAN, 4 x Ethernet",
        "Protection rating 防護等級: IP65",
        "Certification 認證: CE, FCC, BSMI"]},
}

_BROKEN_FOR = {"zh-t": "broken-cjk", "zh-s": "broken-cjk", "ja": "broken-cjk", "ko": "broken-cjk"}


def _base_page(pdf: Pdf, key: str, *, broken: Sequence[int] = (), broken_title: bool = False,
               mode_font: Optional[str] = None):
    """Write base text ``key`` on a new page. Lines whose index is in ``broken`` (and the title
    with ``broken_title``) are drawn with the broken font of the base's script."""
    spec = BASES[key]
    size = spec.get("size", 10.5)
    font = spec["font"]
    bad = mode_font or _BROKEN_FOR.get(font, "broken")
    page = pdf.page(spec.get("page", A4))
    left = 60 if page.rect.width < 700 else 70
    title_size = size * 1.9 if page.rect.width < 700 else size * 1.6
    title_font = bad if broken_title else ("bold" if font in ("sans", "serif", "mono") else font)
    pdf.text(page, left, 90, spec["title"], title_size, title_font)
    y = 90 + title_size * 1.8
    for index, line in enumerate(spec["lines"]):
        if line:
            pdf.text(page, left, y, line, size, bad if index in broken else font)
        y += size * 1.55
    return page


def _scan_png(lines: Sequence[str], size=(1240, 1754), font_px=24, top=150, left=100):
    image = Image.new("L", size, 255)
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=font_px)
    boxes, y = [], top
    for line in lines:
        if not line.strip():
            y += font_px
            boxes.append(None)
            continue
        box = draw.textbbox((left, y), line, font=font)
        if box[2] > size[0] - 40:
            raise ValueError(f"scan line too wide: {line!r}")
        draw.text((left, y), line, fill=0, font=font)
        boxes.append(box)
        y = box[3] + font_px
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue(), boxes, size


def _ocr_layer_pdf(path: Path, key: str, noise: Callable[[str], str], *, under_image: bool) -> Path:
    """A scan of base ``key`` with an OCR text layer of the noisy lines on each scanned line:
    invisible (render mode 3) over the picture, or painted under it (``under_image``, as
    "text under the page image" exports do)."""
    spec = BASES[key]
    lines = [spec["title"]] + list(spec["lines"])
    png, boxes, size = _scan_png(lines)
    pdf = Pdf()
    page = pdf.page(A4)
    sx, sy = A4[0] / size[0], A4[1] / size[1]
    if not under_image:
        page.insert_image(page.rect, stream=png)
    for line, box in zip(lines, boxes):
        if box is None:
            continue
        noisy = noise(line)
        left, top, _, bottom = box
        size_pt = (bottom - top) * sy * 1.1
        page_text_width = _font("helv").text_length(noisy, fontsize=size_pt)
        if left * sx + page_text_width > A4[0]:
            size_pt *= (A4[0] - left * sx - 5) / page_text_width
        pdf.text(page, left * sx, bottom * sy, noisy, size_pt, "sans", render_mode=0 if under_image else 3)
    if under_image:
        page.insert_image(page.rect, stream=png)
    return pdf.save(path)


_OCR_CONFUSIONS = {"m": "rn", "l": "1", "1": "l", "I": "l", "O": "0", "0": "O", "e": "c", "c": "e", "a": "o",
                   "o": "a", "h": "b", "b": "h", "u": "v", "n": "h", "S": "5", "5": "S", "B": "8", "g": "q",
                   "t": "f", "i": "!", "r": "t", "s": "z"}


def _ocr_noise(rate: float, seed: int) -> Callable[[str], str]:
    """OCR-like errors at ``rate`` per letter or digit: confusions, random characters,
    dropped and inserted characters, merged words."""
    rnd = random.Random(seed)

    def noise(line: str) -> str:
        out = []
        for char in line:
            if char.isalnum() and rnd.random() < rate:
                roll = rnd.random()
                if roll < 0.5 and char in _OCR_CONFUSIONS:
                    out.append(_OCR_CONFUSIONS[char])
                elif roll < 0.8:
                    out.append(rnd.choice(_NOISE_POOL.replace(BACKSLASH, "")))
                elif roll < 0.9:
                    continue
                else:
                    out.append(char + rnd.choice("il1.,'"))
            elif char == " " and rnd.random() < rate / 4:
                continue
            else:
                out.append(char)
        return "".join(out) or line
    return noise


def _stray_glyphs_pdf(path: Path) -> Path:
    """An invoice whose logo glyph and one bullet come from a font without ToUnicode (2 x U+FFFD)."""
    pdf = Pdf()
    page = _base_page(pdf, "invoice_en")
    pdf.text(page, 480, 90, "*", 26, "broken")
    pdf.text(page, 60, 720, "•", 10.5, "broken")
    pdf.text(page, 72, 720, "Thank you for your business.", 10.5, "sans")
    return pdf.save(path, {"F9": None})


def _rating_cards_pdf(path: Path) -> Path:
    """A product page with rows of star icons mapped to private-use code points (icon font)."""
    pdf = Pdf()
    page = pdf.page(A4)
    y = 90
    pdf.text(page, 60, y, "Customer reviews", 20, "bold")
    y += 40
    for name, rating in [("Cordless drill DX-18", "4.6 out of 5 (1,204 reviews)"),
                         ("Impact driver ID-20", "4.3 out of 5 (388 reviews)"),
                         ("Oscillating tool MT-12", "4.8 out of 5 (97 reviews)")]:
        pdf.text(page, 60, y, name, 14, "bold")
        pdf.text(page, 60, y + 22, "*****", 13, "broken")
        pdf.text(page, 60, y + 42, rating, 10.5, "sans")
        y += 80
    return pdf.save(path, {"F9": _pua_char})


def _cid_literal_pdf(path: Path) -> Path:
    """A form re-typeset from a broken extraction: its visible text is "(cid:NN)" runs."""
    pdf = Pdf()
    page = pdf.page(A4)
    y = 90
    for line in [BASES["form_en"]["title"]] + BASES["form_en"]["lines"][:5]:
        encoded = _cid(line)
        while encoded:
            cut = encoded[:78]
            if len(encoded) > 78 and ")" in cut:
                cut = cut[:cut.rindex(")") + 1]
            pdf.text(page, 60, y, cut, 9, "sans")
            encoded = encoded[len(cut):].lstrip()
            y += 14
        y += 4
    return pdf.save(path)


# (item suffix, kind, label, base key, builder) for generated pages; a builder gets (path, key)
def _broken_all(fn):
    return lambda path, key: _save_broken(path, key, fn, lines="all")


def _save_broken(path, key, fn, lines="all", title=True):
    pdf = Pdf()
    spec = BASES[key]
    broken = range(len(spec["lines"])) if lines == "all" else lines
    _base_page(pdf, key, broken=broken, broken_title=title)
    resource = "F8" if spec["font"] in _BROKEN_FOR else "F9"
    return pdf.save(path, {resource: fn})


def _legible(path, key):
    pdf = Pdf()
    _base_page(pdf, key)
    return pdf.save(path)


GENERATED = [
    # legible pages in varied genres, scripts and layouts
    *[(f"{key}-legible", "legible_generated", 1, key, _legible) for key in BASES],
    # valid-Unicode garbage the deterministic detector cannot see
    *[(f"{key}-shifted", "garb_shifted_latin", 0, key, _broken_all(_shift3))
      for key in ("invoice_en", "contract_en", "agb_de", "letter_en")],
    *[(f"{key}-permuted", "garb_substituted_latin", 0, key, _broken_all(_perm(seed)))
      for key, seed in (("cashflow_en", 31), ("manual_en", 32), ("cgv_fr", 33), ("report_es", 34), ("code_py", 35))],
    *[(f"{key}-glyph-offset", "garb_glyph_offset", 0, key, _broken_all(_glyph_offset_real))
      for key in ("manual_en", "timetable_en", "slide_en", "form_en")],
    *[(f"{key}-cjk-substituted", "garb_cjk_substituted", 0, key, _broken_all(_cjk_char(k)))
      for key, k in (("notice_zh_tw", 1200), ("manual_zh_cn", 777), ("quality_ja", 3001), ("slide_zh_tw", 451),
                     ("form_zh_tw", 2024), ("privacy_ko", 613))],
    # garbage the detector sees
    *[(f"{key}-pua", "garb_pua", 0, key, _broken_all(_pua_char)) for key in ("invoice_de", "invoice_ru")],
    *[(f"{key}-no-tounicode", "garb_fffd", 0, key, _broken_all(None)) for key in ("code_py", "timetable_en", "invoice_el")],
    *[(f"{key}-mojibake", "garb_mojibake", 0, key, _broken_all(_mojibake_char))
      for key in ("cgv_fr", "invoice_de", "notice_zh_tw")],
    ("form_en-cid-literal", "garb_cid", 0, "form_en", lambda path, key: _cid_literal_pdf(path)),
    # bad OCR layers of scans
    *[(f"{key}-bad-ocr-invisible", "garb_bad_ocr_layer", 0, key,
       (lambda seed: lambda path, key: _ocr_layer_pdf(path, key, _ocr_noise(0.45, seed), under_image=False))(seed))
      for key, seed in (("letter_en", 41), ("manual_en", 42), ("contract_en", 43))],
    *[(f"{key}-bad-ocr-under-image", "garb_bad_ocr_layer", 0, key,
       (lambda seed: lambda path, key: _ocr_layer_pdf(path, key, _ocr_noise(0.5, seed), under_image=True))(seed))
      for key, seed in (("invoice_en", 44), ("cashflow_en", 45), ("letter_en", 46))],
    # partly garbled: a slide's big title (a large share of its words), or every other line
    ("slide_en-title-garbled", "garb_partial_title", 0, "slide_en",
     lambda path, key: _save_broken(path, key, _perm(51), lines=(0,), title=True)),
    ("slide_zh_tw-title-garbled", "garb_partial_title", 0, "slide_zh_tw",
     lambda path, key: _save_broken(path, key, _cjk_char(888), lines=(0,), title=True)),
    ("form_en-title-garbled", "garb_partial_title", 0, "form_en",
     lambda path, key: _save_broken(path, key, _shift3, lines=(0, 1, 2), title=True)),
    ("contract_en-half-lines", "garb_mixed_half", 0, "contract_en",
     lambda path, key: _save_broken(path, key, _shift3, lines=range(0, 15, 2), title=True)),
    ("agb_de-half-lines", "garb_mixed_half", 0, "agb_de",
     lambda path, key: _save_broken(path, key, _perm(52), lines=range(1, 11, 2), title=False)),
    ("cashflow_en-half-lines", "garb_mixed_half", 0, "cashflow_en",
     lambda path, key: _save_broken(path, key, _glyph_offset_real, lines=range(0, 12, 2), title=True)),
    # legible despite a few bad glyphs or a slightly imperfect OCR layer
    ("invoice_en-two-stray-glyphs", "legible_stray_glyphs", 1, "invoice_en", lambda path, key: _stray_glyphs_pdf(path)),
    ("reviews-icon-stars", "legible_stray_glyphs", 1, "reviews", lambda path, key: _rating_cards_pdf(path)),
    ("letter_en-light-ocr-invisible", "legible_ocr_layer", 1, "letter_en",
     lambda path, key: _ocr_layer_pdf(path, key, _ocr_noise(0.03, 61), under_image=False)),
    ("manual_en-light-ocr-under-image", "legible_ocr_layer", 1, "manual_en",
     lambda path, key: _ocr_layer_pdf(path, key, _ocr_noise(0.03, 62), under_image=True)),
]


# --- measuring pages as the pipeline does -------------------------------------------------


class _NoOCR:
    """Stands in for an OCR provider so the pipeline consults the legibility judge (it is
    never called: only page routes are computed)."""


def _page_rows(page, measure):
    """The text-layer rows of ``page`` the judge text is built from (as measure_page does):
    per line, the (text, size, font) of the painted spans, or of the invisible OCR layer on a
    searchable scan."""
    from doc2mark.pipelines import pdf_routing
    trace = None if pdf_routing.char_flags_mark_painting() else frozenset(pdf_routing.invisible_trace_origins(page))
    visible, invisible, lines = pdf_routing._span_rows(page.get_text("dict", flags=pdf_routing.TEXT_FLAGS), trace)
    if measure.signals.searchable_scan:
        boxes = set(measure.layer_rects)
        source = {id(row) for row in invisible if row[0].strip() and row[2] in boxes}
    else:
        source = {id(row) for row in visible}
    return [[(row[0], row[1], row[3]) for row in line if id(row) in source] for line in lines]


def _join(rows_by_line) -> str:
    return "\n".join(joined for joined in ("".join(text for text, _, _ in line) for line in rows_by_line) if joined)


def _stats_fields(spans) -> dict:
    from doc2mark.core.strategy import MIN_JUDGED_CHARS, text_layer_stats
    stats = text_layer_stats(spans)
    return {"detector": {"garbled": stats.garbled, "chars": stats.chars, "garbage_glyphs": stats.garbage_glyphs,
                         "garbage_ratio": round(stats.garbage_ratio, 4)},
            "judged": not stats.garbled and stats.chars >= MIN_JUDGED_CHARS}


def measure_pdf(path: Path) -> List[dict]:
    """Per page: judge text, spans, detector stats, route, whether the pipeline asks the judge."""
    from doc2mark.core.strategy import text_layer_stats
    from doc2mark.pipelines.pymupdf_advanced_pipeline import PDFLoader
    asked: Dict[int, List[str]] = defaultdict(list)
    current = [0]
    loader = PDFLoader(path, ocr=_NoOCR(), legibility_judge=lambda text: asked[current[0]].append(text))
    pages = []
    try:
        for number in range(len(loader.doc)):
            current[0] = number
            route = loader._page_route(number)
            measure = loader._page_measure(number)
            rows = _page_rows(loader.doc.load_page(number), measure)
            text = measure.text or ""
            if _join(rows) != text:
                raise AssertionError(f"{path} p{number + 1}: rebuilt text differs from measure_page's")
            spans = [[t, s, f] for line in rows for t, s, f in line]
            stats, layer = text_layer_stats(spans), measure.signals.text_layer
            if (stats.garbled, stats.chars, stats.garbage_glyphs) != (layer.garbled, layer.chars, layer.garbage_glyphs):
                raise AssertionError(f"{path} p{number + 1}: spans do not reproduce the detector stats")
            if asked[number] and asked[number] != [text]:
                raise AssertionError(f"{path} p{number + 1}: the judge was asked about another text")
            pages.append({"text": text, "spans": spans, "rows": rows, "route": f"{route[0]}:{route[1]}",
                          "pipeline_asks": bool(asked[number])})
    finally:
        loader.close()
    return pages


def _item(id_, group, source, kind, label, text, spans, family=None, **extra) -> dict:
    item = {"id": id_, "group": group, "family": family or group, "source": source, "kind": kind, "label": label,
            "split": None, "text": text, "spans": spans}
    item.update(_stats_fields(spans))
    item.update(extra)
    return item


def _string_spans(text: str) -> list:
    return [[line, 11.0, "Helvetica"] for line in text.split("\n")]


def _transform_rows(rows_by_line, factory, line_wise: bool):
    """Apply a spike transform span by span (one transform instance, reading order)."""
    rows_by_line = [line for line in rows_by_line if "".join(t for t, _, _ in line)]
    if line_wise:
        seed = factory()
        return [[(_perm(seed + index)(t) if index % 2 else t, s, f) for t, s, f in line]
                for index, line in enumerate(rows_by_line)]
    fn = factory()
    return [[(fn(t), s, f) for t, s, f in line] for line in rows_by_line]


def _transform_text(text: str, factory, line_wise: bool) -> str:
    if line_wise:
        seed = factory()
        return "\n".join(_perm(seed + i)(line) if i % 2 else line for i, line in enumerate(text.split("\n")))
    return factory()(text)


# sample_pdf.pdf is sample_document.docx printed to PDF, and sample_text.txt their plain-text
# sibling from the same sample generator: one family.
_SAMPLE_FAMILY = dict.fromkeys(("sample_pdf.pdf", "sample_document.docx", "sample_text.txt"), "sample:sample-document")


def build_legibility(tmp: Path) -> List[dict]:
    import docx
    items: List[dict] = []
    bases: Dict[str, dict] = {}   # base key -> {"rows": ..., "text": ..., "group": ...}
    # real text layers of every committed sample PDF page
    for rel in ("sample_pdf.pdf", "test-table.pdf", "complex-tables/complex_table_test.pdf"):
        name = Path(rel).name
        for number, page in enumerate(measure_pdf(SAMPLES / rel), 1):
            group = f"sample:{name}#p{number}"
            family = _SAMPLE_FAMILY.get(name, f"sample:{name}")
            bases[f"pdf:{name}#p{number}"] = {"rows": page["rows"], "group": group, "family": family}
            items.append(_item(f"lg-real-{Path(rel).stem}-p{number}", group, f"sample_documents/{rel}#p{number}",
                               "real_pdf_page", 1, page["text"], page["spans"], family=family,
                               route=page["route"], pipeline_asks=page["pipeline_asks"]))
    # spike A non-deck, non-PDF legible items (committed documents read here)
    dx = docx.Document(str(SAMPLES / "sample_document.docx"))
    fx = docx.Document(str(SAMPLES / "fail-1.docx"))
    strings = {
        "docx:sample_document.docx": ("real_sample_docx", "\n".join(p.text for p in dx.paragraphs if p.text.strip()),
                                      "sample_documents/sample_document.docx (paragraphs)"),
        "docx:fail-1.docx": ("real_fail1_docx", "\n".join(p.text for p in fx.paragraphs if p.text.strip())[:1500],
                             "sample_documents/fail-1.docx (paragraphs, first 1500 chars)"),
        "txt:sample_text.txt": ("real_sample_text", (SAMPLES / "sample_text.txt").read_text(encoding="utf-8")[:1500],
                                "sample_documents/sample_text.txt"),
    }
    for key, (spike_id, text, source) in strings.items():
        name = key.split(":", 1)[1]
        bases[key] = {"text": text, "group": f"sample:{name}", "family": _SAMPLE_FAMILY.get(name, f"sample:{name}")}
        items.append(_item(f"lg-spike-{spike_id}", bases[key]["group"], f"spikeA:{spike_id} <- {source}",
                           "real_document_text", 1, text, _string_spans(text), family=bases[key]["family"], route=None,
                           pipeline_asks=None))
    for spike_id, text in SPIKE_A_HARD.items():
        bases[f"hard:{spike_id}"] = {"text": text, "group": f"spikeA:{spike_id}", "family": f"spikeA:{spike_id}"}
        items.append(_item(f"lg-spike-{spike_id}", f"spikeA:{spike_id}", f"spikeA:{spike_id}", "hard_legible_string", 1,
                           text, _string_spans(text), route=None, pipeline_asks=None))
    # spike A garbage on the same bases (the deck-derived ones are only in deck_items)
    for spike_id, kind, key, factory, line_wise in SPIKE_A_GARBAGE:
        base = bases[key]
        if "rows" in base:
            rows = _transform_rows(base["rows"], factory, line_wise)
            text, spans = _join(rows), [[t, s, f] for line in rows for t, s, f in line]
        else:
            text = _transform_text(base["text"], factory, line_wise)
            spans = _string_spans(text)
        items.append(_item(f"lg-spike-{spike_id}", base["group"], f"spikeA:{spike_id} <- {key}", kind, 0, text, spans,
                           family=base["family"], route=None, pipeline_asks=None))
    # generated PDFs
    for suffix, kind, label, key, builder in GENERATED:
        path = builder(tmp / f"lg-{suffix}.pdf", key)
        [page] = measure_pdf(path)
        items.append(_item(f"lg-gen-{suffix}", f"gen:{key}", f"generated:{suffix}", kind, label, page["text"],
                           page["spans"], route=page["route"], pipeline_asks=page["pipeline_asks"]))
    ids = [item["id"] for item in items]
    if len(set(ids)) != len(ids):
        raise ValueError("duplicate legibility ids")
    return items


# ======================================================================== boilerplate

_EN_WORDS = ("growth margin customer region product platform investment strategy capital supply demand pricing "
             "inventory logistics partner contract service quality risk audit compliance digital cloud security "
             "revenue cost forecast budget market share plant order shipment warranty").split()
_DE_WORDS = ("Umsatz Kunde Lieferung Vertrag Qualität Anlage Werkzeug Auftrag Rechnung Bestand Planung Prüfung "
             "Wartung Leistung Termin Bericht Ergebnis Kosten Maschine Standort").split()
_ZH_POOL = "本公司於年度持續投資核心平台並維持營運成本穩定管理階層認為現行策略能使企業具備競爭優勢與長期成長動能"
_JA_POOL = "当社は本年度も品質の向上と生産性の改善に取り組み新製品の開発を計画通りに進めることができた"


def _body(pdf: Pdf, page, lang: str, seed: str, top: float, bottom: float, size: float) -> None:
    rnd = random.Random(_hash(seed))
    y, left = top, 60 if page.rect.width < 700 else 70
    while y < bottom:
        if lang in ("zh", "ja"):
            pool = _ZH_POOL if lang == "zh" else _JA_POOL
            chars = int((page.rect.width - 2 * left) / size) - 2
            pdf.text(page, left, y, "".join(rnd.choice(pool) for _ in range(chars)) + "。", size,
                     "zh-t" if lang == "zh" else "ja")
        else:
            words = _DE_WORDS if lang == "de" else _EN_WORDS
            line, width = [], 0.0
            limit = page.rect.width - 2 * left - 20
            while True:
                word = rnd.choice(words)
                extra = _font("helv").text_length(word + " ", fontsize=size)
                if width + extra > limit:
                    break
                line.append(word)
                width += extra
            pdf.text(page, left, y, " ".join(line), size, "sans")
        y += size * 1.45


# A line placed on pages: (y baseline, x, text(p, n) -> str | None, size, font, kind, label).
# y < 0 counts from the bottom edge.
LineSpec = Tuple[float, float, Callable[[int, int], Optional[str]], float, str, str, int]


def _const(text: str, on=None):
    return lambda p, n: text if on is None or p in on else None


BP_DOCS: List[dict] = []


def _doc(name, lang, pages, lines, size=A4, body_size=10.5, body_top=None, body_bottom=None, family=None):
    """One generated document. ``family`` names the template it is a variant of (its
    translations and brand or name substitutions share it, and so share a split)."""
    BP_DOCS.append({"name": name, "lang": lang, "pages": pages, "lines": lines, "size": size,
                    "body_size": body_size, "body_top": body_top, "body_bottom": body_bottom,
                    "family": family or _BP_FAMILY.get(name, name)})


# What must remain in the Markdown of a line repeated on several pages: ONE copy (keep only the
# first: branding, logo text, taglines, marks, copyright, contact lines, print stamps, a brand with
# a page number) or ALL copies (a reader needs it where it appears: per-page labels, titles, unit
# notes, disclaimers, table header cells, identifiers that belong to the content). label = 1 / 0.
ONE, ALL = 1, 0
COPIES = {ONE: "one", ALL: "all"}

# Documents written as variants of one template (translations, brand or name substitutions).
_BP_FAMILY = {
    **dict.fromkeys(("statement-header", "de-bilanz", "zh-statement", "ja-statement"), "statement-header"),
    **dict.fromkeys(("brand-header", "zh-brand"), "brand-header"),
    **dict.fromkeys(("classification-row", "ja-report", "zh-cn-report"), "brand-and-mark"),
    **dict.fromkeys(("confidential-footer", "de-vertraulich", "zh-confidential", "slides-confidential"),
                    "confidential-mark"),
    **dict.fromkeys(("copyright-footer", "de-copyright", "zh-copyright", "ja-copyright"), "copyright-footer"),
    **dict.fromkeys(("disclaimer-footer", "de-hinweis", "zh-disclaimer", "forward-looking", "notes-footer"),
                    "disclaimer-footer"),
    **dict.fromkeys(("letterhead", "de-letterhead", "url-footer"), "letterhead"),
    **dict.fromkeys(("account-header", "customer-header", "patient-header"), "identifier-header"),
    **dict.fromkeys(("lesson-footer", "zh-lesson", "de-lektion"), "lesson-label"),
    **dict.fromkeys(("exhibit-header", "exhibit-plain"), "exhibit-label"),
    **dict.fromkeys(("invoice-ids", "daily-log"), "per-page-id"),
    **dict.fromkeys(("slides-number-brand", "de-folien", "ja-slides"), "slide-brand-number"),
    **dict.fromkeys(("slides-slide-of", "slides-slide-plain", "de-folie"), "slide-number"),
    **dict.fromkeys(("slides-tagline", "event-footer"), "slide-tagline"),
    **dict.fromkeys(("section-few-pages", "slides-section-few", "two-page-title"), "few-pages-title"),
    **dict.fromkeys(("brand-two-pages", "two-page-brand"), "few-pages-brand"),
    **dict.fromkeys(("table-header-rows", "de-tabellenkopf"), "table-header"),
}
_PAGE_OF = lambda p, n: f"Page {p} of {n}"  # noqa: E731 (a page number: the rule removes it, never asked)

# --- English reports --------------------------------------------------------------------------
_doc("brand-header", "en", 6, [(40, 60, _const("Northwind Traders"), 9, "sans", "brand_name", ONE),
                               (-30, 270, _PAGE_OF, 9, "sans", "page_number", ONE)])
_doc("confidential-footer", "en", 5, [(-40, 60, _const("Confidential – Internal Use Only"), 8, "sans",
                                       "confidentiality_mark", ONE),
                                      (-40, 500, lambda p, n: str(p), 8, "sans", "page_number", ONE)])
_doc("copyright-footer", "en", 7, [(-32, 60, _const("© 2026 Fabrikam, Inc. All rights reserved."), 8, "sans",
                                    "copyright", ONE)])
_doc("url-footer", "en", 4, [(-32, 60, _const("www.northwind.example | +1 (425) 555-0100 | info@northwind.example"), 8,
                              "sans", "contact_line", ONE)])
_doc("statement-header", "en", 5, [(38, 230, _const("Contoso Ltd."), 11, "bold", "brand_name", ONE),
                                   (54, 200, _const("Consolidated Balance Sheet"), 11, "bold", "statement_title", ALL),
                                   (69, 225, _const("(in thousands of USD)"), 9, "sans", "unit_note", ALL),
                                   (-30, 290, lambda p, n: str(p), 9, "sans", "page_number", ONE)])
_doc("disclaimer-footer", "en", 6, [(-36, 60, _const("Past performance is not a reliable indicator of future results."),
                                     8, "sans", "disclaimer", ALL)])
_doc("notes-footer", "en", 4, [(-36, 60, _const("The accompanying notes are an integral part of these financial "
                                                "statements."), 8, "sans", "legal_note", ALL)])
_doc("draft-mark", "en", 5, [(38, 60, _const("DRAFT – for internal review"), 9, "bold", "marking", ONE)])
_doc("print-stamp", "en", 3, [(-28, 60, _const("Printed 2026-09-12 14:03 by jdoe"), 7, "sans", "print_timestamp", ONE)])
_doc("account-header", "en", 6, [(40, 60, _const("Account No. 4471-0098-22 · Statement period August 2026"), 9, "sans",
                                  "account_identifier", ALL)])
_doc("customer-header", "en", 4, [(40, 60, _const("Customer: Globex Corporation (ID 88213)"), 9, "sans",
                                   "account_identifier", ALL)])
_doc("section-few-pages", "en", 6, [(40, 60, _const("Adatum Research"), 9, "sans", "brand_name", ONE),
                                    (80, 60, _const("Summary of Findings", on=(1, 4)), 16, "bold", "section_title",
                                     ALL)])
_doc("brand-two-pages", "en", 8, [(40, 60, _const("Lamna Healthcare", on=(2, 6)), 9, "sans", "brand_name", ONE)])
_doc("letterhead", "en", 4, [(34, 60, _const("Fabrikam Legal LLP"), 10, "bold", "brand_name", ONE),
                             (48, 60, _const("1200 Harbor Blvd, Suite 400 · Seattle, WA 98101"), 8, "sans",
                              "contact_line", ONE),
                             (60, 60, _const("Tel +1 206 555 0199 · fabrikam-legal.example"), 8, "sans", "contact_line",
                              ONE)])
_doc("lesson-footer", "en", 6, [(-30, 72, lambda p, n: f"Lesson · {p}", 9, "sans", "per_page_label", ALL),
                                (-30, 500, lambda p, n: str(p), 9, "sans", "page_number", ONE)])
_doc("exhibit-header", "en", 5, [(40, 60, lambda p, n: f"Exhibit – {p + 3}", 10, "sans", "per_page_label", ALL)])
_doc("exhibit-plain", "en", 5, [(40, 60, lambda p, n: f"Exhibit {p}", 10, "sans", "per_page_label", ALL)])
_doc("invoice-ids", "en", 5, [(40, 60, lambda p, n: f"Invoice · {1000 + p}", 10, "sans", "per_page_id", ALL),
                              (56, 60, lambda p, n: f"Invoice No. INV-{7000 + p}", 9, "sans", "per_page_id", ALL)])
_doc("daily-log", "en", 4, [(40, 60, lambda p, n: f"Site log – {p + 7} September 2026", 10, "sans", "per_page_date",
                             ALL)])
_doc("step-of", "en", 6, [(40, 60, lambda p, n: f"Step {p} of {n}", 10, "sans", "per_page_label", ALL)])
_doc("classification-row", "en", 5, [(38, 60, _const("Contoso Ltd."), 9, "sans", "brand_name", ONE),
                                     (38, 470, _const("INTERNAL"), 9, "bold", "marking", ONE)])
_doc("brand-title-row", "en", 6, [(44, 60, _const("Contoso Labs"), 9, "sans", "brand_name", ONE),
                                  (46, 300, lambda p, n: ["Overview", "Market", "Pipeline", "Operations", "Finance",
                                                          "Outlook"][p - 1] + " review", 14, "bold", "section_title",
                                   ALL)])
_doc("header-no-gap", "en", 5, [(66, 60, _const("Northwind Traders"), 9, "sans", "brand_name", ONE)], body_top=78)
_doc("unit-note-footer", "en", 5, [(-36, 60, _const("All amounts in EUR thousands unless otherwise stated."), 8, "sans",
                                    "unit_note", ALL)])
_doc("forward-looking", "en", 3, [(-36, 60, _const("This report contains forward-looking statements that involve "
                                                   "risks."), 8, "sans", "disclaimer", ALL)])
_doc("chapter-headers", "en", 6, [(40, 60, lambda p, n: "Chapter 1 – Market Review" if p <= 3 else
                                   "Chapter 2 – Operations", 9, "sans", "section_title", ALL)])
_doc("patient-header", "en", 4, [(40, 60, _const("Patient: Jane Roe · MRN 00412877 · DOB 1979-03-02"), 9, "sans",
                                  "account_identifier", ALL)])
_doc("two-page-brand", "en", 2, [(40, 60, _const("Adatum Corporation"), 9, "sans", "brand_name", ONE),
                                 (-32, 60, _const("adatum.example · +44 20 7946 0000"), 8, "sans", "contact_line",
                                  ONE)])
_doc("two-page-title", "en", 2, [(40, 60, _const("Quarterly Operations Report"), 9, "sans", "section_title", ALL)])
_doc("event-footer", "en", 5, [(-30, 60, _const("Contoso Summit 2026 · Seattle · 14-16 October"), 8, "sans",
                                "brand_tagline", ONE)], size=SLIDE, body_size=18)
_doc("slides-confidential", "en", 6, [(-24, 36, _const("Contoso Labs Confidential"), 10, "sans",
                                       "confidentiality_mark", ONE),
                                      (-24, 900, lambda p, n: str(p), 10, "sans", "page_number", ONE)],
     size=SLIDE, body_size=18)
_doc("slides-number-brand", "en", 6, [(-24, 36, lambda p, n: f"{p} | Contoso Labs", 10, "sans", "page_number_label",
                                       ONE)], size=SLIDE, body_size=18)
_doc("slides-slide-of", "en", 5, [(-24, 860, lambda p, n: f"Slide {p} / {n}", 10, "sans", "page_number_label",
                                   ONE)], size=SLIDE, body_size=18)
_doc("slides-slide-plain", "en", 5, [(-24, 880, lambda p, n: f"Slide {p}", 10, "sans", "page_number_label", ONE)],
     size=SLIDE, body_size=18)
_doc("slides-tagline", "en", 7, [(-24, 36, _const("Innovation that scales"), 10, "sans", "brand_tagline", ONE)],
     size=SLIDE, body_size=18)
_doc("slides-section-few", "en", 8, [(44, 60, _const("Roadmap", on=(2, 6)), 26, "bold", "section_title", ALL)],
     size=SLIDE, body_size=18)
_doc("table-header-rows", "en", 5, [(40, 60, _const("Contoso Bank"), 9, "sans", "brand_name", ONE),
                                    (96, 60, _const("Date"), 10, "bold", "table_header_row", ALL),
                                    (96, 200, _const("Description"), 10, "bold", "table_header_row", ALL),
                                    (96, 400, _const("Amount"), 10, "bold", "table_header_row", ALL),
                                    (96, 480, _const("Balance"), 10, "bold", "table_header_row", ALL)],
     body_top=110)
# --- German --------------------------------------------------------------------------------------
_doc("de-letterhead", "de", 4, [(36, 60, _const("Müller & Söhne GmbH · Industriestraße 12 · 70565 Stuttgart"), 9,
                                 "sans", "contact_line", ONE),
                                (-34, 60, _const("Geschäftsführer: Hans Müller · Amtsgericht Stuttgart HRB 12345"), 7,
                                 "sans", "company_imprint", ONE)])
_doc("de-vertraulich", "de", 5, [(40, 60, _const("Streng vertraulich"), 9, "bold", "confidentiality_mark", ONE),
                                 (-30, 270, lambda p, n: f"Seite {p} von {n}", 9, "sans", "page_number", ONE)])
_doc("de-bilanz", "de", 4, [(38, 230, _const("Fabrikam AG"), 11, "bold", "brand_name", ONE),
                            (54, 190, _const("Konzernbilanz zum 31. Dezember 2025"), 11, "bold", "statement_title",
                             ALL),
                            (69, 240, _const("(in Tausend EUR)"), 9, "sans", "unit_note", ALL)])
_doc("de-hinweis", "de", 5, [(-34, 60, _const("Alle Angaben ohne Gewähr. Irrtümer und Änderungen vorbehalten."), 8,
                              "sans", "disclaimer", ALL)])
_doc("de-copyright", "de", 6, [(-32, 60, _const("© 2026 Fabrikam AG. Alle Rechte vorbehalten."), 8, "sans",
                                "copyright", ONE)])
_doc("de-folien", "de", 6, [(-24, 36, lambda p, n: f"Fabrikam · {p}", 10, "sans", "page_number_label", ONE)],
     size=SLIDE, body_size=18)
_doc("de-tabellenkopf", "de", 4, [(96, 60, _const("Datum"), 10, "bold", "table_header_row", ALL),
                                  (96, 200, _const("Beschreibung"), 10, "bold", "table_header_row", ALL),
                                  (96, 400, _const("Betrag"), 10, "bold", "table_header_row", ALL),
                                  (96, 480, _const("Saldo"), 10, "bold", "table_header_row", ALL)], body_top=110)
# --- Chinese ------------------------------------------------------------------------------------
_doc("zh-brand", "zh", 6, [(40, 60, _const("北辰精密工業股份有限公司"), 9, "zh-t", "brand_name", ONE),
                           (-30, 280, lambda p, n: f"第 {p} 頁", 9, "zh-t", "page_number", ONE)])
_doc("zh-confidential", "zh", 5, [(-34, 60, _const("機密文件 請勿外流"), 8, "zh-t", "confidentiality_mark", ONE)])
_doc("zh-statement", "zh", 4, [(38, 240, _const("北辰精密工業股份有限公司"), 11, "zh-t", "brand_name", ONE),
                               (54, 260, _const("合併資產負債表"), 11, "zh-t", "statement_title", ALL),
                               (69, 262, _const("單位：新台幣千元"), 9, "zh-t", "unit_note", ALL)])
_doc("zh-disclaimer", "zh", 5, [(-34, 60, _const("本資料僅供參考，不構成任何投資建議。"), 8, "zh-t", "disclaimer", ALL)])
_doc("zh-copyright", "zh", 4, [(-32, 60, _const("版權所有 © 2026 北辰精密"), 8, "zh-t", "copyright", ONE)])
_doc("zh-cn-report", "zh", 5, [(40, 60, _const("华东新能源科技有限公司"), 9, "zh-s", "brand_name", ONE),
                               (-34, 60, _const("内部资料 注意保密"), 8, "zh-s", "confidentiality_mark", ONE)])
_doc("zh-lesson", "zh", 5, [(40, 60, lambda p, n: f"第{p}課 基礎會計", 10, "zh-t", "per_page_label", ALL)])
# --- Japanese ------------------------------------------------------------------------------------
_doc("ja-report", "ja", 5, [(40, 60, _const("株式会社ミナト精機"), 9, "ja", "brand_name", ONE),
                            (40, 480, _const("社外秘"), 9, "ja", "confidentiality_mark", ONE)])
_doc("ja-statement", "ja", 4, [(54, 250, _const("連結貸借対照表"), 11, "ja", "statement_title", ALL),
                               (69, 250, _const("（単位：百万円）"), 9, "ja", "unit_note", ALL)])
_doc("ja-copyright", "ja", 6, [(-32, 60, _const("Copyright © 2026 Minato Seiki Co., Ltd. All Rights Reserved."), 7,
                                "sans", "copyright", ONE)])
_doc("ja-slides", "ja", 6, [(-24, 36, lambda p, n: f"ミナト精機 | {p}", 10, "ja", "page_number_label", ONE)],
     size=SLIDE, body_size=18)

# --- Lines the rule keeps and asks about: rows it cannot peel off (attached_to_content), lines
# on too few pages (few_pages), numbered labels (numbered_label); every template in its own family
# unless it is a variant of another.
_ROW_TITLES = {"zh": ["營運概況", "財務摘要", "倉儲網路", "車隊管理", "永續發展", "未來展望"],
               "ja": ["事業概要", "売上推移", "生産体制", "品質保証", "今後の計画"]}
_doc("zh-brand-title-row", "zh", 6, [(44, 60, _const("遠翔物流"), 9, "zh-t", "brand_name", ONE),
                                     (46, 260, lambda p, n: _ROW_TITLES["zh"][p - 1], 14, "zh-t", "section_title", ALL)],
     family="brand-title-row")
_doc("ja-brand-title-row", "ja", 5, [(44, 60, _const("ひかり電機"), 9, "ja", "brand_name", ONE),
                                     (46, 250, lambda p, n: _ROW_TITLES["ja"][p - 1], 14, "ja", "section_title", ALL),
                                     (44, 480, _const("社外秘"), 9, "ja", "confidentiality_mark", ONE)],
     family="brand-title-row")
_NOTE_TITLES = {"en": ["Revenue", "Leases", "Income taxes", "Inventories", "Provisions"],
                "de": ["Umsatzerlöse", "Leasingverhältnisse", "Ertragsteuern", "Vorräte", "Rückstellungen"],
                "zh": ["營業收入", "租賃", "所得稅", "存貨", "負債準備"],
                "ja": ["売上収益", "リース", "法人所得税", "棚卸資産", "引当金"]}
for _lang, _face, _unit in (("en", "sans", "(in USD thousands)"), ("de", "sans", "(Angaben in TEUR)"),
                            ("zh", "zh-t", "（單位：新臺幣百萬元）"), ("ja", "ja", "（単位：千円）")):
    _doc(f"{_lang}-notes-unit-row", _lang, 5,
         [(46, 60, (lambda titles: lambda p, n: titles[p - 1])(_NOTE_TITLES[_lang]), 12,
           "bold" if _face == "sans" else _face, "section_title", ALL),
          (46, 400, _const(_unit), 9, _face, "unit_note", ALL)], family="notes-unit-row")
_doc("zh-price-table", "zh", 4, [(96, 60, _const("品名"), 10, "zh-t", "table_header_row", ALL),
                                 (96, 300, _const("單價"), 10, "zh-t", "table_header_row", ALL),
                                 (96, 450, _const("數量"), 10, "zh-t", "table_header_row", ALL)],
     body_top=110, family="price-table-header")
_doc("ja-price-table", "ja", 4, [(96, 60, _const("品番"), 10, "ja", "table_header_row", ALL),
                                 (96, 300, _const("単価"), 10, "ja", "table_header_row", ALL),
                                 (96, 450, _const("数量"), 10, "ja", "table_header_row", ALL)],
     body_top=110, family="price-table-header")
for _name, _lang, _face, _text in (("de-konto-no-gap", "de", "sans", "Konto 4471 0098 22 · Kontoauszug August 2026"),
                                   ("zh-account-no-gap", "zh", "zh-t", "帳號 0123-456-789 對帳單期間 2026年8月"),
                                   ("ja-account-no-gap", "ja", "ja", "口座番号 1234567 2026年8月分お取引明細")):
    _doc(_name, _lang, 5, [(66, 60, _const(_text), 9, _face, "account_identifier", ALL)], body_top=78,
         family="account-no-gap")
for _name, _lang, _face, _text in (("de-header-no-gap", "de", "sans", "Keller Präzisionstechnik GmbH"),
                                   ("zh-header-no-gap", "zh", "zh-t", "華岳建設股份有限公司"),
                                   ("ja-header-no-gap", "ja", "ja", "さくら精工株式会社")):
    _doc(_name, _lang, 5, [(66, 60, _const(_text), 9, _face, "brand_name", ONE)], body_top=78, family="header-no-gap")
for _name, _lang, _face, _text in (("footer-no-gap-mark", "en", "sans", "Trey Research Confidential"),
                                   ("de-footer-no-gap-mark", "de", "sans", "Vertraulich – nur für den internen Gebrauch"),
                                   ("zh-footer-no-gap-mark", "zh", "zh-t", "機密資料 未經授權不得複製"),
                                   ("ja-footer-no-gap-mark", "ja", "ja", "社外秘 無断転載禁止")):
    _doc(_name, _lang, 5, [(-60, 60, _const(_text), 8, _face, "confidentiality_mark", ONE)], body_bottom=770,
         family="footer-no-gap-mark")
for _name, _lang, _face, _text in (("continued-footer", "en", "sans", "(continued on next page)"),
                                   ("de-continued-footer", "de", "sans", "(Fortsetzung auf der nächsten Seite)"),
                                   ("zh-continued-footer", "zh", "zh-t", "（續下頁）"),
                                   ("ja-continued-footer", "ja", "ja", "（次ページへ続く）")):
    _doc(_name, _lang, 5, [(-60, 60, _const(_text, on=(1, 2, 3, 4)), 8, _face, "continuation_note", ALL)],
         body_bottom=770, family="continued-footer")
_doc("footer-no-gap-contact", "en", 5, [(-60, 60, _const("Tailspin Toys · 1 Toy Street, Leeds LS1 4AP · "
                                                         "tailspintoys.example"), 8, "sans", "contact_line", ONE)],
     body_bottom=770, family="footer-no-gap-contact")
_doc("de-footer-no-gap-contact", "de", 5, [(-60, 60, _const("Weber Antriebstechnik GmbH · Hafenstraße 9 · 28217 Bremen"),
                                            8, "sans", "contact_line", ONE)], body_bottom=770,
     family="footer-no-gap-contact")
_doc("de-lektion", "de", 5, [(-30, 72, lambda p, n: f"Lektion – {p}", 9, "sans", "per_page_label", ALL)])
_doc("unit-footer", "en", 6, [(-30, 72, lambda p, n: f"Unit · {p}", 9, "sans", "per_page_label", ALL)], family="unit-label")
_doc("zh-unit-footer", "zh", 5, [(-30, 72, lambda p, n: f"單元 · {p}", 9, "zh-t", "per_page_label", ALL)],
     family="unit-label")
_doc("ja-unit-footer", "ja", 5, [(-30, 72, lambda p, n: f"講義 · {p}", 9, "ja", "per_page_label", ALL)], family="unit-label")
_doc("report-brand-number", "en", 6, [(40, 400, lambda p, n: f"Woodgrove Bank | {p}", 9, "sans", "page_number_label",
                                       ONE)], family="report-brand-number")
_doc("zh-report-brand-number", "zh", 5, [(40, 420, lambda p, n: f"聯華電訊 | {p}", 9, "zh-t", "page_number_label", ONE)],
     family="report-brand-number")
_doc("de-folie", "de", 5, [(-24, 860, lambda p, n: f"Folie {p} / {n}", 10, "sans", "page_number_label", ONE)],
     size=SLIDE, body_size=18)
_doc("de-chapter-pair", "de", 8, [(40, 60, _const("Anhang B – Messprotokolle", on=(6, 7)), 11, "bold", "section_title",
                                   ALL)], family="chapter-title-pair")
_doc("zh-chapter-pair", "zh", 8, [(40, 60, _const("第四章 風險管理", on=(3, 4)), 11, "zh-t", "section_title", ALL)],
     family="chapter-title-pair")
_doc("ja-chapter-pair", "ja", 7, [(40, 60, _const("第5章 設備投資計画", on=(2, 3)), 11, "ja", "section_title", ALL)],
     family="chapter-title-pair")
_doc("unit-note-few", "en", 8, [(-36, 60, _const("(in thousands of USD, unaudited)", on=(3, 4, 5)), 8, "sans",
                                 "unit_note", ALL)])
_doc("de-konto-few", "de", 6, [(40, 60, _const("Kontoauszug Nr. 9/2026 · Konto 5520 1187 03", on=(2, 5)), 9, "sans",
                                "account_identifier", ALL)])
_doc("de-mark-pair", "de", 8, [(40, 60, _const("Nur für den Dienstgebrauch", on=(3, 4)), 9, "bold",
                                "confidentiality_mark", ONE)])
_doc("zh-copyright-pair", "zh", 8, [(-32, 60, _const("版權所有 翻印必究", on=(7, 8)), 8, "zh-t", "copyright", ONE)])
_doc("ja-notice-pair", "ja", 6, [(-32, 60, _const("無断転載を禁じます", on=(1, 6)), 8, "ja", "copyright", ONE)])
_doc("website-pair", "en", 6, [(-32, 60, _const("www.lucernepublishing.example", on=(1, 6)), 8, "sans", "contact_line",
                                ONE)])
_doc("stamp-pair", "en", 5, [(-28, 60, _const("Printed 2026-09-18 09:12 by mlopez", on=(2, 3)), 7, "sans",
                              "print_timestamp", ONE)])
_doc("zh-two-page-quote", "zh", 2, [(40, 60, _const("東岳貿易股份有限公司"), 9, "zh-t", "brand_name", ONE),
                                    (40, 400, _const("報價單 Q-2026-118"), 9, "zh-t", "document_id", ALL)],
     family="two-page-quote")
_doc("ja-two-page-quote", "ja", 2, [(40, 60, _const("あおば商事株式会社 営業部"), 9, "ja", "brand_name", ONE),
                                    (40, 400, _const("見積書 No. 2026-0457"), 9, "ja", "document_id", ALL)],
     family="two-page-quote")

# Deck-like documents (16:9, the brand line under the logo text, beside a numbered per-slide label).
DECKS = [
    # (name, lang, brand line, logo text, slide titles)
    ("deck-fabrikam", "en", "by Fabrikam Robotics", "FABRIKAM",
     ["Why automate", "Cells and robots", "Vision system", "Safety", "Service plans", "References"]),
    ("deck-zh", "zh", "by 北辰雲端科技", "BC",
     ["導入挑戰", "平台架構", "資安防護", "報表中心", "權限管理", "成功實績", "價格方案"]),
    ("deck-ja", "ja", "by ミナト精機", "MINATO", ["会社概要", "製品紹介", "品質管理", "導入事例", "サポート"]),
    ("deck-de", "en", "by Müller Automation", "MÜLLER",
     ["Ausgangslage", "Lösung", "Architektur", "Projektplan", "Kosten", "Referenzen", "Kontakt", "Anhang"]),
]
_DECK_LANG = {"deck-de": "de"}   # German titles drawn with the Latin font


def _deck_pdf(path: Path, lang: str, brand: str, logo: str, titles: Sequence[str], registry: dict,
              placed: dict) -> Path:
    pdf = Pdf()
    font = {"zh": "zh-t", "ja": "ja"}.get(lang, "sans")
    for number, title in enumerate(titles, 1):
        page = pdf.page((1440, 810))
        page.draw_rect(pymupdf.Rect(40, 30, 88, 74), color=(0.1, 0.3, 0.6), fill=(0.1, 0.3, 0.6))
        pdf.text(page, 98, 56, logo, 19, "bold")
        pdf.text(page, 98, 71, brand, 8.6, font)
        label = f"{number:02d} / {title}"
        width = len(label) * 20 if font != "sans" else _font("helv").text_length(label, fontsize=20)
        pdf.text(page, 1440 - 52 - width, 62, label, 20, font)
        pdf.text(page, 60, 190, title, 48, font if font != "sans" else "bold")
        _body(pdf, page, lang, f"{path.name}-{number}", 290, 700, 20)
        registry[logo] = ("logo_text", ONE)
        registry[brand] = ("brand_tagline", ONE)
        registry[label] = ("per_page_label", ALL)
        for text in (logo, brand, label):
            placed.setdefault(text, set()).add(number)
    return pdf.save(path)


def _bp_pdf(path: Path, spec: dict, registry: dict, placed: dict) -> Path:
    pdf = Pdf()
    width, height = spec["size"]
    body_size = spec["body_size"]
    top = spec["body_top"] or (130 if height > 600 else 120)
    bottom = spec["body_bottom"] or (height - 130 if height > 600 else height - 110)
    for p in range(1, spec["pages"] + 1):
        page = pdf.page(spec["size"])
        for y, x, text_fn, size, font, kind, label in spec["lines"]:
            text = text_fn(p, spec["pages"])
            if text is None:
                continue
            pdf.text(page, x, y if y > 0 else height + y, text, size, font)
            registry[text] = (kind, label)
            placed.setdefault(text, set()).add(p)
        _body(pdf, page, spec["lang"], f"{spec['name']}-{p}", top, bottom, body_size)
    return pdf.save(path)


def _capture(path: Path, registry: dict, oracle: bool) -> Tuple[List[Tuple[str, dict]], dict]:
    """The questions the pipeline asks a recording judge (None, or the intended answer with
    ``oracle``), and the converted document."""
    from doc2mark.pipelines.pymupdf_advanced_pipeline import pdf_to_simple_json
    calls = []

    def judge(text, context):
        calls.append((text, dict(context)))
        if not oracle:
            return None
        return 1.0 if _lookup(registry, text)[1] == ONE else 0.0

    output = pdf_to_simple_json(str(path), extract_images=False, ocr_images=False, show_progress=False,
                                boilerplate_judge=judge)
    return calls, output


def _nfkc(text: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", text).split())


def _copies_left(output: dict, placed: dict) -> Dict[str, int]:
    """Per placed line: on how many of its pages it is still content (not typed text:header or
    text:footer, i.e. in the Markdown)."""
    pages: Dict[int, List[str]] = defaultdict(list)
    for entry in output["content"]:
        if entry.get("type") not in ("text:header", "text:footer") and isinstance(entry.get("content"), str):
            pages[entry["page"]].append(entry["content"])
    text_of = {page: _nfkc(" ".join(parts)) for page, parts in pages.items()}
    return {text: sum(1 for page in on if _nfkc(text) in text_of.get(page, "")) for text, on in placed.items()}


def _lookup(registry: dict, text: str) -> Tuple[str, int]:
    if text in registry:
        return registry[text]
    norm = " ".join(unicodedata.normalize("NFKC", text).split())
    for placed, value in registry.items():
        if " ".join(unicodedata.normalize("NFKC", placed).split()) == norm:
            return value
    raise KeyError(f"the pipeline asked about a line no generator placed: {text!r}")


def _masked(text: str) -> str:
    return re.sub(r"[0-9]+", "#", _nfkc(text))


def build_boilerplate(tmp: Path, report: Optional[list] = None) -> List[dict]:
    """The questions the pipeline asks about repeated lines (reasons attached_to_content,
    few_pages, numbered_label), labelled with how many copies must remain in the Markdown.
    ``report`` receives every designed line it does not ask about, with the copies the rule
    leaves in the output."""
    sys.path.insert(0, str(ROOT))
    from tests.e2e import builders_judge
    jobs = []
    for spec in BP_DOCS:
        registry, placed = {}, {}
        path = _bp_pdf(tmp / f"bp-{spec['name']}.pdf", spec, registry, placed)
        jobs.append((spec["name"], spec["family"], spec["lang"], f"generated:{spec['name']}", path, registry, placed))
    for name, lang, brand, logo, titles in DECKS:
        registry, placed = {}, {}
        path = _deck_pdf(tmp / f"bp-{name}.pdf", lang, brand, logo, titles, registry, placed)
        jobs.append((name, "deck", _DECK_LANG.get(name, lang), f"generated:{name}", path, registry, placed))
    slides = set(range(1, len(builders_judge.DECK_SLIDES) + 1))
    registry = {builders_judge.BRAND_LINE: ("brand_tagline", ONE), builders_judge.LOGO_TEXT: ("logo_text", ONE)}
    placed = {builders_judge.BRAND_LINE: set(slides), builders_judge.LOGO_TEXT: set(slides)}
    for number, (title, _) in enumerate(builders_judge.DECK_SLIDES, 1):
        registry[f"{number:02d} / {title}"] = ("per_page_label", ALL)
        placed[f"{number:02d} / {title}"] = {number}
    jobs.append(("e2e-deck", "deck", "en", "tests/e2e/builders_judge.py:deck_pdf",
                 builders_judge.deck_pdf(tmp / "bp-e2e-deck.pdf"), registry, placed))
    jobs.append(("sample-pdf", "sample:sample_pdf.pdf", "en", "sample_documents/sample_pdf.pdf",
                 SAMPLES / "sample_pdf.pdf", {}, {}))
    items, never = [], []
    for name, family, lang, source, path, registry, placed in jobs:
        asked, left = {}, {}
        for run in ("none", "oracle"):
            calls, output = _capture(path, registry, oracle=run == "oracle")
            if run == "none":
                left = _copies_left(output, placed)
            for text, context in calls:
                if context.get("reason") not in ("attached_to_content", "few_pages", "numbered_label"):
                    raise AssertionError(f"{name}: unexpected reason {context.get('reason')!r} for {text!r}")
                key = (text, json.dumps(context, sort_keys=True, ensure_ascii=False))
                asked.setdefault(key, []).append(run)
        covered = set()
        for index, ((text, context_json), runs) in enumerate(asked.items(), 1):
            kind, label = _lookup(registry, text)
            context = json.loads(context_json)
            items.append({"id": f"bp-{name}-{index}", "group": f"bp:{name}", "family": f"bp:{family}", "source": source,
                          "kind": kind, "lang": lang, "label": label, "expected_copies": COPIES[label], "split": None,
                          "asked": True, "text": text, "context": context,
                          "asked_in": "both" if len(set(runs)) == 2 else runs[0]})
            if context["reason"] == "numbered_label":   # one question for the whole numbered series
                covered |= {placed_text for placed_text, on in placed.items()
                            if _masked(placed_text) == _masked(text) and on & set(context["pages"])}
            covered |= {placed_text for placed_text in registry if _nfkc(placed_text) == _nfkc(text)}
        for text, (kind, label) in registry.items():
            if text not in covered:
                never.append({"doc": name, "kind": kind, "expected_copies": COPIES[label], "text": text,
                              "pages": len(placed.get(text, ())), "copies_left": left.get(text)})
    if report is not None:
        report.extend(never)
    return items


# ========================================================================== external set
#
# Fresh items written by the PR #22 reviewer (not by the generators above), never used for
# calibration: .executors/d2m-review-jev/probes/fresh_items.py, embedded verbatim below.


def _review_pr22_items():
    import codecs

    def rot13(s): return codecs.encode(s, "rot13")
    def rev_lines(s): return "\n".join(l[::-1] for l in s.split("\n"))
    def mojibake(s, enc_from, enc_to):
        return s.encode(enc_from).decode(enc_to, errors="ignore")

    EN_MEMO = ("Operations update - Week 38\nThe Kaohsiung warehouse shipped 4,812 orders this week, up 6 percent.\n"
               "Two forklifts are scheduled for maintenance on Thursday morning.\nPlease submit overtime requests by Friday noon.")
    ZH_NOTICE = "重大訊息公告\n主旨：本公司董事會決議通過第二季合併財務報告。\n說明：合併營業收入新台幣三十二億元，較去年同期成長百分之十二。"
    EL_TEXT = "Τιμολόγιο αρ. 88\nΕκδότης: Αιγαίο Λογισμικό Α.Ε.\nΣύνολο πληρωμής: 1.240,00 EUR\nΠληρωμή εντός 30 ημερών."
    RU_TEXT = "Счёт-фактура № 45 от 12.09.2026\nПродавец: ООО «Северный ветер»\nИтого к оплате: 118 000,00 руб."

    LEGIBILITY = [  # (id, label 1=legible, text)
        ("f-lg-lorem", 1, "Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor incididunt ut labore et dolore magna aliqua.\nUt enim ad minim veniam, quis nostrud exercitation ullamco laboris nisi ut aliquip ex ea commodo consequat.\nDuis aute irure dolor in reprehenderit in voluptate velit esse cillum dolore eu fugiat nulla pariatur."),
        ("f-lg-pem", 1, "Appendix C - Server certificate\n-----BEGIN CERTIFICATE-----\nMIIDdzCCAl+gAwIBAgIEAgAAuTANBgkqhkiG9w0BAQUFADBaMQswCQYDVQQGEwJJ\nRTESMBAGA1UEChMJQmFsdGltb3JlMRMwEQYDVQQLEwpDeWJlclRydXN0MSIwIAYD\nVQQDExlCYWx0aW1vcmUgQ3liZXJUcnVzdCBSb290MB4XDTAwMDUxMjE4NDYwMFoX\n-----END CERTIFICATE-----\nFingerprint (SHA-1): D4:DE:20:D0:5E:66:FC:53:FE:1A:50:88:2C:78:DB:28:52:CA:E4:74"),
        ("f-lg-hexdump", 1, "Firmware image header (offset 0x0000)\n00000000  7f 45 4c 46 02 01 01 00  00 00 00 00 00 00 00 00  |.ELF............|\n00000010  03 00 3e 00 01 00 00 00  60 10 00 00 00 00 00 00  |..>.....`.......|\n00000020  40 00 00 00 00 00 00 00  e8 3a 00 00 00 00 00 00  |@........:......|\nChecksum: 0x5A3C  Build: 2026.09.14-r3"),
        ("f-lg-medical", 1, "Progress note 09/12\nPt c/o SOB x3d, worse on exertion. Hx HTN, DM2, CKD3a.\nVitals: BP 152/94, HR 104, RR 22, SpO2 93% RA, T 37.9\nA/P: r/o CAP vs CHF exac. CXR, CBC w/ diff, BMP, BNP, trop q6h. Start ceftriaxone 1 g IV qd + azithro 500 mg PO."),
        ("f-lg-welsh", 1, "Cyngor Sir Ceredigion\nMae'r cyngor yn gwahodd ceisiadau am grantiau cymunedol ar gyfer prosiectau sy'n gwella llwybrau cerdded a chyfleusterau chwarae.\nDylid cyflwyno pob cais erbyn dydd Gwener, 31 Hydref 2026."),
        ("f-lg-vietnamese", 1, "HỢP ĐỒNG DỊCH VỤ\nĐiều 5. Quyền và nghĩa vụ của các bên\nBên A có trách nhiệm thanh toán đầy đủ và đúng hạn theo quy định tại Điều 3 của hợp đồng này.\nBên B cam kết bảo mật mọi thông tin liên quan đến khách hàng."),
        ("f-lg-thai", 1, "ประกาศบริษัท\nบริษัทจะปิดทำการในวันจันทร์ที่ 13 ตุลาคม 2569 เนื่องในวันหยุดราชการ\nกรุณาติดต่อฝ่ายบริการลูกค้าในวันทำการถัดไป"),
        ("f-lg-legal-cites", 1, "II. ARGUMENT\nA. Plaintiff states a claim under 42 U.S.C. § 1983.\nSee Monell v. Dep't of Soc. Servs., 436 U.S. 658, 690-91 (1978); cf. Pembaur v. City of Cincinnati, 475 U.S. 469, 480 (1986).\nId. at 481 n.9. Accord Fed. R. Civ. P. 12(b)(6)."),
        ("f-lg-chem", 1, "Safety Data Sheet - Section 3\nIbuprofen: 2-(4-isobutylphenyl)propanoic acid, C13H18O2, CAS 15687-27-1, EC 239-784-6\nMagnesium stearate: CAS 557-04-0, 0.5-1.0 % w/w\nHPMC E5: hydroxypropyl methylcellulose, CAS 9004-65-3"),
        ("f-lg-parts", 1, "PICK LIST PL-20260912-07\nSKU 4471-AX  QTY 12  BIN C-14-03  LOT 26J0912\nSKU 88-0192-R QTY 4   BIN A-02-11  LOT 26H3301\nSKU TX9-4L-BK QTY 40  BIN F-07-02  LOT 26G1187\nPicked by: KWL  Checked: ____"),
        ("f-lg-rot13", 0, rot13(EN_MEMO)),
        ("f-lg-reversed", 0, rev_lines(EN_MEMO)),
        ("f-lg-greek-cp1252", 0, mojibake(EL_TEXT, "cp1253", "cp1252")),
        ("f-lg-russian-cp1252", 0, mojibake(RU_TEXT, "cp1251", "cp1252")),
        ("f-lg-big5-as-gbk", 0, mojibake(ZH_NOTICE, "big5", "gbk")),
        ("f-lg-sjis-as-cp1252", 0, mojibake("第3章 品質管理体制\n当社は全工程において品質マネジメントシステムを運用している。", "shift_jis", "cp1252")),
    ]

    BOILERPLATE = [  # (id, label 1=chrome, line, context)
    ]
    def ctx(zone, n, total, fs, body, reason):
        pages = list(range(1, n + 1))
        return {"zone": zone, "pages": pages, "repeated_on": pages, "page_count": total, "font_size": fs,
                "body_font_size": body, "reason": reason}
    BOILERPLATE += [
        ("f-bp-notes-title", 0, "Notes to the Consolidated Financial Statements", ctx("header", 12, 14, 9.0, 10.0, "first_occurrence")),
        ("f-bp-unit-ntd", 0, "Amounts in NT$ thousands, except per-share data", ctx("header", 6, 6, 8.0, 10.0, "first_occurrence")),
        ("f-bp-prepared-for", 0, "Prepared for: Industrial Development Administration, Ministry of Economic Affairs", ctx("footer", 9, 10, 8.0, 10.5, "first_occurrence")),
        ("f-bp-project-rev", 0, "Project: Riverside Tower - Structural Calculations - Rev. B", ctx("header", 20, 22, 9.0, 10.0, "first_occurrence")),
        ("f-bp-court-case", 0, "Case 2:24-cv-01187-JLR Document 45 Filed 09/12/26", ctx("header", 18, 18, 10.0, 12.0, "first_occurrence")),
        ("f-bp-patient", 0, "陳大文 MRN 0048123 出生日期 1962-05-14", ctx("header", 4, 4, 9.0, 10.5, "first_occurrence")),
        ("f-bp-spec-section", 0, "Section 23 05 00 - Common Work Results for HVAC", ctx("footer", 7, 7, 8.0, 10.0, "first_occurrence")),
        ("f-bp-table-cont", 0, "Table 3 - Test results by batch (continued)", ctx("header", 3, 9, 10.0, 10.0, "attached_to_content")),
        ("f-bp-safety", 0, "WARNING: Disconnect power before servicing. Read all instructions before use.", ctx("footer", 10, 10, 8.0, 10.0, "first_occurrence")),
        ("f-bp-effective", 0, "Policy HR-07 Rev. 3.2 - Effective 1 July 2026", ctx("footer", 8, 8, 8.0, 10.5, "first_occurrence")),
        ("f-bp-exam", 0, "Student: Lin Yu-Chen   Exam: Physics 101 Midterm", ctx("header", 5, 5, 9.0, 11.0, "first_occurrence")),
        ("f-bp-contract-no", 0, "Contract No. C-2026-0457 between Contoso Ltd. and Fabrikam AG", ctx("footer", 15, 16, 8.0, 10.5, "first_occurrence")),
        ("f-bp-chapter-few", 0, "Chapter 4: Thermal Management", ctx("header", 2, 12, 9.0, 10.0, "few_pages")),
        ("f-bp-brand-pharma", 1, "Contoso Pharmaceuticals", ctx("header", 8, 8, 9.0, 10.5, "first_occurrence")),
        ("f-bp-proprietary", 1, "PROPRIETARY AND CONFIDENTIAL", ctx("footer", 12, 12, 7.5, 10.0, "first_occurrence")),
        ("f-bp-website", 1, "www.fabrikam.example", ctx("footer", 6, 6, 8.0, 10.5, "first_occurrence")),
        ("f-bp-recycled", 1, "Printed on 100% recycled paper", ctx("footer", 4, 4, 7.0, 10.0, "first_occurrence")),
        ("f-bp-zh-contact", 1, "台北市信義區松仁路100號 | 電話 (02) 2720-1234 | www.beichen.example", ctx("footer", 5, 5, 8.0, 10.5, "first_occurrence")),
        ("f-bp-copyright-short", 1, "© 2026 Contoso. Proprietary.", ctx("footer", 9, 9, 7.0, 10.0, "first_occurrence")),
        ("f-bp-slide-brand", 1, "Northwind | Investor Day 2026", ctx("footer", 14, 15, 10.0, 20.0, "attached_to_content")),
    ]

    NON_CONTENT = [  # (id, label 1=non-content, text)
        ("f-nc-qr", 0, "Scan the QR code to download the app"),
        ("f-nc-bsod", 0, "Your PC ran into a problem and needs to restart. We're just collecting some error info.\nStop code: INACCESSIBLE_BOOT_DEVICE"),
        ("f-nc-prop65", 0, "WARNING: This product can expose you to chemicals known to the State of California to cause cancer."),
        ("f-nc-fig-caption", 0, "Figure 5. No significant difference was observed between the groups (p = 0.42)."),
        ("f-nc-reconnecting", 0, "Connection lost. Reconnecting..."),
        ("f-nc-no-results", 0, "No results found for \"quarterly report 2019\""),
        ("f-nc-left-blank", 0, "This page has been intentionally left blank."),
        ("f-nc-access-denied", 0, "Access denied\nYou do not have permission to view this file. Request access from the owner."),
        ("f-nc-out-of-stock", 0, "Out of stock"),
        ("f-nc-printer", 0, "Unable to connect to printer.\nCheck the USB cable and try again."),
        ("f-nc-handwritten", 0, "Call Mr. Wang re: shipment - ext. 214"),
        ("f-nc-table", 0, "| Item | Qty |\n|---|---|\n| Bolt M8 x 40 | 200 |\n| Washer M8 | 400 |"),
        ("f-nc-feuerwehr", 0, "Feuerwehrzufahrt - Parken verboten"),
        ("f-nc-zh-maintenance", 0, "系統維護中，暫停服務，造成不便敬請見諒。"),
        ("f-nc-captcha", 0, "Please verify you are human.\nI'm not a robot"),
        ("f-nc-black", 1, "The image you've provided appears to be completely black; there's no discernible text."),
        ("f-nc-corrupted", 1, "Sorry - the file seems to be corrupted and I couldn't open it."),
        ("f-nc-nothing", 1, "Nothing to transcribe here."),
        ("f-nc-zh-clearer", 1, "这张图片的内容我无法识别，请提供更清晰的版本。"),
        ("f-nc-blurred", 1, "I can only see a blurred image; unfortunately no characters are recognizable."),
        ("f-nc-fr-none", 1, "Il n'y a aucun texte visible sur cette image."),
        ("f-nc-ja-none", 1, "該当するテキストは見つかりませんでした。"),
        ("f-nc-only-photo", 1, "The image only contains a photograph without any text, so there is nothing to transcribe."),
        ("f-nc-it-refusal", 1, "Mi dispiace, non riesco a leggere il testo di questa immagine."),
    ]
    return LEGIBILITY, BOILERPLATE, NON_CONTENT


_EXT_LEGIBILITY_KIND = {"f-lg-rot13": "garb_shifted_latin", "f-lg-reversed": "garb_reversed",
                        "f-lg-greek-cp1252": "garb_mojibake", "f-lg-russian-cp1252": "garb_mojibake",
                        "f-lg-big5-as-gbk": "garb_mojibake", "f-lg-sjis-as-cp1252": "garb_mojibake"}
_EXT_BOILERPLATE_KIND = {
    "f-bp-notes-title": ("statement_title", "en"), "f-bp-unit-ntd": ("unit_note", "en"),
    "f-bp-prepared-for": ("document_id", "en"), "f-bp-project-rev": ("document_id", "en"),
    "f-bp-court-case": ("document_id", "en"), "f-bp-patient": ("account_identifier", "zh"),
    "f-bp-spec-section": ("section_title", "en"), "f-bp-table-cont": ("continuation_note", "en"),
    "f-bp-safety": ("disclaimer", "en"), "f-bp-effective": ("document_id", "en"),
    "f-bp-exam": ("account_identifier", "en"), "f-bp-contract-no": ("document_id", "en"),
    "f-bp-chapter-few": ("section_title", "en"), "f-bp-brand-pharma": ("brand_name", "en"),
    "f-bp-proprietary": ("confidentiality_mark", "en"), "f-bp-website": ("contact_line", "en"),
    "f-bp-recycled": ("print_note", "en"), "f-bp-zh-contact": ("contact_line", "zh"),
    "f-bp-copyright-short": ("copyright", "en"), "f-bp-slide-brand": ("brand_tagline", "en"),
}


def build_external() -> Dict[str, List[dict]]:
    """The reviewer's fresh items in the schemas of the three sets (split "external")."""
    from doc2mark.ocr.refusal import MAX_JUDGE_CHARS, _normalize, matches_non_content_pattern
    legibility_items, boilerplate_items, non_content_items = _review_pr22_items()
    source = "review-pr22"
    legibility = [_item(id_, id_, source, _EXT_LEGIBILITY_KIND.get(id_, "hard_legible_string"), label, text,
                        _string_spans(text), family=f"{source}:{id_}", route=None, pipeline_asks=None)
                  for id_, label, text in legibility_items]
    boilerplate = []
    for id_, label, text, context in boilerplate_items:
        asked = context["reason"] != "first_occurrence"
        copies = COPIES[label] if asked else "one"
        kind, lang = _EXT_BOILERPLATE_KIND[id_]
        boilerplate.append({"id": id_, "group": id_, "family": f"{source}:{id_}", "source": source, "kind": kind,
                            "lang": lang, "label": 1 if copies == "one" else 0, "expected_copies": copies,
                            "split": "external", "asked": asked, "text": text, "context": dict(context),
                            "reviewer_label": label})
    non_content = []
    for id_, label, text in non_content_items:
        answer = _normalize(text)
        fires = matches_non_content_pattern(text)
        non_content.append({"id": id_, "group": id_, "family": f"{source}:{id_}", "source": source,
                            "kind": "no_content" if label else "content", "label": label, "split": "external",
                            "text": text, "pattern": fires,
                            "judged": bool(answer) and not fires and len(answer) <= MAX_JUDGE_CHARS})
    for items in (legibility, boilerplate, non_content):
        for item in items:
            item["split"] = "external"
    return {"legibility": legibility, "boilerplate": boilerplate, "non_content": non_content}


# ============================================================================ deck slice


def deck_items(path) -> Dict[str, List[dict]]:
    """The uncommitted slice built from a private deck read in place (never write it into the
    repository): every page's text layer (label 1), the spike A transforms of six of its pages
    (label 0), and the repeated lines the pipeline asks about, labelled by their shape (a
    "by <company>" line and a short logo text are chrome; a numbered "NN / title" label is
    content; anything else is left out and listed under "unlabelled"). All items are "test"."""
    _doc2mark()
    path = Path(path)
    legibility = []
    pages = measure_pdf(path)
    for number, page in enumerate(pages, 1):
        legibility.append(_item(f"deck-p{number}", f"deck#p{number}", f"deck#p{number}", "real_deck_slide", 1,
                                page["text"], page["spans"], route=page["route"], pipeline_asks=page["pipeline_asks"]))
    transforms = [(4, "garb_cjk_substituted", lambda: _cjk_offset(1200)), (19, "garb_cjk_substituted", lambda: _cjk_offset(777)),
                  (26, "garb_cjk_substituted", lambda: _cjk_offset(3001)), (23, "garb_pua", lambda: _pua),
                  (6, "garb_fffd_partial", lambda: _fffd(0.2, 4)), (15, "garb_mojibake", lambda: _mojibake)]
    for number, kind, factory in transforms:
        if number > len(pages):
            continue
        rows = _transform_rows(pages[number - 1]["rows"], factory, False)
        spans = [[t, s, f] for line in rows for t, s, f in line]
        legibility.append(_item(f"deck-p{number}-{kind}", f"deck#p{number}", f"deck#p{number} (spike A transform)",
                                kind, 0, _join(rows), spans, route=None, pipeline_asks=None))
    boilerplate, unlabelled = [], []
    asked = {}
    for text, context in _capture(path, {}, oracle=False)[0]:
        asked.setdefault((text, json.dumps(context, sort_keys=True, ensure_ascii=False)), None)
    for index, (text, context_json) in enumerate(asked, 1):
        if re.match(r"^by\s+\S", text):
            kind, label = "brand_tagline", ONE
        elif re.fullmatch(r"[A-Z]{1,6}", text):
            kind, label = "logo_text", ONE
        elif re.match(r"^\d{1,2}\s*/\s*\S", text):
            kind, label = "per_page_label", ALL
        else:
            unlabelled.append({"text": text, "context": json.loads(context_json)})
            continue
        boilerplate.append({"id": f"deck-bp-{index}", "group": "deck", "family": "deck", "source": "deck", "kind": kind,
                            "lang": "zh", "label": label, "expected_copies": COPIES[label], "split": "test",
                            "asked": True, "text": text, "context": json.loads(context_json), "asked_in": "none"})
    for item in legibility:
        item["split"] = "test"
    return {"legibility": legibility, "boilerplate": boilerplate, "unlabelled": unlabelled}


# ================================================================================== main


HOOKS = ("legibility", "boilerplate", "non_content")


def build_all(only: Optional[str] = None, report: Optional[list] = None) -> Dict[str, List[dict]]:
    _doc2mark()
    sets: Dict[str, List[dict]] = {}
    external = build_external()
    with tempfile.TemporaryDirectory(prefix="judge-sets-") as tmp:
        if only in (None, "non_content"):
            sets["non_content"], sets["non_content_ambiguous"] = build_non_content()
        if only in (None, "legibility"):
            sets["legibility"] = build_legibility(Path(tmp))
        if only in (None, "boilerplate"):
            sets["boilerplate"] = build_boilerplate(Path(tmp), report)
    for hook in HOOKS:
        if hook in sets:
            assign_splits(sets[hook], hook, [item for item in external[hook] if item.get("asked", True)])
            sets[f"external/{hook}"] = external[hook]
    return sets


def _print_similarity() -> None:
    for hook, info in SIMILARITY.items():
        print(f"{hook}: {info['families']} families in {info['clusters']} clusters (largest {info['largest_cluster']} "
              f"items; {info['forced_test']} items forced to test by an external near-duplicate); max cross-split "
              f"char-5-gram Jaccard {info['max_jaccard']:.3f}, difflib ratio {info['max_ratio']:.3f}; pairs above "
              f"0.6: {len(info['above_0.6'])}")
        for jaccard, ratio, a, b in info["above_0.6"]:
            print(f"    {a} ~ {b}: Jaccard {jaccard:.3f}, ratio {ratio:.3f}")
        for split in ("train", "test"):
            if f"external_{split}" in info:
                jaccard, ratio, above = info[f"external_{split}"]
                print(f"    external vs {split}: max Jaccard {jaccard:.3f}, ratio {ratio:.3f}; above 0.6: {len(above)}")
                for jac, rat, a, b in above:
                    print(f"        {a} ~ {b}: Jaccard {jac:.3f}, ratio {rat:.3f}")


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", type=Path, default=ROOT / "tests" / "data" / "judge")
    parser.add_argument("--only", choices=HOOKS)
    parser.add_argument("--check", action="store_true", help="rebuild and compare with --out, byte for byte")
    parser.add_argument("--deck", type=Path, help="build the uncommitted deck slice (prints counts)")
    args = parser.parse_args(argv)
    if args.deck:
        slice_ = deck_items(args.deck)
        for name in ("legibility", "boilerplate"):
            print(summary(f"deck {name}", slice_[name]))
        for item in slice_["boilerplate"]:
            print(item["kind"], item["label"], item["context"]["reason"])
        print(f"unlabelled asked lines: {len(slice_['unlabelled'])}")
        return 0
    never: list = []
    sets = build_all(args.only, never)
    status = 0
    for name, items in sets.items():
        path = args.out / f"{name}.jsonl"
        if args.check:
            expected = "".join(_json_line(item) + "\n" for item in items).encode("utf-8")
            same = path.exists() and path.read_bytes() == expected
            print(f"{path}: {'identical' if same else 'DIFFERS'}")
            status |= 0 if same else 1
        else:
            write_jsonl(path, items)
            print(f"wrote {path} ({len(items)} items)")
        if name != "non_content_ambiguous":
            print(summary(name, items))
        print()
    _print_similarity()
    if never:
        repeated = [line for line in never if line["pages"] >= 2]
        one_copy = [line for line in repeated if line["copies_left"] == 1]
        print(f"\nDesigned lines the pipeline does not ask about: {len(never)} texts; repeated on 2+ pages: "
              f"{len(repeated)}, of which the rule keeps exactly one copy: {len(one_copy)} "
              f"({Counter(line['expected_copies'] for line in one_copy)})")
        for line in never:
            print(f"  {line['doc']}: [{line['kind']}, {line['expected_copies']}] {line['text']!r} "
                  f"on {line['pages']} pages, {line['copies_left']} left")
    return status


if __name__ == "__main__":
    sys.exit(main())
