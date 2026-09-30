"""Reproduce the code issues the docs audit found (listed in the PR under "Code issues found").

Each probe prints what the code does; nothing is fixed here. Run it in the E2E image from the
repository root (``eval/docs_audit/run_on_spark.sh`` runs it too)::

    python eval/docs_audit/probes.py
"""
from __future__ import annotations

import contextlib
import io
import logging
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
SAMPLES = ROOT / "sample_documents"
logging.disable(logging.WARNING)


def cli(*args: str, cwd: Path) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, "-m", "doc2mark", *args], cwd=cwd, capture_output=True, text=True)


def probe(title: str):
    def wrap(fn):
        print(f"\n## {title}")
        try:
            fn()
        except Exception as exc:  # a probe must never stop the others
            print(f"probe raised {type(exc).__name__}: {exc}")
        return fn
    return wrap


work = Path(tempfile.mkdtemp(prefix="d2m-probes-"))


@probe("CLI folder run: the default --pattern '*' also matches sub-folders, which fail the run")
def _():
    docs = work / "docs"
    (docs / "sub").mkdir(parents=True)
    shutil.copy(SAMPLES / "sample_text.txt", docs / "a.txt")
    shutil.copy(SAMPLES / "sample_text.txt", docs / "sub" / "b.txt")
    r = cli(str(docs), "-o", str(work / "out1"), "-q", cwd=work)
    print("exit", r.returncode, "| stderr:", r.stderr.strip()[-200:])
    r = cli(str(docs), "-r", "-o", str(work / "out2"), "-q", "--skip-errors", cwd=work)
    print("with --skip-errors: exit", r.returncode, "| written:", sorted(p.name for p in (work / "out2").iterdir()))


@probe("CLI folder output is flat by file stem: same-stem files overwrite each other")
def _():
    docs = work / "stems"
    (docs / "2024").mkdir(parents=True)
    shutil.copy(SAMPLES / "sample_text.txt", docs / "report.txt")
    shutil.copy(SAMPLES / "sample_document.md", docs / "2024" / "report.md")
    r = cli(str(docs), "-r", "--pattern", "report.*", "-o", str(work / "out3"), "-q", cwd=work)
    print("exit", r.returncode, "| inputs: 2 | written:", sorted(p.name for p in (work / "out3").iterdir()))


@probe("CLI --sort size keeps the SMALLEST files, the --help epilog says 'Process 10 largest files'")
def _():
    from doc2mark.cli import filter_files
    docs = work / "sizes"
    docs.mkdir()
    for name, size in (("small.txt", 10), ("big.txt", 5000)):
        (docs / name).write_text("x" * size)
    print([p.name for p in filter_files(list(docs.iterdir()), max_files=1, sort_by="size")])
    r = cli("--help", cwd=work)
    print([line.strip() for line in r.stdout.splitlines() if "largest" in line])


@probe("CLI --preserve-structure and --timeout are parsed but never used")
def _():
    source = (ROOT / "doc2mark" / "cli.py").read_text()
    print("args.preserve_structure used:", "args.preserve_structure" in source)
    print("args.timeout used only as future.result(timeout=...) after as_completed:",
          source.count("args.timeout"), "use(s)")


@probe("batch_process skips upper-case extensions and .htm (load() accepts both)")
def _():
    from doc2mark import UnifiedDocumentLoader
    docs = work / "cases"
    docs.mkdir()
    shutil.copy(SAMPLES / "sample_pdf.pdf", docs / "upper.PDF")
    (docs / "page.htm").write_text("<h1>Hello</h1>")
    shutil.copy(SAMPLES / "sample_text.txt", docs / "plain.txt")
    loader = UnifiedDocumentLoader(ocr_provider=None)
    results = loader.batch_process(docs, save_files=False, show_progress=False)
    print("batch_process found:", sorted(Path(k).name for k in results))
    print("load() works on both:", loader.load(docs / "upper.PDF").metadata.format, loader.load(docs / "page.htm").metadata.format)


@probe("batch results: tables_found is always 0 (ProcessedDocument.tables is never filled)")
def _():
    from doc2mark import UnifiedDocumentLoader
    loader = UnifiedDocumentLoader(ocr_provider=None)
    results = loader.batch_process_files([SAMPLES / "complex-tables" / "complex_table_test.pdf"], save_files=False,
                                         show_progress=False)
    info = next(iter(results.values()))
    doc = loader.load(SAMPLES / "complex-tables" / "complex_table_test.pdf")
    print("tables_found:", info["metadata"]["tables_found"], "| table items:",
          sum(i["type"] == "table" for i in doc.json_content), "| extra tables_count:", doc.metadata.extra.get("tables_count"))


@probe("convenience functions pass **kwargs to load(), not to the loader: table_style= raises")
def _():
    from doc2mark import load
    try:
        load(str(SAMPLES / "sample_pdf.pdf"), table_style="markdown_grid")
    except TypeError as exc:
        print("TypeError:", exc)


@probe("ChunkingConfig.include_page_markers has no effect; every text:section is '##' in chunks")
def _():
    from doc2mark import ChunkingConfig, chunk_content
    items = [{"type": "text:title", "content": "T", "page": 1, "level": 1},
             {"type": "text:section", "content": "Deep", "page": 1, "level": 4},
             {"type": "text:normal", "content": "body", "page": 2}]
    a = chunk_content(items, ChunkingConfig(include_page_markers=True))
    b = chunk_content(items, ChunkingConfig(include_page_markers=False))
    print("same output:", [c.content for c in a] == [c.content for c in b], "|", [c.content for c in a])


@probe("batch_process_files is importable from doc2mark but missing from doc2mark.__all__")
def _():
    import doc2mark
    print("hasattr:", hasattr(doc2mark, "batch_process_files"), "| in __all__:", "batch_process_files" in doc2mark.__all__)


@probe("docstrings rendered in the API reference that contradict the code")
def _():
    import inspect
    import doc2mark
    from doc2mark import UnifiedDocumentLoader
    print("load(): 'requires extract_images=True' in docstring:", "requires extract_images=True" in (doc2mark.load.__doc__ or ""))
    print("UnifiedDocumentLoader.load note 'only work with Office and PDF formats':",
          "only work with Office and PDF formats" in (UnifiedDocumentLoader.load.__doc__ or ""))
    print("CLI epilog 'GPT-4V':", "GPT-4V" in inspect.getsource(__import__("doc2mark.cli").cli))



@probe(".tsv always fails; the CSV delimiter argument is ignored")
def _():
    from doc2mark import UnifiedDocumentLoader
    loader = UnifiedDocumentLoader(ocr_provider=None)
    (work / "t.tsv").write_text("a\tb\n1\t2\n")
    try:
        loader.load(work / "t.tsv")
    except Exception as exc:
        print("tsv:", type(exc).__name__, str(exc)[:120])
    (work / "t.csv").write_text("a\tb\n1\t2\n")
    print("same content as .csv:", loader.load(work / "t.csv").content.splitlines()[:3])
    (work / "semi.csv").write_text("a;b\n1;2\n")
    doc = loader.load(work / "semi.csv", delimiter=",")
    print("delimiter=',' on a ';' file ->", repr(doc.metadata.delimiter))


@probe("PowerPoint: layout placeholder prompts leak into every slide; slide_count is unreliable")
def _():
    from doc2mark import UnifiedDocumentLoader
    doc = UnifiedDocumentLoader(ocr_provider=None).load(SAMPLES / "sample_presentation.pptx")
    print("'Click to edit Master title style' occurrences:", doc.content.count("Click to edit Master title style"))
    print("page_count:", doc.metadata.page_count, "| slide_count:", doc.metadata.slide_count)


@probe("PowerPoint: the text of a plain shape (rectangle) is emitted twice")
def _():
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    from doc2mark import UnifiedDocumentLoader
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(1))
    shape.text_frame.text = "UNIQUE BOX TEXT"
    prs.save(work / "shape.pptx")
    content = UnifiedDocumentLoader(ocr_provider=None).load(work / "shape.pptx").content
    print("'UNIQUE BOX TEXT' occurrences:", content.count("UNIQUE BOX TEXT"))


@probe("Excel: sheet_names and total_cells are not set")
def _():
    from doc2mark import UnifiedDocumentLoader
    doc = UnifiedDocumentLoader(ocr_provider=None).load(SAMPLES / "sample_spreadsheet.xlsx")
    print("sheet_names:", doc.metadata.sheet_names, "| total_cells:", doc.metadata.total_cells, "| page_count:", doc.metadata.page_count)


@probe("an unknown table_style: PDF conversion fails, Office silently falls back to the basic converter")
def _():
    from doc2mark import UnifiedDocumentLoader
    loader = UnifiedDocumentLoader(ocr_provider=None, table_style="MARKDOWN_GRID")
    doc = loader.load(SAMPLES / "complex-tables" / "complex_table_test.docx")
    print("docx: colspan in output:", "colspan" in doc.content, "| json_content:", doc.json_content is not None)
    try:
        loader.load(SAMPLES / "complex-tables" / "complex_table_test.pdf")
        print("pdf: converted")
    except Exception as exc:
        print("pdf:", type(exc).__name__, str(exc)[:100])


@probe("Markdown: a file that starts with a --- rule loses its first block to frontmatter")
def _():
    from doc2mark import UnifiedDocumentLoader
    (work / "rule.md").write_text("---\nIntro paragraph\n---\n# Title\n\nBody\n")
    doc = UnifiedDocumentLoader(ocr_provider=None).load(work / "rule.md")
    print("content:", repr(doc.content[:40]), "| frontmatter:", repr(doc.metadata.frontmatter))


@probe("ProcessedDocument.text removes '# ' anywhere")
def _():
    from doc2mark import OutputFormat, UnifiedDocumentLoader
    (work / "c.txt").write_text("I write C# code.\n")
    print(repr(UnifiedDocumentLoader(ocr_provider=None).load(work / "c.txt", output_format=OutputFormat.TEXT).content))


@probe("chunker: footnotes starting with the same 20 characters collide; a title-less document keeps its first section as a parent")
def _():
    from doc2mark import chunk_content
    items = [{"type": "text:section", "content": "A", "page": 1}, {"type": "text:normal", "content": "a", "page": 1},
             {"type": "text:section", "content": "B", "page": 1}, {"type": "text:normal", "content": "b", "page": 1},
             {"type": "text:footnote", "content": "* Source: company filings 2024", "page": 1},
             {"type": "text:footnote", "content": "* Source: company filings 2025", "page": 1}]
    chunks = chunk_content(items)
    print("hierarchies:", [c.section_hierarchy for c in chunks])
    joined = "\n".join(c.content for c in chunks)
    print("footnotes kept:", joined.count("company filings"), "of 2")


@probe("OCR facade and loader: settings that do not reach the Vertex AI provider")
def _():
    from doc2mark import OCR, UnifiedDocumentLoader
    try:
        OCR("vertex_ai", project="my-gcp-project")
    except TypeError as exc:
        print("OCR('vertex_ai', project=...):", "TypeError:", exc)
    print("OCR('vertex_ai', model='gemini-2.0-flash').model ->", OCR("vertex_ai", model="gemini-2.0-flash")._provider.model)
    gemini = UnifiedDocumentLoader(ocr_provider="gemini", model="gemini-2.0-flash", location="europe-west4").ocr
    print("loader 'gemini': model", gemini.model, "| location", gemini.location)


@probe("OCR cache key ignores OCRConfig.task")
def _():
    from doc2mark import OCRConfig, Task
    from doc2mark.ocr.cache import build_ocr_cache_key
    from doc2mark.ocr.openai import OpenAIOCR
    a = OpenAIOCR(api_key="k", config=OCRConfig(task=Task.AUTO))
    b = OpenAIOCR(api_key="k", config=OCRConfig(task=Task.RECEIPT))
    print("same key for task=auto and task=receipt:", build_ocr_cache_key(a, b"img") == build_ocr_cache_key(b, b"img"))


@probe("non_content_judge values outside [0, 1] are not rejected")
def _():
    from doc2mark.ocr.refusal import non_content_reason
    print("judge returning 1.5 ->", non_content_reason("Totals by region: see chart", judge=lambda text: 1.5))
    print("judge returning None ->", non_content_reason("Totals by region: see chart", judge=lambda text: None))


@probe("loader top_p / frequency_penalty / presence_penalty never reach the OpenAI request")
def _():
    import os
    from tests.e2e import fake_openai as fo
    from doc2mark import UnifiedDocumentLoader
    from tests.e2e import pdfgen
    with fo.FakeOpenAI() as fake:
        fake.script(structured=[fo.page("HELLO")], free_form=[fo.text("HELLO")])
        os.environ.update(fake.env)
        loader = UnifiedDocumentLoader(ocr_provider="openai", top_p=0.1, frequency_penalty=0.5, presence_penalty=0.5)
        pdfgen.image_pdf(work / "p.pdf", "HELLO")
        loader.load(work / "p.pdf", ocr_images=True)
        body = fake.requests[0]
        print("request keys:", sorted(k for k in body if k != "messages"))


@probe("create_ocr_cache('redis') without a URL quietly returns a memory cache")
def _():
    from doc2mark import create_ocr_cache
    print(type(create_ocr_cache("redis")).__name__)


@probe("batch_process with extract_images=True reports a PDF with pictures as failed (base64 str written in 'wb' mode)")
def _():
    from doc2mark import UnifiedDocumentLoader
    loader = UnifiedDocumentLoader(ocr_provider=None)
    results = loader.batch_process_files([SAMPLES / "sample_pdf.pdf"], output_dir=work / "imgout", extract_images=True,
                                         show_progress=False)
    info = next(iter(results.values()))
    print(info["status"], "|", info.get("error"), "| .md written:", (work / "imgout" / "sample_pdf.md").exists())


shutil.rmtree(work, ignore_errors=True)
